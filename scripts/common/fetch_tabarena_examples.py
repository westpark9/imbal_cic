"""Download inspectable examples, without running or changing an IDS experiment.

Keep UCI originals and the official TabArena/BeyondArena release distinct.
The latter is not asserted byte-identical to the OpenML v0.1 release.
"""

import argparse
import hashlib
import io
import json
import re
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HF_REV = "2ecfe882ccfb814fc27c4de10a64ceefd5d7655c"
GIT_REV = "4b055decdd272314cf44078d98eafc37b105d785"
HF_BASE = f"https://huggingface.co/datasets/TabArena/BeyondArena/resolve/{HF_REV}"
META_URL = (
    f"https://raw.githubusercontent.com/autogluon/tabarena/{GIT_REV}/packages/"
    "tabarena/src/tabarena/benchmark/task/metadata/data/curated_tabarena_dataset_metadata.csv"
)
EXAMPLES = [
    ("credit_g", "credit-g", "019d7367-39a7-7335-8448-2dc9b2872361", "금융 · 이진분류 · 수치/범주 혼합"),
    ("maternal_health_risk", "maternal_health_risk", "019d7368-bc0d-723e-a4f9-abfd36e085e9", "의료 · 3개 위험군 분류 · 소수 수치 특징"),
    ("concrete_compressive_strength", "concrete_compressive_strength", "019d7367-0f12-7c25-b411-de9c711bd4ec", "재료 · 연속값 회귀 · 수치 특징"),
    ("mic", "MIC", "019d7444-22cb-7d04-940f-6fa8b0436ea2", "의료 · 불균형 다중분류 · 결측값/다수 특징"),
]
FILES = [
    "dataset.parquet", "container_metadata.json", "dataset_metadata.dataset-mold-v1.json",
    "task_metadata.predictive-ml-task-mold-v1.json", "dtypes.json",
    "experiment_metadata.predictive-ml-splits-mold-v1.json",
]


def download(url, dest):
    response = requests.get(url, timeout=(10, 45))
    response.raise_for_status()
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(response.content)
    return {"url": url, "sha256": hashlib.sha256(response.content).hexdigest(), "bytes": len(response.content)}


def fetch_example(entry, output, curated):
    name, arena_name, uuid, reason = entry
    folder = output / name
    sources = {}
    for filename in FILES:
        dest = folder / "official_release" / filename
        sources[str(dest.relative_to(output))] = download(f"{HF_BASE}/{name}/{uuid}/{filename}", dest)
    meta = json.loads((folder / "official_release/dataset_metadata.dataset-mold-v1.json").read_text())
    task = json.loads((folder / "official_release/task_metadata.predictive-ml-task-mold-v1.json").read_text())
    urls = re.findall(r"https://archive\.ics\.uci\.edu/[^\s]+?\.zip", meta["download_description"])
    if len(set(urls)) != 1:
        raise ValueError(f"Ambiguous original source for {name}: {urls}")
    archive = folder / "uci_original/source.zip"
    sources[str(archive.relative_to(output))] = download(urls[0], archive)
    with zipfile.ZipFile(archive) as z:
        for member in z.infolist():
            target = archive.parent / member.filename
            if not target.resolve().is_relative_to(archive.parent.resolve()):
                raise ValueError(f"Unsafe archive path: {member.filename}")
        z.extractall(archive.parent)
    df = pd.read_parquet(folder / "official_release/dataset.parquet")
    target = task["target_column_name"]
    assert target in df.columns and not df[target].isna().any(), name
    assert df.columns.is_unique, name
    df.to_csv(folder / "data.csv", index=False, encoding="utf-8-sig")
    schema = pd.DataFrame([
        {"column": c, "role": "target" if c == target else "feature", "dtype": str(df[c].dtype),
         "missing": int(df[c].isna().sum()), "nunique": int(df[c].nunique()),
         "example_values": " | ".join(str(x) for x in df[c].dropna().unique()[:5])}
        for c in df.columns
    ])
    schema.to_csv(folder / "schema.csv", index=False, encoding="utf-8-sig")
    row = curated.loc[curated.dataset_name == arena_name]
    if len(row) != 1:
        raise ValueError(f"No unique TabArena metadata row: {arena_name}")
    official = json.loads(row.to_json(orient="records"))[0]
    result = {
        "name": name, "selection_reason": reason, "rows": len(df), "input_features": len(df.columns) - 1,
        "target": target, "task": task["problem_type"], "metric": task["objective_metric_name"],
        "missing_cells": int(df.isna().sum().sum()), "duplicate_rows": int(df.duplicated().sum()),
        "license": meta["license"], "curation_comments": meta["curation_comments"],
        "upstream_url": meta["original_dataset_source_download_link"], "tabarena_v0_1_metadata": official,
        "matches_v0_1_row_count": len(df) == int(official["num_instances"]),
        "matches_v0_1_feature_count": len(df.columns) == int(official["num_features"]),
        "openml_content_identity": "UNVERIFIED: OpenML unavailable; HF release is not assumed identical",
        "hf_revision": HF_REV, "sources": sources,
    }
    if "classification" in task["problem_type"]:
        result["target_counts"] = {str(k): int(v) for k, v in df[target].value_counts().items()}
    else:
        result["target_summary"] = df[target].describe().to_dict()
    (folder / "manifest.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: result[k] for k in ["name", "rows", "input_features", "matches_v0_1_row_count", "matches_v0_1_feature_count"]}), flush=True)
    return result, df, schema


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="17365D")
        c.font = Font(color="FFFFFF", bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(c.value, str):
                c.data_type = "s"
    for col in ws.columns:
        width = min(55, max(14, max(len(str(c.value or "")) for c in list(col)[:30]) + 2))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    provenance = download(META_URL, output / "_metadata/curated_tabarena_dataset_metadata.csv")
    curated = pd.read_csv(output / "_metadata/curated_tabarena_dataset_metadata.csv")
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda e: fetch_example(e, output, curated), EXAMPLES))
    wb = Workbook()
    overview = wb.active
    overview.title = "안내_데이터목록"
    overview.append(["데이터", "선정 이유", "배포본 행 수", "입력 특징 수", "목표 열", "과제", "분류 빈도 또는 회귀 요약", "v0.1 행/특징 수 일치", "원 출처", "공식 정제 설명"])
    schemas = wb.create_sheet("열_설명")
    schemas.append(["데이터", "열 이름", "역할", "자료형", "결측 수", "고유값 수", "값 예시"])
    for result, df, schema in results:
        overview.append([result["name"], result["selection_reason"], len(df), len(df.columns) - 1, result["target"], result["task"],
                         json.dumps(result.get("target_counts", result.get("target_summary")), ensure_ascii=False),
                         f'{result["matches_v0_1_row_count"]} / {result["matches_v0_1_feature_count"]} (값 동일성은 미검증)',
                         result["upstream_url"], result["curation_comments"]])
        ws = wb.create_sheet(result["name"][:31])
        ws.append(list(df.columns))
        for record in df.head(200).itertuples(index=False, name=None):
            ws.append([None if pd.isna(v) else v for v in record])
        for record in schema.itertuples(index=False, name=None):
            schemas.append([result["name"], *record])
    for ws in wb:
        style_sheet(ws)
    wb.save(output / "TabArena_대표데이터.xlsx")
    (output / "안내.txt").write_text(
        "TabArena 수록 데이터 4종 열람 자료\n\n"
        "TabArena_대표데이터.xlsx: 데이터 개요·열 설명과 각 데이터 앞 200행. data.csv는 전체 행.\n"
        "각 데이터/uci_original/: UCI에서 받은 원본 ZIP과 해제 파일.\n"
        "각 데이터/official_release/: TabArena 팀의 공식 BeyondArena 배포본 Parquet 및 처리·분할 메타데이터.\n"
        "각 데이터/data.csv: 위 공식 배포본을 값 변경·보간·표본추출 없이 CSV로 변환.\n"
        "각 데이터/schema.csv, manifest.json: 목표 열·자료형·결측·분포·출처·라이선스·체크섬.\n\n"
        "주의: TabArena v0.1 OpenML 서버는 이 환경에서 연결 시간 초과였다.\n"
        "동일 이름의 공식 BeyondArena 배포본과 UCI 원본을 확보했다. v0.1과 행·열 수를 비교했으나\n"
        "OpenML 파일과 행별 값·순서·분할까지 동일하다고 검증한 것은 아니다. 차이는 manifest에 명시했다.\n"
        "이 자료는 원본 구조 열람용이며 IDS 실험에 투입하거나 벤치마크 재현에 사용하지 않았다.\n",
        encoding="utf-8",
    )
    (output / "_metadata/download_manifest.json").write_text(json.dumps({
        "retrieved_at_utc": datetime.now(timezone.utc).isoformat(), "tabarena_git_revision": GIT_REV,
        "hf_revision": HF_REV, "curated_metadata_source": provenance,
        "datasets": [r for r, _, _ in results], "experiments_run": False,
    }, ensure_ascii=False, indent=2) + "\n")


if __name__ == "__main__":
    main()
