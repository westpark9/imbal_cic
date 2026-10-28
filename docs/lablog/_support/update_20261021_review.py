"""Reconcile the user's literature workbook with the papers read on 2026-10-21."""
from pathlib import Path
from copy import copy
import json, hashlib, datetime, csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

ROOT=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/lablog.py').exists())
E=ROOT/'docs/lablog/evidence/references/20261021'
PATH=E.parent/'tabpfn_papers_survey.xlsx'
RUN=ROOT/'tabpfn/results/v7/exp86/20261021_232457_attack_fewshot_s43_47_0a849540'

def main():
    w=openpyxl.load_workbook(PATH);changes=[];before=hashlib.sha256(PATH.read_bytes()).hexdigest()
    original_values=list(w['Lit TabPFN'].values)
    def setcell(s,r,c,value):
        old=s.cell(r,c).value
        if old != value:
            changes.append(dict(sheet=s.title,cell=s.cell(r,c).coordinate,before=old,after=value));s.cell(r,c,value)
    def update(s,r,values):
        for c,v in values.items():setcell(s,r,c,v)
    s=w['20261021_추가문헌']
    update(s,3,{11:'§3.1: 사전학습 자료(합성/정제 실표/웹 표/KB) × 전이(ICL/PFN, FT encoder, text generation, weight generation). §4.5 context 선택·학습.',
        16:'최종 포함 146편·기반 모델 42개; synthetic 26/42, ICL/PFN 28/42. 검색 기록 1,148건과 포함 목록을 구분.',
        17:'수집 기준일 2026-09-15. 정적 single-table 분류/회귀가 핵심 범위. 보안 응용 중 검색 기록에만 있는 논문을 포함 문헌이라고 부르지 않음.',
        22:'사용자 제공 PDF 전문 확인: 20261021/Tabular Foundation Models: A Systematic Surve.pdf. §3.1 p.5; §4 pp.11–14; §5 pp.14–15; §6.3 p.19.'})
    update(s,2,{16:'데이터 구성·평가 근거를 원문 pp.3–4,6–8에서 확인. known의 균형 라벨 자료와 known+unknown 비라벨 long-tail; 군집 대응에는 Hungarian 사용.',
        17:'p.7 Table 1 CTU-13 2.8M vs 본문 1.3M, long-tail 배분식/정규화, header/payload 설명의 구현 대조 필요. 독립 held-out 및 전처리 중복 통제는 미확정.',
        22:'원문·공식 배포처 대조 완료. 자세한 근거와 PCAP 후속안: docs/lablog/reports/20261021_references.tex. 저자 코드 공개 API 404, IoT-23 PCAP/라벨 공개, CTU-13 전체 정상 PCAP 비공개, ISCX2014 공식 다운로드 불가.'})
    update(s,5,{11:'v3 전체 구조 유지; Fourier 수치+ECDF rank encoding, classification/regression 공동 checkpoint, CLS 4→8·폭512→1024. 별도 quantile/robust/SVD 전처리 제거. default8/Fast4 estimators.',
        16:'§3.1–3.4 pp.9–11을 확인. 모델 용량·입력 표현·기본 estimator 변경을 현재 v3와 구분. Plus/Thinking 세부 구현은 비공개.',
        22:'20261021/tabpfn35_2609.17895.pdf v2 확인: §3 pp.9–11, §4 pp.11–12. 이번 EXP86은 v3 유지.'})
    update(s,6,{22:'20261021/streams_2610.05352.pdf §2–4 pp.2–5 및 Appendix D 검토. prequential과 고정 test 구분, 원문 비용을 우리 v3에 직접 일반화하지 않음.'})
    title='Ensembling Tabular Foundation Models - A Diversity Ceiling And A Calibration Trap'
    found=[r for r in range(2,s.max_row+1) if s.cell(r,3).value==title]
    r=found[0] if found else s.max_row+1
    vals=['20261021 추가','TFM-Ensemble-Diversity',title,
        '6개 TFM의 오류 중복·보정·추가 계산을 분석해 앙상블 필요성을 검증했다.','앙상블·보정·비용','프리프린트','2026-05-18',None,
        '모델 수를 늘리는 것이 예측 다양성과 계산 대비 품질 이득을 보장하는지 불명확했다.',
        '앙상블의 이득과 오류 중복·calibration 비용을 분리 평가.',
        '가중평균, greedy selection, stacking, temperature blending, cascade, 3-seed ensemble. 내부 validation/OOF 사용.',
        '153 OpenML classification datasets; TabArena subset의 보정 지표.',
        'TabPFNv2.5/v2.6, TabICL/v2, LimiX, OrionMSPv1.5',
        '80/20 stratified split, train 내부75/25 validation; OOF folds는 방법별 구분.',
        'Accuracy, weighted F1, AUC, log-loss, fit time; ECE/AURC/worst-group 진단.',
        '§4–5: 오류 상관이 높았고 일부 앙상블의 추가 비용 대비 품질 이득이 제한적이었다.',
        '특정 pool·153과제·분할 결과이며 모든 앙상블 실패를 입증하지 않음. TabFM+의 입력 view ensemble과 구별.',
        'arXiv:2605.18696',2026,'https://arxiv.org/pdf/2605.18696',
        '현재 expert의 상호 보완 오류와 oracle 상한부터 확인. 같은 global context·총 라벨·보존량·추론예산에서 결합기/보정의 기여를 분리.',
        '20261021/2605.18696.pdf §3–5 pp.3–4. Systematic Survey 최종 포함 문헌(analysis_theory).']
    update(s,r,{i+1:v for i,v in enumerate(vals)})

    s=w['20261021_근거와요청목록']
    update(s,6,{2:'사용자 제공 전문 검토 완료',3:'PDF §3.1 p.5; §4 pp.11–14; §6.3 p.19',5:'2축 분류와 연구 확장 분야를 사용. 검색 기록1148과 최종포함146 구분.',6:'서베이 원문 추가 요청 해소. IDS 응용은 최종 포함 여부를 별도 표시.'})
    update(s,8,{2:'원문 구조·전처리 절 검토 완료',3:'§3.1–3.4 pp.9–11',6:'v3.5 실행/모델 교체는 이번 범위 아님.'})
    update(s,10,{6:'원문 §1–2 IID 범위·반복분할·지표 검토. OpenML/HF 행별 내용 동일성은 미검증 유지.'})
    update(s,12,{1:'OpenReview 첨부 = In-Context Data Distillation with TabPFN',2:'기존 Lit TabPFN!61에 연결; 중복 추가하지 않음',
        3:'5_In_Context_Data_Distillation.pdf pp.1–4',4:None,
        5:'고정 PFN으로 distilled X를 최적화. 라벨 수와 context 압축 크기는 다름. tuned/default XGB 구분.',
        6:'필수 원문 요청 해소; 저자 공식 구현은 이번 확인에서 미확인.',7:'https://arxiv.org/abs/2402.06971'})
    extra=[['LotBoNC 원문·공식 통계 대조','검토 완료 / 재현 미확정','PDF p.7 Table1·§5.1 및 공식 IoT-23/CTU-13',None,
            'CTU-13 flow 수 본문/표 불일치; IoT-23 공식 20 captures를 20 attack types로 해석하지 않음.',
            '정확한 클래스 ID·flow timeout·표본 배분·분할 seed·저자 코드 필요.','https://www.stratosphereips.org/datasets-iot23'],
           ['TabArena 원본 4종','다운로드·검증 완료','data/tabarena_examples/_metadata/download_manifest.json',None,
            'UCI 원본 + 공식 HF parquet/metadata + CSV + 열람용 Excel; 이진/다중/회귀/결측불균형 사례.',
            'IDS 실행에는 사용하지 않음.','data/tabarena_examples/TabArena_대표데이터.xlsx']]
    for row in extra:
        found=[r for r in range(2,s.max_row+1) if s.cell(r,1).value==row[0]];r=found[0] if found else s.max_row+1
        update(s,r,{i+1:v for i,v in enumerate(row)})

    s=w['20261021_선행실험안']
    run_state=json.loads((RUN/'results.json').read_text())
    update(s,2,{3:f"승인 후 EXP86 실행 중 ({len(run_state['jobs'])}/100 완료; 최신 상태는 실행 경로 progress.json)"})
    for r in [4,6,7,8,14,15]:setcell(s,r,3,'확정·실행에 반영')
    setcell(s,9,3,'코드 점검 및 표본/가중치/예측 검증에 반영')
    setcell(s,18,3,'2026-10-21 23:24 KST 시작; 자동 보고서 생성·임시 컴파일 검증 확인')
    setcell(s,19,3,'docs/lablog/reports/20261021_exp86.tex 자동 갱신')
    r=s.max_row+1
    if not any(s.cell(i,1).value=='실행 경로' for i in range(1,r)):
        update(s,r,{1:'실행 경로',2:str(RUN.relative_to(ROOT)),3:'v7/EXP86, 5 seeds; 로그·예측·표본 ID·해시 보존'})

    review='20261021_선별검토'
    if review in w:del w[review]
    s=w.create_sheet(review)
    headers=['문헌/단위','서베이 분류·포함 관계','확인한 기제','라벨/학습/추론 비용','현재 연구의 gap·Discussion','근거 페이지·파일','원문 검토 수준']
    s.append(headers)
    entries=[
        ['ICD','최종포함: context_management','모델θ 고정, distilled 특징 학습','큰 labeled train + 1000 distilled context; 역전파/최적화 비용',
         '작은 저장 context와 극소수 실제 라벨 가용성을 구별. EXP86의 전제를 대신하지 못함.','ICD PDF §3 p.3, §4 p.4; 기존61행','원문 해당 절 검토'],
        ['LoCalPFN','최종포함: context_management','query kNN context + end-to-end FT','검색 pool·가중치 FT·query별 context 처리 비용',
         '지역성 이점과 FT 기여를 분리. 공유 context보다 처리 비용이 늘 수 있음.','localpfn_neurips2024.pdf §2.2–2.4 pp.3–5; 기존36행','원문 해당 절 검토'],
        ['BoostPFN','기존 관련 모델 보충','오분류/잔차에 따라 context 표본 가중치·앙상블 계수 갱신','큰 train pool·여러 PFN pass·계수 선택',
         '대표성 개선과 추가 계산 효과 분리. 작은 각 context가 작은 전체 라벨 예산을 뜻하지 않음.','boostpfn_2503.01256.pdf §4 pp.3–4; 기존27행','원문 해당 절 검토'],
        ['DistPFN','최종포함: analysis_theory','prior와 posterior 비율 보정; FT 요구 없음','개별 query 또는 배치 평균 posterior; 원문 기본은 배치 평균',
         '균형 context·고정 test에서의 prior 변화와 특징 구분 실패를 구별. oversampling은 label-shift 평가 구성.','distpfn_2605.04363.pdf §4 pp.4–5; 기존21행','원문 해당 절 검토'],
        ['PFN-Imb','최종포함: context_management','downsampling·thresholding 등 비교','11 binary tasks, balanced test, TabPFN2.5',
         '다중공격·자연 test분포의 tail 개선으로 일반화하지 않음. 정상 FPR tradeoff 필요.','2605.21742.pdf §II–III pp.2–3; 기존71행','원문 해당 절 검토'],
        ['TabFM+','사용자 요청 보충: 모델·앙상블','32 views + NNLS 검증 가중치·보정','backbone 고정이나 validation·후처리 학습·여러pass 사용',
         '극소수 라벨에서 validation 확보비용과 입력 다양성 이점을 따로 검증.','tabfm_2609.37959.pdf §4 pp.8–9','원문 해당 절 검토'],
        ['TFM-Ensemble-Diversity','최종포함: analysis_theory','6 TFM/6 결합법의 오류중복·보정·비용 분석','153 classification tasks; validation/OOF 추가 계산',
         '우리 expert 오류 상보성·oracle 확인 후 gating 설계. 전체 앙상블 실패로 확대하지 않음.','2605.18696.pdf §3–5 pp.3–4','원문 해당 절 검토'],
        ['StreamTFM-Study','서베이 기준일 이후 보충','prequential, label delay, FIFO/DualFIFO/CURE','메모리·backbone·batch·serving cost 통제',
         '현재 고정 test/라벨 공개와 차이. v1 등의 비용 결과를 v3로 일반화하지 않음.','streams_2610.05352.pdf §2–4 pp.2–5','원문 해당 절 검토'],
        ['SmartCity Forensics','서베이 검색 기록에만 포함: 보안 응용 보충','ToN 장치 telemetry, PFN triage + RF/GB 제안','stratified80/20+5fold; PFN 준비비용 측정 범위 불명확',
         '우리 NF-ToN 특징·공격family와 다름. cascade 전체효과와 모델별 표 구분.','2604.11394.pdf §3–4 pp.4–7; 기존120행','원문 해당 절 검토'],
        ['RF-TabPFN / IIoT / FM-CyberEval','보안 응용 보충 후보','기존 workbook 요약 유지','기존 발췌/스니펫 범위를 넘는 새 주장 없음',
         '데이터 단위·few-shot 수·split/전처리·노출 라벨·전체 cascade 평가를 전문으로 대조해야 함.','기존79/88/77행','이번 전문 재검토 미완료'],
        ['LotBoNC','사용자 요청 보안 데이터 구성','균형 known labelled + long-tail known/unknown unlabelled','packet/burst encoder; SSL pretraining + fine-tuning',
         '소수 신규 라벨 분류와 novel-category discovery는 다름. PCAP 구성 재현에 참고.','../20261015/lotbonc.pdf §§4.2,5 pp.3–4,6–8','원문 및 공식 출처 대조'],
        ['TabArena','벤치마크','IID 분할/과제별 지표/선별 데이터','default/HPO/ensemble 조건 구별',
         '라벨 예산·보존 예산·시간순 적응 질문은 별도. 실제 표 4종을 열람 가능하게 보존.','tabarena_2506.16791.pdf §1–2 pp.2,4–6; 기존81행','원문·배포본 검토']]
    for row in entries:s.append(row)
    for name in ['20261021_추가문헌','20261021_근거와요청목록','20261021_선행실험안',review]:
        s=w[name];s.freeze_panes='D2' if name.endswith('추가문헌') else 'B2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244763');c.alignment=Alignment(wrap_text=True,vertical='top')
        for row in s.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
            s.row_dimensions[row[0].row].height=105 if name.endswith('추가문헌') else 86
        for col in s.columns:s.column_dimensions[col[0].column_letter].width=24 if col[0].column==1 else 48
        s.row_dimensions[1].height=34
    # Preserve original notes, add explicit review pointers rather than overwrite user prose.
    original=w['Lit TabPFN']
    for row,comment in {
        21:'2026-10-21 원문 §4 pp.4–5 재확인: inverse-frequency oversampling은 label-shift 실험 자료 구성. DistPFN 자체는 posterior 보정이며 FT 없음. 원문 기본은 배치 평균 posterior 사용. 선별검토 시트 참조.',
        61:'사용자 제공 OpenReview 5_In_Context_Data_Distillation.pdf와 동일 논문. 원문 pp.1–4 검토. 큰 train을 써서 1000 distilled context를 최적화하므로 실제 라벨 n=1–20과 구분.',
        71:'2026-10-21 원문 확인: TabPFN2.5, 11 binary tasks, test 각 클래스500, imbalance/threshold 비교. 다중공격 자연분포 평가와 구분.',
        81:'2026-10-21 원문 §1–2 및 실제 원본4종 검토. data/tabarena_examples/TabArena_대표데이터.xlsx, 상세 보고서20261021_references.tex 참조.',
        120:'2026-10-21 원문 §§3–4 확인: ToN 장치 telemetry로 우리 NF-ToN과 다름. Table3의 training0을 context 준비0으로 해석하지 않음.'}.items():
        c=original.cell(row,3);old=c.comment.text if c.comment else ''
        if comment not in old:c.comment=Comment((old+'\n\n'+comment).strip(),'Codex review')
    tmp=PATH.with_name(PATH.stem+'.tmp.xlsx');w.save(tmp);tmp.replace(PATH)
    # Validate archive and unchanged original sheet values (comments are intentionally added).
    final=openpyxl.load_workbook(PATH);assert list(final['Lit TabPFN'].values)==original_values
    audit=dict(updated_at=datetime.datetime.now().astimezone().isoformat(),before_sha256=before,
        after_sha256=hashlib.sha256(PATH.read_bytes()).hexdigest(),changed_cells=changes,
        added_review_sheet=review,original_sheet_values_preserved=True,
        original_comments_added_to_title_rows=[21,61,71,81,120],new_paper=title,icd_duplicate_added=False)
    (E/'review_workbook_update.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2,default=str)+'\n')
    print(json.dumps(dict(sheets=final.sheetnames,changed_cells=len(changes),review_rows=len(entries),sha256=audit['after_sha256']),ensure_ascii=False))

if __name__=='__main__':main()
