"""Consolidate the user's imbal / TabPFN / IDS workbook, preserving source notes.

Run from the project root. Inputs are archived once before replacing the active
workbooks; the TabPFN-only workbook becomes a projection of the integrated book.
"""
from __future__ import annotations

from collections import Counter
from copy import copy
import hashlib
import json
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
import zipfile

import openpyxl

from reformat_literature_20261027 import (
    HEADERS, GROUPS, EXTRA_CHAPTERS, MUTED, collect_tabpfn, render_sheet, date_key,
)

IMBAL_GROUPS = [
    (1, "01 원본 유형 1 — C열 분류 유지"),
    (1, "02 원본 유형 2 — C열 분류 유지"),
    (1, "03 원본 유형 3 — C열 분류 유지"),
    (2, "01 서베이·분류 체계"), (2, "02 데이터·라벨 노이즈"), (2, "03 OOD 탐지"),
    (3, "01 손실·마진·재가중·분류기 보정"),
    (3, "02 샘플링·생성·데이터 증강"),
    (3, "03 표현·대조학습·구조 설계"),
    (3, "04 전문가·앙상블·지속학습"),
    (3, "05 GNN·그래프 불균형"),
    (3, "06 회귀·다중라벨·검출 등 확장"),
    (3, "07 능동학습·프라이버시·강건성"),
]
SOTA_GROUPS = {
    "IMMAX": 1, "LSR": 1, "PRIME": 6, "DP": 7, "DIRECT": 7, "DCE": 4,
    "IP2SL": 2, "MORE": 6, "LLM-AutoDA": 2, "EPIC": 2, "DiffMatch": 6,
    "TTC": 3, "LFN": 2, "ADR": 2, "TAET": 7, "TinyChange": 4, "MGS": 3,
    "NodeImport": 5, "ORD": 2, "DBM": 1, "BCE3S": 1, "LOS": 1,
    "ConMix": 3, "LTRL": 3, "DaSC": 7, "RRB": 6, "SIMBA": 5,
    "AD-GSMOTE": 5, "MATI": 4, "GraphMuGu": 5, "DILDL": 6,
    "LT-SpikeFormer": 3, "CFAMG": 2, "OVO-OVR": 1, "Uni-GNN": 5,
    "IP-DPP": 2, "IBCL": 1, "OSB-SMOTE": 2, "ImGDA": 5, "RLAL": 3,
    "GraphIAM": 5,
}
IDS_GROUPS = [
    (1, "01 IDS 개론·서베이"), (1, "02 데이터셋·평가 구성"),
    (2, "01 Transformer·트래픽 사전학습"),
    (2, "02 GNN·그래프 표현학습 (Transformer 결합 포함)"),
    (2, "03 표 데이터·특징 기반 ML·앙상블"),
    (2, "04 CNN·DNN·ANN·Random Neural Network"),
    (2, "05 RNN·시퀀스 모델"), (2, "06 Autoencoder·재구성"),
    (2, "07 대조·자기지도 표현학습 (기반 구조 확인 필요)"),
    (2, "08 Provenance·규칙·시그니처 (GNN과 구별)"),
    (2, "09 통계·산업 공정 모델"),
    (3, "01 설명가능성·강건성·보조 평가"),
    (3, "02 기반 모델 확인 필요"),
]
IDS_MODEL = {
    "CoLD": 3, "Entente": 2, "GeCo": 9, "CertTA": 0, "Feco": 7,
    "CAPTAIN": 8, "GraphIDS": 2, "E-GRACL": 2, "BS-GAT": 2, "XMF-GNN": 2,
    "AOC-IDS": 6, "RIDS": 5, "SSID": 4, "FlowTransformer": 1,
    "Cloud-Transformer (목록명)": 1, "GTID": 1, "xNIDS": 0, "DISTDET": 8,
    "GIPS": 8, "RUIDS": 1, "NEGSC": 2, "TrafficFormer": 1,
    "TS-IDS": 2, "CIDS": 7, "Anomal-E": 2, "E-GraphSAGE": 2,
}
IDS_INTRO = {
    "ankit_2019": (1, 1), "bayesh_2024": (3, 2), "bandarupalli_2025": (2, 4),
    "vinayakumar_2019": (2, 4), "lee_2019": (2, 4), "zhang_2020": (2, 4),
    "sharafaldin2018toward": (1, 2), "zhou2020building": (2, 3),
    "liu2019machine": (1, 1), "ring2019survey": (1, 1),
    "gharib2016evaluation": (1, 2), "lashkari2017characterization": (1, 2),
    "hewapathirana2025comparative": (2, 3), "abdulhammed2019features": (2, 3),
    "moustafa2015unsw": (1, 2), "tavallaee2009detailed": (1, 2),
    "kdd_cup_1999_data_130": (1, 2), "moustafa2017big": (2, 9),
}
SUMMARIES = {
    "mixup": "입력과 라벨을 선형 보간해 학습 자료를 증강했다.",
    "OLTR": "미지 클래스가 존재하는 open-world 환경의 long-tail 인식을 다뤘다.",
    "Class-Balanced": "유효 표본 수에 기반한 클래스 균형 손실을 제안했다.",
    "LDAM": "라벨 분포에 따른 마진 손실로 불균형 학습을 다뤘다.",
    "BS\nBALMS": "long-tail 인식을 위한 balanced meta-softmax를 제안했다.",
    "BBN": "두 분기 네트워크와 누적 학습으로 long-tail 인식을 다뤘다.",
    "cRT\nt-norm\nLWS": "표현 학습과 분류기 학습을 분리해 long-tail 인식을 평가했다.",
    "LFME": "여러 전문가의 지식을 자기주도 속도 조절 방식으로 증류했다.",
    "Remix": "클래스 불균형을 고려한 mixup 증강을 제안했다.",
    "LEAP-CF\nLEAP-AF": "학습 가능한 임베딩 증강으로 long-tail 표현 학습을 다뤘다.",
    "RIDE": "분포에 반응하는 다양한 전문가를 라우팅해 long-tail 인식을 다뤘다.",
    "MiSLAS": "long-tail 인식의 calibration 개선을 다뤘다.",
    "LADE": "라벨 분포를 분리해 long-tail 시각 인식을 다뤘다.",
    "ACE": "상호 보완적인 전문가를 결합해 long-tail 인식을 다뤘다.",
    "Hybrid-PSC\nHybrid-SC": "대조학습을 결합한 하이브리드 네트워크를 제안했다.",
    "DRO-LT": "분포 강건성 손실을 long-tail 학습에 적용했다.",
    "Logit adj": "logit 조정으로 long-tail 학습의 클래스 편향을 다뤘다.",
    "MLLT": "균등·재균형 샘플링의 협력 학습으로 다중라벨 long-tail 인식을 다뤘다.",
    "PaCo": "파라미터를 갖는 대조학습을 제안했다.",
    "SADE": "다양한 전문가를 자기지도 방식으로 집계했다.",
    "Weight balancing": "가중치 균형을 통해 long-tail 인식을 다뤘다.",
    "NCL": "중첩 협력 학습으로 long-tail 시각 인식을 다뤘다.",
    "ResLT": "잔차 학습으로 long-tail 인식을 다뤘다.",
    "BalPoE": "보정된 전문가의 균형 product 결합을 제안했다.",
    "SHIKE": "자기 이질적 통합과 지식 추출로 long-tail 인식을 다뤘다.",
    "MDCS": "다양한 전문가와 일관성 자기증류를 결합했다.",
    "LGLA": "국소·전역 logit 조정으로 long-tail 학습을 다뤘다.",
    "CSA": "재샘플링이 long-tail 학습에 미치는 영향을 분석했다.",
    "PRL": "하이퍼네트워크로 다양한 전문가를 생성해 head–tail 선호를 제어했다.",
    "LSC": "test 분포를 모르는 long-tail 인식에서 라벨 시프트 보정을 학습했다.",
    "NCMC": "소수 클래스에 여러 중심을 두는 neural collapse와 분류 규칙을 분석했다.",
    "HTC-LTPLL": "head·tail 분류기의 협력으로 long-tail 부분라벨 학습을 다뤘다.",
    "H2T": "head 특징을 tail로 융합해 long-tail 시각 인식을 다뤘다.",
    "GBG": "long-tail 학습을 다목적 최적화 문제로 다뤘다.",
    "DeiT-LT": "long-tail 자료에서 Vision Transformer 학습에 증류를 적용했다.",
    "DirMixE": "test 분포를 모르는 조건의 계층적 라벨 분포 변화를 다뤘다.",
    "MOOSF": "다목적 최적화에 기반한 전략 융합으로 long-tail 학습을 다뤘다.",
    "BEM": "균형·엔트로피 기반 혼합으로 long-tail 준지도학습을 다뤘다.",
    "DELTA": "long-tail 온라인 지속학습에서 학습 요소를 분리했다.",
    "Meta-CALA": "클래스와 속성을 고려한 logit 조정을 제안했다.",
    "Noisy-Ostracods (목록명)": "실제 라벨·특징 노이즈와 불균형이 있는 세분류 데이터셋을 제시했다.",
    "ImOOD": "불균형 ID 분포에서 tail과 OOD가 혼동되는 문제를 다뤘다.",
    "Zhang_2023": "deep long-tail 학습 방법을 정리한 서베이다.",
    "Zhang_2025": "long-tail 학습을 체계적으로 검토한 서베이다.",
    "ankit_2019": "침입탐지 데이터셋의 발전을 정리했다.",
    "bayesh_2024": "AI 기법을 결합한 침입탐지·식별 시스템을 제시했다. 세부 모델은 원본 미기재.",
    "bandarupalli_2025": "CIC-IDS2017에서 심층 신경망 기반 침입탐지를 다뤘다.",
    "vinayakumar_2019": "심층학습 기반 지능형 침입탐지를 다뤘다.",
    "lee_2019": "이벤트 프로파일과 인공신경망을 이용한 위협 탐지를 다뤘다.",
    "zhang_2020": "SMOTE·Gaussian mixture model을 결합한 CNN 침입탐지를 다뤘다.",
    "sharafaldin2018toward": "침입탐지 데이터셋 생성과 공격 트래픽 특성화를 다뤘다.",
    "zhou2020building": "특징 선택과 앙상블 분류기를 결합한 침입탐지를 다뤘다.",
    "liu2019machine": "IDS의 머신러닝·딥러닝 방법을 정리한 서베이다.",
    "ring2019survey": "네트워크 침입탐지 데이터셋을 정리한 서베이다.",
    "gharib2016evaluation": "침입탐지 데이터셋을 평가하는 틀을 제시했다.",
    "lashkari2017characterization": "시간 기반 특징으로 Tor 트래픽을 특성화했다.",
    "hewapathirana2025comparative": "CSE-CIC-IDS2018에서 두 단계 머신러닝 침입탐지를 비교했다.",
    "abdulhammed2019features": "머신러닝 침입탐지에 사용하는 특징의 차원 축소를 비교했다.",
    "moustafa2015unsw": "UNSW-NB15 침입탐지 데이터셋을 제시했다.",
    "tavallaee2009detailed": "KDD CUP 99 데이터셋을 상세 분석했다.",
    "kdd_cup_1999_data_130": "KDD CUP 99 데이터셋 참고 항목이다.",
    "moustafa2017big": "유한 Dirichlet 혼합 모델의 통계적 의사결정으로 침입탐지를 다뤘다.",
}
PRIMARY = {
    "PRL": "https://proceedings.neurips.cc/paper_files/paper/2024/hash/0e21697f3b9385f6c4100aad94c26bed-Abstract.html",
    "NCMC": "https://proceedings.neurips.cc/paper_files/paper/2024/hash/78a2ed8e7da1829d64fb3561adbd706b-Abstract-Conference.html",
    "ImOOD": "https://proceedings.neurips.cc/paper_files/paper/2024/hash/c554c1305b8f4f993db4738a9c633d14-Abstract-Conference.html",
    "IP-DPP": "https://arxiv.org/abs/2510.08836",
    "CoLD": "https://www.ndss-symposium.org/ndss-paper/cold-collaborative-label-denoising-framework-for-network-intrusion-detection/",
    "RUIDS": "https://www.sciencedirect.com/science/article/pii/S016740482300041X",
}


def norm(value):
    return re.sub(r"[^\w]", "", unicodedata.normalize("NFKC", str(value)).casefold())


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def note(rec, col, text):
    rec["notes"][col] = rec["notes"].get(col, "") + "\n" + text


def group_name(groups, chapter, number):
    return next(name for ch, name in groups if ch == chapter and name.startswith(f"{number:02d} "))


def source_record(sheet, row, mapping, abb=None):
    values = [sheet.cell(row, c).value if c else None for c in mapping]
    if abb:
        values[0] = abb
    notes = {i: sheet.cell(row, c).comment.text for i, c in enumerate(mapping)
             if c and sheet.cell(row, c).comment}
    rec = {"values": values, "notes": notes, "strike": bool(sheet.cell(row, mapping[1]).font.strike),
           "origin": f"references.xlsx:{sheet.title}!{row}", "origins": [f"{sheet.title}!{row}"],
           "source_cells": [], "chapter": 0, "group": "", "kind": ""}
    note(rec, 0, f"원본: {rec['origin']}")
    note(rec, 2, "첨부 원본의 서지 정보를 유지했다. 별도 근거가 적힌 항목 외에는 이번에 게재 상태를 재검증하지 않았다.")
    a = sheet.cell(row, 1)
    if a.fill.patternType == "solid":
        rec["a_fill"] = copy(a.fill)
        note(rec, 0, "A열 색은 첨부 원본에서 보존했다.")
    if sheet.title == "Lit. imbal":
        rec["a_font"] = copy(a.font)
        note(rec, 0, f"원본 C열 유형: {sheet.cell(row, 3).value}; 원본 인용키: {sheet.cell(row, 2).value}")
    # Preserve ancillary run notes even though the requested schema has no column for them.
    extra = [(c.coordinate, c.value) for c in sheet[row] if c.value is not None and c.column > max(mapping)]
    if extra:
        note(rec, 15, "원본 보조 메모: " + json.dumps(extra, ensure_ascii=False))
    return rec


def merge_records(records):
    merged, by_abb, by_title, events = [], {}, {}, []
    for rec in records:
        abb, title = norm(rec["values"][0]), norm(rec["values"][1])
        old = by_title.get(title) or by_abb.get(abb)
        if old is None:
            merged.append(rec)
            by_abb[abb], by_title[title] = rec, rec
            continue
        events.append({"retained": old["values"][0], "merged_from": rec["origin"]})
        old["origins"] += rec["origins"]
        for col, value in enumerate(rec["values"]):
            if value is not None and value != "":
                if old["values"][col] is None or old["values"][col] == "":
                    old["values"][col] = value
                elif norm(old["values"][col]) != norm(value):
                    note(old, col, f"중복 원본 {rec['origin']}의 기록(기존 본문 유지): {value}")
        for col, text in rec["notes"].items():
            note(old, col, text)
        by_abb[abb], by_title[title] = old, old
    return merged, events


def collect_imbal(workbook):
    records = []
    s = workbook["Lit. imbal"]
    aliases = {16: "Noisy-Ostracods (목록명)", 17: "ImOOD", 36: "NCMC"}
    for row in range(3, 47):
        if not s.cell(row, 4).value:
            continue
        abb = aliases.get(row) or s.cell(row, 1).value or s.cell(row, 2).value
        rec = source_record(s, row, [1,4,15,16,0,5,0,6,7,8,9,10,11,12,13,14,0], abb)
        category = s.cell(row, 3).value
        if category in [1, 2, 3]:
            rec["chapter"], rec["group"] = 1, group_name(IMBAL_GROUPS, 1, category)
        else:
            number = 3 if abb == "ImOOD" else 2 if row == 16 else 1
            rec["chapter"], rec["group"] = 2, group_name(IMBAL_GROUPS, 2, number)
        records.append(rec)
    mapping = [1,2,13,14,0,3,0,4,5,6,7,8,9,10,11,12,15]
    for name, start, end in [("Lit Imbal SOTA",2,46), ("Lit IDS  SOTA",32,56)]:
        s = workbook[name]
        for row in range(start, end):
            if not s.cell(row, 2).value:
                continue
            rec = source_record(s, row, mapping)
            # SOTA alternating fills were decoration, not the original imbal A-column coding.
            rec.pop("a_fill", None)
            abb = rec["values"][0]
            if abb in {"PRL", "NCMC", "ImOOD"}:
                rec["chapter"], rec["group"] = 2, group_name(IMBAL_GROUPS, 2, 3)
            else:
                rec["chapter"], rec["group"] = 3, group_name(IMBAL_GROUPS, 3, SOTA_GROUPS[abb])
            records.append(rec)
    records, events = merge_records(records)
    by_abb = {r["values"][0]: r for r in records}
    for abb, code in [("PRL", "https://github.com/DataLab-atom/PRL"), ("ImOOD", "https://github.com/alibaba/imood")]:
        rec = by_abb[abb]
        note(rec, 5, f"중복 정리 중 코드 링크가 다른 논문과 뒤섞인 것을 확인했다. 공식 원문에 따라 교정: {PRIMARY[abb]}\n원래 값: {rec['values'][5]}")
        rec["values"][5] = code
    rec = by_abb["NCMC"]
    note(rec, 5, "원본 SOTA의 DataLab-atom/PRL은 다른 논문의 코드다. NCMC 코드로 사용하지 않으며 원래 값은 메모에 보존한다. " + str(rec["values"][5]) + "\n" + PRIMARY["NCMC"])
    rec["values"][5] = None
    for abb in ["PRL", "NCMC", "ImOOD"]:
        if not by_abb[abb]["values"][16]:
            by_abb[abb]["values"][16] = PRIMARY[abb]
            note(by_abb[abb], 16, "공식 논문집 상세 페이지. 해당 페이지에서 PDF에 접근할 수 있다.")
    for abb in ["IP2SL", "IP-DPP"]:
        note(by_abb[abb], 0, "IP2SL/IP-DPP 항목은 기제 설명이 유사하지만 서로 다른 제목으로 기록되어 있다. 동일 논문 여부가 확정되지 않아 임의 통합하지 않았다. 확인된 IP-DPP 원문: " + PRIMARY["IP-DPP"])
    return records, events


def collect_ids(workbook):
    records = []
    s = workbook["Lit IDS  SOTA"]
    for row in range(3, 29):
        abb = s.cell(row, 1).value or "Cloud-Transformer (목록명)"
        rec = source_record(s, row, [1,2,13,14,0,3,0,4,5,6,7,8,9,10,11,12,15], abb)
        number = IDS_MODEL[abb]
        rec["chapter"] = 3 if number == 0 else 2
        rec["group"] = group_name(IDS_GROUPS, rec["chapter"], number or 1)
        note(rec, 0, "기반 모델 분류는 원본 제목·Key methods에 근거한다. GNN 결합 모델과 provenance 규칙 기반을 구별했다.")
        if abb in PRIMARY:
            note(rec, 0, "분류 보조 확인: " + PRIMARY[abb])
        if not s.cell(row, 1).value:
            note(rec, 0, "Abb.가 없어 목록용 이름을 붙였다. 논문의 공식 약어가 아니다.")
        records.append(rec)
    s = workbook["Lit IDS"]
    for row in range(3, s.max_row + 1):
        if not s.cell(row, 3).value:
            continue
        abb = s.cell(row, 2).value
        rec = source_record(s, row, [2,3,14,15,0,4,0,5,6,7,8,9,10,11,12,13,0], abb)
        if row == 17:
            rec["values"][1], rec["values"][5], rec["values"][16] = s.cell(row,4).value, None, s.cell(row,3).value
            note(rec, 1, "원본 C17에 논문 URL, D17에 제목이 있어 올바른 열로 옮겼다. A17 메모: " + str(s.cell(row,1).value))
        rec["chapter"], number = IDS_INTRO[abb]
        rec["group"] = group_name(IDS_GROUPS, rec["chapter"], number)
        if rec["values"][3] is None:
            match = re.search(r"(?:19|20)\d{2}", abb)
            if match:
                rec["values"][3] = int(match.group())
                note(rec, 3, "원본의 인용키에 적힌 연도를 정렬용으로 옮겼다. 공식 게재연도 재검증은 하지 않았다.")
        note(rec, 0, "원본 인용키를 Abb. 식별자로 유지했다. 기반 모델은 제목과 원본 설명에서 식별되는 범위로 분류했다.")
        records.append(rec)
    return merge_records(records)


def prepare_values(records):
    for rec in records:
        values, abb = rec["values"], rec["values"][0]
        if not values[6]:
            base = SUMMARIES.get(abb) or values[8] or values[9] or values[1]
            base = " ".join(str(base).split())
            values[6] = base if len(base) <= 180 else base[:177] + "…"
            note(rec, 6, "원본 제목·Goal·Key methods를 요약했다. 추가 성능 주장을 만들지 않았다.")
        if not values[4]:
            values[4] = "미기재"
            note(rec, 4, "첨부 파일에 정확한 게재일이 없어 날짜를 추정하지 않았다.")
        if isinstance(values[5], str):
            old = values[5]
            text = old.strip()
            if text.startswith("github.com/") or text.startswith("shizhen-zhao.github.io/"):
                values[5] = "https://" + text
            elif "liamdm/FlowTransformer" in text:
                values[5] = "https://github.com/liamdm/FlowTransformer"
            elif not re.search(r"https?://", text):
                note(rec, 5, "코드 URL이 아닌 원본 메모: " + text)
                values[5] = None
            if old != values[5]:
                note(rec, 5, "원본 코드 필드: " + old)


def validate(path, locations, records_by_sheet):
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
    w = openpyxl.load_workbook(path)
    assert list(w.sheetnames) == list(records_by_sheet)
    for name, records in records_by_sheet.items():
        s = w[name]
        assert s.max_column == 17 and s.freeze_panes == "B4"
        assert [s.cell(3,c).value for c in range(1,18)] == HEADERS
        assert len(locations[name]) == len(records)
        assert len({norm(r['values'][0]) for r in records}) == len(records)
        grouping = {}
        for rec in records:
            abb = rec["values"][0]
            row = locations[name][abb]["row"]
            assert [s.cell(row,c).value for c in range(1,18)] == rec["values"], (name,abb,"value mismatch")
            if rec.get("a_fill"):
                assert s.cell(row,1).fill == rec["a_fill"], (name,abb,"A color lost")
            if rec["strike"]:
                assert s.cell(row,2).font.strike
            grouping.setdefault((rec["chapter"],rec["group"]), []).append((row,rec))
        for seq in grouping.values():
            seq.sort()
            keys=[(-date_key(r["values"][3]),-date_key(r["values"][4]),r["values"][0]) for _,r in seq]
            assert keys == sorted(keys), (name,"sort")
            for (ra,a),(rb,b) in zip(seq,seq[1:]):
                assert s.cell(ra,2).fill.fgColor.rgb != s.cell(rb,2).fill.fgColor.rgb, (name,ra,rb,"banding")
    return w


def main():
    root = Path.cwd()
    refs = root / "docs/lablog/evidence/references"
    ev = refs / "20261027"
    source_paths = {"references.xlsx": refs/"references.xlsx", "tabpfn_papers_survey.xlsx": refs/"tabpfn_papers_survey.xlsx"}
    archives = ev / "before_three_parts"
    archives.mkdir(exist_ok=True)
    active_hashes = {name:sha(path) for name,path in source_paths.items()}
    for name, path in source_paths.items():
        if not (archives/name).exists():
            shutil.copy2(path, archives/name)
    source = openpyxl.load_workbook(archives/"references.xlsx")
    reviewed = openpyxl.load_workbook(archives/"tabpfn_papers_survey.xlsx")
    plan = json.loads((ev/"literature_update.json").read_text())
    plan["_inventory"] = str(ev/"inventory_before.json")
    imbal, imbal_merges = collect_imbal(source)
    ids, ids_merges = collect_ids(source)
    tab = collect_tabpfn(reviewed, plan)
    # Use the older attachment's annotations where they differ; keep newer
    # reviewed data in comments so the user's writing is not silently erased.
    source_tab = source["Lit TabPFN"]
    old_by_abb = {source_tab.cell(r,2).value:r for r in range(1,source_tab.max_row+1)
                  if isinstance(source_tab.cell(r,1).value,(int,float))}
    changed_from_attachment = []
    narrative_mapping = {6:4,7:9,8:10,9:11,10:12,11:13,12:14,13:15,14:16,15:17}
    for rec in tab:
        abb = rec["values"][0]
        if abb in old_by_abb:
            row = old_by_abb[abb]
            for out_col,old_col in narrative_mapping.items():
                value = source_tab.cell(row,old_col).value
                if value is not None and value != rec["values"][out_col]:
                    note(rec,out_col,"최신 파일의 이전 서술(첨부 원본의 사용자 메모를 우선): " + str(rec["values"][out_col]))
                    rec["values"][out_col] = value
                    changed_from_attachment.append({"abb":abb,"column":HEADERS[out_col]})
        if abb == "RF-TabPFN":
            rec["chapter"],rec["group"] = 2,"01 IDS 모델·보안 데이터 구성"
            note(rec,0,"사용자 분류 기준에 따라 RF–TabPFN 결합·특징 선택 IDS 방법을 Ch.2로 이동했다. 전체 방법의 우위를 새로 검증한 것은 아니다.")
        if abb in {"PFN-Imb", "Plugin-BB", "AdapTable"}:
            rec["group"] = "08 불균형 보정·결정규칙 비교"
        if abb == "TFM-Surv":
            rec["group"] = "06 생성·태스크 확장"
        if abb == "Rockburst-TabPFN":
            rec["group"] = "02 다른 도메인 응용"
        if abb in MUTED:
            note(rec,15,"사용자 선별 의견: 기존 모델 비교 중심 문헌으로 표시. TabNet·SMOTE 처리의 정확한 관계는 전문 대조 없이 새 사실로 확정하지 않았다.")
    prepare_values(imbal)
    prepare_values(ids)
    result = openpyxl.Workbook()
    result.remove(result.active)
    books = {"1. imbal":imbal,"2. TabPFN":tab,"3. IDS":ids}
    locations = {}
    locations["1. imbal"] = render_sheet(result,"1. imbal",imbal,
        chapter_names={1:"기존 문헌 | 원본 C열 접근 유형",2:"서베이·데이터·OOD",3:"최근 SOTA | 접근 방법별"},groups=IMBAL_GROUPS,
        title="1. imbal | 기존 유형 · long-tail SOTA · OOD",
        subtitle="유형별 연도 → 게재일 최신순. Abb.만 고정. A열의 원본 강조색·서식과 C열 유형 1/2/3을 보존했다. 빈 날짜는 미기재로 두고, 합친 중복 문헌의 다른 서술은 셀 메모에 남겼다.")
    locations["2. TabPFN"] = render_sheet(result,"2. TabPFN",tab,
        groups=GROUPS+[(3,"08 불균형 보정·결정규칙 비교")])
    locations["3. IDS"] = render_sheet(result,"3. IDS",ids,
        chapter_names={1:"IDS 소개 | 서베이·데이터셋",2:"IDS 방법 | 기반 모델·입력 표현별",3:"보조 평가·미확정 구조"},groups=IDS_GROUPS,
        title="3. IDS | 소개·데이터셋 / Transformer·GNN·표 데이터 등 기반 모델 / 보조 평가",
        subtitle="유형별 최신순; Abb.만 고정. 원본 A색 유지: 분홍=detect만, 노랑=binary만, 초록=multi-class(원본 분류, 재검증 아님). 표 기반 TabPFN 방법은 2. TabPFN의 Ch.2에도 정리했다.")
    result["1. imbal"].sheet_properties.tabColor="805D38"
    result["2. TabPFN"].sheet_properties.tabColor="244763"
    result["3. IDS"].sheet_properties.tabColor="28644F"
    result.active=1
    projection=openpyxl.Workbook()
    projection.remove(projection.active)
    projection_locs=render_sheet(projection,"Lit TabPFN",tab,groups=GROUPS+[(3,"08 불균형 보정·결정규칙 비교")])
    stage=Path(tempfile.mkdtemp(prefix="literature_three_parts_"))
    main_path,tab_path=stage/"references.xlsx",stage/"tabpfn_papers_survey.xlsx"
    result.save(main_path)
    projection.save(tab_path)
    checked=validate(main_path,locations,books)
    checked_tab=validate(tab_path,{"Lit TabPFN":projection_locs},{"Lit TabPFN":tab})
    for abb in MUTED:
        r=locations["2. TabPFN"][abb]["row"]
        assert checked["2. TabPFN"].cell(r,1).font.strike
        assert checked["2. TabPFN"].cell(r,2).font.color.rgb.endswith("858585")
    for abb,ch in EXTRA_CHAPTERS.items():
        assert locations["2. TabPFN"][abb]["chapter"]==ch
    assert not {"FM-Industrial-Security","Conformal-TabPFN-IDS"} & set(locations["2. TabPFN"])
    assert len(imbal)==85 and len(ids)==44 and len(tab)==135
    for name,path in source_paths.items():
        assert sha(path)==active_hashes[name],"Active workbook changed during preparation."
    checked.close();checked_tab.close()
    for staged,target in [(main_path,source_paths["references.xlsx"]),(tab_path,source_paths["tabpfn_papers_survey.xlsx"])]:
        with tempfile.NamedTemporaryFile(dir=refs,suffix=".xlsx",delete=False) as f:
            temp=Path(f.name)
        shutil.copy2(staged,temp)
        temp.replace(target)
    records_json={name:[{k:v for k,v in r.items() if k not in {"source_cells","a_fill","a_font"}} for r in rows] for name,rows in books.items()}
    (ev/"three_parts_records.json").write_text(json.dumps(records_json,ensure_ascii=False,indent=2)+"\n")
    original_cells={s.title:[[{"cell":c.coordinate,"value":c.value,"comment":c.comment.text if c.comment else None,
                                  "hyperlink":c.hyperlink.target if c.hyperlink else None} for c in row if c.value is not None]
                            for row in s if any(c.value is not None for c in row)] for s in source}
    (ev/"three_parts_original_cells.json").write_text(json.dumps(original_cells,ensure_ascii=False,indent=2,default=str)+"\n")
    report={"date":"2026-10-27","status":"complete","workbook":str(source_paths['references.xlsx'].relative_to(root)),
            "tabpfn_projection":str(source_paths['tabpfn_papers_survey.xlsx'].relative_to(root)),
            "input_sha256":{n:sha(archives/n) for n in source_paths},"output_sha256":{n:sha(p) for n,p in source_paths.items()},
            "counts":{n:len(v) for n,v in books.items()},"total_rows":sum(map(len,books.values())),
            "count_scope":"파트 간 중복은 허용한 문헌 행 수. 서로 다른 분야에서 참고하는 동일 논문이 있을 수 있다.",
            "merged_duplicates":imbal_merges+ids_merges,"original_imbal_type_counts":dict(Counter(r['group'] for r in imbal if r['chapter']==1)),
            "original_imbal_highlight_count":sum('a_fill' in r for r in imbal),
            "tabpfn_chapters":dict(Counter(r['chapter'] for r in tab)),"removed_unsolicited_papers":["FM-Industrial-Security","Conformal-TabPFN-IDS"],
            "muted":sorted(MUTED),"acceptance_updates_retained":len(plan['updates']),"user_attachment_prose_retained":changed_from_attachment,
            "classification_sources":PRIMARY,"locations":locations,
            "validation":["ZIP integrity + reload of both workbooks","exact 17-column schema","Abb. only frozen",
                          "all output cell values match reviewed mapping","original A fills preserved","category/date order checked",
                          "adjacent rows alternate color","muted papers + chapter placements checked","new unsolicited sheets/papers removed"],
            "unresolved":["IP2SL/IP-DPP 제목이 달라 동일 문헌 여부 미확정: 별도 행 유지 및 상호 메모",
                          "원본에 없는 정확한 게재일·성능·코드는 추정하지 않음",
                          "IDS 인용키 연도는 정렬용으로 옮겼으며 게재연도 재검증과 구별",
                          "TabPFN 원본의 단독 A137 PDF 링크는 문헌 행이 없어 원본 셀 JSON·보관 파일에 보존"]}
    (ev/"three_parts_validation.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({k:v for k,v in report.items() if k not in {'locations','merged_duplicates','classification_sources'}},ensure_ascii=False,indent=2))


if __name__=="__main__":
    main()
