"""Rebuild the TabPFN literature sheet with the user's 17-column schema.

This reads the reviewed workbook and creates a new file; it does not search for
papers, infer acceptance, or import an unconfirmed historical workbook.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
import hashlib
import json
from pathlib import Path
import re
import unicodedata
import zipfile

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = ["Abb.", "Title", "Venue", "연도", "게재일", "코드", "한줄요약",
           "Motivation", "Goal", "Key methods", "Eval-Data",
           "Eval-Comparison Models", "Eval-Experiments", "Metrics",
           "Key findings", "Limitations", "PDF Link"]
OLD_COLUMNS = [2, 3, 18, 19, 7, 8, 4, 9, 10, 11, 12, 13, 14, 15, 16, 17, 20]
WIDTHS = [27, 68, 43, 9, 30, 36, 48, 50, 47, 60, 40, 43, 48, 37, 58, 58, 48]
EXTRA_CHAPTERS = {"LotBoNC": 2, "TFM-Systematic-Survey": 1,
                  "TabFM / TabFM+": 2, "TabPFN-3.5": 1,
                  "StreamTFM-Study": 2, "TFM-Ensemble-Diversity": 2}
GROUPS = [
    (1, "01 공식 모델·확장"), (1, "02 서베이"),
    (2, "01 IDS 모델·보안 데이터 구성"),
    (2, "02 컨텍스트 선택·압축"), (2, "03 스트림·컨텍스트 관리"),
    (2, "04 아키텍처·사전학습·추론 효율"),
    (2, "05 파인튜닝·테스트 시점 학습"),
    (2, "06 앙상블·보정"), (2, "07 불균형·라벨 시프트"),
    (2, "08 생성·증강"), (2, "09 그 밖의 확장·분석"),
    (3, "01 IDS 응용·모델 비교"), (3, "02 다른 도메인 응용"),
    (3, "03 벤치마크·분석"), (3, "04 이론·해석·불확실성"),
    (3, "05 LLM·다른 기반 모델"), (3, "06 생성·태스크 확장"),
    (3, "07 서베이·도구·포지션"),
]
CHAPTER_NAMES = {
    1: "Ch.1 | TabPFN 원본·공식 확장·서베이",
    2: "Ch.2 | IDS 관련 방법·TabPFN 개선·컨텍스트 및 앙상블",
    3: "Ch.3 | 도메인 응용·기존 모델 비교·간접 관련 연구",
}
CHAPTER_COLORS = {1: "244763", 2: "28644F", 3: "805D38"}
MUTED = {"FM-CyberEval", "FM-IDS-FLLM"}


def group_for(abb, chapter, kind):
    if chapter == 1:
        return "02 서베이" if "서베이" in kind else "01 공식 모델·확장"
    if chapter == 2:
        if abb in {"LotBoNC", "WFE-Tab"} or "IDS" in kind:
            return "01 IDS 모델·보안 데이터 구성"
        if "스트림" in kind or "스트리밍" in kind or "컨텍스트 관리" in kind:
            return "03 스트림·컨텍스트 관리"
        if abb == "CIA-DD" or "컨텍스트 선택" in kind or "컨텍스트 압축" in kind or "컨텍스트 구성" in kind or "검색 컨텍스트" in kind:
            return "02 컨텍스트 선택·압축"
        if abb in {"TabFM / TabFM+", "TFM-Ensemble-Diversity"} or "앙상블" in kind or "배깅" in kind:
            return "06 앙상블·보정"
        if "불균형" in kind or "라벨 시프트" in kind:
            return "07 불균형·라벨 시프트"
        if "생성" in kind or "증강" in kind:
            return "08 생성·증강"
        if "파인튜닝" in kind or "테스트타임 트레이닝" in kind:
            return "05 파인튜닝·테스트 시점 학습"
        if any(token in kind for token in ["아키텍처", "고차원", "attention", "사전학습", "추론 가속"]):
            return "04 아키텍처·사전학습·추론 효율"
        return "09 그 밖의 확장·분석"
    if "IDS" in kind or abb in MUTED:
        return "01 IDS 응용·모델 비교"
    if "도메인" in kind or any(token in kind for token in ["금융", "임상", "의료", "보험", "신용"]):
        return "02 다른 도메인 응용"
    if "서베이" in kind or "도구" in kind or "포지션" in kind:
        return "07 서베이·도구·포지션"
    if "벤치마크" in kind or "분석" in kind:
        return "03 벤치마크·분석"
    if "이론" in kind or "해석" in kind or "불확실성" in kind:
        return "04 이론·해석·불확실성"
    if "생성" in kind or "태스크" in kind:
        return "06 생성·태스크 확장"
    return "05 LLM·다른 기반 모델"


def date_key(value):
    m = re.search(r"(20\d{2})(?:-(\d{2}))?(?:-(\d{2}))?", str(value))
    return int("".join(p or "00" for p in m.groups())) if m else 0


def normalize_venue(source):
    value = str(source or "")
    # The complete source, DOI and original title remain in cell comments.
    return re.split(r"; (?:arXiv|medRxiv)| / arXiv:", value, maxsplit=1)[0]


def add_comment(cell, text):
    old = cell.comment.text + "\n\n" if cell.comment else ""
    cell.comment = Comment(old + text, "Literature review")
    cell.comment.width = 560
    cell.comment.height = 300


def collect_tabpfn(source, plan):
    updates = {p["abb"]: p for p in plan["updates"]}
    inventory = json.loads((plan_path := Path(plan["_inventory"])).read_text())
    initial = {p["abb"]: p for p in inventory}
    records = []
    for name in ["Lit TabPFN", "20261021_추가문헌"]:
        sheet = source[name]
        for row in range(1, sheet.max_row + 1):
            if name == "Lit TabPFN" and not isinstance(sheet.cell(row, 1).value, (int, float)):
                continue
            if name != "Lit TabPFN" and row == 1:
                continue
            abb = sheet.cell(row, 2).value
            if not abb:
                continue
            cells = [sheet.cell(row, col) for col in OLD_COLUMNS]
            values = [c.value for c in cells]
            kind = str(sheet.cell(row, 5).value or "")
            chapter = int(sheet.cell(row, 1).value) if name == "Lit TabPFN" else EXTRA_CHAPTERS[abb]
            item = updates.get(abb)
            note = []
            if item:
                values[1] = item.get("published_title", values[1])
                values[2], values[3], values[4], values[16] = item["source"], item["year"], item["venue_date"], item["url"]
                note += [item["note"], "채택 확인 근거: " + "\n".join(item["evidence"])]
            elif abb == "TabPFN-2.5":
                values[2] = initial[abb]["source"]
            if not item:
                original_date = str(values[4] or "")
                if re.fullmatch(r"\d{4}-\d{2}-\d{2}", original_date):
                    values[4] = original_date + " (최초 공개)"
                elif original_date.endswith("*"):
                    values[4] = original_date[:-1] + " (기존 근사 기록)"
                note.append("기존 날짜의 의미(최초 공개/근사 기록)를 유지했다. 공식 게재일로 새로 검증한 값은 채택 갱신 12편에 한정한다.")
            full_source = values[2]
            values[2] = normalize_venue(full_source)
            notes = {i: c.comment.text for i, c in enumerate(cells) if c.comment}
            notes[0] = notes.get(0, "") + f"\n원본: {name}!{row}\n원래 유형: {kind}\n출처등급: {item['grade'] if item else sheet.cell(row, 6).value}"
            notes[2] = notes.get(2, "") + "\n원래 Source: " + str(full_source) + "\n" + "\n".join(note)
            notes[4] = notes.get(4, "") + "\n기존 공개일: " + str(sheet.cell(row, 7).value)
            if item:
                notes[1] = notes.get(1, "") + "\n기존 제목: " + str(sheet.cell(row, 3).value)
                notes[16] = notes.get(16, "") + "\n기존 preprint 링크: " + str(initial[abb]["url"])
            if name == "20261021_추가문헌":
                notes[15] = notes.get(15, "") + "\nDiscussion: " + str(sheet.cell(row, 21).value or "") + "\n검토 범위: " + str(sheet.cell(row, 22).value or "")
            records.append({"values": values, "chapter": chapter, "kind": kind,
                            "group": group_for(abb, chapter, kind), "notes": notes,
                            "strike": bool(cells[1].font.strike) or abb in MUTED,
                            "origin": f"{name}!{row}", "source_cells": cells})
    # Explicitly retained by the user's later instruction, despite removing the
    # rest of the unsolicited three-paper addition.
    p = next(p for p in plan["new_papers"] if p["abb"] == "FM-IDS-FLLM")
    fields = ["abb", "title", "source", "year", "date", "code", "summary",
              "motivation", "goal", "methods", "datasets", "baselines",
              "experiments", "metrics", "findings", "limitations", "url"]
    records.append({"values": [p.get(k) for k in fields], "chapter": 3, "kind": p["type"],
                    "group": "01 IDS 응용·모델 비교", "strike": True,
                    "notes": {0: "사용자가 이번 요청에서 명시적으로 지정한 회색·취소선 문헌.",
                              2: "\n".join(p["evidence"]),
                              15: p["level"] + "\n" + p["discussion"]},
                    "origin": "20261027/literature_update.json:FM-IDS-FLLM", "source_cells": []})
    assert len(records) == 135
    assert len({r["values"][0] for r in records}) == 135
    return records


def render_sheet(workbook, name, records, *, chapter_names=None, groups=None,
                 chapter_colors=None, title=None, subtitle=None):
    chapter_names = chapter_names or CHAPTER_NAMES
    groups = groups or GROUPS
    chapter_colors = chapter_colors or CHAPTER_COLORS
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 80
    sheet.sheet_properties.tabColor = "244763"
    sheet.freeze_panes = "B4"  # Abb. only; Title scrolls.
    for col, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row, text in [(1, title or "2. TabPFN | 원본·서베이 / 개선·IDS 방법 / 도메인 응용·모델 비교"),
                      (2, subtitle or "각 유형 안에서 연도 → 게재일 최신순. Abb.만 고정. 회색·취소선은 기존 표시와 사용자 지정 문헌. 날짜의 최초 공개/근사 여부와 채택 근거는 셀 메모에 보존.")]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=17)
        c = sheet.cell(row, 1, text)
        c.font = Font(name="맑은 고딕", size=12 if row == 1 else 10, bold=row == 1, color="FFFFFF" if row == 1 else "244763")
        c.fill = PatternFill("solid", fgColor="244763" if row == 1 else "E8EFF5")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 29
    for col, title in enumerate(HEADERS, 1):
        c = sheet.cell(3, col, title)
        c.font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="334E68")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[3].height = 36
    locations, order = {}, []
    for chapter in chapter_names:
        if not any(r["chapter"] == chapter for r in records):
            continue
        row = sheet.max_row + 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=17)
        cell = sheet.cell(row, 1, chapter_names[chapter])
        cell.fill = PatternFill("solid", fgColor=chapter_colors[chapter])
        cell.font = Font(name="맑은 고딕", size=12, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[row].height = 30
        for ch, group in groups:
            selected = [r for r in records if r["chapter"] == chapter and r["group"] == group]
            if ch != chapter or not selected:
                continue
            selected.sort(key=lambda r: (-date_key(r["values"][3]), -date_key(r["values"][4]), r["values"][0]))
            row = sheet.max_row + 1
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=17)
            cell = sheet.cell(row, 1, group + f" | {len(selected)}편")
            cell.fill = PatternFill("solid", fgColor="DFE7EE")
            cell.font = Font(name="맑은 고딕", size=11, bold=True, color=chapter_colors[chapter])
            cell.alignment = Alignment(vertical="center")
            sheet.row_dimensions[row].height = 25
            for index, record in enumerate(selected):
                sheet.append(record["values"])
                row, abb = sheet.max_row, record["values"][0]
                fill = ("F3F3F3" if index % 2 == 0 else "E5E5E5") if record["strike"] else ("FFFFFF" if index % 2 == 0 else "EEF2FF")
                for col, cell in enumerate(sheet[row], 1):
                    cell.font = Font(name="맑은 고딕", size=11, color="858585" if record["strike"] else "243746", strike=record["strike"])
                    cell.fill = PatternFill("solid", fgColor=fill)
                    cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right" if col == 4 else "left")
                    cell.border = Border(bottom=Side(style="hair", color="D9E1E8"))
                    if col - 1 in record["notes"]:
                        add_comment(cell, record["notes"][col - 1].strip())
                    if col == 1 and record.get("a_fill") is not None:
                        cell.fill = copy(record["a_fill"])
                    if col == 1 and record.get("a_font") is not None:
                        cell.font = copy(record["a_font"])
                    if col in [6, 17] and isinstance(cell.value, str):
                        match = re.search(r"https?://[^\s]+", cell.value)
                        if match:
                            cell.hyperlink = match.group().rstrip(".,;)")
                            font = copy(cell.font)
                            font.underline = "single"
                            font.color = "858585" if record["strike"] else "0563C1"
                            cell.font = font
                if abb in MUTED:
                    add_comment(sheet.cell(row, 1), "사용자 요청(2026-10-27): 기존 모델 비교 중심 문헌으로 Ch.3에 회색·취소선 표시. 이는 채택 여부를 부정하는 표시가 아니다.")
                max_lines = 1
                for col, val in enumerate(record["values"], 1):
                    lines = str(val or "").splitlines() or [""]
                    estimate = sum(max(1, int(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line) / (WIDTHS[col - 1] * 0.85)) + 1) for line in lines)
                    max_lines = max(max_lines, estimate)
                sheet.row_dimensions[row].height = min(409, max(84, 16 * max_lines + 10))
                sheet.row_dimensions[row].outlineLevel = 1
                locations[abb] = {"sheet": name, "row": row, "chapter": chapter, "group": group, "origin": record["origin"]}
                order.append(abb)
    sheet.auto_filter.ref = f"A3:Q{sheet.max_row}"
    sheet.print_title_rows = "1:3"
    sheet.sheet_properties.outlinePr.summaryBelow = False
    assert len(order) == len(records), "Unmapped category"
    return locations


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    evidence = args.root / "docs/lablog/evidence/references/20261027"
    plan = json.loads((evidence / "literature_update.json").read_text())
    plan["_inventory"] = str(evidence / "inventory_before.json")
    source = openpyxl.load_workbook(args.input)
    records = collect_tabpfn(source, plan)
    result = openpyxl.Workbook()
    result.remove(result.active)
    locations = render_sheet(result, "Lit TabPFN", records)
    result.save(args.output)
    with zipfile.ZipFile(args.output) as z:
        assert z.testzip() is None
    checked = openpyxl.load_workbook(args.output)
    sheet = checked["Lit TabPFN"]
    assert [sheet.cell(3, c).value for c in range(1, 18)] == HEADERS
    assert sheet.max_column == 17 and sheet.freeze_panes == "B4"
    by_abb = {r["values"][0]: r for r in records}
    for abb, info in locations.items():
        row = info["row"]
        assert [sheet.cell(row, c).value for c in range(1, 18)] == by_abb[abb]["values"]
        if abb in MUTED:
            assert sheet.cell(row, 1).font.strike and sheet.cell(row, 2).font.strike
    for abb, chapter in EXTRA_CHAPTERS.items():
        assert locations[abb]["chapter"] == chapter
    assert not {"FM-Industrial-Security", "Conformal-TabPFN-IDS"} & set(locations)
    audit = {"date": "2026-10-27", "input": str(args.input), "output": str(args.output),
             "input_sha256": hashlib.sha256(args.input.read_bytes()).hexdigest(),
             "output_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
             "status": "TabPFN reformatted; original imbal/IDS workbook confirmation pending",
             "papers": len(records), "chapters": dict(Counter(r["chapter"] for r in records)),
             "headers": HEADERS, "freeze_panes": "B4", "locations": locations,
             "removed_unsolicited_papers": ["FM-Industrial-Security", "Conformal-TabPFN-IDS"],
             "explicitly_retained_and_muted": sorted(MUTED), "acceptance_updates_retained": len(plan["updates"]),
             "validation": ["ZIP integrity", "reopened workbook", "17 requested columns",
                            "Abb. only frozen", "135 unique records", "chapter placement", "two gray/struck papers",
                            "all 17 output values equal reviewed source mapping"]}
    (evidence / "reformat_validation.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
    # Preserve historical supplementary notes outside the simplified workbook.
    old_notes = {s.title: [list(row) for row in s.values] for s in source if s.title not in {"Lit TabPFN"}}
    (evidence / "before_reformat_supplementary_sheets.json").write_text(json.dumps(old_notes, ensure_ascii=False, indent=2, default=str) + "\n")
    print(json.dumps({k: v for k, v in audit.items() if k != "locations"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
