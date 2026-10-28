"""Regroup all imbalanced-learning references without old/SOTA partitions."""
from copy import copy
import hashlib
import json
from pathlib import Path
import re
import shutil
import tempfile
import zipfile

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange


# Apply the same primary approach/topic to every paper, independent of year.
# A paper with several mechanisms appears once, under its primary approach.
GROUPS = [
    ("01 서베이·데이터셋", ["Zhang_2025", "Zhang_2023", "Noisy-Ostracods (목록명)"]),
    ("02 샘플링·생성·데이터 및 특징 증강", [
        "mixup", "Remix", "LEAP-CF\nLEAP-AF", "CSA", "H2T", "BEM",
        "IP2SL", "IP-DPP", "LLM-AutoDA", "EPIC", "LFN", "ADR", "ORD",
        "ConMix", "DiffMatch", "CFAMG", "OSB-SMOTE",
    ]),
    ("03 손실·재가중·로짓 및 분류기 보정", [
        "Class-Balanced", "LDAM", "BS\nBALMS", "cRT\nt-norm\nLWS", "MiSLAS",
        "LADE", "DRO-LT", "Logit adj", "Weight balancing", "LSC", "GBG",
        "MOOSF", "Meta-CALA", "IMMAX", "LSR", "MORE", "DBM", "BCE3S",
        "LOS", "OVO-OVR", "IBCL",
    ]),
    ("04 표현·대조학습·프로토타입", [
        "Hybrid-PSC\nHybrid-SC", "PaCo", "NCMC", "DELTA", "PRIME", "TTC",
        "TinyChange", "MGS", "LTRL", "DILDL",
    ]),
    ("05 다중 분기·전문가·앙상블", [
        "BBN", "LFME", "RIDE", "ACE", "MLLT", "SADE", "NCL", "ResLT",
        "BalPoE", "SHIKE", "MDCS", "LGLA", "PRL", "HTC-LTPLL", "DirMixE",
        "DCE", "MATI", "RRB",
    ]),
    ("06 구조 설계·지식 증류", ["DeiT-LT", "LT-SpikeFormer", "RLAL"]),
    ("07 GNN·그래프 불균형", [
        "NodeImport", "SIMBA", "AD-GSMOTE", "GraphMuGu", "Uni-GNN", "ImGDA", "GraphIAM",
    ]),
    ("08 OOD·Open-world 인식", ["OLTR", "ImOOD"]),
    ("09 라벨 노이즈·강건성·능동학습·프라이버시", ["DP", "DIRECT", "TAET", "DaSC"]),
]


def normalize_current_schema(workbook):
    """Apply the already requested date/comment removal to an older backup."""
    for sheet in workbook:
        headers = [c.value for c in sheet[3]]
        if "게재일" in headers:
            col = headers.index("게재일") + 1
            dims = {i: copy(sheet.column_dimensions[get_column_letter(i)])
                    for i in range(1, sheet.max_column + 1)}
            merges = [CellRange(str(area)) for area in sheet.merged_cells.ranges]
            for area in list(sheet.merged_cells.ranges):
                sheet.unmerge_cells(str(area))
            sheet.delete_cols(col)
            for key in list(sheet.column_dimensions):
                del sheet.column_dimensions[key]
            for i, dim in dims.items():
                if i == col:
                    continue
                c = i if i < col else i - 1
                dim.index = get_column_letter(c)
                dim.min = dim.max = c
                sheet.column_dimensions[dim.index] = dim
            for area in merges:
                if area.min_col > col:
                    area.min_col -= 1
                if area.max_col >= col:
                    area.max_col -= 1
                sheet.merge_cells(str(area))
            sheet.auto_filter.ref = f"A3:P{sheet.max_row}"
        for row in sheet:
            for cell in row:
                if cell.comment:
                    cell.comment = None
                if cell.hyperlink:
                    cell.hyperlink.ref = cell.coordinate
        if sheet.title in ["2. TabPFN", "Lit TabPFN"]:
            sheet.cell(2, 1).value = "각 유형 안에서 연도 최신순. Abb.만 고정. 회색·취소선은 기존 표시와 사용자 지정 문헌."


def year(value):
    found = re.search(r"(?:19|20)\d{2}", str(value or ""))
    return int(found.group()) if found else 0


def snapshot(sheet):
    return {
        "cells": [(c.coordinate, c.value, copy(c._style), c.hyperlink.target if c.hyperlink else None)
                  for row in sheet for c in row],
        "merged": sorted(map(str, sheet.merged_cells.ranges)),
        "freeze": sheet.freeze_panes,
        "filter": sheet.auto_filter.ref,
    }


def main():
    root = Path.cwd()
    target = root / "docs/lablog/evidence/references/references.xlsx"
    fallback = Path("/tmp/references_remove_date_comments_wmqz5aid/references.xlsx")
    source = target if target.exists() else fallback
    if not source.exists():
        source = Path("/tmp/literature_three_parts_1qdej4x3/references.xlsx")
    if not source.exists():
        raise FileNotFoundError("The integrated workbook or its pre-edit backup is required.")
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    backup = Path(tempfile.mkdtemp(prefix="imbal_before_regroup_")) / "references.xlsx"
    shutil.copy2(source, backup)
    workbook = openpyxl.load_workbook(source)
    normalize_current_schema(workbook)
    other = {s.title: snapshot(s) for s in workbook if s.title != "1. imbal"}
    old = workbook["1. imbal"]
    records = {old.cell(r, 1).value: [old.cell(r, c) for c in range(1, 17)]
               for r in range(4, old.max_row + 1) if old.cell(r, 2).value is not None}
    members = [abb for _, abbs in GROUPS for abb in abbs]
    assert len(members) == len(set(members)) == 85
    assert set(members) == set(records), (set(records) - set(members), set(members) - set(records))
    headers = [c.value for c in old[3]]
    title_style, note_style = copy(old.cell(1,1)._style), copy(old.cell(2,1)._style)
    group_style = copy(old.cell(4,1)._style)
    sheet = workbook.create_sheet("imbal_regrouped", workbook.index(old))
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = old.sheet_view.zoomScale
    sheet.sheet_properties.tabColor = copy(old.sheet_properties.tabColor)
    sheet.freeze_panes = "B4"
    sheet.sheet_properties.outlinePr.summaryBelow = False
    for key, dim in old.column_dimensions.items():
        sheet.column_dimensions[key] = copy(dim)
    for r, value, style in [
        (1, "1. imbal | 접근 방법·연구 주제별 통합 분류", title_style),
        (2, "발표 시기와 관계없이 주된 접근 방법·연구 주제로 분류. 복합 방법은 한 유형에 배치하고, 유형 안에서는 연도 최신순. Abb. 고정과 A열 강조색 유지.", note_style),
    ]:
        sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=16)
        sheet.cell(r,1,value)._style = style
        sheet.row_dimensions[r].height = old.row_dimensions[r].height
    for c, text in enumerate(headers,1):
        sheet.cell(3,c,text)._style = copy(old.cell(3,c)._style)
    sheet.row_dimensions[3].height = old.row_dimensions[3].height
    locations = {}
    for group, abbs in GROUPS:
        r = sheet.max_row + 1
        sheet.merge_cells(start_row=r,start_column=1,end_row=r,end_column=16)
        sheet.cell(r,1,f"{group} | {len(abbs)}편")._style = copy(group_style)
        sheet.row_dimensions[r].height = 30
        order = sorted(abbs, key=lambda abb: (-year(records[abb][3].value),abb))
        for idx, abb in enumerate(order):
            original = records[abb]
            r = sheet.max_row + 1
            for c, prior in enumerate(original,1):
                cell = sheet.cell(r,c,prior.value)
                cell._style = copy(prior._style)
                if c != 1 or prior.fill.fgColor.type != "rgb" or prior.fill.fgColor.rgb[-6:] != "CFE2F3":
                    cell.fill = PatternFill("solid",fgColor="FFFFFF" if idx%2==0 else "EEF2FF")
                if prior.hyperlink:
                    cell.hyperlink = copy(prior.hyperlink)
                    cell.hyperlink.ref = cell.coordinate
            sheet.row_dimensions[r].height = old.row_dimensions[original[0].row].height
            sheet.row_dimensions[r].outlineLevel = 1
            locations[abb] = {"row":r,"group":group,"source_row":original[0].row}
    workbook.remove(old)
    sheet.title = "1. imbal"
    sheet.auto_filter.ref = f"A3:P{sheet.max_row}"
    workbook.active = 0
    with tempfile.NamedTemporaryFile(dir=target.parent,suffix=".xlsx",delete=False) as f:
        staged = Path(f.name)
    try:
        workbook.save(staged)
        with zipfile.ZipFile(staged) as z:
            assert z.testzip() is None
            assert not any("comment" in n.lower() for n in z.namelist())
        checked = openpyxl.load_workbook(staged)
        for name, state in other.items():
            assert snapshot(checked[name]) == state, f"Unrelated sheet changed: {name}"
        updated = checked["1. imbal"]
        assert updated.max_column == 16 and updated.freeze_panes == "B4"
        assert [c.value for c in updated[3]] == headers
        highlighted = 0
        for abb, location in locations.items():
            r = location["row"]
            assert [updated.cell(r,c).value for c in range(1,17)] == [c.value for c in records[abb]]
            for c, prior in enumerate(records[abb],1):
                assert copy(updated.cell(r,c).font) == copy(prior.font), (abb,c,"font")
                if prior.hyperlink:
                    assert updated.cell(r,c).hyperlink.target == prior.hyperlink.target
            if records[abb][0].fill.fgColor.type == "rgb" and records[abb][0].fill.fgColor.rgb[-6:] == "CFE2F3":
                assert copy(updated.cell(r,1).fill) == copy(records[abb][0].fill)
                highlighted += 1
        for group, abbs in GROUPS:
            ordered = sorted(abbs,key=lambda a:locations[a]["row"])
            years = [year(records[a][3].value) for a in ordered]
            assert years == sorted(years,reverse=True)
            for left,right in zip(ordered,ordered[1:]):
                assert updated.cell(locations[left]["row"],2).fill != updated.cell(locations[right]["row"],2).fill
        assert not any(c.comment for s in checked for row in s for c in row)
        checked.close()
        assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash, "Input changed during editing"
        staged.replace(target)
    finally:
        staged.unlink(missing_ok=True)
    report = {"source":str(source),"restored_from_backup":source != target,"output":str(target),
              "source_sha256":original_hash,"output_sha256":hashlib.sha256(target.read_bytes()).hexdigest(),
              "backup":str(backup),"papers":85,"groups":{g:len(a) for g,a in GROUPS},
              "preserved_highlights":highlighted,"locations":locations,
              "validation":["85 paper rows unchanged", "single common classification independent of publication year",
                            "descending years per group", "alternating fills and original A highlights",
                            "16 columns; no dates or comments; Abb. freeze", "other two sheets unchanged"]}
    evidence = target.parent/"20261027/imbal_unified_classification.json"
    evidence.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({k:v for k,v in report.items() if k != "locations"},ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
