"""Build the research supplement from cached primary-source arXiv metadata.

The XML inputs are the arXiv API responses saved during the 2026-09-09 search.
Rebuilding also works from the exported arxiv_metadata_checked.csv alone.
Only bibliographic fields, original screening notes, and URLs are exported.
"""
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent
CACHE = Path("/tmp/imbalcic-research-0909")
NS = {"a": "http://www.w3.org/2005/Atom"}

NOTES = [
    ("2609.07956", "신규", "HINT", "높음", "Workshop (arXiv comment: ECML PKDD 2026)", "검색 context와 선택적 TabPFN 호출을 결합. kNN confidence 기반 offload; 별도 사후 verifier는 제시하지 않음.", "6개 stream의 첫 20k; IDS tail 결과로 일반화 불가."),
    ("2609.05955", "신규", "LoGIC", "높음", "Preprint", "구조·특징·피복 검색과 cluster별 공유 context. label context와 unlabeled halo 예산 분리.", "그래프 node classification. NetFlow 행 분류와 입력 구조가 다름."),
    ("2609.03880", "신규", "Xiaomi-TabLDM", "중간", "Technical report / preprint", "공유 expert와 routed expert를 가진 내부 sparse MoE, context와 test-time compute 확장.", "사전학습 모델 내부의 FFN MoE. frozen PFN 외부 context expert routing과 구별."),
    ("2609.04540", "신규", "Mitra-v2", "중간", "Technical report / preprint", "synthetic prior 확장, 장문 context, 공개 backbone 비교 후보.", "일반 benchmark 성능은 CIC2018 tail 성능의 증거가 아님."),
    ("2609.06912", "신규", "Structural coverage", "중간", "Preprint", "합성 pretraining prior의 구조적 피복과 downstream 성능의 관계를 조사.", "inference context 선택 연구와 다른 층위; 연관성을 인과적 보장으로 해석하지 않음."),
    ("2608.16319", "기존 항목 버전 갱신", "TabPFN-Rel", "중간", "Preprint", "xlsx에 있던 항목의 v2가 9월 7일 제출됨을 확인.", "버전 갱신 사실을 확인했으며 본문의 어떤 기법이 v2에서 새로 생겼는지는 주장하지 않음."),
    ("2605.27254", "기존 조사 누락 보완", "LUCoS", "높음", "Preprint", "unsupervised PFN latent geometry에서 medoid context 선택. 거리 표현의 타당성 대조군.", "low-label cold-start 67 datasets; 100k IDS의 효과는 검증 필요."),
    ("2607.17962", "기존 조사 누락 보완", "Context topology", "중간", "Preprint", "내부 표현의 topology와 context-level reliability 간 연관성.", "합성 geometry 진단 연구; 표본별 gate의 보장이나 완성된 선택기로 인용하지 않음."),
    ("2603.14324", "기존 조사 누락 보완", "L2D with advice", "매우 높음", "Preprint", "expert와 추가 정보의 결합 action 공간을 최적화하는 일관성 이론.", "현재 scorer/verifier 구조가 그 논문의 반례라는 뜻은 아님. 제안 구조의 이론적 출발점."),
    ("2506.20650", "기존 조사 누락 보완", "Multi-expert routing", "높음", "ICML 2025 (PMLR verified)", "고정 expert의 정확도·비용을 함께 최적화하는 surrogate 및 H-consistency.", "유한 표본·시간 이동·CIC2018에서 자동 보장은 아님."),
    ("2602.17144", "기존 조사 누락 보완", "When More Experts Hurt", "중간", "Preprint metadata checked", "multi-expert deferral 목적의 underfitting 문제와 PiCCE 제안.", "공동 학습 classifier의 결과를 frozen TabPFN에 그대로 적용하지 않음."),
    ("2110.01052", "이론 보완", "Learn then Test", "높음", "Primary paper", "보정 정책의 위험 통제를 다중 가설 검정으로 설계.", "고정 후보 정책과 표본 가정 필요. 시간 상관 데이터에 iid 보장을 그대로 붙이지 않음."),
    ("2202.13415", "이론 보완", "Beyond exchangeability", "높음", "Annals of Statistics 2023", "시간 이동 아래 conformal coverage 손실과 가중 보정의 조건을 다룸.", "임의 concept drift에 대해 무조건적 finite-sample 보장을 주지 않음."),
    ("2608.12989", "기존 조사 재검토", "BAPS", "높음", "Preprint", "class balance·대표성·경계·밀도·다양성의 context 구성. 가장 직접적인 구성 baseline.", "512 prototype 중심; IDS multi-class·충돌·시간 분할과 차이."),
    ("2606.11473", "기존 조사 재검토", "CRUMB", "높음", "Preprint", "query clustering + MMD 기반 context matching + batch 공유.", "test X를 묶는 transductive batch 설정. X matching만으로 tail label 피복은 보장되지 않음."),
    ("2405.16156", "기존 조사 재검토", "MixturePFN", "매우 높음", "ICLR 2025 (기존 서베이)", "local context bank·nearest-centroid routing·배포 검색에 맞춘 adaptation.", "context expert bank 자체의 신규성을 주장하기 어려운 직접 선행연구."),
    ("2607.26628", "기존 조사 재검토", "Context sampling", "높음", "Preprint", "context 크기·random draw·coverage·선택 비용을 비교.", "15개 small OpenML; 큰 context가 모든 IDS에서 더 낫다는 정리가 아님."),
    ("2607.25532", "기존 조사 재검토", "Spurious routing", "높음", "Preprint", "합성 composite feature 환경에서 spurious signal과 context 크기 상호작용, environment 구성 완화책.", "정리는 ridge ICL, TabPFN은 경험적 증거. expert 선택 routing과 용어를 구별."),
    ("2605.21742", "기존 조사 재검토", "PFN imbalance", "높음", "IEEE ITW 2026 (기존 서베이)", "binary imbalance에서 threshold·sampling·calibration 분리 비교.", "다중클래스 macro-F1와 동일한 목표는 아님. threshold baseline 필요성을 뒷받침."),
    ("2605.04363", "기존 조사 재검토", "DistPFN", "높음", "ICML 2026 (기존 서베이)", "label-shift posterior adjustment; local repro 0907~0908 기록과 함께 비교.", "우리 latest n=4에서 tail 이득 없음; n별 효과와 지표 차이를 구별."),
    ("2406.05207", "기존 조사 재검토", "LoCalPFN", "높음", "NeurIPS 2024 (기존 서베이)", "retrieval 및 검색 배포 조건에 맞춘 fine-tuning.", "local 재현은 v1, ours는 v3이므로 wrapper 기여를 분리하려면 같은 backbone 필요."),
    ("2402.11137", "기존 조사 재검토", "TuneTables", "중간", "NeurIPS 2024 (기존 서베이)", "context 최적화·압축을 통한 PFN 확장.", "context 최적화라는 포괄적 주장 자체는 신규성이 낮음."),
    ("2605.13986", "기존 조사 재검토", "TabPFN-3", "높음", "Technical report / preprint", "현 backbone은 1M rows까지 확장. 버전별 context 제약을 구분.", "실용적인 비용·정확도 optimum은 물리적 최대 길이와 다름."),
]


def main():
    metadata = {}
    saved = OUT / "arxiv_metadata_checked.csv"
    if saved.exists():
        for rec in pd.read_csv(saved, keep_default_na=False).to_dict("records"):
            found = re.search(r"(\d{4}\.\d{4,5})(v\d+)?$", rec["id"])
            if found:
                metadata[found[1]] = rec
    for file in CACHE.glob("*metadata.xml"):
        for entry in ET.fromstring(file.read_text()).findall("a:entry", NS):
            rec = {key: " ".join(entry.findtext("a:" + key, default="", namespaces=NS).split())
                   for key in ["id", "title", "published", "updated"]}
            found = re.search(r"(\d{4}\.\d{4,5})(v\d+)?$", rec["id"])
            if found:
                metadata[found[1]] = rec
    recent = CACHE / "arxiv_recent.xml"
    entries = ET.fromstring(recent.read_text()).findall("a:entry", NS) if recent.exists() else []
    for entry in entries:
        rec = {key: " ".join(entry.findtext("a:" + key, default="", namespaces=NS).split())
               for key in ["id", "title", "published", "updated"]}
        metadata[re.search(r"(\d{4}\.\d{4,5})v", rec["id"])[1]] = rec
    rows = []
    for aid, status, abb, relevance, venue, idea, limits in NOTES:
        rec = metadata[aid]
        rows.append(dict(status=status, abbreviation=abb, title=rec["title"],
            first_submitted_utc=rec["published"], last_updated_utc=rec["updated"],
            version=rec["id"].split("v")[-1], relevance=relevance, source_status=venue,
            connection_to_project=idea, claim_boundary=limits,
            url="https://arxiv.org/abs/" + aid, reviewed_on="2026-09-09"))
    data = pd.DataFrame(rows)
    data.to_csv(OUT / "literature_supplement.csv", index=False)
    xlsx = OUT / "literature_supplement.xlsx"
    data.to_excel(xlsx, index=False, sheet_name="20260909 supplement")
    book = load_workbook(xlsx)
    sheet = book.active
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24475C")
    for col, width in {"A":24,"B":24,"C":65,"D":24,"E":24,"F":10,"G":12,"H":42,"I":70,"J":70,"K":42,"L":16}.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 64
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[10].hyperlink = row[10].value
        row[10].style = "Hyperlink"
    book.save(xlsx)
    allmeta = pd.DataFrame(metadata.values()).sort_values("updated", ascending=False)
    allmeta.to_csv(OUT / "arxiv_metadata_checked.csv", index=False)
    print("Supplement:", len(data), "papers. Metadata records:", len(allmeta))


if __name__ == "__main__":
    main()
