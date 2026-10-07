#!/usr/bin/env python3
"""Intent-first Korean deck. Existing EXP61/63 evidence only; no experiments launched."""
import hashlib, json, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from slide_style import *

BASE = HERE / 'archive/20261002_before_intent_revision/docs/slides/kor/1002_labmeeting_ko_draft.pptx'
prs = Presentation(BASE)
old = list(prs.slides)
assert len(old) == 12
sota = json.loads((ROOT/'docs/research/20260930/exp61_report_data.json').read_text())
cap = json.loads((ROOT/'docs/research/20260930/exp63_capability_report_data.json').read_text())

# 1. Research intent, distinguished from demonstrated results.
s = rebuild(old[0], '연구 목표 · 라벨이 적은 공격을 잘 탐지하는 IDS', '01 연구 의도')
subtitle(s, '제한된 context 안에서 희소 공격의 구분에 필요한 사례를 구성 · 클래스별 성능과 비용으로 검증')
box(s,.62,1.55,12.09,1.12,fill=HEAD)
text(s,.84,1.72,11.65,.72,'공격의 라벨 사례가 적을 때, 무엇을 context에 담아야\nBenign과 다른 공격 사이에서 그 공격을 더 잘 구분할 수 있는가?',size=22,bold=True,ink=NAVY,inset=0)
for x,title,body in [(.62,'대상 · 희소 공격','학습 사례가 적은 공격의 탐지 개선\n전체 클래스의 일률적 개선은 부차 목표'),(4.75,'제약 · 라벨과 메모리','제한된 라벨 획득량과 context 크기\nContext 구성·모델 학습·추론 비용'),(8.88,'보존 · 기존 탐지 능력','Benign 오탐과 기존 공격 성능 확인\n희소 공격 개선과 함께 평가')]:
 card(s,x,3.00,3.83,1.88,title)
 text(s,x+.17,3.65,3.49,1.04,body,size=14,inset=0)
text(s,.62,5.32,12.09,.38,'아이디어: Global의 공통 사례 + 공격 구분에 필요한 전문 사례를 context로 구성',size=18,bold=True,ink=NAVY,inset=0)
text(s,.62,5.90,12.09,.53,'현재 확인: 고정 test에서 일부 클래스 개선     다음 검증: 희소 공격 성능과 구성·추론 비용',size=14,ink=GREY,inset=0)
footer(s,'검증할 주장: 제한된 라벨·context 예산에서 희소 공격 탐지 개선')
notes(s,'사용자 피드백 2026-10-02. 목표와 이미 입증된 결과를 분리한다. 새 공격은 첫 실험에서 반드시 unseen label이라는 뜻이 아니다. OOD는 후속 확장으로 정의. TabPFN의 희소 공격 성능 및 비용 우위는 아직 미검증. 최종 사용자 확인: 소수 라벨 확보 전후 적응은 필수 목표가 아님. EXP64–68을 합의된 연구 설계나 이번 발표의 주된 근거로 채택하지 않는다.')

# 2. Why this mechanism is a reasonable research subject, not a cost claim.
s = rebuild(old[1], '왜 TabPFN인가 · 적은 사례를 context로 활용하는 분류기', '01 연구 의도')
subtitle(s,'XGBoost는 강한 비교 기준 · TabPFN의 선택 근거는 사전학습 지식과 소수 context 사례의 활용 가능성')
card(s,.62,1.58,5.90,2.85,'XGBoost · 주어진 데이터로 학습')
text(s,.82,2.23,5.50,1.84,'주어진 라벨 데이터로 결정 규칙 학습\n\n현재 고정 test에서 강한 성능과 빠른 추론\n희소 공격에서도 반드시 비교할 기준',size=16,inset=0)
card(s,6.81,1.58,5.90,2.85,'TabPFN · 사례를 조건으로 추론')
text(s,7.01,2.23,5.50,1.84,'사전학습된 backbone과 라벨 context로 추론\n\n가중치 재학습 없이 사례 구성 변경 가능\n적은 공격 사례의 효과적인 활용이 가설',size=16,inset=0)
box(s,.62,4.78,12.09,1.64,fill=HEAD)
text(s,.84,4.94,11.65,.38,'비교해야 할 질문',size=18,bold=True,ink=NAVY,inset=0)
text(s,.84,5.47,11.65,.78,'동일한 라벨 정보에서, 어떤 context 구성이 희소 공격의 구분을 돕는가?\nXGBoost / Global TabPFN / 제안 context·expert 구성의 성능과 비용 비교',size=15,inset=0)
footer(s,'TabPFN의 희소 공격 성능·비용 이점은 가설 · XGBoost와 동등한 정보로 검증')
notes(s,'XGBoost는 전체 재학습만 가능한 모델이 아님: https://xgboost.readthedocs.io/en/stable/python/examples/continuation.html . TabPFN context 변경은 backbone gradient update가 없지만 전처리, 캐시, calibration, validation 및 추론 비용은 발생. 학습형 임베딩·context 최적화는 별도 최적화 비용을 포함. 현재 EXP61의 shared-GPU 시간은 실제 온라인 갱신 비용의 증거가 아니다.')

# Existing editable architecture / audit / distributions / context pages are retained.
for s,section in [(old[4],'02 접근'),(old[2],'03 데이터 검토'),(old[3],'03 데이터 검토'),(old[5],'04 현재 근거')]:
 for sh in s.shapes:
  if sh.has_text_frame and abs(sh.left-Inches(9.05))<Inches(.04) and sh.top>Inches(7):
   set_text(sh.text_frame,section,size=10,ink='D6E4F3',align=PP_ALIGN.RIGHT)
old[4].notes_slide.notes_text_frame.text += '\n이 구조는 지금까지 구현한 출발점이다. 순차 도입 시나리오의 Global/expert/S/V 갱신 여부는 미확정이며 필수 조건 아님. 새 라벨을 반영한다고 기존 S/V가 자동으로 유효해지는 것은 아니다.'

# 7. Compact evidence, avoiding a leaderboard-led argument.
s = new_slide(prs,'현재 근거 · residual context는 일부 공격에서 효과', '04 현재 근거')
subtitle(s,'정제 데이터 · 고정 test · seed 43 · 대표 K = 4 · 순차 공격 도입 시나리오는 아직 설계 단계')
rows=[]
for ds,name in [('cic2018','CIC2018'),('toniot','ToN')]:
 d=sota['datasets'][ds]['results']; p=cap['datasets'][ds]['banks']['4']['policy']
 rows.append([name,f'{d["xgb"]["macro_f1"]:.4f}',f'{d["global_raw"]["macro_f1"]:.4f}',f'{p["macro_f1"]:.4f}',f'{100*p["proposed"]/p["rows"]:.2f}',f'{100*p["accepted"]/p["rows"]:.2f}'])
table(s,.62,1.64,[1.45,2.09,2.09,2.09,2.19,2.18],['데이터','XGB 100k F1','Global F1','현재 모델 F1','Expert 호출 (%)','출력 채택 (%)'],rows,row_h=.60,font=13)
text(s,.62,3.62,12.09,.32,'F1은 Macro-F1 · 현재 모델은 expert/route 추가 학습 데이터 사용 · 같은 100k 예산의 우위 주장 불가',size=12.5,ink=GREY,inset=0)
card(s,.62,4.18,5.90,1.94,'관찰 · 무엇이 달라졌나')
text(s,.82,4.86,5.50,1.10,'ToN Scanning F1: 0.0394 → 0.8587\nCIC2018: 호출 0건으로 Global 예측 유지\n모든 tail의 개선을 확인한 상태는 아님',size=14.5,inset=0)
card(s,6.81,4.18,5.90,1.94,'다음 질문 · 어떤 정보가 필요한가')
text(s,7.01,4.86,5.50,1.10,'공통·전문 사례의 구성으로 얻은 개선인가?\n라벨이 적은 다른 공격에도 효과가 있는가?\nBenign 오탐과 구성·추론 비용의 변화는?',size=14.5,inset=0)
footer(s,'현재 결과는 context 구성 연구의 출발점 · 희소 공격 전반의 개선은 추가 검증')
notes(s,'Sources: EXP61 report_data and EXP63 K4 policy. K4 is the existing representative configuration, not test-best selection. Logical call rate applies the policy to cached expert predictions; it is not measured conditional latency. Full SOTA and per-expert class F1 are in the appendix. No claim that experts fail due solely to benign overlap or that a class is impossible.')
summary=s

# 8. Sequential arrivals are an evaluation axis, not necessarily adaptation.
s = rebuild(old[11], '평가 방향 · 공격을 순차적으로 도입해 희소 공격 탐지 확인', '05 다음 검증')
subtitle(s,'설계안 · 공격 도입 순서·혼합 비율을 통제 · 보유 라벨 수와 context 갱신 여부는 별도 설정')
steps=[('초기 배경','Benign과 일부 공격이\n함께 존재하는 트래픽'),('공격 도입','라벨이 적은 공격이\n순차적으로 출현'),('혼합 변화','기존·신규 도입 공격의\n혼합 비율 변화'),('성능 관찰','공격별 탐지와 오탐\n기존 공격 성능 추적')]
for i,(title,body) in enumerate(steps):
 x=.62+i*3.08
 card(s,x,1.62,2.85,1.72,title)
 text(s,x+.15,2.29,2.55,.83,body,size=14,inset=0)
 if i<3:text(s,x+2.86,2.17,.20,.40,'›',size=24,bold=True,ink=BLUE,inset=0)
card(s,.62,3.72,5.90,2.65,'먼저 볼 것 · 라벨이 적은 공격의 구분')
text(s,.82,4.38,5.50,1.77,'동일한 학습 정보·공격 도입 순서로 비교\n공격별 Precision / Recall / F1과 Benign 오탐\n구간별 혼합 조건의 영향 확인\nContext 갱신 여부는 별도 설계 항목',size=14.5,inset=0)
card(s,6.81,3.72,5.90,2.65,'후속 확장 · OOD와 새로운 공격')
text(s,7.01,4.38,5.50,1.77,'평가 흐름에 뒤늦게 등장한 공격과\n학습에 없던 공격은 서로 다른 조건\n먼저 소수 라벨 조건의 탐지 성능 검증\n이후 미지 공격·분포 밖 입력으로 확장',size=14.5,inset=0)
footer(s,'평가의 중심은 희소 공격 탐지 · 순차 도입과 온라인 갱신은 별도 축')
notes(s,'사용자 확인 2026-10-02: 실제 수집 시간을 반드시 따를 필요 없음. 소수 라벨 확보 전후 적응 비교도 필수 아님. 목표는 전체 성능 개선이 아니라 소수 라벨 공격의 성능 개선. 본 순차 도입 흐름은 개념 설계안이며 구체 공격 순서·비율·구간 수는 미확정. EXP68을 채택하지 않음. 기본 분류 평가에서 attack은 소수 학습 라벨이 존재할 수 있으므로 나중에 test에 출현한다고 OOD는 아님. 라벨 수와 공격 도입 시점은 분리 통제. 향후 온라인 갱신을 추가할 경우에만 predict-before-label 및 라벨 지연을 엄수. 고정 context와 동일 표본에서는 평가 순서만 바꿔도 예측은 동일. 순차 시나리오는 구간별 혼합 조건의 평가이며 순서만으로 적응 효과를 주장하지 않는다.')

# 9. Hypotheses, not verdicts on unagreed screening runs.
s = rebuild(old[10], 'Context 설계 후보 · 부족한 정보에 대한 가설부터 검증', '05 다음 검증')
subtitle(s,'Anchor + residual block은 하나의 출발점 · 같은 라벨·메모리 예산에서 구성 요소의 역할을 확인')
rows=[['1  Benign 다양화','정상 트래픽의 여러 동작을 대표','공격 recall 유지 시 Benign 오탐 감소'],['2  Tail 합성','희소 공격의 관측 범위를 보완','동일 개수 실제 사례·단순 복제와 비교'],['3  표현 학습·임베딩','유사한 입력에서도 클래스 구분 강화','원본 feature 대비 분리력과 추가 학습 비용'],['4  Context 최적화','적은 행에 유용한 정보를 집중','동일 context 크기의 선택 방식과 비교']]
table(s,.62,1.62,[2.55,4.45,5.09],['접근','검증할 가설','확인할 변화'],rows,row_h=.74,font=14)
text(s,.62,5.56,12.09,.37,'2  LITO (ICLR 2024)   ·   3  Supervised contrastive prototype embedding   ·   4  In-context data distillation',size=11.8,ink=GREY,inset=0)
text(s,.62,6.12,12.09,.42,'표현 학습·합성·최적화의 비용도 합산 · 성능이 오른 이유를 분리할 수 있는 비교 조건 설계',size=15,bold=True,ink=NAVY,inset=0)
footer(s,'가설 → 정보·비용을 통제한 비교 → 효과 확인의 순서로 context 설계')
notes(s,'후속 방법은 사용자 요청의 네 접근. 기존 EXP64–67 결과는 원본에 보존하되 이번 주장에 자동 채택하지 않는다. 세부 screening 조건과 방법 충실성 검토 후 사용 여부 판단. LITO: https://proceedings.iclr.cc/paper_files/paper/2024/file/5d54d2df6ec8f7b920aa0fec9a6d1b2e-Paper-Conference.pdf ; embedding: https://github.com/mlopezm/Supervised-contrastive-learning-over-prototype-label-embeddings ; ICD: https://arxiv.org/abs/2402.06971 (ICLR 2024 ME-FoMo workshop; not main conference). https://valthom.github.io/')

# 10. Close on claim/evidence, not another configuration.
s = new_slide(prs,'다음 단계 · 희소 공격 개선이라는 주장에 맞춰 검증', '05 다음 검증')
subtitle(s,'핵심 주장 후보: 제한된 라벨·context 예산에서 희소 공격을 더 잘 구분하고 기존 탐지 능력 유지')
rows=[['① 목표와 조건','대상 희소 공격·보유 라벨 수·순차 도입 방식·context 예산 정의'],['② 단순 비교부터','같은 라벨 데이터로 XGB / Global / 제안 context 구성 비교'],['③ 제안의 역할 확인','Context 구성의 기여 확인 후 expert·S/V 기여 분리'],['④ OOD 확장','새 분포·미지 공격의 인지와 라벨 확보 후 적응을 분리']]
table(s,.62,1.68,[2.65,9.44],['순서','확인할 내용'],rows,row_h=.76,font=16)
box(s,.62,5.87,12.09,.66,fill=HEAD)
text(s,.84,5.98,11.65,.44,'판단 기준: 라벨이 적은 어떤 공격을 더 잘 구분하며, 오탐과 기존 성능을 보존하는 데 드는 비용은?',size=16,bold=True,ink=NAVY,inset=0)
footer(s,'주장 → 희소 공격 중심 평가 → 구성 요소 기여 확인 → OOD 확장')
notes(s,'현재 고정 test EXP61/63 결과는 탐색 근거. 사용자 확인: 실제 timestamp와 소수 라벨 전후 적응 비교는 필수 아님. 희소 공격 탐지가 중심. 임의의 라벨 수/구간 수/효과 목표를 확정하지 않음. 문서 분업: raw results immutable, LaTeX selected evidence, decision log for claims/protocol revisions.')
closing=s

# Main 10 slides plus 4 detailed appendix pages.
order=[old[0],old[1],old[4],old[2],old[3],old[5],summary,old[11],old[10],closing,old[6],old[7],old[8],old[9]]
ids={id(prs.slides[i]):el for i,el in enumerate(prs.slides._sldIdLst)}
ordered_ids=[ids[id(s)] for s in order]
lst=prs.slides._sldIdLst
for el in list(lst):lst.remove(el)
for el in ordered_ids:lst.append(el)
for i,s in enumerate(prs.slides,1):
 for sh in s.shapes:
  if sh.has_text_frame and abs(sh.left-Inches(12.23))<Inches(.05) and abs(sh.top-Inches(7.01))<Inches(.05):set_text(sh.text_frame,str(i),size=11,bold=True,ink='FFFFFF',align=PP_ALIGN.RIGHT)
  if i>10 and sh.has_text_frame and abs(sh.left-Inches(9.05))<Inches(.05) and sh.top>Inches(7):set_text(sh.text_frame,'부록 · 현재 실험 근거',size=10,ink='D6E4F3',align=PP_ALIGN.RIGHT)
prs.core_properties.title='10월 2일 랩미팅 · 희소 공격 탐지를 위한 context 설계'
prs.core_properties.comments='2026-10-02 intent revision: Korean only; EXP61/63 exploratory evidence; temporal protocol draft, OOD extension; EXP68 not adopted.'
prs.save(OUTPUT)

# Verify geometry, numeric alignment, text style; PDF QA follows rendering.
issues=[];dump=[];num=0
for i,s in enumerate(prs.slides,1):
 ts=[];entries=[]
 for sh in s.shapes:
  if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+20 or sh.top+sh.height>prs.slide_height+20:issues.append((i,'outside',sh.name))
  if sh.top<Inches(6.9) and sh.top+sh.height>Inches(6.9):issues.append((i,'footer',sh.name))
  if sh.has_text_frame and sh.text.strip():entries.append(sh.text);ts.append(sh)
  if sh.has_table:
   entries.extend(' | '.join(c.text for c in row.cells) for row in sh.table.rows)
   for ri,row in enumerate(sh.table.rows):
    for ci,c in enumerate(row.cells):
     if ri and ci and NUM.fullmatch(c.text):
      num+=1
      assert all(p.alignment==PP_ALIGN.RIGHT for p in c.text_frame.paragraphs),(i,ri,ci,c.text)
 for j,a in enumerate(ts):
  for b in ts[j+1:]:
   if min(a.left+a.width,b.left+b.width)-max(a.left,b.left)>Inches(.005) and min(a.top+a.height,b.top+b.height)-max(a.top,b.top)>Inches(.005):issues.append((i,'text overlap',a.text[:24],b.text[:24]))
 joined='\n'.join(entries)
 assert not re.search(r'합니다|입니다|습니다|[—–]|학습 0 s|개선 불가|context 최적화만|첫 10%',joined),(i,'wording')
 dump.append(f'## {i}\n{joined}\n')
assert not issues,issues
QA.mkdir(exist_ok=True)
(QA/'slide_text.md').write_text('\n'.join(dump))
(QA/'validation.json').write_text(json.dumps({'slides':len(prs.slides),'main':10,'appendix':4,'generator':'docs/slides/revise_1002_ko_intent.py','numeric_cells_right_aligned':num,'geometry_issues':issues,'temporal_status':'design_pending; EXP68 not adopted','sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest()},ensure_ascii=False,indent=2))
print(OUTPUT, 'slides', len(prs.slides))
