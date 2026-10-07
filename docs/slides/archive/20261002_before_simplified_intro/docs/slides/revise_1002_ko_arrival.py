#!/usr/bin/env python3
"""Korean revision: new-attack incorporation first, static SOTA as current status.

Base is the current user deck archived immediately before this revision.
EXP70 results are deliberately not inserted while the campaign is running.
"""
import csv
import hashlib
import json
import re
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from slide_style import *

BASE=HERE/'archive/20261002_before_new_attack_revision/docs/slides/kor/1002_labmeeting_ko_draft.pptx'
QA=HERE/'kor/1002_arrival_qa'
QA.mkdir(exist_ok=True)
prs=Presentation(BASE);old=list(prs.slides)
assert len(old)==14
sota=json.loads((ROOT/'docs/research/20260930/exp61_report_data.json').read_text())
cap=json.loads((ROOT/'docs/research/20260930/exp63_capability_report_data.json').read_text())
allocation=json.loads((ROOT/'tabpfn/configs/exp70_allocation.json').read_text())
LABELS=dict(benign='Benign',bot='Bot',brute_force='Brute force',ddos='DDoS',dos='DoS',
    infiltration='Infiltration',web_attacks='Web attack',backdoor='Backdoor',injection='Injection',
    mitm='MitM',password='Password',ransomware='Ransomware',scanning='Scanning',xss='XSS')


def arrow(s,x1,y1,x2,y2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=color(BLUE);sh.line.width=Pt(1.7)
    tip=etree.SubElement(sh.line._get_or_add_ln(),qn('a:tailEnd'));tip.set('type','triangle')
    return sh


def node(s,x,y,w,h,title,body,fill=CARD):
    box(s,x,y,w,h,fill=fill,border=LINE)
    text(s,x+.15,y+.13,w-.30,.38,title,size=17,bold=True,ink=NAVY,inset=0)
    text(s,x+.15,y+.68,w-.30,h-.83,body,size=15,inset=0)


def retitle(s,title,section):
    for sh in s.shapes:
        if not sh.has_text_frame:continue
        if sh.top<Inches(.96) and sh.text.strip():set_text(sh.text_frame,title,size=28,bold=True,ink=NAVY,inset=0)
        if sh.left>Inches(9) and sh.top>Inches(7) and sh.width>Inches(1):
            set_text(sh.text_frame,section,size=10,ink='D6E4F3',align=PP_ALIGN.RIGHT)


# 1. Operational problem and the user's actual central question.
s=rebuild(old[0],'새로운 공격 사례를 반영하는 IDS','01 연구 아이디어')
subtitle(s,'공격 C의 라벨 사례 확보 이후 · 기존 공격과 새 공격을 함께 분류하는 탐지기 갱신')
for x,title,body in [(.62,'기존 탐지','Benign · 공격 A · 공격 B\n기존 클래스 분류'),(4.79,'새 공격 사례 확보','분석·라벨링된 공격 C\n탐지기에 반영할 사례'),(8.96,'탐지기 갱신','Benign · 공격 A · B · C\n기존 탐지와 새 공격 대응')]:
    node(s,x,1.78,3.75,1.72,title,body)
arrow(s,4.40,2.65,4.75,2.65);arrow(s,8.57,2.65,8.92,2.65)
box(s,.62,4.04,12.09,1.85,fill=HEAD)
text(s,.88,4.23,11.57,1.43,
    '새로운 공격의 라벨 사례가 확보됐을 때,\nTabPFN은 context 확장만으로 XGBoost 재학습보다\n효과적으로 새로운 공격을 분류할 수 있는가?',size=23,bold=True,ink=NAVY,inset=0)
text(s,.62,6.22,12.09,.36,'새 공격 분류 · 기존 탐지 유지 · 갱신과 추론에 드는 비용',size=17,ink=BLUE,bold=True,inset=0)
footer(s,'연구 대상: 새로운 공격 사례를 탐지기에 반영하는 방식')
notes(s,'발표의 중심 질문은 사용자가 명시한 신규 공격 라벨 확보 이후의 context 확장 대 XGB 재학습 비교다. Tail/context 최적화나 평균 F1 순위가 발표의 중심을 대체하지 않게 설명한다. 라벨 확보 이전 미지 공격 인지는 본 평가와 별개의 후속 OOD 주제다. 기술적 갱신 메커니즘은 설명하되 성능·비용 우위를 이미 입증했다고 말하지 않는다.')

# 2. Editable mechanism diagram, not a leaderboard.
s=rebuild(old[1],'모델 재학습과 context 확장을 통한 갱신','01 연구 아이디어')
subtitle(s,'동일한 기존 사례와 신규 공격 사례 제공 · 갱신 방식의 차이에 대한 비교')
for y,method,mid_title,mid_body,right_title,right_body,fill in [
    (1.70,'XGBoost','누적 데이터로 재학습','분류 규칙·모델 파라미터 갱신','갱신된 XGB','기존 공격 + 새 공격 분류',CARD),
    (3.67,'TabPFN','Context 확장','Backbone 고정\n전처리·캐시 재구성','동일한 TabPFN','확장된 context로 분류',HEAD)]:
    text(s,.62,y+.10,1.55,.40,method,size=20,bold=True,ink=NAVY,inset=0)
    node(s,2.31,y,2.90,1.54,'동일한 누적 사례','기존 사례 + 새 공격 사례',fill)
    node(s,5.74,y,3.41,1.54,mid_title,mid_body,fill)
    node(s,9.67,y,3.04,1.54,right_title,right_body,fill)
    arrow(s,5.24,y+.77,5.69,y+.77);arrow(s,9.18,y+.77,9.62,y+.77)
text(s,.62,5.70,12.09,.45,'TabPFN의 구조적 특성: backbone 가중치 재학습 없이 사례 구성 변경',size=19,bold=True,ink=NAVY,inset=0)
text(s,.62,6.24,12.09,.34,'효용 검증: 새 공격 성능·기존 탐지 유지와 전처리·캐시·추론 비용의 동시 비교',size=14.5,ink=GREY,inset=0)
footer(s,'가중치 재학습과 사례 갱신의 차이 · 실제 효용은 성능과 비용으로 검증')
notes(s,'그림의 XGBoost는 본 비교에서 채택한 누적 데이터 전체 재학습 전략이다. XGBoost의 모든 사용 방식에 재학습만 가능하다는 주장이 아니다. TabPFN fit에는 전처리·캐시 구성이 있으며 first prediction 지연과 전체 추론 비용도 측정한다. EXP70은 Global TabPFN 비교이며 S/V 전체의 무학습 갱신을 검증한 실험은 아니다.')

# 3. Preserve the existing, user-reviewed audit visualization.
audit=old[3];retitle(audit,'입력·라벨 품질 검토 · 동일 feature의 상충 라벨','02 데이터 검토')
audit.notes_slide.notes_text_frame.text+='\n서사 연결: 새 공격 반영 방식의 효과를 비교하기 위한 데이터 기반 점검. 라벨 상충의 원인을 모두 시간 구간 라벨링으로 단정하지 않는다.'

# 4. Explicit whole-data before/after versus fixed test support.
s=rebuild(old[4],'정제 이후 클래스 분포와 고정 test','02 데이터 검토')
subtitle(s,'정제 전후: 전체 데이터 · Test: 정제 후 평가 표본 · 정제 후 전체 표본 수 내림차순')
distribution_checks={}
for ds,title,x in [('cic2018','CIC2018',.62),('toniot','ToN',6.81)]:
    dist=list(csv.DictReader((HERE/f'kor/1002_intro_assets/{ds}_class_distribution.csv').open()))
    dist=sorted(dist,key=lambda r:-int(r['after']))
    support={c['name']:c['support'] for c in sota['datasets'][ds]['results']['global_raw']['classes']}
    rows=[[LABELS[r['class']],f"{int(r['before']):,}",f"{int(r['after']):,}",f"{support[r['class']]:,}"] for r in dist]
    rows.append(['합계',f"{sum(int(r['before']) for r in dist):,}",f"{sum(int(r['after']) for r in dist):,}",f'{sum(support.values()):,}'])
    text(s,x,1.51,5.90,.38,title,size=19,bold=True,ink=NAVY,inset=0)
    sh=table(s,x,1.99,[1.41,1.61,1.55,1.33],['클래스','정제 전 전체','정제 후 전체','Test'],rows,row_h=.372,font=11.2)
    highlight(sh,len(rows),list(range(4)),ink=NAVY)
    distribution_checks[ds]=dict(order=[r['class'] for r in dist],test_total=sum(support.values()),rows=rows)
box(s,.62,5.69,5.90,.80,fill=HEAD)
text(s,.81,5.80,5.52,.56,'Tail: 고정 test에서 표본 수가 적은 클래스\nContext에 제공한 학습 사례 수와 별도 표기',size=14,bold=True,ink=NAVY,inset=0)
footer(s,'평가에서의 클래스 희소성과 context의 학습 사례 수 구분')
notes(s,'정제 전후 집계는 전체 train+validation+test이다. Test 열은 별도 부분집합이며 이를 정제 후 전체에 다시 더하지 않는다. 원본: clean split class_counts.csv와 EXP61 global_raw 클래스 support. 사용자 확정: tail은 고정 test 표본 수 기준으로 해석. 학습 샘플 배분이 균등하더라도 평가의 tail은 유지된다. 특정 tail threshold는 임의로 추가하지 않는다.')

# 5. Existing abstract architecture remains, with precise frozen/learned scope.
architecture=old[2];retitle(architecture,'Context 기반 Global·expert 모델 구조','03 제안 구조')
for sh in architecture.shapes:
    if not sh.has_text_frame:continue
    if '학습 데이터의 역할 분리' in sh.text:
        set_text(sh.text_frame,'공통 context와 전문 context 활용 · TabPFN backbone 고정 · Scorer·Verifier 별도 학습',size=13,ink=GREY,inset=0)
    if sh.top>Inches(7) and sh.left<Inches(1) and 'Backbone' in sh.text:
        set_text(sh.text_frame,'공통·전문 사례로 분류기 구성 · S/V로 예측 교정 선택',size=11,bold=True,ink='FFFFFF',inset=0)
notes(architecture,'현재 구현한 구조는 Global/expert/route 풀을 분리하고, anchor+residual block으로 expert context를 구성하며 S/V를 별도 학습한다. TabPFN backbone이 고정이라는 말은 S/V까지 학습 없이 갱신된다는 뜻이 아니다. 순차 도입 EXP70에서는 이 구조의 기반인 XGB 대 Global TabPFN의 사례 갱신부터 검증하며, 완전판의 신규 공격 갱신 규칙과 비용은 후속 검증이다.')

# 6. Static scorecard as current status; keep actual full-model evidence from EXP63.
s=rebuild(old[12],'정제 데이터의 현재 성능 · SOTA 비교','04 현재 성능')
subtitle(s,'전체 공격 클래스를 학습한 정적 조건 · 동일 정제 test · Seed 43 · Macro-F1 (%)')
rows=[];sota_values=[]
for method,title in [('global_raw','Global TabPFN v3'),('xgb','XGBoost'),('boostpfn','BoostPFN'),('localpfn','LoCalPFN · FT'),('distpfn','DistPFN · v3'),('xgb_full','XGBoost · 전체 train')]:
    vals=[100*sota['datasets'][ds]['results'][method]['macro_f1'] for ds in ['cic2018','toniot']]
    rows.append([title,*[f'{v:.2f}' for v in vals],'공통 100,000행' if method!='xgb_full' else '전체 train'])
    sota_values.append(dict(method=method,cic2018=vals[0],toniot=vals[1]))
vals=[100*cap['datasets'][ds]['banks']['4']['policy']['macro_f1'] for ds in ['cic2018','toniot']]
rows.append(['제안 모델 · K = 4 · S/V',*[f'{v:.2f}' for v in vals],'Global + Expert·Route 풀'])
sh=table(s,.62,1.62,[4.30,1.55,1.55,4.69],['방법','CIC2018','ToN','학습 정보'],rows,row_h=.48,font=14)
highlight(sh,7,list(range(4)),fill=HEAD,ink=NAVY)
text(s,.62,5.63,12.09,.33,'제안 모델: 기존 anchor + residual block 결과 · Expert·Route의 추가 학습 정보 사용',size=13,ink=GREY,inset=0)
text(s,.62,6.10,12.09,.47,'현재 성능 수준의 확인 · 신규 공격 반영 효과와 갱신·추론 비용은 순차 도입으로 검증',size=17,bold=True,ink=NAVY,inset=0)
footer(s,'정적 분류 성능의 현재 수준 · 신규 공격 반영 방식은 별도 검증')
notes(s,'수치는 완료된 EXP61/63, 대표 K4를 유지. EXP70의 부분 결과는 넣지 않는다. 제안 모델은 Global에 더해 Expert·Route 데이터로 학습하므로 같은 100k 정보 예산의 승리로 주장하지 않는다. EXP70의 Global 결과가 완료돼도 이를 제안 완전판 수치로 바꾸지 않는다. PFN backbones: Global/DistPFN v3, BoostPFN/LoCalPFN v1, LoCalPFN 미세조정 포함. SOTA 표는 현재 정적 성능을 보고하는 역할이며 연구 목표를 평균 F1 순위 경쟁으로 대체하지 않는다. 운영 비용은 부록의 측정 조건을 함께 설명한다.')
sota_slide=s

# 7. One concrete validation slide; detailed quotas are an appendix.
s=rebuild(old[7],'신규 공격의 순차 도입을 통한 검증','05 갱신 방식 검증')
subtitle(s,'EXP70 실행 중 · XGBoost 재학습과 Global TabPFN context 확장 · 동일한 누적 사례 제공')
cic=[('초기 · 2/14~16','Benign\nBrute force · DoS'),('2/20~21','DDoS'),('2/22','Web attack'),('2/28','Infiltration'),('3/2','Bot')]
ton=[('초기 · 4/23~24','Benign\nScanning · DoS'),('4/25','Injection\nDDoS'),('4/26','Password'),('4/27','XSS'),('4/28','Ransomware\nBackdoor'),('4/29','MitM')]
for label,items,ty,cy in [('CIC2018 · 2018년',cic,1.53,2.02),('ToN · 2019년',ton,3.38,3.87)]:
    text(s,.62,ty,12.09,.35,label,size=18,bold=True,ink=NAVY,inset=0)
    gap=.18;w=(12.09-gap*(len(items)-1))/len(items)
    for i,(date,classes) in enumerate(items):
        x=.62+i*(w+gap);box(s,x,cy,w,1.08,fill=HEAD if i==0 else CARD,border=LINE)
        text(s,x+.10,cy+.09,w-.20,.23,date,size=11.8,ink=GREY,inset=0)
        text(s,x+.10,cy+.42,w-.20,.56,classes,size=14,bold=True,ink=NAVY,inset=0,align=PP_ALIGN.CENTER)
        if i+1<len(items):arrow(s,x+w+.02,cy+.54,x+w+gap-.02,cy+.54)
box(s,.62,5.38,12.09,.68,fill=HEAD)
text(s,.84,5.51,11.65,.40,'최종 각 100,000행 · 등장 전 클래스 학습 제외 · 기존 test 유지 · Seed 43',size=16,bold=True,ink=NAVY,inset=0)
text(s,.62,6.33,12.09,.30,'원본의 클래스 최초 출현 순서 보존 · 확인 대상: 새 공격 분류 / 기존 탐지 유지 / 갱신·추론 비용',size=13,ink=GREY,inset=0)
notes(s,'실행 프로토콜: docs/research/20261002/exp70_protocol.md. 승인 배분대로 CIC 초기84820, ToN80712에서100000으로 누적. 클래스별 전체 test에서 평가하며 raw 시간 흐름 또는 실제 네트워크 환경 변화 재현으로 주장하지 않는다. 새 공격의 라벨 사례를 제공한 뒤 분류한다. Tail만 순차적으로 등장하는 설계가 아니다. EXP70은 Global 비교이고 expert·S/V 갱신까지 자동 검증된 것은 아니다. 결과를 채워 넣지 않아도 확정된 검증 설계로 완결된 장이다.')

# 8. Close on system-level utility, as explicitly requested.
s=rebuild(old[9],'후속 검증 · 새로운 공격을 반영하는 시스템','05 갱신 방식 검증')
subtitle(s,'새 공격 사례를 반영하는 방식의 효용 확인 · 성능·기존 탐지 유지·운영 비용의 연결')
for x,title,body in [(.62,'새 공격 반영','동일한 신규 사례 제공\n갱신 후 공격별 분류 성능\n희소 클래스의 탐지 변화'),(4.79,'기존 탐지 유지','동일한 기존 표본 비교\n교정·훼손과 정상 오탐\nContext 구성의 영향 확인'),(8.96,'운영 비용','사례 추가 시 갱신 비용\n다음 갱신까지의 추론 비용\n처리량·메모리의 비교')]:
    node(s,x,1.78,3.75,2.22,title,body)
box(s,.62,4.42,12.09,1.04,fill=HEAD)
text(s,.86,4.60,11.61,.67,'연구의 판단 기준: 새 공격을 얼마나 잘 반영하고,\n기존 탐지 성능을 어느 비용으로 유지하는가',size=21,bold=True,ink=NAVY,inset=0)
text(s,.62,5.88,12.09,.40,'Context 구성 개선 → Expert·S/V의 역할과 갱신 검증 → 운영 조건별 효과 확인',size=17,bold=True,ink=BLUE,inset=0)
text(s,.62,6.42,12.09,.26,'F1은 시스템 효용을 확인하는 지표 중 하나 · 미지 공격 인지는 이후 OOD 확장',size=13,ink=GREY,inset=0)
notes(s,'사용자 피드백의 핵심을 마지막에 다시 강조: 실험마다 전체 F1 순위를 올리는 수치 경쟁이 연구의 목적이 아니다. 새로운 공격의 라벨 사례를 확보했을 때 이를 탐지기에 반영하는 방식의 성능과 운영 효용을 검증한다. Context 후보는 구현 수단이며 부록에 둔다. S/V 학습과 갱신을 포함한 완전판 비용은 별도 측정해야 한다. 이 장은 계획이므로 하단 밴드에 같은 슬로건을 반복하지 않는다.')

# Appendix: retain detailed evidence, and update context candidates' role.
for slide in [old[10],old[11],old[13]]:
    retitle(slide,next(sh.text for sh in slide.shapes if sh.has_text_frame and sh.top<Inches(.96) and sh.text.strip()),'부록 · 상세 근거')
candidate=old[8];retitle(candidate,'Context 구성의 후속 후보','부록 · 구현 후보')
for sh in candidate.shapes:
    if sh.has_text_frame and sh.top>Inches(1) and sh.top<Inches(1.4):
        set_text(sh.text_frame,'새 공격 반영과 기존 탐지 유지를 위한 구성 후보 · 동일 정보량·context 예산에서 효과 분리',size=13,ink=GREY,inset=0)
    if sh.has_text_frame and sh.left<Inches(1) and sh.top>Inches(7):set_text(sh.text_frame,'',size=11)
candidate.notes_slide.notes_text_frame.text+='\n2026-10-02 후속 확정: 네 가지 방법은 연구 목적 자체가 아니라 새 공격 반영 및 기존 탐지 유지를 위한 구현 후보. 기존 상대사례 실험은 본문이나 부록에 추가하지 않는다.'

alloc_slide=new_slide(prs,'최종 100,000행의 단계별 배분','부록 · EXP70 조건')
subtitle(alloc_slide,'Benign 75,000행 + 공격 25,000행 · 가용 수를 고려한 균등 배분 · 전체 clean train에서 추출')
allocation_rows={}
groups={'cic2018':[['benign','brute_force','dos'],['ddos'],['web_attacks'],['infiltration'],['bot']],
        'toniot':[['benign','scanning','dos'],['injection','ddos'],['password'],['xss'],['ransomware','backdoor'],['mitm']]}
for ds,title,x in [('cic2018','CIC2018',.62),('toniot','ToN',6.81)]:
    text(alloc_slide,x,1.58,5.90,.36,title,size=19,bold=True,ink=NAVY,inset=0)
    total=0;rows=[]
    for i,group in enumerate(groups[ds]):
        added=sum(allocation[ds][c] for c in group);total+=added
        desc='Benign · Brute force · DoS' if ds=='cic2018' and not i else ('Benign · Scanning · DoS' if not i else ' · '.join(LABELS[c] for c in group))
        rows.append(['초기' if not i else str(i),desc,f'{added:,}',f'{total:,}'])
    assert total==100000
    table(alloc_slide,x,2.10,[.52,2.74,1.25,1.39],['단계','추가 클래스','추가 수','누적 수'],rows,row_h=.51,font=11.5)
    allocation_rows[ds]=rows
text(alloc_slide,.62,6.04,12.09,.40,'Web attack 450행 · Ransomware 2,151행은 가용 사례 전체 사용',size=16,bold=True,ink=NAVY,inset=0)
text(alloc_slide,.62,6.49,12.09,.27,'XGB와 TabPFN에 동일한 사례 제공 · 이전 소규모 실험 사례를 포함해 증설 · Test 고정',size=13,ink=GREY,inset=0)
notes(alloc_slide,'정확한 클래스별 배분은 tabpfn/configs/exp70_allocation.json. CIC 공격은Web450과나머지각4910. ToN Ransomware2151, Backdoor2857, 나머지각2856. 기존 실험에서Global풀이train의절반이었던 것과달리 전체cleantrain사용. 최종예산100000이며 신규클래스당100000이아니다. 같은날복수공격그룹은합산. 최초출현순서표이므로 클래스표본수정렬의예외다.')

# Final order: 8 main slides, 5 appendix slides. The separate context-results page is removed.
order=[old[0],old[1],audit,old[4],architecture,sota_slide,old[7],old[9],old[10],old[11],old[13],candidate,alloc_slide]
ids={id(prs.slides[i]):el for i,el in enumerate(prs.slides._sldIdLst)}
new_ids=[ids[id(s)] for s in order]
for el in list(prs.slides._sldIdLst):prs.slides._sldIdLst.remove(el)
for el in new_ids:prs.slides._sldIdLst.append(el)
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if sh.has_text_frame and abs(sh.left-Inches(12.23))<Inches(.05) and abs(sh.top-Inches(7.01))<Inches(.05):
            set_text(sh.text_frame,str(i),size=11,bold=True,ink='FFFFFF',align=PP_ALIGN.RIGHT)
prs.core_properties.title='신규 공격 사례를 반영하는 context 기반 IDS'
prs.core_properties.comments='Korean only; 8 main + 5 appendix. Existing EXP61/63 static results; EXP70 protocol only. New-attack incorporation is the central research question.'
prs.save(OUTPUT)

# Numeric/source checks and layout bounds; PDF visual QA is performed after rendering.
issues=[];dump=[];numeric=0
for i,s in enumerate(prs.slides,1):
    content=[];texts=[]
    for sh in s.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+20 or sh.top+sh.height>prs.slide_height+20:issues.append((i,'outside',sh.name))
        if sh.top<Inches(6.9) and sh.top+sh.height>Inches(6.9):issues.append((i,'footer overlap',sh.name))
        if sh.has_text_frame and sh.text.strip():content.append(sh.text);texts.append(sh)
        if sh.has_table:
            for ri,row in enumerate(sh.table.rows):
                content.append(' | '.join(c.text for c in row.cells))
                for ci,c in enumerate(row.cells):
                    if ri and ci and NUM.fullmatch(c.text):
                        numeric+=1
                        assert all(p.alignment==PP_ALIGN.RIGHT for p in c.text_frame.paragraphs),(i,ri,ci)
    for j,a in enumerate(texts):
        for b in texts[j+1:]:
            if min(a.left+a.width,b.left+b.width)-max(a.left,b.left)>Inches(.005) and min(a.top+a.height,b.top+b.height)-max(a.top,b.top)>Inches(.005):issues.append((i,'text boxes overlap',a.text[:20],b.text[:20]))
    joined='\n'.join(content)
    assert not re.search(r'합니다|입니다|습니다|[—–]|상대사례|라벨이 적은 공격이\n순차적으로|갱신 여부는 별도|순차 공격 도입 시나리오는 아직 설계',joined),(i,'stale wording')
    dump.append(f'## {i}\n{joined}\n')
assert not issues,issues
eng=json.loads((HERE/'archive/20261002_before_new_attack_revision/english_sha256.json').read_text())
for p,h in eng.items():assert hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==h
(QA/'slide_text.md').write_text('\n'.join(dump))
(QA/'validation.json').write_text(json.dumps(dict(slides=len(prs.slides),main=8,appendix=5,
    numeric_cells_right_aligned=numeric,geometry_issues=issues,english_unchanged=True,
    distributions=distribution_checks,sota_values=sota_values,proposed_model_k4_percent=vals,
    allocation_rows=allocation_rows,generator=str(Path(__file__).relative_to(ROOT)),
    sha256=hashlib.sha256(OUTPUT.read_bytes()).hexdigest()),ensure_ascii=False,indent=2))
print(OUTPUT,'main=8 appendix=5 numeric cells=',numeric)
