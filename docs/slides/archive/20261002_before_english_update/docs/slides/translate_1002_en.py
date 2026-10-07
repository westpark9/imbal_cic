#!/usr/bin/env python3
"""kor/1002_labmeeting_ko_draft.pptx → eng/1002_labmeeting_en_draft.pptx.

Layout, formatting and numbers stay as they are; Korean text runs (slides, tables,
notes) are replaced through a phrase dictionary, and the four slide-1 matrix images are
redrawn with an English annotation from the saved Korean assets. Korean phrases missing
from the dictionary are reported and fail the run.
"""
import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
IN = HERE / 'kor/1002_labmeeting_ko_draft.pptx'
OUT = HERE / 'eng/1002_labmeeting_en_draft.pptx'
KO_ASSETS = HERE / 'kor/1002_intro_assets'
EN_ASSETS = HERE / 'eng/1002_intro_assets'
QA = HERE / 'eng/1002_draft_qa'
DATASETS = [('cse_cic_ids2018', 'CIC2018'), ('ton_iot', 'ToN-IoT'),
            ('bot_iot', 'BoT-IoT'), ('unsw_nb15', 'UNSW-NB15')]
LABELS = dict(benign='Benign', bot='Bot', brute_force='Brute force', ddos='DDoS', dos='DoS',
              infiltration='Infiltration', web_attacks='Web attack', backdoor='Backdoor',
              injection='Injection', mitm='MitM', password='Password', ransomware='Ransomware',
              scanning='Scanning', xss='XSS', reconnaissance='Recon', theft='Theft',
              analysis='Analysis', exploits='Exploits', fuzzers='Fuzzers', generic='Generic',
              shellcode='Shellcode', worms='Worms')
SHORT = dict(benign='Ben', bot='Bot', brute_force='Brute', ddos='DDoS', dos='DoS',
             infiltration='Inf', web_attacks='Web', backdoor='Back', injection='Inj',
             mitm='MitM', password='Pass', ransomware='Ran', scanning='Scan', xss='XSS',
             reconnaissance='Recon', theft='Theft', analysis='Anal', exploits='Expl',
             fuzzers='Fuzz', generic='Gen', shellcode='Shell', worms='Worm')

# ── phrase dictionary (exact, stripped match on a run group) ────────────────────────────
T = {
    # chrome
    '초안 · TBD': 'Draft · TBD', '데이터 검토': 'Data review', '모델 소개': 'Model',
    'SOTA 비교': 'SOTA comparison', 'Expert 평가': 'Expert evaluation',
    '구성요소 비교': 'Component ablation', '최종 모델 비교': 'Final comparison', '결론': 'Conclusion',
    '해석': 'Takeaway', '발표 결론': 'Conclusion',
    # slide 1
    '동일 feature를 공유하는 클래스': 'Classes sharing identical feature vectors',
    '행: 원래 라벨     열: 동일 feature의 다른 라벨     셀: 행 클래스 전체 대비 비율 (%)':
        'Row: original label     Column: other label on the same feature vector     Cell: share of all rows in the row class (%)',
    '빈칸: 0%   ·: 1% 미만   회색: 동일 클래스   같은 행의 중복 집계로 합산 불가':
        'Blank: 0%   ·: below 1%   Grey: same class   A row can count in several cells, so cells do not add up',
    'Web attack: 1,764 + 432 − 중복 429 = 1,767행 제거 · 771행 잔존':
        'Web attack: 1,764 + 432 − 429 overlap = 1,767 rows removed · 771 remain',
    # slide 2
    '정제 이후 클래스 분포': 'Class distribution after cleaning',
    '전체 데이터 기준 (train + validation + test) · 행 수와 데이터셋 내 클래스 비중 (%)':
        'All data (train + validation + test) · row counts and within-dataset class share (%)',
    '클래스': 'Class', '정제 전': 'Before', '정제 후': 'After', '비중 전 → 후': 'Share before → after',
    '합계': 'Total',
    'Infiltration 75.1% 감소': 'Infiltration down 75.1%',
    'Web attack 69.6% 감소 · 정제 후 771행': 'Web attack down 69.6% · 771 rows after cleaning',
    '정제 후에도 남는 클래스 불균형 · CIC2018 Web attack 771행, ToN Scanning 39,069행':
        'Imbalance remains after cleaning · CIC2018 Web attack 771 rows, ToN Scanning 39,069 rows',
    # slide 3
    'Global과 residual expert 구조': 'Global model and residual expert structure',
    '학습 데이터의 역할 분리 · Global 구성 / Expert 구성 / S·V 학습':
        'Separate roles of the training data · Global construction / expert construction / S·V training',
    'Global 풀 · 공통 context': 'Global pool · shared context',
    'Expert 풀 · residual과 expert': 'Expert pool · residuals and experts',
    'Route 풀 · S/V 학습': 'Route pool · S/V training',
    '오프라인 · 모델 구성': 'Offline · model construction',
    'Global 구성': 'Build Global', '공통 context로': 'Base classifier from', '기본 분류기 구성': 'the shared context',
    'Residual 계산': 'Compute residuals', '입력 표현·예측 확률': 'Embedding · probabilities',
    '정답과의 차이·오류 크기': 'gap to truth · error size',
    'Residual 군집': 'Cluster residuals', '실패 특성에 따라': 'Split into K clusters', 'K개 군집으로 분할': 'by failure pattern',
    'Expert 구성': 'Build experts', '공통 anchor': 'Shared anchor +', '+ 군집별 특화 block': 'cluster-specific block',
    'S/V 학습': 'Train S/V', 'Global·expert의': 'Learn from Global/expert', '교정·훼손 사례 학습': 'corrections and damage',
    '온라인 · 입력별 예측': 'Online · per-input prediction',
    'Global 예측': 'Global prediction', '기본 클래스 예측': 'Base class prediction',
    'Expert 선택': 'Selects an expert,', '호출 여부 판단': 'decides whether to call',
    'Expert 추론': 'Expert inference', '선택한 context로 분류': 'Classify with the selected context',
    'Expert 예측의': 'Decides whether to', '채택 여부 판단': 'accept the expert',
    '최종 예측': 'Final prediction', '채택 시 expert 예측': 'Expert prediction if accepted',
    '호출 또는 채택 조건 미충족 → Global 예측 유지': 'Call or acceptance condition not met → keep the Global prediction',
    'Backbone 고정 · context로 expert 구성 · S/V로 예측 교체 여부 결정':
        'Backbone frozen · experts built from contexts · S/V decide whether to replace the prediction',
    # slide 4
    'Global의 클래스별 성능과 expert의 병목': 'Global per-class performance and the expert bottleneck',
    '정제 데이터 · seed 43 · 기존 residual context의 진단 · S/V 적용 전':
        'Cleaned data · seed 43 · diagnosis of the existing residual contexts · before S/V',
    '기존 e3 · Infiltration': 'Existing e3 · Infiltration', '기존 e3 · Injection': 'Existing e3 · Injection',
    '담당영역 F1  0.440 → 0.967': 'Own residual region F1  0.440 → 0.967',
    '담당영역 F1  0.159 → 0.939': 'Own residual region F1  0.159 → 0.939',
    '입력유사군 Precision 12.4% · FP 7,935 → 21,535': 'Input-nearest group precision 12.4% · FP 7,935 → 21,535',
    '입력유사군 Precision 8.7% · FP 581 → 14,222': 'Input-nearest group precision 8.7% · FP 581 → 14,222',
    '담당영역: 정답 residual 배정 · 입력유사군: 입력 표현과 Global 확률로 배정 · 화살표: Global → Expert':
        'Own region: assigned by true residual · Input-nearest group: assigned by input embedding and Global probabilities · Arrow: Global → Expert',
    '담당 오류를 고치는 능력과 유사한 다른 클래스를 구분하는 능력을 함께 개선':
        'Improve both: correcting own-region errors and separating similar rows of other classes',
    # slide 5
    '정제 데이터의 SOTA 분류 성능': 'SOTA classification on cleaned data',
    'Seed 43 · 동일 C0 100,000행 · CIC2018 test 304만 행 / ToN test 227만 행':
        'Seed 43 · same C0 100,000 rows · CIC2018 test 3.04M rows / ToN test 2.27M rows',
    '방법': 'Method', '학습 풀 행 수 · CIC / ToN': 'Training pool rows · CIC / ToN',
    'XGBoost · 전체 train': 'XGBoost · full train', '866만 / 467만': '8.67M / 4.67M',
    '전체 train XGB: 추가 학습량을 허용한 참고 기준 · Global: expert·S/V 적용 전 기본 분류기':
        'Full-train XGB: reference with extra training data · Global: base classifier before experts and S/V',
    'CIC2018 · 클래스별 차이 · 공통 100k': 'CIC2018 · per-class differences · shared 100k',
    'ToN · 클래스별 차이 · 공통 100k': 'ToN · per-class differences · shared 100k',
    'Infiltration 최고 F1 0.3055 · DistPFN': 'Best Infiltration F1 0.3055 · DistPFN',
    'Web attack 최고 F1 0.3321 · XGB': 'Best Web attack F1 0.3321 · XGB',
    'MitM 최고 F1 0.2363 · XGB': 'Best MitM F1 0.2363 · XGB',
    'Scanning 최고 F1 0.3281 · BoostPFN': 'Best Scanning F1 0.3281 · BoostPFN',
    '동일 10만 행 기준 최고 Macro-F1 · CIC2018 DistPFN 0.7839 · ToN XGB 0.6983':
        'Best Macro-F1 on the same 100k rows · CIC2018 DistPFN 0.7839 · ToN XGB 0.6983',
    # slide 6
    'SOTA 학습·추론 비용': 'SOTA training and inference cost',
    'RTX 4090 24GB · 최대 2개 작업 병렬 · 자원 경합을 포함한 경과 시간':
        'RTX 4090 24GB · up to 2 jobs in parallel · elapsed time including resource contention',
    'CIC 학습 (초)': 'CIC fit (s)', 'CIC 추론 (초)': 'CIC predict (s)',
    'ToN 학습 (초)': 'ToN fit (s)', 'ToN 추론 (초)': 'ToN predict (s)',
    'LoCalPFN 학습에 validation 포함 · 전체 test 추론에 kNN 검색 포함':
        'LoCalPFN fit includes validation · full-test inference includes kNN search',
    'DistPFN은 Global fit·추론 공유 · prior 보정 추가 시간 약 0.19초':
        'DistPFN shares the Global fit and inference · prior correction adds about 0.19 s',
    'GPU: 1초 간격 프로세스별 최대값 · 단독 실행 속도 비교와 구분':
        'GPU: per-process peak sampled every 1 s · not a standalone speed comparison',
    'LoCalPFN 전체 test 추론 · CIC2018 4.81시간 · ToN 3.10시간':
        'LoCalPFN full-test inference · CIC2018 4.81 h · ToN 3.10 h',
    # slide 7
    '담당 영역과 입력 유사군의 분류 능력': 'Expert ability: own region vs input-nearest group',
    '관점 A · 담당 residual 영역': 'View A · own residual region',
    '관점 B · 입력 유사군': 'View B · input-nearest group',
    '대표 expert · 클래스: TBD': 'Representative expert · class: TBD',
    '지표': 'Metric', '교정: TBD': 'Corrections: TBD', '훼손: TBD': 'Damage: TBD',
    # slide 8
    'Scorer와 verifier의 기여': 'Contribution of scorer and verifier',
    '구성': 'Configuration', 'S·V 없음': 'No S·V', 'S만': 'S only', 'V만': 'V only',
    '호출 · 채택': 'Calls · accepted', '교정 · 훼손': 'Corrections · damage',
    'Context · gate 설정': 'Context · gate settings',
    # slide 9
    '제안 모델의 최종 성능과 비용': 'Final performance and cost of the proposed model',
    '검증된 expert·S/V 구성의 최종 비교': 'Final comparison of the validated expert and S/V configuration',
    '학습 정보': 'Training info', '호출 비용': 'Call cost',
    'Expert · S/V 없음': 'Expert · no S/V', 'Expert · S만': 'Expert · S only', 'Expert · V만': 'Expert · V only',
    '비교 기준: 5장의 SOTA 성능 · 6장의 측정 비용': 'Baselines: SOTA performance (slide 5) · measured cost (slide 6)',
    'Expert·route의 추가 학습 행 수와 조건부 추론 비용을 포함한 비교':
        'Includes the extra training rows for experts and routing, and the conditional inference cost',
    # slide 10
    '확인된 효과와 다음 개선': 'Confirmed effects and next improvements',
    'Context 구성': 'Context composition', 'Expert 구분 능력': 'Expert discrimination',
    '최종 모델 성능': 'Final model performance', '다음 개선': 'Next improvements',
}

NOTES = {
    1: ('Share of rows whose exact feature vector also appears under another label; not a model confusion matrix\n'
        'Denominator: all rows of the row class, over the full uncleaned data\n'
        'Even when A and B share vectors with each other, the cell shares are asymmetric because class sizes differ\n'
        'A row shared with several other labels counts in several cells, so row sums are not removal rates\n'
        'Web attack: 1,335 rows shared with Benign only, 3 with Infiltration only, 429 with both, 771 unshared\n'
        'Data and figure sources: docs/slides/kor/1002_intro_assets/ (English figures: docs/slides/eng/1002_intro_assets/)\n'
        'Column abbreviations follow the full class names in row order'),
    2: ('Per-class rows before/after cleaning are the train/val/test totals of the current fixed split class_counts.csv\n'
        'CIC2018 total 20,115,529 → 14,755,417; ToN total 27,520,260 → 10,807,288\n'
        'Share denominators are the total rows before and after cleaning, respectively\n'
        'CIC2018 Infiltration 188,152 → 46,761; Web attack 2,538 → 771\n'
        'Removal criteria omitted at the user\'s request'),
    3: ('Offline/online two-stage flow from the user-edited 0904 deck, restructured to match the current implementation\n'
        'Residuals are computed on the expert pool; contexts are not restricted to Global misclassifications\n'
        'Residuals use the input embedding, Global probabilities, the gap between truth and probabilities, and the error size\n'
        'A shared anchor supplies class examples; each expert adds a cluster-specific block\n'
        'The scorer selects an expert and decides whether to call it; the verifier decides whether to accept its prediction\n'
        'The current context comparisons do not yet apply the S/V of the new bank\n'
        'This slide introduces the overall design and claims neither active S/V nor a performance gain'),
    4: ('Global per-class F1 is the full-test result of the seed 43 Global confirmed in EXP59 and EXP61\n'
        'The existing e3 example uses the EXP59 residual context\n'
        'Own residual region: label-assisted diagnosis; input-nearest group: fixed assignment from input embedding and Global probabilities\n'
        'The two scopes contain different rows, so F1 is not compared across scopes\n'
        'CIC2018 e3 Infiltration: precision 12.4% in the input-nearest group despite high recall; ToN e3 Injection: precision 8.7%\n'
        'Direction: context composition that keeps the correction targets while reducing FP on other classes'),
    5: ('EXP61 completed results, updated 2026-09-30\n'
        'Conflict-free fixed split and exactly the same 100,000 training IDs as the Global C0\n'
        'Test: CIC2018 3,042,473 rows, ToN 2,271,723 rows, full-test evaluation\n'
        'Full-train XGB: CIC2018 8,666,430 rows, ToN 4,669,119 rows, a separate training budget\n'
        'Global/DistPFN backbone v3, BoostPFN/LoCalPFN backbone v1\n'
        'LoCalPFN selects the fine-tuned checkpoint by validation AUC; validation CIC 30,194 / ToN 41,744 rows\n'
        'BoostPFN T=50, context 500; LoCalPFN context 1000, FT 21 epochs × 30 steps\n'
        'Same training IDs; backbones and validation usage are not identical\n'
        'Per-class best values are among the shared-100k methods, excluding full-train XGB\n'
        'The full proposed expert/S/V system is not included in this table\n'
        'Single seed; the current test is the already observed development holdout\n'
        'docs/research/20260930/exp61_report_data.json'),
    6: ('Fit, full-test prediction and per-process GPU usage measured in the same EXP61 run\n'
        'Two workers share the GPU, so the numbers are not standalone speed ratios between models\n'
        'LoCalPFN validation is included in fit: CIC 3,129.97 s / ToN 5,694.41 s\n'
        'DistPFN fit/predict share the Global computation and are not counted twice\n'
        'Global uses the separately saved backbone predictions; DistPFN adds only about 0.19 s of prior post-processing\n'
        'Inference throughput, per-stage RAM and the input-loading peak are in the HTML detail tables and cost.json\n'
        'Model RAM covers preprocessing, model loading, fit, validation and predict, excluding raw PKL loading\n'
        'Actual single-GPU wall time is about 6.65 h, distinct from the sum of parallel job times\n'
        'docs/research/20260930/exp61_report_data.json'),
    7: 'Body, numbers, figures and takeaway TBD\nClassification ability in own region and input-nearest group',
    8: 'Body, numbers, figures and takeaway TBD\nContribution of scorer and verifier',
    9: ('New S/V ablation and final proposed-model results on the current seed 43, K=4 bank: TBD\n'
        'EXP61 baseline performance is on slide 5 and cost on slide 6\n'
        'Extra training labels for experts and routing are reported separately from the Global C0 100k budget'),
    10: 'Body, numbers, figures and takeaway TBD\nConfirmed effects and next improvements',
}

# Font-size overrides for English text that would otherwise wrap (slide, Korean key) → pt
RESIZE = {
    (3, '공통 context로'): 10.6, (3, '기본 분류기 구성'): 10.6,
    (3, '입력 표현·예측 확률'): 10.6, (3, '정답과의 차이·오류 크기'): 10.6,
    (3, '실패 특성에 따라'): 10.6, (3, 'K개 군집으로 분할'): 10.6,
    (3, '공통 anchor'): 10.6, (3, '+ 군집별 특화 block'): 10.6,
    (3, 'Global·expert의'): 10.6, (3, '교정·훼손 사례 학습'): 10.6,
    (3, '기본 클래스 예측'): 10.6, (3, 'Expert 선택'): 10.6, (3, '호출 여부 판단'): 10.6,
    (3, '선택한 context로 분류'): 10.6, (3, 'Expert 예측의'): 10.6, (3, '채택 여부 판단'): 10.6,
    (3, '채택 시 expert 예측'): 10.6,
    (4, '담당영역: 정답 residual 배정 · 입력유사군: 입력 표현과 Global 확률로 배정 · 화살표: Global → Expert'): 10.0,
    (1, '행: 원래 라벨     열: 동일 feature의 다른 라벨     셀: 행 클래스 전체 대비 비율 (%)'): 12.0,
}

hang = re.compile('[가-힣]')
missing, replaced = set(), 0


def fmt(r):
    f = r.font
    return (f.size.pt if f.size else None, f.bold,
            str(f.color.rgb) if (f.color is not None and f.color.type is not None) else None)


def translate_paragraph(p, slide_no):
    global replaced
    groups = []
    for r in p.runs:
        k = fmt(r)
        if groups and groups[-1][0] == k:
            groups[-1][1].append(r)
        else:
            groups.append([k, [r]])
    for _, runs in groups:
        text = ''.join(r.text for r in runs)
        if not hang.search(text):
            continue
        key = text.strip()
        if key not in T:
            missing.add(text)
            continue
        lead = text[:len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()):]
        runs[0].text = lead + T[key] + trail
        rp = runs[0]._r.get_or_add_rPr()
        rp.set('lang', 'en-US')
        if (slide_no, key) in RESIZE:
            runs[0].font.size = Pt(RESIZE[(slide_no, key)])
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        replaced += 1


def text_frames(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            yield sh.text_frame
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield c.text_frame


def redraw_matrices():
    """Same figure as intro_1002.prepare_data, drawn from the saved Korean assets with an English annotation."""
    EN_ASSETS.mkdir(parents=True, exist_ok=True)
    for path in (Path.home() / '.local/share/fonts/Pretendard').glob('*.otf'):
        font_manager.fontManager.addfont(path)
    plt.rcParams.update({'font.family': 'Pretendard', 'svg.fonttype': 'none', 'axes.unicode_minus': False})
    prov = json.loads((KO_ASSETS / 'provenance.json').read_text())
    cmap = LinearSegmentedColormap.from_list('KENTECH', ['#F2F7FB', '#80B6DA', '#1875B4', '#00306C'])
    out = {}
    for ds, title in DATASETS:
        m = pd.read_csv(KO_ASSETS / f'{ds}_pair_percent.csv', index_col=0)
        classes = list(m.index)
        assert classes == prov[ds]['class_order'] and list(m.columns) == classes
        matrix = m.values
        fig = plt.figure(figsize=(5.90, 2.25), facecolor='white')
        ax = fig.add_axes([0.215, 0.035, 0.765, 0.71])
        ax.imshow(matrix, cmap=cmap, vmin=0, vmax=100, aspect='auto', interpolation='nearest')
        ax.set_xticks(range(len(classes)), [SHORT[c] for c in classes], fontsize=9.3)
        ax.set_yticks(range(len(classes)), [LABELS[c] for c in classes], fontsize=9.5)
        ax.tick_params(axis='both', length=0, pad=4)
        ax.xaxis.tick_top()
        for spine in ax.spines.values():
            spine.set_visible(False)
        for i in range(len(classes)):
            for j in range(len(classes)):
                if i == j:
                    ax.add_patch(Rectangle((j - .5, i - .5), 1, 1, facecolor='#E9EDF1', edgecolor='white', linewidth=.6))
                    continue
                v = matrix[i, j]
                if v:
                    label = f'{v:.1f}' if v >= 1 else '·'
                    ax.text(j, i, label, ha='center', va='center', fontsize=8.5 if len(classes) >= 10 else 9.3,
                            color='white' if v >= 55 else '#163F60', weight='bold' if v >= 50 else 'normal')
        ax.set_xticks(np.arange(-.5, len(classes), 1), minor=True)
        ax.set_yticks(np.arange(-.5, len(classes), 1), minor=True)
        ax.grid(which='minor', color='white', linewidth=.6)
        ax.tick_params(which='minor', bottom=False, left=False, top=False)
        fig.text(.015, .945, title, color='#00306C', fontsize=14.5, weight='bold', va='center')
        fig.text(.985, .945, f"Conflicting rows {prov[ds]['conflict_fraction'] * 100:.2f}%",
                 color='#5D6773', fontsize=10.3, va='center', ha='right')
        fig.savefig(EN_ASSETS / f'{ds}_matrix.png', dpi=260)
        fig.savefig(EN_ASSETS / f'{ds}_matrix.svg')
        plt.close(fig)
        out[ds] = EN_ASSETS / f'{ds}_matrix.png'
    (EN_ASSETS / 'README.md').write_text(
        'English redraw of kor/1002_intro_assets/*_matrix.png from the saved *_pair_percent.csv and provenance.json; '
        'numbers unchanged, only the annotation text differs.\n')
    return out


def replace_picture(slide, pic, path):
    sp = pic._element
    parent = sp.getparent()
    idx = parent.index(sp)
    rid = sp.blip_rId
    new = slide.shapes.add_picture(str(path), pic.left, pic.top, pic.width, pic.height)
    parent.remove(new._element)
    parent.insert(idx, new._element)
    parent.remove(sp)
    slide.part.drop_rel(rid)


def main():
    OUT.parent.mkdir(exist_ok=True)
    prs = Presentation(IN)
    assert len(prs.slides) == 10
    for i, s in enumerate(prs.slides, 1):
        for tf in text_frames(s):
            for p in tf.paragraphs:
                translate_paragraph(p, i)
        s.notes_slide.notes_text_frame.text = NOTES[i]
    if missing:
        print('MISSING translations:')
        for m in sorted(missing):
            print('   ', repr(m))
        sys.exit(1)
    # slide 1 matrices: match the four pictures to datasets by position
    pngs = redraw_matrices()
    slots = {(0.62, 1.47): 'cse_cic_ids2018', (6.81, 1.47): 'ton_iot', (0.62, 3.89): 'bot_iot', (6.81, 3.89): 'unsw_nb15'}
    s1 = prs.slides[0]
    pics = [sh for sh in s1.shapes if sh.shape_type == 13]
    assert len(pics) == 4
    for pic in pics:
        key = min(slots, key=lambda k: abs(k[0] - pic.left / 914400) + abs(k[1] - pic.top / 914400))
        replace_picture(s1, pic, pngs[slots[key]])
    prs.core_properties.title = 'October 2 lab meeting · English draft'
    prs.core_properties.subject = 'Data review, context composition, expert capability, S/V, comparison with existing methods'
    prs.core_properties.comments = 'Slides 1–6 written, 7–10 TBD pending model experiments; translated from the Korean draft'
    prs.save(OUT)
    print('saved', OUT, 'replacements', replaced)

    # dump + residual-Korean check + fit estimate
    check = Presentation(OUT)
    dump, leftovers = [], []
    for i, s in enumerate(check.slides, 1):
        entries = []
        for sh in s.shapes:
            if sh.has_text_frame:
                entries.append(sh.text_frame.text)
            if sh.has_table:
                entries.extend(' | '.join(c.text for c in r.cells) for r in sh.table.rows)
        joined = '\n'.join(entries)
        if hang.search(joined) or hang.search(s.notes_slide.notes_text_frame.text):
            leftovers.append(i)
        dump.append(f'## {i}\n{joined}\n')
    QA.mkdir(exist_ok=True)
    (QA / 'slide_text.md').write_text('\n'.join(dump))
    assert not leftovers, f'Korean text left on slides {leftovers}'
    print('\n=== text fit estimate (need > box)')
    for si, sl in enumerate(check.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() and Emu(sh.height).inches > 0.2:
                need, box = need_h(sh), Emu(sh.height).inches
                if need > box * 0.99:
                    print(f'  slide {si}: need={need:.2f} box={box:.2f}  {sh.text_frame.text[:70]!r}')


def char_w(ch, sz):
    o = ord(ch)
    if ch == ' ':
        return 0.30 * sz
    if 0xAC00 <= o <= 0xD7A3 or 0x3000 <= o <= 0x9FFF:
        return 1.0 * sz
    if ch in '·•✓✕→↑↓×∪−':
        return 0.9 * sz
    if ch.isdigit():
        return 0.58 * sz
    if ch.isupper():
        return 0.68 * sz
    if ch.isalpha():
        return 0.54 * sz
    if ch in '.,:;\'"()[]|/':
        return 0.32 * sz
    return 0.55 * sz


def lines(s, width_in, sz):
    avail = max(width_in * 72, 1)
    n, cur = 1, 0.0
    for word in re.split(r'(\s+)', s):
        w = sum(char_w(c, sz) for c in word)
        if cur + w > avail and word.strip():
            n += 1
            cur = w
        else:
            cur += w
    return n


def need_h(sh):
    tf = sh.text_frame
    ml, mr = (tf.margin_left or 0) / 914400, (tf.margin_right or 0) / 914400
    mt, mb = (tf.margin_top or 0) / 914400, (tf.margin_bottom or 0) / 914400
    w = Emu(sh.width).inches - ml - mr
    h = 0.0
    for p in tf.paragraphs:
        txt = ''.join(r.text for r in p.runs)
        sz = max([r.font.size.pt for r in p.runs if r.font.size] or [11])
        h += (lines(txt, w, sz) if txt else 1) * sz * 1.25 / 72
    return h + mt + mb


if __name__ == '__main__':
    main()
