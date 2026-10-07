"""Data-backed slides 1–4 for the October 2 lab meeting."""
from pathlib import Path
import csv
import hashlib
import json

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.patches import Rectangle
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / 'docs/slides/kor/1002_intro_assets'
AUDIT = ROOT / 'docs/research/20260911/dataset_quality_audit'
DATASETS = [
    ('cse_cic_ids2018', 'CIC2018'), ('ton_iot', 'ToN-IoT'),
    ('bot_iot', 'BoT-IoT'), ('unsw_nb15', 'UNSW-NB15'),
]
LABELS = dict(benign='Benign', bot='Bot', brute_force='Brute force', ddos='DDoS', dos='DoS',
              infiltration='Infiltration', web_attacks='Web attack', backdoor='Backdoor',
              injection='Injection', mitm='MitM', password='Password', ransomware='Ransomware',
              scanning='Scanning', xss='XSS', reconnaissance='Recon', theft='Theft',
              analysis='Analysis', exploits='Exploits', fuzzers='Fuzzers', generic='Generic',
              shellcode='Shellcode', worms='Worms')
SHORT = dict(benign='Ben', bot='Bot', brute_force='Brute', ddos='DDoS', dos='DoS',
             infiltration='Inf', web_attacks='Web', backdoor='Back', injection='Inj',
             mitm='MitM', password='Pass', ransomware='Ran', scanning='Scan', xss='XSS',
             reconnaissance='Recon', theft='Theft', analysis='Anal', exploits='Expl',
             fuzzers='Fuzz', generic='Gen', shellcode='Shell', worms='Worm')


def prepare_data():
    ASSETS.mkdir(exist_ok=True)
    for path in (Path.home()/'.local/share/fonts/Pretendard').glob('*.otf'):
        font_manager.fontManager.addfont(path)
    plt.rcParams.update({'font.family':'Pretendard', 'svg.fonttype':'none', 'axes.unicode_minus':False})
    directed = json.loads((AUDIT/'pair_conflicts.json').read_text())
    undirected = json.loads((AUDIT/'pair_conflicts_unordered.json').read_text())
    summary = pd.read_csv(AUDIT/'audit_summary.csv').set_index('dataset')
    cmap = LinearSegmentedColormap.from_list('KENTECH', ['#F2F7FB','#80B6DA','#1875B4','#00306C'])
    sources, all_cells = {}, []
    for ds,title in DATASETS:
        data = directed[ds]['all']
        classes = list(data['class_rows'])
        matrix = np.zeros((len(classes),len(classes)),dtype=float)
        counts = np.zeros_like(matrix,dtype=np.int64)
        pairs = {(p['cls'],p['other']):p for p in data['pairs']}
        assert len(pairs)==2*len(undirected[ds]['pairs'])
        assert sum(data['class_rows'].values())==int(summary.loc[ds,'rows'])
        for p in undirected[ds]['pairs']:
            for a,b,key in [(p['a'],p['b'],'a'),(p['b'],p['a'],'b')]:
                assert pairs[a,b]['rows_of_cls_sharing']==p[f'rows_{key}']
        for (a,b),p in pairs.items():
            value=p['rows_of_cls_sharing']/data['class_rows'][a]
            assert abs(value-p['frac_of_cls'])<1e-12
            counts[classes.index(a),classes.index(b)]=p['rows_of_cls_sharing']
            matrix[classes.index(a),classes.index(b)]=100*value
        pd.DataFrame(matrix,index=classes,columns=classes).to_csv(ASSETS/f'{ds}_pair_percent.csv')
        for i,a in enumerate(classes):
            for j,b in enumerate(classes):
                if i!=j:
                    all_cells.append(dict(dataset=ds,row_class=a,column_class=b,
                        class_rows=data['class_rows'][a],shared_rows=int(counts[i,j]),percent=float(matrix[i,j])))
        fig=plt.figure(figsize=(5.90,2.25),facecolor='white')
        ax=fig.add_axes([0.215,0.035,0.765,0.71])
        ax.imshow(matrix,cmap=cmap,vmin=0,vmax=100,aspect='auto',interpolation='nearest')
        ax.set_xticks(range(len(classes)),[SHORT[c] for c in classes],fontsize=9.3)
        ax.set_yticks(range(len(classes)),[LABELS[c] for c in classes],fontsize=9.5)
        ax.tick_params(axis='both',length=0,pad=4)
        ax.xaxis.tick_top()
        for spine in ax.spines.values():spine.set_visible(False)
        for i in range(len(classes)):
            for j in range(len(classes)):
                if i==j:
                    ax.add_patch(Rectangle((j-.5,i-.5),1,1,facecolor='#E9EDF1',edgecolor='white',linewidth=.6))
                    continue
                v=matrix[i,j]
                if v:
                    label=f'{v:.1f}' if v>=1 else '·'
                    ax.text(j,i,label,ha='center',va='center',fontsize=8.5 if len(classes)>=10 else 9.3,
                            color='white' if v>=55 else '#163F60',weight='bold' if v>=50 else 'normal')
        ax.set_xticks(np.arange(-.5,len(classes),1),minor=True)
        ax.set_yticks(np.arange(-.5,len(classes),1),minor=True)
        ax.grid(which='minor',color='white',linewidth=.6)
        ax.tick_params(which='minor',bottom=False,left=False,top=False)
        fig.text(.015,.945,title,color='#00306C',fontsize=14.5,weight='bold',va='center')
        fig.text(.985,.945,f"전체 상충 행 {summary.loc[ds,'frac_mixed']*100:.2f}%",
                 color='#5D6773',fontsize=10.3,va='center',ha='right')
        fig.savefig(ASSETS/f'{ds}_matrix.png',dpi=260)
        fig.savefig(ASSETS/f'{ds}_matrix.svg')
        plt.close(fig)
        sources[ds]=dict(class_order=classes,rows=int(summary.loc[ds,'rows']),
                         conflict_rows=int(summary.loc[ds,'rows_in_mixed_label_groups']),
                         conflict_fraction=float(summary.loc[ds,'frac_mixed']))
    pd.DataFrame(all_cells).to_csv(ASSETS/'all_pair_cells.csv',index=False)
    tables={}
    for ds,dirname in [
        ('cic2018','cic2018_conflict_free_fixed_split_20260915_180000'),
        ('toniot','toniot_conflict_free_fixed_split_20260916_021730')]:
        base=ROOT/'data/derived'/dirname
        manifest=json.loads((base/'manifest.json').read_text())
        counts=pd.read_csv(base/'class_counts.csv').groupby('class',sort=False)[['before','after']].sum()
        assert counts.before.sum()==manifest['rows_before'] and counts.after.sum()==manifest['rows_after']
        key='cse_cic_ids2018' if ds=='cic2018' else 'ton_iot'
        assert counts.before.to_dict()==directed[key]['all']['class_rows']
        counts['before_share']=100*counts.before/counts.before.sum()
        counts['after_share']=100*counts.after/counts.after.sum()
        counts.to_csv(ASSETS/f'{ds}_class_distribution.csv')
        tables[ds]=counts
    inputs=[AUDIT/'pair_conflicts.json',AUDIT/'pair_conflicts_unordered.json',AUDIT/'audit_summary.csv']
    sources['definition']='Cell(A,B): number of A rows with an exact feature vector also labeled B / all A rows. All original rows, before cleaning. Pair counts may overlap.'
    sources['input_sha256']={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}
    (ASSETS/'provenance.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2)+'\n')
    return tables,cmap


def footer(h,s,value):
    h.text(s,.62,7.035,8.20,.30,value,size=11,bold=True,ink='FFFFFF',inset=0)


def line(h,s,x1,y1,x2,y2,ink=None,head=False,width=1.6):
    shape=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    shape.line.color.rgb=h.color(ink or h.BLUE)
    shape.line.width=Pt(width)
    if head:
        end=etree.SubElement(shape.line._get_or_add_ln(),qn('a:tailEnd'))
        end.set('type','triangle');end.set('w','sm');end.set('len','sm')
    return shape


def node(h,s,x,y,w,title,sub,fill=None):
    h.box(s,x,y,w,1.06,fill=fill or h.CARD,border=h.LINE)
    h.text(s,x+.06,y+.10,w-.12,.32,title,size=15,bold=True,ink=h.NAVY,align=PP_ALIGN.CENTER,inset=.02)
    h.text(s,x+.08,y+.49,w-.16,.40,sub,size=11.3,ink=h.GREY,align=PP_ALIGN.CENTER,inset=.02)


def build(h):
    tables,cmap=prepare_data()
    s=h.slide('동일 feature를 공유하는 클래스', '데이터 검토')
    h.text(s,.62,1.01,12.09,.33,'행: 원래 라벨     열: 동일 feature의 다른 라벨     셀: 행 클래스 전체 대비 비율 (%)',
           size=13,ink=h.GREY,inset=0)
    for (ds,title),(x,y) in zip(DATASETS,[(.62,1.47),(6.81,1.47),(.62,3.89),(6.81,3.89)]):
        s.shapes.add_picture(str(ASSETS/f'{ds}_matrix.png'),Inches(x),Inches(y),width=Inches(5.90),height=Inches(2.25))
    h.text(s,.62,6.34,7.15,.31,'빈칸: 0%   ·: 1% 미만   회색: 동일 클래스   같은 행의 중복 집계로 합산 불가',size=10.3,ink=h.GREY,inset=0)
    for i,v in enumerate([0,25,50,75,100]):
        rgba=cmap(v/100);hexval=''.join(f'{round(255*c):02X}' for c in rgba[:3])
        h.box(s,8.42+i*.66,6.39,.45,.14,fill=hexval,rounded=False)
        h.text(s,8.39+i*.66,6.57,.51,.19,str(v),size=8.3,ink=h.GREY,align=PP_ALIGN.CENTER,inset=0)
    h.text(s,11.81,6.36,.90,.27,'%',size=10,ink=h.GREY,inset=0)
    footer(h,s,'Web attack: 1,764 + 432 − 중복 429 = 1,767행 제거 · 771행 잔존')
    s.notes_slide.notes_text_frame.text=(
        '모델의 예측 혼동행렬이 아니라 동일 feature에서 복수 라벨이 관찰되는 비율\n'
        '분모는 원래 행 클래스의 모든 표본이며 정제 전 전체 데이터 기준\n'
        'A와 B가 서로 feature를 공유하더라도 클래스 크기가 달라 셀 비율은 비대칭\n'
        '한 행이 여러 다른 라벨과 공유되면 여러 셀에 집계되므로 행 합은 제거 비율이 아님\n'
        'Web attack: Benign만 1,335행, Infiltration만 3행, 양쪽 모두 429행, 겹침 없는 771행\n'
        '데이터 및 그림 원본: docs/slides/kor/1002_intro_assets/\n'
        '열 약어는 행의 전체 클래스명과 같은 순서')

    s=h.slide('정제 이후 클래스 분포', '데이터 검토')
    h.text(s,.62,1.01,12.09,.33,'전체 데이터 기준 (train + validation + test) · 행 수와 데이터셋 내 클래스 비중 (%)',
           size=13,ink=h.GREY,inset=0)
    for ds,title,x in [('cic2018','CIC2018',.62),('toniot','ToN-IoT',6.81)]:
        d=tables[ds]
        h.text(s,x,1.51,5.90,.39,title,size=18,bold=True,ink=h.NAVY,inset=0)
        def pct(v):return f'{v:.3f}' if v<.1 else f'{v:.2f}'
        rows=[[LABELS[c],f'{int(v.before):,}',f'{int(v.after):,}',f'{pct(v.before_share)} → {pct(v.after_share)}']
              for c,v in d.iterrows()]
        rows.append(['합계',f'{int(d.before.sum()):,}',f'{int(d.after.sum()):,}','100 → 100'])
        shape=h.table(s,x,2.03,[1.43,1.38,1.38,1.71],['클래스','정제 전','정제 후','비중 전 → 후'],rows,row_h=.354,font=11.0)
        for r in range(1,len(rows)+1):
            is_total=r==len(rows)
            is_tail=rows[r-1][0] in (['Infiltration','Web attack'] if ds=='cic2018' else ['Scanning'])
            for c in range(4):
                cell=shape.table.cell(r,c)
                if is_total or is_tail:
                    cell.fill.fore_color.rgb=h.color(h.HEAD if is_total else 'EEF5FB')
                for p in cell.text_frame.paragraphs:
                    if c>0:p.alignment=PP_ALIGN.RIGHT
                    for run in p.runs:
                        if is_total or is_tail:run.font.bold=True
                cell.margin_left=cell.margin_right=Inches(.08)
                cell.margin_top=cell.margin_bottom=Inches(.025)
        if ds=='cic2018':
            h.box(s,x,5.55,5.90,.72,fill=h.CARD)
            h.text(s,x+.17,5.61,5.56,.55,'Infiltration 75.1% 감소\nWeb attack 69.6% 감소 · 정제 후 771행',
                   size=13,ink=h.NAVY,inset=0)
    footer(h,s,'정제 후에도 남는 클래스 불균형 · CIC2018 Web attack 771행, ToN Scanning 39,069행')
    s.notes_slide.notes_text_frame.text=(
        '정제 전후 클래스별 행 수는 현재 fixed-split의 class_counts.csv를 train/val/test 전체 합산\n'
        'CIC2018 합계 20,115,529 → 14,755,417; ToN 합계 27,520,260 → 10,807,288\n'
        '구성비 분모는 각각 정제 전 전체 행과 정제 후 전체 행\n'
        'CIC2018 Infiltration 188,152 → 46,761; Web attack 2,538 → 771\n'
        '사용자 요청에 따라 제거 기준 설명 제외')

    s=h.slide('Global과 residual expert 구조', '모델 소개')
    h.text(s,.62,1.02,12.09,.31,'학습 데이터의 역할 분리 · Global 구성 / Expert 구성 / S·V 학습',size=13,ink=h.GREY,inset=0)
    for x,w,label in [(.62,3.87,'Global 풀 · 공통 context'),(4.73,3.87,'Expert 풀 · residual과 expert'),(8.84,3.87,'Route 풀 · S/V 학습')]:
        sh=h.box(s,x,1.50,w,.46,fill=h.HEAD)
        h.set_text(sh.text_frame,label,size=12,bold=True,ink=h.NAVY,align=PP_ALIGN.CENTER,inset=.03)
    h.text(s,.62,2.18,12.09,.28,'오프라인 · 모델 구성',size=13,bold=True,ink=h.GREY,inset=0)
    xs=[.62,3.09,5.56,8.03,10.50];w=2.21;y=2.69
    offline=[('Global 구성','공통 context로\n기본 분류기 구성'),
             ('Residual 계산','입력 표현·예측 확률\n정답과의 차이·오류 크기'),
             ('Residual 군집','실패 특성에 따라\nK개 군집으로 분할'),
             ('Expert 구성','공통 anchor\n+ 군집별 특화 block'),
             ('S/V 학습','Global·expert의\n교정·훼손 사례 학습')]
    for x,(title,sub) in zip(xs,offline):node(h,s,x,y,w,title,sub)
    for a,b in zip(xs,xs[1:]):line(h,s,a+w,y+.53,b,y+.53,head=True)
    h.text(s,.62,4.03,12.09,.28,'온라인 · 입력별 예측',size=13,bold=True,ink=h.GREY,inset=0)
    y=4.53
    online=[('Global 예측','기본 클래스 예측'),('Scorer','Expert 선택\n호출 여부 판단'),
            ('Expert 추론','선택한 context로 분류'),('Verifier','Expert 예측의\n채택 여부 판단'),
            ('최종 예측','채택 시 expert 예측')]
    for x,(title,sub) in zip(xs,online):node(h,s,x,y,w,title,sub,fill=h.HEAD if title in ['Scorer','Verifier'] else None)
    for a,b in zip(xs,xs[1:]):line(h,s,a+w,y+.53,b,y+.53,head=True)
    sx,vx,fx=xs[1]+w/2,xs[3]+w/2,xs[4]+w/2
    for x in [sx,vx]:line(h,s,x,5.59,x,6.10,ink=h.MUTED,head=False,width=1.3)
    line(h,s,sx,6.10,fx,6.10,ink=h.MUTED,width=1.3)
    line(h,s,fx,6.10,fx,5.59,ink=h.MUTED,head=True,width=1.3)
    h.text(s,.62,6.27,12.09,.31,'호출 또는 채택 조건 미충족 → Global 예측 유지',
           size=13,ink=h.GREY,align=PP_ALIGN.CENTER,inset=0)
    footer(h,s,'Backbone 고정 · context로 expert 구성 · S/V로 예측 교체 여부 결정')
    s.notes_slide.notes_text_frame.text=(
        '0904 사용자 편집본의 오프라인/온라인 두 단계 흐름도를 현재 구현에 맞춰 재구성\n'
        'Expert 풀에서 residual을 계산하되 Global 오분류만으로 context를 제한하지 않음\n'
        'Residual에는 입력 표현, Global 확률, 정답과 확률의 차이, 오류 크기 사용\n'
        '공통 anchor로 클래스 사례를 제공하고 군집별 특화 block으로 각 expert 구성\n'
        'Scorer는 expert 선택 및 사전 호출 판단; verifier는 expert 예측 수락 판단\n'
        '현재 context 비교 결과에는 새 bank의 S/V를 아직 적용하지 않음\n'
        '본 장은 전체 모델 설계 소개이며 S/V 활성화나 성능 개선을 주장하지 않음')
    build_global(h)


def build_global(h):
    source=ROOT/'tabpfn/results/20260929_exp61_sota_local_s43'
    capability=json.loads((ROOT/'tabpfn/results/20260928_exp59_residual_oracle_s43/expert_capability/capability.json').read_text())
    s=h.slide('Global의 클래스별 성능과 expert의 병목', '모델 소개')
    h.text(s,.62,1.01,12.09,.33,'정제 데이터 · seed 43 · 기존 residual context의 진단 · S/V 적용 전',size=13,ink=h.GREY,inset=0)
    for ds,title,x in [('cic2018','CIC2018',.62),('toniot','ToN-IoT',6.81)]:
        df=pd.read_csv(source/'class_metrics.csv').query("dataset==@ds and method=='global_raw'").rename(columns={'name':'class_name'})
        df=df.sort_values('f1',ascending=False)
        fig,ax=plt.subplots(figsize=(5.9,3.0));fig.subplots_adjust(left=.22,right=.92,top=.84,bottom=.15)
        pos=np.arange(len(df));bars=ax.barh(pos,df.f1,color=['#B36B56' if v<.5 else '#327CAD' for v in df.f1],height=.65)
        ax.set_yticks(pos,[LABELS[c] for c in df.class_name],fontsize=10)
        ax.invert_yaxis();ax.set_xlim(0,1.12)
        ax.set_xticks([0,.5,1],['0','0.5','1.0'],fontsize=9)
        ax.set_xlabel('Global class F1',fontsize=10,color='#5D6773')
        ax.grid(axis='x',alpha=.17);ax.set_axisbelow(True);ax.tick_params(axis='y',length=0)
        for spine in ax.spines.values():spine.set_visible(False)
        for i,v in enumerate(df.f1):ax.text(v+.017,i,f'{v:.3f}',va='center',fontsize=10,color='#203044')
        fig.text(.015,.955,title,fontsize=14.5,weight='bold',color='#00306C',va='center')
        fig.text(.985,.955,f'Macro-F1 {df.f1.mean():.4f}',fontsize=11,color='#5D6773',va='center',ha='right')
        for ext in ['png','svg']:fig.savefig(ASSETS/f'{ds}_global_f1.{ext}',dpi=260)
        plt.close(fig)
        s.shapes.add_picture(str(ASSETS/f'{ds}_global_f1.png'),Inches(x),Inches(1.51),width=Inches(5.90),height=Inches(3.0))
        e=next(e for e in capability['datasets'][ds]['experts'] if e['expert']==3)
        cl='infiltration' if ds=='cic2018' else 'injection'
        a=next(c for c in e['residual']['classes'] if c['name']==cl)
        b=next(c for c in e['observable']['classes'] if c['name']==cl)
        h.box(s,x,4.82,5.90,1.32,fill=h.CARD)
        h.text(s,x+.15,4.90,5.60,.28,f'기존 e3 · {LABELS[cl]}',size=14,bold=True,ink=h.NAVY,inset=0)
        h.text(s,x+.15,5.26,5.60,.29,
               f"담당영역 F1  {a['global_metrics']['f1']:.3f} → {a['expert_metrics']['f1']:.3f}",size=13,ink=h.INK,inset=0)
        h.text(s,x+.15,5.64,5.60,.29,
               f"입력유사군 Precision {b['expert_metrics']['precision']*100:.1f}% · FP {b['global_metrics']['FP']:,} → {b['expert_metrics']['FP']:,}",
               size=12.0,ink=h.INK,inset=0)
    h.text(s,.62,6.36,12.09,.32,'담당영역: 정답 residual 배정 · 입력유사군: 입력 표현과 Global 확률로 배정 · 화살표: Global → Expert',
           size=10.8,ink=h.GREY,inset=0)
    footer(h,s,'담당 오류를 고치는 능력과 유사한 다른 클래스를 구분하는 능력을 함께 개선')
    s.notes_slide.notes_text_frame.text=(
        'Global 클래스별 F1은 EXP61에서 확인한 seed 43 Global의 전체 test 결과\n'
        '기존 expert e3 사례는 EXP59의 기존 residual context\n'
        '담당 residual 영역은 정답 사용 진단, 입력 유사군은 입력·Global 확률 기반의 고정 배정\n'
        '두 범위의 표본은 다르므로 F1을 범위 사이에 직접 비교하지 않음\n'
        'CIC2018 e3 Infiltration은 입력유사군에서 높은 recall에도 precision 12.4%; ToN e3 Injection은 precision 8.7%\n'
        '해결 방향은 교정 대상을 유지하면서 상대 클래스 FP를 줄이는 context 구성')
