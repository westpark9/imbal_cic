#!/usr/bin/env python3
"""Revise the explanation and scenario; preserve the supplied result graphs.

One scene description produces editable PPTX shapes and browser-rendered PDF.
The original result-graph pixels and values are reused without reanalysis.
"""
from pathlib import Path
import hashlib
import html
import json
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = next(p for p in Path(__file__).resolve().parents if (p/'AGENTS.md').exists())
SLIDES = ROOT/'docs/slides'
QA = SLIDES/'_support/qa/20261020_meeting'
QA.mkdir(parents=True, exist_ok=True)
SOURCE = SLIDES/'eng/26S_Meeting.pdf'
ARCH = SLIDES/'_support/qa/kor/1015_labmeeting_qa/v3_architecture_page.png'
PAPER = 'https://arxiv.org/html/2605.13986v1'
W, H = 13.333333, 7.5
NAVY, BLUE, INK, GREY, PALE = '00306C', '007BC6', '20262E', '617080', 'F1F5F9'

def text(x,y,w,h,s,size=22,bold=False,color=INK,align="left"):
    h=max(h,len(s.splitlines())*size*1.14/72*1.06)
    return dict(kind='text',x=x,y=y,w=w,h=h,text=s,size=size,bold=bold,color=color,align=align)

def box(x,y,w,h,fill=PALE):
    return dict(kind='box',x=x,y=y,w=w,h=h,fill=fill)

def arrow(x,y,w=.35,h=.28):
    return dict(kind='arrow',x=x,y=y,w=w,h=h,fill=BLUE)

def arch():
    return dict(kind='image',x=.62,y=1.08,w=3.95,h=5.88,path=str(ARCH),crop=[.20,.082,.19,.285])

def make_scenes(lang, include_results=False):
    def t(ko,en): return ko if lang=='ko' else en
    scenes=[]
    def page(title,items,notes,message='',section=''):
        n=len(scenes)+1
        footer=[box(0,7.05,W,.45,fill=NAVY),text(.62,7.14,10.65,.23,message,14,False,'FFFFFF'),
                text(11.45,7.15,1.25,.22,f'{section}   {n}',12,False,'FFFFFF',align='right')]
        scenes.append(dict(title=title,notes=notes,items=[text(.60,.33,12.1,.57,title,30,True,NAVY)]+items+footer))

    items=[text(.62,1.10,12.0,.35,t('침입 탐지 시스템(IDS) · 통신 기록의 특징으로 정상·공격 유형 분류',
        'Intrusion detection system (IDS) · Classify network flows as benign or attack types'),21),
        text(.62,1.72,12.0,.62,t('처음 확보한 라벨에 없는 공격의 등장\n고정된 분류 점수만으로 파악하기 어려운 신규 사례 반영 능력',
        'An attack absent from the initial labelled data may appear later\nA fixed test score alone does not show how well a detector can adapt'),21,True,NAVY)]
    cards=[(t('운영 시작','At deployment'),t('정상 + 공격 A·B\n일부 유형의 라벨만 확보','Benign + attacks A and B\nLabels for a limited set of types')),
           (t('추가 조사 이후','After investigation'),t('새 공격 C의 라벨 소량 확보\n기존 탐지기에 반영할 필요','A few labelled cases of C\nNew labels for the detector')),
           (t('갱신 후 확인','After updating'),t('C 탐지 품질과 갱신 비용\n기존 정상·A·B 구분 유지','Detection of C and update cost\nQuality on earlier classes'))]
    for i,(title,body) in enumerate(cards):
        x=.62+i*4.22
        items += [box(x,2.84,3.65,1.58),text(x+.18,3.02,3.29,.35,title,21,True,BLUE),
                  text(x+.18,3.63,3.29,.55,body,17)]
        if i<2:items.append(arrow(x+3.81,3.48))
    items += [text(.65,4.90,1.85,.36,'XGBoost',22,True),
        text(2.65,4.90,9.90,.42,t('새 사례를 포함한 누적 라벨 데이터로 모델 재학습',
            'Refit the model on all available labels, including the new cases'),20),
        text(.65,5.57,1.85,.36,'TabPFN',22,True,BLUE),
        text(2.65,5.57,9.90,.42,t('사전학습 모델에 기존·신규 사례를 context로 제공',
            'Supply existing and new cases as context to the pretrained model'),20),
        text(.65,6.36,12,.32,t('이번 평가 범위 · 새 공격의 라벨 확보 이후 적응',
            'Evaluation scope · Adaptation after labels for a new attack become available'),17,color=GREY)]
    page(t('소수 신규 공격 사례에 대응하는 IDS','Adapting IDS with few newly labelled attacks'),items,
         'IDS is a classifier of flow records, not a name for the dataset. Begin with the operational gap: only some attack labels are available initially; investigation later supplies a few examples of another attack. '
         'The research question is how much newly labelled data and update computation each method needs, while tracking existing-class quality. '
         'Do not infer a universal need for retraining from a weekday collection schedule. Full cumulative XGB refitting is our chosen baseline. '
         'This experiment evaluates adaptation after labelling, not unknown-attack detection before labelling or generalization across networks.',
         t('비교 질문 · 소수 신규 사례에서 얻는 탐지 품질과 갱신 비용','Research question · Detection quality and update cost with few new labelled cases'), '01')

    items=[text(.62,1.10,12.0,.34,t('합성 표 과제로 예측 규칙 사전학습 · TabPFN-3, arXiv 2026',
        'Prediction rules pretrained on synthetic tabular tasks · TabPFN-3, arXiv 2026'),19,color=GREY),
        text(.62,1.67,12.0,.63,t('사전학습된 것은 사례에서 예측하는 방법\n현재 IDS의 정상·공격 클래스에 해당하는 사례는 추론 시 입력',
        'Pretraining learns how to predict from examples\nExamples defining the current IDS classes are supplied at prediction time'),22,True,NAVY),
        box(.62,2.62,3.23,1.00),text(.82,2.78,2.85,.30,'Context',20,True,BLUE),
        text(.82,3.22,2.85,.25,t('기존·신규 특징 + 라벨','Existing/new flows + labels'),16),
        box(.62,3.87,3.23,.87),text(.82,4.00,2.85,.30,'Query',20,True,BLUE),
        text(.82,4.40,2.85,.25,t('분류할 기록 · 정답 미입력','Unlabelled flow to classify'),16),
        arrow(4.07,3.04,.47,.32),arrow(4.07,4.15,.47,.32),
        box(4.80,3.10,3.55,1.46,fill=NAVY),text(5.00,3.35,3.15,.39,'TabPFN',26,True,'FFFFFF'),
        text(5.00,3.98,3.15,.29,t('사전학습 가중치 고정','Pretrained weights fixed'),18,color='FFFFFF'),
        arrow(8.57,3.66,.47,.32),box(9.30,3.10,3.38,1.46),
        text(9.50,3.40,2.97,.33,t('클래스별 확률','Class probabilities'),21,True,BLUE),
        text(9.50,4.00,2.97,.30,t('정상 · 공격 A · 공격 C …','Benign · Attack A · Attack C …'),16),
        text(.65,5.13,2.65,.35,'In-context learning',20,True,BLUE),
        text(3.45,5.13,9.05,.36,t('context와 query의 관계를 이용한 예측 · 가중치 갱신 없음',
        'Prediction conditioned on the supplied context; weights unchanged'),19),
        text(.65,5.84,2.40,.35,'Fine-tuning',20,True,BLUE),
        text(3.45,5.84,9.05,.66,t('IDS 사례를 context/query로 나누고 query 정답 손실을 역전파\n별도의 가중치 추가 학습 · 이번 비교에서는 미수행',
        'Split IDS examples into context/query and backpropagate query loss\nAdditional weight training; not performed in this comparison'),18)]
    page(t('TabPFN · 사전학습 가중치와 입력 사례의 역할','TabPFN · Pretrained weights and supplied examples'),items,
         PAPER+' Sections 1, 2.2, 2.5. Synthetic tasks concern pretraining, not our IDS evaluation rows. '
         'Context lives outside the weights. Updating a context does not write its rows into the model parameters. '
         'IDS fine-tuning would optimize the model using losses on labelled queries, but the displayed no-FT comparisons do not do this.',
         t('가중치를 고정해도 필요한 context · 현재 분류 문제를 알려 주는 라벨 사례',
           'Context supplies the current classification task even when pretrained weights are fixed'),'02')

    items=[arch()]
    rows=[('1',t('입력값 표현 · Cell embedding','Encode values · Cell embedding'),
              t('패킷 길이 등 각 값을 벡터로 변환\n결측 여부도 함께 입력','Encode each value, such as packet length\nInclude missing-value indicators')),
          ('2',t('열 정보 요약 · Feature embedding','Summarize columns · Feature embedding'),
              t('여러 기록에서 각 특징의 분포 파악\ncontext 라벨을 함께 참조','Summarize each feature across flows\nUse the labels in the context')),
          ('3',t('행 정보 통합 · Feature aggregation','Combine features · Feature aggregation'),
              t('한 기록의 특징들을 하나의 표현으로 통합\n사례 간 관계 비교에 사용할 행 표현 생성','Combine a flow’s features into a representation\nPrepare it for comparison with other flows'))]
    for i,(num,title,body) in enumerate(rows):
        y=1.40+i*1.74
        items += [text(4.98,y,7.73,.45,num+'  '+title,22,True,BLUE),text(5.33,y+.68,7.25,.87,body,20)]
    page(t('v3 구조 · 통신 기록의 표현','v3 architecture · Representing network flows'),items,
        PAPER+' Figure 5, Section 2.1. Original model diagram without a visible caption as requested. '
        'These are inference representations; they are not a sequence of IDS weight-training stages.',section='02')

    items=[arch(),text(4.98,1.36,7.73,.45,t('4  사례 참조 · In-context learning','4  Consult examples · In-context learning'),22,True,BLUE),
        text(5.33,2.03,7.23,.78,t('query와 context 사례의 관계 계산\nattention으로 참고할 사례의 비중 결정',
        'Compare the query with context examples\nUse attention to weight their relevance'),20),
        text(4.98,3.02,7.73,.45,t('5  확률 출력 · Many-class decoder','5  Predict probabilities · Many-class decoder'),22,True,BLUE),
        text(5.33,3.69,7.23,.78,t('context 라벨을 가중합해 확률 계산\n정상·기존 공격·신규 공격의 사례 모두 필요',
        'Weight and combine the context labels\nRetain benign, existing and new attack examples'),20),
        box(4.98,5.11,7.70,1.57),text(5.20,5.31,7.25,.35,t('가중치 학습 없이도 발생하는 갱신 비용','Update computation without weight training'),21,True,NAVY),
        text(5.20,5.96,7.23,.48,t('context 전처리 · 표현 계산 · 캐시 준비\nXGB 재학습보다 저렴한지는 실측으로 비교',
        'Context preprocessing · Representations · Cache\nMeasure the cost against XGB refitting'),18)]
    page(t('v3 구조 · 사례 참조와 갱신 비용','v3 architecture · Context and update computation'),items,
        PAPER+' Sections 2.2 and 2.4.2. The many-class decoder uses attention-weighted context labels. '
        'Updating the context may require rebuilding preprocessing and cache states. Frozen weights do not imply a free update.',
        t('가중치 추가 학습의 부재와 낮은 갱신 비용을 구별한 비교',
          'No additional weight training does not by itself establish low update cost'),'02')

    items=[text(.62,1.10,12.0,.67,t('한 행 = 실제 수집된 통신 기록의 특징 + 정상/공격 라벨\n각 벤치마크를 별도로 정제·분리한 뒤 같은 출처 안에서 사례 선택',
        'One row = features of a recorded network flow + its benign/attack label\nClean and split each benchmark separately; select examples within that benchmark'),20)]
    for j,name in enumerate(['CIC2018','ToN-IoT']):
        y=2.26+j*1.40
        items += [box(.62,y,2.07,1.06),text(.80,y+.33,1.73,.34,name,22,True,NAVY),
            arrow(2.90,y+.38,.40,.28),
            box(3.54,y,4.18,1.06),text(3.73,y+.15,3.80,.32,t('Train · 학습 후보 집합','Train · Candidate examples'),19,True,BLUE),
            text(3.73,y+.63,3.80,.29,t('초기 사례 → 신규 공격 사례 선택','Select initial cases, then new attacks'),16),
            box(8.09,y,4.59,1.06),text(8.28,y+.15,4.21,.32,t('Test · 고정 평가 집합','Test · Fixed evaluation flows'),19,True,BLUE),
            text(8.28,y+.63,4.21,.29,t(('3,042,473' if j==0 else '2,271,723')+'개 · 갱신에 미사용',
                ('3,042,473' if j==0 else '2,271,723')+' rows · Excluded from updates'),16)]
    items += [text(.65,5.12,12,.39,t('CIC 예시 · 초기 1,000개 + 같은 CIC train의 신규 DDoS 100개 = 제공 라벨 1,100개',
        'CIC example · Initial 1,000 + 100 DDoS rows from the same CIC train = 1,100 labels'),20,True,NAVY),
        text(.65,5.85,12,.66,t('이 1,100개로 XGB 재학습 / PFN context 구성\n“현재 데이터셋”의 의미 = 현재까지 제공한 라벨 사례 집합',
        'Use these same 1,100 rows for XGB refitting or as PFN context\n“Current dataset” means the labelled examples supplied so far'),19)]
    page(t('평가 데이터 구성 · 출처와 제공 사례의 구분','Evaluation data · From benchmark rows to labelled examples'),items,
        'Data trace: EXP77 data manifests -> EXP71 cumulative run data -> conflict-free fixed-split manifests. '
        'The stored suite may contain multiple benchmark keys, but the experiment selects one dataset independently; no CIC attacks are injected into ToN or vice versa. '
        'The 100-row example explains the old setting; the next page explicitly states it has no operational or optimality justification. '
        'Original row IDs are disjoint across train/test, but duplicate feature vectors can exist; this is not a claim of feature-disjoint holdout. '
        'Stages control which labels become available and how many; the researcher does not synthesize the network-feature rows.',
        t('CIC와 ToN의 독립 평가 · 같은 출처의 행을 선택해 라벨 공개 단계만 구성',
          'Independent CIC and ToN evaluations; selected real rows with a designed label-release schedule'),'03')

    items=[text(.62,1.10,12,.35,t('아래 수량 = 해당 단계에 새로 제공한 학습/context 사례 · 공격 도입 단계순',
        'Counts = newly supplied training/context examples at each stage · Ordered by class arrival'),18,color=GREY)]
    # Aligned numerical columns; stage order is the explicit ordering criterion.
    widths=[.63,2.61,1.00,1.60]
    def schedule(x,title,rows):
        out=[text(x,1.75,5.84,.36,title,23,True,BLUE)]
        heads=t(['단계','새로 제공한 클래스','클래스당','누적 합계'],['Stage','Newly supplied class','Per class','Total so far'])
        xs=[x]
        for w in widths[:-1]:xs.append(xs[-1]+w)
        y=2.28
        out.append(box(x,y,sum(widths),.46,fill=NAVY))
        for c,(xx,w,head) in enumerate(zip(xs,widths,heads)):
            out.append(text(xx+.06,y+.10,w-.12,.25,head,14,True,'FFFFFF',align='right' if c>=2 else 'left'))
        y+=.46
        for i,(stage,names,counts,total) in enumerate(rows):
            h=.49 if '\n' not in names else .68
            if names.count('\n')==2:h=.95
            out.append(box(x,y,sum(widths),h,fill=PALE if i%2==0 else 'FFFFFF'))
            for c,value in enumerate([stage,names,counts,total]):
                out.append(text(xs[c]+.08,y+.10,widths[c]-.16,h-.14,value,16,False,INK,align='right' if c>=2 else 'left'))
            y+=h
        return out
    items += schedule(.62,'CIC2018',[
        ('1','Benign\nBrute force\nDoS','750\n125\n125','1,000'),
        ('2','DDoS','n','1,000 + n'),('3','Web attacks','n','1,000 + 2n'),
        ('4','Infiltration','n','1,000 + 3n'),('5','Bot','n','1,000 + 4n')])
    items += schedule(6.88,'ToN-IoT',[
        ('1','Benign\nDoS\nScanning','750\n125\n125','1,000'),
        ('2','DDoS\nInjection','n\nn','1,000 + 2n'),('3','Password','n','1,000 + 3n'),
        ('4','XSS','n','1,000 + 4n'),('5','Backdoor\nRansomware','n\nn','1,000 + 6n'),('6','MitM','n','1,000 + 7n')])
    items += [text(.65,6.62,12,.28,t('n = 100 예시 · CIC 총 1,400개 / ToN 총 1,700개 · 하루 공격량이 아닌 확보 라벨 수',
        'At n = 100: CIC 1,400 / ToN 1,700 total labels; n is a label budget, not daily traffic volume'),16,color=GREY)]
    page(t('평가 시나리오 · 신규 공격당 라벨 수 n','Evaluation scenario · n labels per new attack class'),items,
        'Counts and class arrival order from EXP77/EXP71 manifests. Initial total 750+125+125=1000. '
        'The n rows are real labelled training examples, not daily attack frequencies or test counts. '
        'Both methods receive the same IDs. Cumulative final sizes: CIC 1000+4n, ToN 1000+7n. '
        'Old displayed results sweep n=10/20/50/100/200/400, and the old cost snapshot uses n=100. '
        '100 is an illustrative setting, not a statistical midpoint, optimum, or measured operational frequency. '
        'The revised n=5/10/20/30/50 buffer plan is design-only and has no results yet.',
        t('기존 결과 n = 10·20·50·100·200·400 · 100개는 최적값이 아닌 설명용 조건',
          'Completed sweep: n = 10, 20, 50, 100, 200, 400; 100 is illustrative, not an optimum'),'03')

    if include_results:
        items=[text(.62,1.10,12,.37,t('초기 1,000개 · 신규 10–400개/클래스 · 두 모델 모두 누적 · 마지막 모델 전체 test',
            'Initial 1,000 · 10 to 400 labels/new class · Cumulative data · Final model on full test'),18,color=GREY),
            dict(kind='image',x=.22,y=1.80,w=12.89,h=4.07,path=str(QA/'original-5.png'),crop=[.017,.286,.017,.171]),
            text(.65,6.13,12,.64,t('전체 점수로 관찰한 라벨 수에 따른 변화\n신규 공격별 탐지와 기존 클래스의 갱신 전후 성능은 별도 확인 필요',
              'Overall quality as the number of available labels changes\nNew-attack detection and before/after retention require separate measurements'),19)]
        page(t('기존 누적 실험 · 라벨 수에 따른 전체 성능','Completed cumulative study · Quality versus labels'),items,
             'The graph is an unchanged crop of supplied PDF page 5; every plotted value is preserved. '
             'Only the framing changed: removed the 400-row winner callout and distinguished overall F1 from evidence of retention. '
             'This is completed 10–400 cumulative evidence, not the new 5–50 fixed-buffer experiment.',
             t('새 버퍼 비교는 설계 단계 · 신규 5·10·20·30·50개를 전량 투입',
               'Planned buffer comparison: supply all 5, 10, 20, 30 or 50 new examples per class'),'04')
        items=[text(.62,1.10,12,.36,t('신규 공격당 100개 예시 · XGB 누적 재학습 / PFN 누적 context · IDS FT 없음',
            'Illustrative n = 100 · XGB cumulative refit / PFN cumulative context · No IDS fine-tuning'),18,color=GREY),
            text(.62,1.63,12,.35,t('갱신 비용 = 2단계 이후 시간 합계 · 초기 준비·전체 추론 시간 제외',
            'Update cost = sum over stages after stage 1; initial preparation and inference excluded'),18,True,NAVY),
            dict(kind='image',x=.90,y=2.13,w=11.53,h=4.65,path=str(QA/'original-6.png'),crop=[.0135,.297,.0135,.005])]
        page(t('기존 누적 실험 · 품질과 후속 갱신 비용','Completed cumulative experiment · Quality and update cost'),items,
             'Unchanged graph crop from supplied PDF page 6. At n=100 final cumulative size is CIC 1400 / ToN 1700. '
             'Existing update times include model/context preparation after the initial stage, and exclude inference. '
             'These values do not establish fixed-buffer retention, minimum sufficient labels, or an end-to-end deployment cost advantage. '
             'Next experiment measures final new-class quality, paired existing-class changes and later update cost at n=5/10/20/30/50.',
             t('이 소규모 누적 조건의 실측 비용 · 고정 버퍼의 기존 성능 유지는 후속 평가',
               'Measured costs for this small cumulative setting; fixed-buffer retention remains to be tested'),'04')
    return scenes

def write_deck(lang,scenes):
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    geometry=[]
    for page_no,scene in enumerate(scenes,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.notes_slide.notes_text_frame.text=scene['notes']
        for item in scene['items']:
            x,y,w,h=[item[k] for k in ('x','y','w','h')]
            assert x>=0 and y>=0 and x+w<=W+.01 and y+h<=H+.01,item
            geometry.append(dict(page=page_no,**item))
            if item['kind']=='text':
                shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=shape.text_frame
                frame.clear();frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.NONE;frame.vertical_anchor=MSO_ANCHOR.TOP
                frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                for i,line in enumerate(item['text'].split('\n')):
                    p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
                    p.space_before=p.space_after=Pt(0);p.line_spacing=1.14
                    p.alignment=PP_ALIGN.RIGHT if item.get('align')=='right' else PP_ALIGN.LEFT
                    r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(item['size']);r.font.bold=item['bold'];r.font.color.rgb=RGBColor.from_string(item['color'])
                    properties=r._r.get_or_add_rPr();properties.set('lang','ko-KR' if lang=='ko' else 'en-US')
                    for tag in ('a:ea','a:cs'):
                        font_element=OxmlElement(tag);font_element.set('typeface','Pretendard');properties.append(font_element)
            elif item['kind']=='image':
                shape=slide.shapes.add_picture(item['path'],Inches(x),Inches(y),width=Inches(w),height=Inches(h))
                shape.crop_left,shape.crop_top,shape.crop_right,shape.crop_bottom=item['crop']
            else:
                shape=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW if item['kind']=='arrow' else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(item['fill']);shape.line.fill.background()
    stem='20261020_Meeting_en' if lang=='en' else '20261020_Meeting_intro_ko'
    dest=SLIDES/('eng' if lang=='en' else 'kor')/stem
    prs.save(dest.with_suffix('.pptx'))
    body=[]
    for scene in scenes:
        tags=[]
        for i in scene['items']:
            style=f'left:{i["x"]}in;top:{i["y"]}in;width:{i["w"]}in;height:{i["h"]}in;'
            if i['kind']=='text':
                style+=f'font-size:{i["size"]}pt;font-weight:{700 if i["bold"] else 400};color:#{i["color"]};text-align:{i.get("align","left")};'
                tags.append(f'<div class="text" style="{style}">{html.escape(i["text"])}</div>')
            elif i['kind']=='image':
                l,t,r,b=i['crop'];iw=i['w']/(1-l-r);ih=i['h']/(1-t-b)
                tags.append(f'<div class="clip" style="{style}"><img src="{Path(i["path"]).as_uri()}" style="position:absolute;width:{iw}in;height:{ih}in;left:{-l*iw}in;top:{-t*ih}in"></div>')
            else:
                style+=f'background:#{i["fill"]};'
                if i['kind']=='arrow':style+='clip-path:polygon(0 25%,60% 25%,60% 0,100% 50%,60% 100%,60% 75%,0 75%);'
                tags.append(f'<div class="shape" style="{style}"></div>')
        body.append('<section><div class="scene">'+''.join(tags)+'</div></section>')
    rendered=QA/f'{lang}_intro.html'
    rendered.write_text('<!doctype html><meta charset="utf-8"><style>'
        '@page{size:10in 5.625in;margin:0}'
        '*{box-sizing:border-box}html,body{margin:0;padding:0}body{font-family:Pretendard,sans-serif;-webkit-print-color-adjust:exact;print-color-adjust:exact}'
        'section{position:relative;width:10in;height:5.625in;page-break-after:always;overflow:hidden}'
        f'.scene{{position:relative;width:{W}in;height:{H}in;transform:scale(.75);transform-origin:top left}}'
        'section:last-child{page-break-after:auto}.text,.shape,.clip{position:absolute}.text{white-space:pre;line-height:1.14}.clip{overflow:hidden}'
        '</style>'+''.join(body))
    rendered_pdf=QA/f'{lang}_intro.pdf'
    command=['/opt/google/chrome/chrome','--headless','--no-sandbox','--disable-dev-shm-usage','--disable-gpu','--no-pdf-header-footer',
        '--allow-file-access-from-files','--run-all-compositor-stages-before-draw','--virtual-time-budget=5000',f'--print-to-pdf={rendered_pdf}',rendered.as_uri()]
    with (QA/f'{lang}_render.log').open('w') as log:
        subprocess.run(command,stdout=log,stderr=log,check=True,timeout=60)
    dest.with_suffix('.pdf').write_bytes(rendered_pdf.read_bytes())
    (QA/f'{lang}_scene.json').write_text(json.dumps(dict(scenes=scenes,geometry=geometry,output=str(dest.relative_to(ROOT))),ensure_ascii=False,indent=2)+'\n')
    return str(dest.relative_to(ROOT))

def main():
    subprocess.run(['pdftoppm','-f','5','-l','6','-scale-to','2400','-png',str(SOURCE),str(QA/'original')],check=True)
    outputs=[write_deck('ko',make_scenes('ko')),write_deck('en',make_scenes('en',True))]
    manifest=dict(source=str(SOURCE.relative_to(ROOT)),source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        revision=2, revised_explanation_pages=6, original_graph_pages=[5,6],new_result_page_numbers=[7,8],
        outputs=outputs,main_tex_changed=False,experiment_values_changed=False,
        changes=['operational gap before method','external context versus pretrained weights','independent data-source flow',
                 'same-benchmark 100-row worked example','stage counts as n; 100 illustrative',
                 'result scope and update-cost definition; winner callout removed'],
        pptx_limitation='Original architecture and result graphs are images; explanatory text and scenario table are editable',
        reference=PAPER)
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
