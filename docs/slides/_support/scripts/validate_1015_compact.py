#!/usr/bin/env python3
"""Check compact slides and historical appendix against measured sources."""
from pathlib import Path
import json,re,hashlib,subprocess,argparse
import numpy as np
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/experiment_paths.py').is_file())
parser=argparse.ArgumentParser();parser.add_argument('--lang',choices=['ko','en'],default='ko');args=parser.parse_args();LANG=args.lang
qa=ROOT/'docs/slides/_support/qa'/('kor' if LANG=='ko' else 'eng')/'1015_labmeeting_qa'
m=json.loads((qa/'manifest.json').read_text())
result=json.loads((ROOT/m['new_source']).read_text())
out=ROOT/m['output'];pdf=out.with_suffix('.pdf');prs=Presentation(out)
assert not m['layout_preview']
assert len(prs.slides)==10 and m['evaluation']=='final_full_test' and m['language']==LANG
assert m['main_slide_count']==8 and m['appendix_slides']==[9,10]
assert m['large_final_slide']==7 and m['small_sota_removed']
assert m['data_scope_by_slide']=={'6':'small_label_1400_1700','7':'100000_train_plus_separate_validation','8':'new_attack_budget_sweep'}
budget=json.loads((ROOT/m['budget_source']).read_text())
budget['jobs'].extend(json.loads((ROOT/m['expanded_budget_source']).read_text())['jobs'])
budget_jobs={(j['dataset'],j['method'],j['new_per_class']):j for j in budget['jobs']}
history=json.loads((ROOT/m['historical_source']).read_text())
large=json.loads((ROOT/m['large_final_source']).read_text())
assert large['state']=='complete' and m['historical_evaluation']=='final_full_test'
large_jobs={(j['dataset'],j['case'],j['method']):j for j in large['jobs']}
mixed=json.loads((ROOT/m['mixed_final_source']).read_text())
assert mixed['state']=='complete' and m['query_order']=='seed43_mixed' and m['query_batch_rows']==8192
assert json.loads((ROOT/m['mixed_final_source']).parent.joinpath('verification.json').read_text())['passed']
large_jobs={(j['dataset'],j['case'],j['method']):j for j in mixed['jobs'] if j['scope']=='large'}
assert m['current_results_include_initial_ids_finetuning'] is False
jobs={(j['dataset'],j['method']):j for j in result['jobs']}
for source in m.get('additional_sources',[]):
    extra=json.loads((ROOT/source).read_text())
    jobs.update({(j['dataset'],j['method']):j for j in extra['jobs']})
jobs.update({(j['dataset'],j['method']):j for j in mixed['jobs'] if j['scope']=='small'})
for (ds,method,n),j in budget_jobs.items():
    if n==100 and method in ['xgb','pfn_pretrained']:j['metrics']=jobs[ds,method]['metrics']
assert m['cost_metric_by_slide']=={'6':'update_seconds','7':None} and m['distpfn_initial_ft'] is False
assert all(r['metric'] in ('f1','update') for r in m['chart_data'])
assert {r['method'] for r in m['chart_data'] if r['slide']==6}=={'xgb','pfn_pretrained'}
assert all(r['metric']=='f1' for r in m['chart_data'] if r['slide']==7)
assert 'pfn_ft_accumulate' not in {r['method'] for r in m['chart_data']}
assert 'distpfn_ft' not in {r['method'] for r in m['chart_data']}
assert 'distpfn' in {r['method'] for r in m['chart_data'] if r['slide']==7}
assert len(m['chart_data'])==52 and not any(r['slide']==11 for r in m['chart_data'])
for ds in ['cic2018','toniot']:
    j=jobs[ds,'distpfn_pretrained'];base=jobs[ds,'pfn_pretrained']
    assert j['initial_finetuning_seconds']==0 and j['update_seconds']==base['update_seconds']
    assert j['initial_finetuning'] is False and j['source_model_sha256']==base['source_model_sha256']
    assert j['test_order']==base['test_order']=='seed43_mixed'
for source,digest in m['source_sha256'].items():
    assert hashlib.sha256((ROOT/source).read_bytes()).hexdigest()==digest
for row in m['chart_data']:
    if row['slide']==8:
        j=budget_jobs[row['dataset'],row['method'],int(row['case'].removeprefix('new_'))]
        v=100*j['metrics']['macro_f1']
    elif row['slide']==7:
        original=large_jobs[row['dataset'],row['case'],row['method']]
        assert original['evaluation']=='final_full_test' and not original['training_rerun']
        v=100*original['metrics']['macro_f1']
    else:
        j=jobs[row['dataset'],row['method']]
        v={'f1':100*j['metrics']['macro_f1'],'update':j['update_seconds']}[row['metric']]
    assert abs(row['value']-v)<1e-10
assert {r['method'] for r in m['chart_data'] if r['slide']==8}=={'xgb','pfn_pretrained'}
assert {r['method'] for r in m['chart_data'] if r['slide']==7}=={'xgb','tabpfn','ours','boostpfn','distpfn'}
assert {r['new_per_class'] for r in m['budget_counts']}=={10,20,50,100,200,400}
for row in m['budget_counts']:
    assert row['total']==budget_jobs[row['dataset'],'pfn_pretrained',row['new_per_class']]['label_budget']
    assert row['total']==1000+row['new_per_class']*(4 if row['dataset']=='cic2018' else 7)
for source,digest in history['source_hashes'].items():
    assert hashlib.sha256((ROOT/source).read_bytes()).hexdigest()==digest
for dataset in ['cic2018','toniot']:
    run=(ROOT/m['new_source']).parent;meta=json.loads((run/'data'/dataset/'manifest.json').read_text())
    source=Path(meta['source'])
    y=np.load(source/'train_y.npy');names=meta['class_names']
    plans=json.loads((run/'data'/dataset/'plans.json').read_text())
    for row in [r for r in m['scenario_data'] if r['dataset']==dataset]:
        plan=plans[row['stage']-1];assert row['cumulative']==len(plan['released'])
        actual={names[c].replace('_',' ').lower():int((y[plan['incoming']]==c).sum()) for c in np.unique(y[plan['incoming']])}
        assert actual=={name.lower():n for name,n in zip(row['classes'],row['new_counts'])}
    rows=[r for r in m['historical_counts'] if r['dataset']==dataset]
    assert [r['natural'] for r in rows]==sorted([r['natural'] for r in rows],reverse=True)
    current={names[c]:int((y[plans[-1]['released']]==c).sum()) for c in np.unique(y[plans[-1]['released']])}
    assert {r['name']:r['current'] for r in rows}==current
    histmeta=json.loads((ROOT/history['baseline']/'data'/dataset/'manifest.json').read_text())
    histdata=Path(histmeta['source']);train_ids=np.load(histdata/'train_ids.npy');train_y=np.load(histdata/'train_y.npy')
    sorter=np.argsort(train_ids)
    for case in ['natural','uniform']:
        job=ROOT/history['baseline']/'jobs'/f'{dataset}_{case}_tabpfn'
        stage=sorted(job.glob('stage_*'))[-1];ids=np.load(stage/'train_ids.npy')
        positions=sorter[np.searchsorted(train_ids[sorter],ids)]
        assert np.array_equal(train_ids[positions],ids)
        actual={name:int((train_y[positions]==c).sum()) for c,name in enumerate(histmeta['class_names'])}
        assert actual=={r['name']:r[case] for r in rows},(dataset,case,actual)
        assert sum(r[case] for r in rows)==100000
        fit=json.loads((stage/'fit_metrics.json').read_text())
        assert dict(zip(fit['class_names'],fit['train_class_counts']))==actual
    # Stage deltas are independently calculated from original input plans.
    large_plans=json.loads((ROOT/history['run']/'data'/dataset/'plans.json').read_text())
    page=9 if dataset=='cic2018' else 10
    table=next(sh.table for sh in prs.slides[page-1].shapes if sh.has_table)
    old_counts={case:{} for case in ['natural','uniform']}
    stages=[r for r in m['large_scenario_data'] if r['dataset']==dataset]
    for row in stages:
        s=row['stage']-1;assert row['page']==page
        ns=row['new_counts']['natural'];assert ns==sorted(ns,reverse=True)
        p=plans[s];small_actual={name:int((y[p['incoming']]==c).sum()) for c,name in enumerate(names)}
        assert [small_actual.get(name,0) for name in row['classes']]==row['new_counts']['small']
        assert row['cumulative']['small']==len(p['released'])
        for case in old_counts:
            ids=large_plans[case][s]['fit'];current={name:int((train_y[ids]==c).sum()) for c,name in enumerate(histmeta['class_names'])}
            delta={name:count-old_counts[case].get(name,0) for name,count in current.items()}
            assert [delta.get(name,0) for name in row['classes']]==row['new_counts'][case]
            assert row['cumulative'][case]==len(ids)
            assert sum(row['new_counts'][case])==sum(delta.values())
            old_counts[case]=current
        assert [table.cell(s+1,c).text for c in range(2,5)]==['\n'.join(f'{n:,}' for n in row['new_counts'][case]) for case in ['small','natural','uniform']]
        assert [table.cell(s+1,c).text for c in range(5,8)]==[f"{row['cumulative'][case]:,}" for case in ['small','natural','uniform']]
    vr=len(stages)+1
    for c,case in enumerate(['natural','uniform'],6):
        stage=sorted((ROOT/history['run']/'jobs'/f'{dataset}_{case}_ours').glob('stage_*'))[-1]
        val=json.loads((stage/'fit_metrics.json').read_text())['validation_rows']
        assert table.cell(vr,c).text==f'{val:,}'
    assert table.cell(vr,5).text=='0'
issues=[];overlaps=[];content_overlaps=[]
for i,slide in enumerate(prs.slides,1):
    texts=[];contents=[]
    for shape in slide.shapes:
        if (shape.has_text_frame and shape.text) or shape.has_table or shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
            contents.append(shape)
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+10000
            assert shape.top+shape.height<=prs.slide_height+10000
        if shape.has_text_frame and shape.text:
            assert not re.search(r'합니다|입니다|습니다|—',shape.text),(i,shape.text)
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+10000
            assert shape.top+shape.height<=prs.slide_height+10000
            texts.append(shape)
        if shape.has_table:
            for row in list(shape.table.rows)[1:]:
                for cell in list(row.cells)[1:]:
                    if re.fullmatch(r'[\d,\s]+',cell.text):
                        assert all(p.alignment==PP_ALIGN.RIGHT for p in cell.text_frame.paragraphs)
    for a,shape in enumerate(texts):
        for other in texts[a+1:]:
            w=min(shape.left+shape.width,other.left+other.width)-max(shape.left,other.left)
            h=min(shape.top+shape.height,other.top+other.height)-max(shape.top,other.top)
            if w>10000 and h>10000:overlaps.append((i,shape.text,other.text))
    for a,shape in enumerate(contents):
        for other in contents[a+1:]:
            w=min(shape.left+shape.width,other.left+other.width)-max(shape.left,other.left)
            h=min(shape.top+shape.height,other.top+other.height)-max(shape.top,other.top)
            if w>10000 and h>10000:content_overlaps.append((i,shape.name,other.name))
assert not overlaps,overlaps
assert not content_overlaps,content_overlaps
assert hashlib.sha256((ROOT/m['source_deck']).read_bytes()).hexdigest()==m['source_deck_sha256']
pages=int(re.search(r'Pages:\s+(\d+)',subprocess.check_output(['pdfinfo',str(pdf)],text=True)).group(1));assert pages==10
text=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
main_text=subprocess.check_output(['pdftotext','-f','1','-l','8','-layout',str(pdf),'-'],text=True)
assert '1014' not in text and '미적용 사전 비교' not in text
if LANG=='ko':
    assert '부록' in text and '별도 검증' in text and '확률 보정' in text
    assert '10·20·50·100·200·400' in text and '최종 모델 전체 test' in text
else:
    assert not re.search('[가-힣]',text), 'Korean text remains in English PDF'
    assert 'Appendix' in text and 'Separate validation' in text and '100k training samples' in text
    assert '10/20/50/100/200/400' in text
    for source in (qa/'compact').glob('*.svg'):
        assert not re.search('[가-힣]',source.read_text()),source
assert '행/초' not in text and '행 / 초' not in text and '배치 검토용' not in text
assert ('최종' in text and '제안 모델' in text) if LANG=='ko' else ('Final' in text and 'Proposed' in text)
assert '준비 + 전체 추론' not in text and '비용은 로그눈금' not in text
page7=subprocess.check_output(['pdftotext','-f','7','-l','7','-layout',str(pdf),'-'],text=True)
assert '후속 갱신' not in page7 and '갱신 합계' not in page7
assert (re.search(r'학습\s*10\s*만\s*개',page7) and 'Global' in page7) if LANG=='ko' else ('100k training samples' in page7 and 'Global' in page7)
assert '52.89' not in page7 and 'expert 호출 0건' not in page7
assert '75,000' in page7 and '25,000' in page7 and '79.00%' in page7
assert not any('초기 FT' in sh.text for slide in list(prs.slides)[5:7] for sh in slide.shapes if sh.has_text_frame)
(qa/'compact/rendered_text.txt').write_text(text)
validation=dict(revision=m['revision'],language=LANG,slides=10,pdf_pages=pages,plot_values_verified=len(m['chart_data']),
    scenario_stages_verified=len(m['scenario_data']),numeric_columns_right_aligned=True,
    initial_ids_finetuning_results=False,evaluation='final_full_test',date_shorthand_removed=True,large_final_result_in_main=7,small_sota_removed=True,
    historical_appendix_pages=[9,10],historical_class_counts_verified=len(m['historical_counts']),
    historical_counts_verified_from_original_train_ids=True,historical_evaluation='final_full_test',
    large_scenario_stages_verified=len(m['large_scenario_data']),
    new_class_budgets=[10,20,50,100,200,400],main_budget_slide=8,
    small_baselines_share_label_budget=True,large_proposed_uses_separate_validation=True,unintended_text_overlap_count=len(overlaps),source_deck_unchanged=True,
    unintended_content_overlap_count=len(content_overlaps),
    pptx_sha256=hashlib.sha256(out.read_bytes()).hexdigest(),pdf_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),
    update_cost_on_slide_6_only=True,sota_slide_performance_only=True,initial_ft_removed_from_slides_6_7=True,
    distpfn_without_finetuning=True,visual_review='pending')
validation.update(query_order='seed43_mixed',query_batch_rows=8192,paired_distpfn_grouping_verified=True)
other=ROOT/'docs/slides/_support/qa'/('eng' if LANG=='ko' else 'kor')/'1015_labmeeting_qa/manifest.json'
if other.exists():
    paired=json.loads(other.read_text())
    if paired['revision']==m['revision']:
        for field in ['chart_data','scenario_data','historical_counts','budget_counts','large_scenario_data','source_sha256','data_scope_by_slide']:
            assert paired[field]==m[field],field
        validation['bilingual_data_parity']=True
        other_prs=Presentation(ROOT/paired['output'])
        for a,b in zip(prs.slides,other_prs.slides):
            ta=[sh.table for sh in a.shapes if sh.has_table];tb=[sh.table for sh in b.shapes if sh.has_table]
            assert len(ta)==len(tb)
            for left,right in zip(ta,tb):
                for i,row in enumerate(left.rows):
                    for j,cell in enumerate(row.cells):
                        if re.fullmatch(r'[\d,\s]+',cell.text):assert right.cell(i,j).text==cell.text
        validation['bilingual_numeric_table_parity']=True
if LANG=='en':
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:assert not re.search('[가-힣]',sh.text)
            if sh.has_table:
                assert not any(re.search('[가-힣]',cell.text) for row in sh.table.rows for cell in row.cells)
    validation['english_translation_complete']=True
for dest in [qa/'validation.json',qa/'compact/validation.json']:dest.write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
