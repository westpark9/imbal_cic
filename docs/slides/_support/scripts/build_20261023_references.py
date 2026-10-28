#!/usr/bin/env python3
"""Korean literature slides, revised on top of the user's annotated PPTX.

Editable text/diagrams in PPTX, matching PDF from the same scene descriptions.
Original paper figures remain attributed raster excerpts. No manuscript edits.
"""
from pathlib import Path
from copy import deepcopy
import hashlib, html, json, re, shutil, subprocess, zipfile
import numpy as np
from PIL import Image, ImageFont, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree

ROOT=next(p for p in Path(__file__).resolve().parents if (p/'AGENTS.md').exists())
SLIDES=ROOT/'docs/slides'
QA=SLIDES/'_support/qa/20261023_references'
ASSETS=QA/'assets'
OUT=SLIDES/'kor/20261023_References_ko'
USER_BASE=SLIDES/'_support/archive/20261023_user_feedback/20261023_References_ko.pptx'
N_SLIDES=8 if USER_BASE.exists() else 10
REPORT=ROOT/'docs/lablog/reports/20261021_references.tex'
FIGS=ROOT/'docs/lablog/figures/references_20261022'
FONT=Path('/home/user/.local/share/fonts/Pretendard/Pretendard-Regular.otf')
BOLD=FONT.with_name('Pretendard-Bold.otf')
W,H=13.333333,7.5
NAVY,BLUE,CYAN,INK,GREY,PALE,TEAL,AMBER='00306C','007BC6','00B8EE','20262E','617080','F1F5F9','227B71','9F6B20'
SOURCES={
 'survey':('Tabular Foundation Models: A Systematic Survey','SSRN preprint, 2026','https://papers.ssrn.com/sol3/papers.cfm?abstract_id=7531718','§3.1, Figure 3, PDF pp.5–6'),
 'pfn35':('TabPFN-3.5: Technical Report','arXiv, 2026','https://arxiv.org/abs/2609.17895','§3.1–3.4, PDF pp.9–11'),
 'local':('Retrieval & Fine-Tuning for In-Context Tabular Models','NeurIPS, 2024','https://papers.nips.cc/paper_files/paper/2024/file/c40daf14d7a6469e65116507c21faeb7-Paper-Conference.pdf','Figure 3 p.5; Table 1 p.7'),
 'boost':('Prior-Fitted Networks Scale to Larger Datasets When Treated as Weak Learners','AISTATS, 2025','https://arxiv.org/abs/2503.01256','Figure 1 p.2; §4; Table 3 p.6'),
 'dist':('Mitigating Label Shift in Tabular In-Context Learning via Test-Time Posterior Adjustment','ICML, 2026','https://arxiv.org/abs/2605.04363','Figure 4 p.4; Algorithm 1 p.5; Table 5 p.6'),
 'tabfm':('TabFM: A Zero-Shot Foundation Model for Tabular Data','arXiv, 2026','https://arxiv.org/abs/2609.37959','Figure 2 p.2; §4 and implementation appendix'),
 'imbalance':('Correcting Class Imbalance in Prior-Data Fitted Networks for Tabular Classification','ITW 표기 원고, 2026','https://arxiv.org/abs/2605.21742','Figure 4 p.4; Table II p.5'),
 'ensemble':('Ensembling Tabular Foundation Models — A Diversity Ceiling And A Calibration Trap','arXiv, 2026','https://arxiv.org/abs/2605.18696','§3–5, Figure 1 p.5'),
 'iot':('Optimizing IoT Intrusion Detection with Tabular Foundation Models for Smart City Forensics','arXiv, 2026','https://arxiv.org/abs/2604.11394','Figure 1 p.5; Tables 2–3 p.7'),
 'lotbonc':('LotBoNC: Novel Botnet Traffic Classification under Long-tailed Distributions','ACM ASIA CCS, 2026','https://doi.org/10.1145/3779208.3785276','§4 pp.3–6; §5 pp.6–8; Appendix A.1 p.15'),
 'streams':('Rethinking Tabular Foundation Models On Data Streams','arXiv, 2026','https://arxiv.org/abs/2610.05352','Tables 2–6, Figures 1 and 7, Appendix A and D'),
 'icd':('In-Context Data Distillation with TabPFN','ICLR ME-FoMo Workshop, 2024','https://arxiv.org/abs/2402.06971','Figure 2 p.3; §3–4 and Table 1 p.4'),
}
SCENES=[]
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def txt(x,y,w,s,size=22,bold=False,color=INK,align='left',url=None):
    return dict(kind='text',x=x,y=y,w=w,h=len(s.splitlines())*size*1.18/72+.035,
                text=s,size=size,bold=bold,color=color,align=align,url=url)
def box(x,y,w,h,fill=PALE,border=None):
    return dict(kind='box',x=x,y=y,w=w,h=h,fill=fill,border=border)
def arrow(x,y,w=.38,h=.25,down=False,color=BLUE):
    return dict(kind='down' if down else 'arrow',x=x,y=y,w=w,h=h,fill=color)
def image(key,x,y,w,h):
    path=ASSETS/(key+'.png');iw,ih=Image.open(path).size
    scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
    return dict(kind='image',x=x+(w-ww)/2,y=y+(h-hh)/2,w=ww,h=hh,path=str(path))
def caption(x,y,w,s,key=None):
    return txt(x,y,w,s,14,color=GREY,url=SOURCES[key][2] if key else None)
def step(n,title,body,y,x=7.05,w=5.65):
    return [box(x,y,.39,.39,BLUE),txt(x,y+.045,.39,str(n),17,True,'FFFFFF','center'),
            txt(x+.56,y-.01,w-.56,title,22,True,NAVY),
            txt(x+.56,y+.48,w-.56,body,20)]
def chip(x,y,w,s,fill=PALE,color=INK,size=17):
    return [box(x,y,w,.46,fill),txt(x+.08,y+.085,w-.16,s,size,True,color,'center')]
def block(x,y,w,h,title,body='',color=NAVY,body_size=19):
    items=[box(x,y,w,h,'FFFFFF','CEDAE5'),box(x,y,w,.44,color),txt(x+.1,y+.075,w-.2,title,19,True,'FFFFFF','center')]
    if body:items.append(txt(x+.12,y+.63,w-.24,body,body_size,False,INK,'center'))
    return items
def source_line(key,label=None,x=.62,y=1.05,w=12.05,size=15):
    title,venue,url,_=SOURCES[key]
    return txt(x,y,w,label or title+' · '+venue,size,color=GREY,url=url)
def page(title,items,summary,section,notes,sources):
    n=len(SCENES)+1
    source_notes='\n\n'.join(f'{SOURCES[k][0]}\n{SOURCES[k][1]}\n{SOURCES[k][3]}\n{SOURCES[k][2]}' for k in sources)
    SCENES.append(dict(title=title,source_keys=sources,notes='발표 설명\n'+notes+'\n\n원문 근거\n'+source_notes,
      items=[txt(.6,.32,12.08,title,30,True,NAVY)]+items+
      [box(0,7.04,W,.46,NAVY),txt(.6,7.15,9.55,summary,14,color='FFFFFF'),
       txt(10.15,7.16,2.15,section,12,color='C9D8EB',align='right'),
       txt(12.48,7.145,.28,str(n),14,color='FFFFFF',align='right')]))

def assets():
    ASSETS.mkdir(parents=True,exist_ok=True)
    keys=['survey','local','boost','tabfm','dist','imbalance','ensemble','lotbonc','streams_cache','icd']
    for key in keys:
        subprocess.run(['pdftoppm','-singlefile','-scale-to','2400','-png',str(FIGS/(key+'.pdf')),str(ASSETS/key)],check=True,capture_output=True)

def scenes():
    # 1: taxonomy and model interface, not a full genealogy lecture.
    a=[source_line('survey','Tabular Foundation Models: A Systematic Survey · 2026'),
       image('survey',.65,1.52,4.1,5.02),
       caption(.7,6.59,4.1,'원문 Figure 3 · 기반 모델 42개','survey'),
       txt(5.15,1.56,7.4,'두 분류 축',23,True,NAVY),
       txt(5.15,2.03,7.4,'사전학습 자료  합성 prior / 실제 표\n전이 방식  ICL / 가중치 FT / 텍스트 생성 /\n                 작은 모델의 가중치 생성',21),
       txt(5.15,3.22,7.4,'TabPFN · 라벨 사례를 입력으로 받는 Transformer',21,True,NAVY)]
    a+=block(5.15,3.8,2.25,1.18,'Context + query','사례 (x, y) + 새 x')
    a+=[arrow(7.55,4.26,.4)]
    a+=block(8.10,3.8,2.22,1.18,'사전학습 모델','가중치 θ 고정')
    a+=[arrow(10.47,4.26,.4)]
    a+=block(11.02,3.8,1.63,1.18,'예측','클래스 확률')
    a+=[txt(5.15,5.28,7.4,'In-context learning (ICL)\n입력 사례를 참고하는 예측 · 가중치 학습과 구별',21),
        box(5.15,6.16,7.49,.55,PALE),
        txt(5.31,6.29,7.1,'v3.5의 변화  수치 표현·용량 확장 · 현재 실험은 v3',18,url=SOURCES['pfn35'][2])]
    page('Tabular foundation model의 분류와 적용 방식',a,
         'ICL에서는 사전학습 가중치와 현재 문제의 라벨 사례가 서로 다른 역할','01 모델과 앙상블',
         '서베이 Figure 3의 두 축만 먼저 읽습니다. 열은 합성 prior와 실제 표, 행은 새 표에 적용하는 방식입니다. 42개 모델의 족보를 모두 외우려는 목적은 아닙니다. Tabular foundation model은 표 데이터의 예측 규칙을 사전학습한 모델입니다. TabPFN은 prior-data fitted network 계열의 Transformer이며, 라벨이 있는 context와 예측할 query를 함께 읽습니다. 여기서 context는 사전학습 가중치 안에 들어 있는 데이터가 아니라 현재 예측에 주는 입력입니다. 가중치 fine-tuning은 별도 최적화입니다. TabPFN-3.5는 수치 인코딩과 용량 등을 개선했지만, 현재 우리 결과는 v3로 측정했습니다. 이후 논문들은 context를 선택하거나, 여러 예측을 합치거나, 출력 확률을 보정하는 서로 다른 지점을 바꿉니다.',
         ['survey','pfn35'])

    # 2: two representative context adaptations, with large original figures.
    a=[txt(.65,1.05,5.75,'LoCalPFN · NeurIPS 2024',23,True,NAVY,url=SOURCES['local'][2]),
       txt(7.0,1.05,5.7,'BoostPFN · AISTATS 2025',23,True,NAVY,url=SOURCES['boost'][2]),
       image('local',.72,1.64,5.55,2.10),image('boost',7.1,1.64,5.4,2.10),
       caption(.72,3.87,5.55,'Query별 이웃 검색과 지역 context를 이용한 FT','local'),
       caption(7.1,3.87,5.4,'오류를 반영한 context 선택과 예측 결합','boost'),
       box(6.66,1.1,.015,5.45,'DCE5ED'),
       txt(.78,4.4,5.55,'① Query와 가까운 사례 검색\n② 지역 context로 TabPFN 예측\n③ 지역 배치로 가중치 FT 추가',21),
       txt(7.14,4.4,5.45,'① 일부 사례로 PFN 약한 예측기 구성\n② 이전 오류에 따라 표본 가중치 변경\n③ 여러 예측기의 출력을 결합',21),
       txt(.78,6.04,5.4,'참고점  검색 효과와 FT 효과의 분리',19,True,TEAL),
       txt(7.14,6.04,5.45,'참고점  실패 사례를 반영한 expert 구성',19,True,TEAL)]
    page('Context 적응 · 지역 검색과 오류 기반 선택',a,
         'LoCalPFN은 검색과 가중치 FT · BoostPFN은 고정 PFN의 context 앙상블','01 모델과 앙상블',
         '두 논문은 context를 바꾼다는 공통점이 있지만 학습 대상은 다릅니다. LoCalPFN은 현재 query 주변의 이웃으로 context를 구성하고, 그 지역적인 입력 방식에 맞춰 본체 가중치도 fine-tuning합니다. 검색만 사용하면 TabPFN-kNN이며 원래 LoCalPFN과 구별해야 합니다. BoostPFN은 전체 라벨 pool의 일부로 약한 PFN 예측기를 만들고, 오분류에 집중하도록 다음 subset의 표본 가중치를 바꿉니다. PFN 본체를 FT하는 대신 여러 context의 예측을 결합합니다. LoCalPFN의 95개 자료 평균 AUC는 0.908, TabPFN 0.867, XGB 0.892였고, BoostPFN의 5,000개 train 조건 평균 AUC는 0.908, XGB 0.849였습니다. 서로 다른 논문과 프로토콜의 수치를 가로로 비교해서는 안 됩니다. 우리에게는 지역 검색의 효과와 오류를 보완하는 expert 구성을 따로 실험할 근거가 됩니다.',
         ['local','boost'])

    # 3: separate TabFM backbone from the TabFM+ view ensemble.
    a=[source_line('tabfm','TabFM: A Zero-Shot Foundation Model for Tabular Data · arXiv 2026'),
       image('tabfm',.67,1.57,5.73,3.65),
       caption(.74,5.3,5.6,'셀 표현 → 행 표현 → context 기반 예측','tabfm')]
    a+=step(1,'같은 모델에 다른 입력 view 제공','특징 변환·순열·교차 특징 등\n32개 view에서 반복 예측',1.65)
    a+=step(2,'Validation으로 결합 가중치 학습','비음수 최소제곱으로 가중치 선택\n학습 가중치 75% + 균등 가중치 25%',3.16)
    a+=step(3,'결합 확률의 사후 보정','본체 가중치는 고정\n결합·보정에는 validation 라벨 사용',4.67)
    a+=chip(.75,5.93,5.65,'TabFM+ = 고정 TabFM + 입력 view 앙상블',fill=PALE,color=NAVY,size=18)
    page('TabFM+ · 같은 모델의 여러 입력 view 결합',a,
         '여러 예측을 얻는 비용과 결합기를 학습할 validation 라벨까지 비교에 포함','01 모델과 앙상블',
         'TabFM은 TabPFN의 버전명이 아니라 별도 사전학습 모델입니다. 그림은 TabFM 기본 모델 구조로, 셀을 표현하고 행을 요약한 뒤 context와 query 관계로 예측합니다. TabFM+는 이 그림의 모델을 32개의 변환된 입력 view에 반복 적용하는 확장입니다. 여러 독립 모델을 새로 사전학습하는 방식은 아닙니다. Validation 정답으로 비음수 최소제곱 결합 가중치를 구하고, 75%의 학습 가중치와 25%의 균등 가중치를 혼합한 뒤 분류 확률을 보정합니다. 원문은 TabArena 51개 과제에서 TabFM+의 개선을 보고했습니다. 우리 연구가 참고할 부분은 서로 다른 관점의 입력이 상호 보완 오류를 만드는지입니다. 다만 본체가 고정돼 있다는 사실이 결합기 학습과 라벨 비용도 없다는 뜻은 아닙니다.',
         ['tabfm'])

    # 4: same scientific question, posterior adjustment and thresholding.
    a=[txt(.68,1.1,5.85,'DistPFN · ICML 2026',22,True,NAVY,url=SOURCES['dist'][2]),
       image('dist',.7,1.69,5.8,1.35),
       txt(.82,3.14,5.52,'Context 클래스 비율과\n예측 분포의 비로 출력 확률 보정',21),
       txt(.82,4.12,5.52,'추가 가중치 FT 없음\n원문 기본값은 query 배치 평균 분포',20,True,TEAL),
       box(6.69,1.1,.015,5.45,'DCE5ED'),
       txt(7.03,1.1,5.65,'클래스 불균형 보정 · ITW 표기 2026',21,True,NAVY,url=SOURCES['imbalance'][2]),
       image('imbalance',7.25,1.72,5.25,3.01),
       caption(7.12,4.89,5.45,'임계값에 따른 다수·소수 클래스 정확도 변화','imbalance'),
       txt(.82,5.35,5.52,'확인 대상\nContext와 test의 클래스 비율 차이',20,True,NAVY),
       txt(7.16,5.35,5.42,'비교 대상\n임계값 조절 / 다수 클래스 축소',20,True,NAVY),
       txt(.82,6.36,11.8,'우리 연구  prior 편향과 특징 구분 실패를 분리한 대조 실험',20,True,TEAL)]
    page('불균형 적응 · 확률 보정과 판정 기준 조절',a,
         '분포 보정의 이득은 클래스별 precision·recall과 정상 오탐으로 확인','01 모델과 앙상블',
         '소수 공격이 약한 이유를 모두 context 개수 부족으로 볼 수는 없습니다. DistPFN은 context의 클래스 prior와 예측 분포를 이용해 출력 확률을 보정합니다. 본체 FT가 필요하지 않으며, 원문 기본 실험은 개별 query가 아니라 query 배치 평균 분포를 보정 계수에 사용했습니다. 253개 분류 자료에서 label shift 강도 β=2일 때 TabPFN-v2 accuracy 0.759가 DistPFN 0.783으로 높아졌습니다. 오른쪽 연구는 이진분류 11개 자료에서 임계값 조절과 downsampling을 비교했습니다. 균형 test의 평균 balanced accuracy는 무보정 0.710, 임계값 0.793, downsampling 0.779였습니다. 이 결과를 다중 공격의 자연 빈도 test에 그대로 옮길 수는 없습니다. 우리에게는 prior가 문제인지 특징의 구분이 문제인지 나누는 대조군이며, 정상 오탐이 함께 늘지 않는지도 봐야 합니다.',
         ['dist','imbalance'])

    # 5: why an ensemble needs complementary errors.
    a=[source_line('ensemble','Ensembling Tabular Foundation Models · arXiv 2026'),
       image('ensemble',.71,1.59,5.67,3.4),
       caption(.77,5.14,5.55,'153개 분류 과제 · 평균 정확도와 준비 비용','ensemble')]
    a+=step(1,'같은 오분류의 반복 여부 확인','Expert를 늘려도 같은 사례를 틀리면\n복구 가능한 오류의 증가 제한',1.65)
    a+=step(2,'결합기보다 oracle 여지부터 확인','어느 expert라도 맞힌 사례의 상한과\n단일 모델의 차이 비교',3.17)
    a+=step(3,'이득이 있을 때 선택·보정 개선','정답으로 고친 사례와 새 오분류 분리\n같은 라벨·보존량·추론 예산으로 비교',4.69)
    a+=[box(.73,5.77,5.72,.92,PALE),txt(.9,5.9,5.36,'IoT 포렌식 응용  PFN 선별 → 트리 앙상블\n개별 모델 점수와 전체 pipeline 효과의 구분',17,color=NAVY,url=SOURCES['iot'][2])]
    page('앙상블 설계 · 모델 수보다 오류의 상호 보완',a,
         '단일 expert의 구분 능력과 결합 정책의 효과를 분리한 검증 필요','01 모델과 앙상블',
         '앙상블 논문은 강한 TFM 6개와 결합 전략 6개를 153개 자료에서 평가했습니다. 그림의 오른쪽으로 갈수록 비용이 크고 위로 갈수록 정확도가 높습니다. 저자들은 높은 오류 유사성과 일부 결합기의 확률 보정 저하를 관찰했습니다. 이는 모든 앙상블이 무의미하다는 결론은 아닙니다. 우리 모델에서는 먼저 어느 expert라도 맞힌 사례를 선택할 수 있다고 가정한 oracle 상한을 봐야 합니다. 그 상한조차 Global과 비슷하면 결합기를 더 복잡하게 만들기보다 expert context나 표현을 바꿀 필요가 있습니다. 상한의 여지가 크면 실제 선택기가 왜 그 수정을 놓치는지 분석합니다. ToN-IoT 장치 telemetry의 포렌식 응용도 PFN 선별 뒤 트리 앙상블 구조를 제안했지만, 개별 모델의 점수와 전체 cascade의 이득은 구분해야 합니다. 현재 제안 모델의 성능을 향상하려면 이 두 효과를 분리하는 진단이 우선입니다.',
         ['ensemble','iot'])

    # 6: dataset provenance, visibility, and imbalance in one explicit diagram.
    a=[source_line('lotbonc','LotBoNC: Novel Botnet Traffic Classification under Long-tailed Distributions · ASIA CCS 2026',size=15),
       txt(.79,1.43,11.7,'각 벤치마크에서 독립 적용',16,True,TEAL)]
    a+=block(.7,1.78,3.1,1.15,'실제 트래픽','IoT-23 / ISCX-Bot / CTU-13',body_size=16.5)
    a+=[arrow(4.01,2.2,.4)]
    a+=block(4.62,1.78,3.45,1.15,'입력 전처리','PCAP → flow → burst / token',body_size=17)
    a+=[arrow(8.3,2.2,.4)]
    a+=block(8.92,1.78,3.68,1.15,'평가 집합 구성','같은 출처 안에서 사례·라벨 분리')
    a+=[txt(.79,3.23,11.9,'실험자가 조절한 부분  라벨 공개 대상과 클래스별 표본 수',22,True,NAVY)]
    a+=block(.78,3.88,5.73,1.78,'라벨 집합 L','알려진 공격만 포함\n클래스별 같은 수의 사례',BLUE)
    a+=block(6.83,3.88,5.73,1.78,'비라벨 집합 U','알려진 공격 + 새로운 공격\n불균형 구성 · 정답은 학습에서 비공개',TEAL)
    # Schematic class blocks show composition, not undocumented row counts.
    for j,color in enumerate([BLUE,BLUE,BLUE,TEAL,TEAL,TEAL]):
        a.append(box(1.22+j*.75,5.86,.62,.25,color))
    for j,color in enumerate([BLUE,BLUE,BLUE,BLUE,TEAL,AMBER]):
        a.append(box(7.38+j*.61,5.86,.48,.25,color))
    a+=[caption(.94,6.27,5.43,'구성 예시 · 파랑·청록 = known / 주황 = novel'),
        txt(6.98,6.23,5.44,'Known / novel 총량 비 γ = 1, 5, 10',18,True,NAVY),
        caption(.94,6.65,11.5,'블록 수는 설명용 예시 · 실제 표본 수와 별개')]
    page('LotBoNC · 원본 트래픽에서 평가 조건을 만드는 과정',a,
         '실제 트래픽의 라벨 가용성과 불균형을 조절해 새로운 botnet 군집 발견 평가','03 평가 시나리오',
         'LotBoNC에서 참고할 첫 부분은 데이터 구성입니다. 원시 PCAP의 TCP/UDP flow를 모으고 burst와 byte bigram token으로 변환했습니다. 사전학습에는 별도 트래픽 출처를 사용했고, 여기에는 평가 구성에 사용한 IoT-23·ISCX-Bot-2014·CTU-13만 표시했습니다. 각 평가 데이터 안에서 알려진 공격의 균형 라벨 집합 L과, 알려진 공격과 새 공격이 섞인 불균형 비라벨 집합 U를 만들었습니다. 두 개의 서로 다른 벤치마크를 합쳐 B를 만드는 설명이 아닙니다. γ는 비라벨 U 안의 known 표본 총량을 novel 표본 총량으로 나눈 값이며, 우리 신규 공격당 n개 라벨 수와 다른 변수입니다. 그림의 블록은 실제 배분 수가 아닌 구성 예시입니다. 우리의 시나리오에도 같은 출처의 어느 행을 언제 라벨 사례로 공개하는지 설명하는 방식을 적용할 수 있습니다. LotBoNC는 새 공격 정답을 주지 않고 군집을 발견하는 과제이고, 우리는 공격별 정답 사례를 제공한 뒤 분류하므로 채점 목표는 다릅니다.',
         ['lotbonc'])

    # 7: one loop, three easy steps, small mathematical semantics only.
    a=[image('lotbonc',.75,1.35,5.8,3.05),
       caption(.9,4.48,5.53,'원문 Figure 4 · 임시 정답 배정과 표현 학습의 반복'),
       box(.8,5.1,5.68,1.38,PALE),
       txt(.98,5.25,5.3,'E-step의 배정 기준',21,True,NAVY),
       txt(.98,5.75,5.3,'예측과 맞는 배정 + 추정 클래스 비율 유지\n예시  10개 사례 · prior 0.6 : 0.3 : 0.1',18)]
    a+=step(1,'E-step · batch 전체의 임시 정답','높은 예측 확률을 우선하되\n클래스별 총량을 추정 prior에 맞춤',1.45)
    a+=step(2,'M-step · 클래스 대표 표현 갱신','배정된 사례로 prototype과 prior 수정\n이전 값과 새 값을 이동평균으로 결합',3.10)
    a+=step(3,'표현이 좋아지도록 가중치 학습','대조 손실 + prototype 손실\nEncoder 가중치 θ 갱신 후 반복',4.75)
    page('LotBoNC · 임시 정답과 클래스 표현의 반복 갱신',a,
         '핵심 참고점은 라벨 공개·불균형 설계 · 고정 TabPFN의 context 교체와는 다른 학습','03 평가 시나리오',
         '그림의 윗부분이 E-step입니다. 각 사례를 단순 argmax로 분류하면 다수 클래스에 몰릴 수 있어, batch 전체의 배정 A를 함께 구합니다. 현재 예측 확률 P에 대한 비용 C=-log P를 작게 하면서, 행 합은 사례별 질량 w, 열 합은 추정 prior r가 되게 하는 최적 수송 문제입니다. 예를 들어 prior가 0.6, 0.3, 0.1이면 10개 사례를 6, 3, 1개에 해당하는 질량으로 배정하되 soft assignment가 가능합니다. Sinkhorn은 행과 열 합을 번갈아 맞춥니다. M-step에서는 클래스 대표 벡터인 prototype과 클래스 비율을 이동평균으로 갱신합니다. Masked-token과 대조 사전학습으로 준비한 encoder를 전체 loss로 실제 업데이트하므로 고정 PFN context 갱신과 다릅니다. Known 분류와 novel 군집 정합도를 평가하며, 군집 번호는 Hungarian matching으로 대응합니다. 원문 OT kernel 부호·행렬 차원과 prototype prior 항 표기에 불일치가 있어 세부 식을 그대로 재현하지 않았고, 보고서 부록에 확인 위치를 남겼습니다. 발표에서는 논문의 일관된 작동 의도를 설명하되 원문 구현이 검증됐다고 말하지 않습니다.',
         ['lotbonc'])

    # 8: prequential protocol then encoding/cache cost.
    a=[source_line('streams','Rethinking Tabular Foundation Models On Data Streams · arXiv 2026')]
    for x,title,body in [( .74,'새 사례 도착','x_t'),(3.83,'먼저 예측','현재 memory 사용'),(6.92,'정답 도착','d개 사례 뒤'),(10.01,'그다음 갱신','Memory M 변경')]:
        a+=block(x,1.63,2.57,1.05,title,body)
    for x in [3.4,6.49,9.58]:a.append(arrow(x,2.03,.31))
    a+=[image('streams_cache',.75,3.29,5.86,1.55),
       caption(.84,5.02,5.61,'매번 context 재계산 / 저장된 표현 재사용'),
       txt(.85,5.68,5.61,'기본 조건  메모리 1,000행 · 즉시 라벨\n실제 stream 5개 + 합성 drift stream 4개',19)]
    a+=step(1,'메모리 정책과 backbone을 분리','같은 모델에서 FIFO와 선택 정책 비교\n단순 FIFO도 강한 기준선',3.24)
    a+=step(2,'행 교체와 context encoding 구분','갱신 빈도를 줄여도 예측마다 재계산하면\n큰 비용은 계속 발생',4.85)
    page('데이터 스트림 · 예측 순서와 context 비용의 분리',a,
         '가중치 갱신이 없어도 context 처리 비용 발생 · cache 재사용 여부의 실측 필요','03 평가 시나리오',
         '이 논문은 미리 고정한 최종 test가 아니라 예측 후 학습하는 prequential 평가입니다. 지금 도착한 사례를 현재까지 라벨을 확보한 memory로 예측하고, 정답이 지연 d 뒤에 도착한 다음에만 그 사례를 사용합니다. 기본 memory는 1,000행, warm-up은 100행, query batch는 1개입니다. 실데이터 5개와 합성 drift stream 4개에서 backbone과 memory 정책을 나눠 비교했습니다. 전체 평균에서는 FIFO가 더 복잡한 정책과 비슷했지만, 라벨 지연이 큰 특정 stream에서는 선택 정책의 이득이 있었습니다. 중요한 비용 구분은 memory의 행을 바꾸는 작업과, 그 내용을 모델이 encoding하는 작업입니다. 원문 구현은 예측마다 context를 다시 처리해, refresh 간격만 늘려서는 큰 비용을 줄이지 못했습니다. 저자들의 작은 causal 모델은 이전 표현을 cache에서 재사용하는 가능성을 보였습니다. 그 결과를 현재 TabPFN-v3의 cache가 공짜로 incremental update된다는 근거로 쓰지는 않습니다. 우리의 다음 비용 계측에서도 선택·전처리·cache 준비·query 추론을 나눠야 합니다.',
         ['streams'])

    # 9: frozen theta and optimized synthetic input, not weight fine-tuning.
    a=[source_line('icd','In-Context Data Distillation with TabPFN · ICLR ME-FoMo Workshop 2024'),
       image('icd',.7,1.66,5.84,2.15),
       caption(.84,3.94,5.56,'왼쪽  모델 학습을 포함한 증류 / 오른쪽  context 직접 최적화'),
       box(.82,4.55,5.67,1.93,PALE),
       txt(1.02,4.76,5.27,'고정하는 것  모델 가중치 θ, context 라벨 y',19,True,NAVY),
       txt(1.02,5.31,5.27,'바꾸는 것  context의 특징값 X',22,True,BLUE),
       txt(1.02,5.91,5.27,'X ← X − α · ∂Loss/∂X',24,True,NAVY)]
    a+=step(1,'작은 합성 context 준비','클래스별 같은 수의 라벨 고정\n특징값은 최적화 가능한 변수',1.69)
    a+=step(2,'실제 train 사례를 query로 예측','현재 합성 context + 고정 TabPFN\n실제 정답과의 오차 계산',3.21)
    a+=step(3,'오차로 context 특징값 수정','Gradient를 입력 X까지 전달\n가중치 θ를 바꾸지 않고 반복',4.73)
    page('ICD · 모델 대신 context의 특징값 학습',a,
         'ICD의 학습 대상은 합성 context의 특징값 · 모델 가중치와 라벨은 고정','04 Context 최적화',
         'ICD는 in-context data distillation의 약자입니다. 핵심은 무엇을 학습하는지입니다. 일반 증류는 합성 데이터로 모델을 학습한 뒤 그 결과로 합성 데이터를 수정하는 두 겹 최적화가 필요합니다. TabPFN은 context를 읽는 forward 자체가 예측 절차이므로 본체를 학습하는 내부 반복을 생략할 수 있습니다. θ는 고정하고 작은 context의 라벨도 고정한 채 특징 X만 변수로 둡니다. 실제 train의 사례를 query로 넣어 정답 y의 음의 로그확률 loss를 계산하고, gradient로 context X를 업데이트합니다. 수식의 α는 context optimizer 학습률이고 ∂Loss/∂X는 어느 방향으로 context를 움직이면 예측 오류가 줄어드는지를 나타냅니다. Gradient가 모델 연산을 통과한다는 사실과 모델 가중치를 optimizer가 변경한다는 사실은 다릅니다. FIFO나 무작위 선택은 실제 행 ID를 바꾸지만, ICD가 만든 점은 실제 flow 행이 아닌 합성 특징값입니다.',
         ['icd'])

    # 10: explicit label pool vs retained context, and result interpretation.
    a=[txt(.75,1.13,5.74,'원문 실험  큰 라벨 집합 → 1,000개 context',22,True,NAVY),
       txt(.78,1.69,5.68,'48개 표 분류 과제 · 전체 2,000∼200,000행\nTrain / validation / test = 80 / 10 / 10',18),
       txt(.79,2.51,5.6,'F1 중앙값 (%)',18,True,NAVY)]
    for i,(name,val,color) in enumerate([('튜닝 XGB',92.1,INK),('TabPFN + ICD',89.9,BLUE),('기본 XGB',89.3,GREY),('무작위 TabPFN',84.4,'A9C7DE')]):
        y=3.11+i*.63
        a+=[txt(.8,y+.02,1.94,name,17,bold=name=='TabPFN + ICD'),
            box(2.83,y,2.72*val/100,.35,color),
            txt(5.67,y+.01,.7,f'{val:.1f}',18,True,color=INK,align='right')]
    a+=[box(2.83,5.57,2.72,.012,'CCD6DF'),
        txt(2.78,5.65,.3,'0',11,color=GREY),
        txt(5.20,5.65,.35,'100',11,color=GREY,align='right'),
        caption(.83,6.06,5.5,'같은 48개 자료의 중앙값 · IDS Macro-F1과 별개','icd')]
    a+=step(1,'압축에 쓴 라벨 수와 보관량 구분','큰 train 전체로 context를 최적화\nContext 1,000개 ≠ 확보 라벨 1,000개',1.44)
    a+=step(2,'반복 최적화 비용까지 포함','Context 증류 + validation 선택\n최종 cache 준비 + query 추론',3.08)
    a+=step(3,'우리 연구의 적용 순서','극소수 신규 라벨의 단계 평가부터\n라벨 이력 증가 시 압축 방법 비교',4.72)
    page('ICD의 적용 조건 · 확보 라벨과 보존 예산',a,
         '현재는 소수 라벨의 활용부터 검증 · 큰 라벨 이력의 압축은 다음 연구 질문','04 Context 최적화',
         'ICD의 실험은 실제 라벨이 극소량인 상황이 아니라 큰 라벨 train을 작은 context로 압축하는 상황입니다. 48개 자료는 전체 2,000개에서 200,000개였고 80/10/10으로 나눴습니다. 최종 context는 총 1,000개지만 그 점들을 최적화할 때 큰 train의 라벨을 계속 사용했습니다. 표의 F1 중앙값은 튜닝 XGB 92.1%, ICD 89.9%, 기본 XGB 89.3%, 무작위 TabPFN 84.4%입니다. ICD가 기본 PFN을 개선했지만 튜닝 XGB를 넘지는 못했습니다. F1 정의와 데이터 집계가 우리 IDS Macro-F1과 같다고 가정하지 않습니다. 작은 context만 보고 라벨 효율이 높다고 주장하면 확보한 정보량을 숨기게 됩니다. 현재 우리의 후속 질문은 실제 새 공격 라벨이 2–5개뿐일 때 정상·기존 공격과 함께 구분할 수 있는지입니다. 큰 라벨 이력이 쌓이면 동일 원본 pool과 보존 예산에서 실제 행 선택과 합성 증류를 비교할 수 있습니다. 이번 슬라이드는 그 연구 순서를 설명하며 새 실험 또는 모델 변경을 확정하지 않습니다.',
         ['icd'])

def revise_user_pptx():
    """Preserve uploaded slide order and edits; edit that presentation in place."""
    prs=Presentation(USER_BASE)
    assert len(prs.slides)==11
    assert '<목차>' in prs.slides[1].shapes[0].text
    while len(prs.slides)>8:
        sid=prs.slides._sldIdLst[-1]
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)

    def replace(slide,old,new,size=None):
        matches=[s for s in slide.shapes if s.has_text_frame and s.text==old]
        assert len(matches)==1,(old,len(matches))
        sh=matches[0];f=sh.text_frame
        source=next(r for p in f.paragraphs for r in p.runs)
        pt=size or source.font.size.pt
        bold=bool(source.font.bold);color=str(source.font.color.rgb)
        align=f.paragraphs[0].alignment
        f.clear();f.word_wrap=False;f.auto_size=MSO_AUTO_SIZE.NONE
        for j,line in enumerate(new.splitlines()):
            p=f.paragraphs[0] if j==0 else f.add_paragraph()
            p.space_before=p.space_after=Pt(0);p.line_spacing=1.18;p.alignment=align
            r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(pt)
            r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
            for tag in ['a:ea','a:cs']:
                e=OxmlElement(tag);e.set('typeface','Pretendard');r._r.get_or_add_rPr().append(e)
        sh.height=Inches(len(new.splitlines())*pt*1.18/72+.035)
        return sh

    replace(prs.slides[0],
        '사전학습 자료:  합성 prior / 실제 표\n전이 방식:  ICL / 가중치 FT / 텍스트 생성 / 작은 모델의 가중치 생성',
        '사전학습 자료  합성 prior / 실제 표\n전이 방식  ICL / 가중치 FT / 텍스트 생성 /\n                 작은 모델의 가중치 생성')
    replace(prs.slides[3],
        '배정된 사례로 prototype(중심)과 prior(빈도) 갱신\n이전 값과 새 값을 이동평균으로 결합',
        'Prototype(클래스 중심)·prior(빈도) 갱신\n이전 값과 새 값을 이동평균으로 결합',19)
    replace(prs.slides[3],
        'Known과 novel이 뒤섞인 라벨 없는 데이터에서, 희귀한 novel이 묻히지 않도록 · 고정 TabPFN의 context 교체와는 다른 학습',
        'Known·novel 혼합 데이터에서 희귀 novel의 군집 학습 · encoder 가중치 갱신')

    # The uploaded page 2 is the user's editorial memo, now a presentation agenda.
    toc=prs.slides[1]
    for shape in list(toc.shapes):
        shape._element.getparent().remove(shape._element)
    items=[txt(.6,.32,12.08,'검토 논문과 활용 목적',30,True,NAVY),
        txt(.82,1.4,11.8,'1  IDS 평가 시나리오·context 운용',23,True,NAVY),
        txt(1.02,2.04,2.8,'LotBoNC',21,True,BLUE,url=SOURCES['lotbonc'][2]),
        txt(3.85,2.04,8.67,'라벨 없는 신규 공격·긴 꼬리 분포의 군집 학습\nKnown·novel 구성 참고 · 시간순 스트림 평가와 구분',19),
        txt(1.02,3.14,2.8,'Rethinking',21,True,BLUE,url=SOURCES['streams'][2]),
        txt(3.85,3.14,8.67,'예측 → 라벨 확보 → 메모리 갱신\n메모리 정책과 처리 비용의 분리',19),
        txt(1.02,4.24,2.8,'In-context distillation',18,True,BLUE,url=SOURCES['icd'][2]),
        txt(3.85,4.24,8.67,'가중치를 고정한 합성 context 최적화\n같은 확보 라벨에서 보관량을 줄이는 방법',19),
        txt(.82,5.51,11.8,'2  SOTA 비교 후보',23,True,NAVY),
        txt(1.02,6.10,11.6,'LoCalPFN·BoostPFN  지역 검색·오류 기반 결합\nDistPFN·불균형 보정  예측 확률·판정 기준 조절',19),
        box(0,7.04,W,.46,NAVY),txt(12.48,7.145,.28,'2',14,color='FFFFFF',align='right')]
    for it in items:
        x,y,w,h=[Inches(it[k]) for k in ('x','y','w','h')]
        if it['kind']=='box':
            sh=toc.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h)
            sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(it['fill']);sh.line.fill.background()
            continue
        sh=toc.shapes.add_textbox(x,y,w,h);f=sh.text_frame;f.clear()
        f.word_wrap=False;f.auto_size=MSO_AUTO_SIZE.NONE;f.vertical_anchor=MSO_ANCHOR.TOP
        f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
        for j,line in enumerate(it['text'].splitlines()):
            p=f.paragraphs[0] if j==0 else f.add_paragraph()
            p.space_before=p.space_after=Pt(0);p.line_spacing=1.18
            p.alignment={'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER}[it['align']]
            r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(it['size'])
            r.font.bold=it['bold'];r.font.color.rgb=RGBColor.from_string(it['color'])
            for tag in ['a:ea','a:cs']:
                e=OxmlElement(tag);e.set('typeface','Pretendard');r._r.get_or_add_rPr().append(e)
            if it.get('url'):r.hyperlink.address=it['url']
    if toc.notes_slide.notes_text_frame is None:
        body=next(sh for sh in prs.slides[0].notes_slide.shapes if sh.has_text_frame and sh.text_frame is not None and sh.text.startswith('발표 설명'))
        cloned=deepcopy(body._element)
        cloned.xpath('.//p:cNvPr')[0].set('id',str(max((sh.shape_id for sh in toc.notes_slide.shapes),default=1)+1))
        toc.notes_slide.shapes._spTree.insert_element_before(cloned,'p:extLst')
    toc.notes_slide.notes_text_frame.text=(
        '발표 설명\n먼저 세 논문을 서로 다른 목적으로 검토합니다. LotBoNC는 라벨 없는 신규 botnet과 '
        '긴 꼬리 분포를 다루는 군집 학습이며, 시간순 스트림 평가가 아닙니다. 라벨 공개 대상과 known/novel '
        '구성 방법을 참고합니다. Rethinking은 예측, 라벨 확보, 메모리 갱신을 분리하고 정확도와 처리 비용을 '
        '함께 비교하므로 순차 평가 설계에 가장 직접적입니다. ICD는 평가 시나리오를 정하는 논문이 아니라, '
        '이미 확보한 라벨 정보를 작은 합성 context에 담는 방법 후보입니다. 다음으로 LoCalPFN, BoostPFN, '
        'DistPFN과 불균형 보정을 비교 후보로 정리합니다. SOTA 후보라는 표기는 우리 IDS 설정에서 우위가 '
        '확인됐다는 뜻이 아닙니다.\n\n원문 근거\n'+
        '\n'.join(SOURCES[k][0]+'\n'+SOURCES[k][2] for k in ['lotbonc','streams','icd','local','boost','dist','imbalance']))

    for i,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text and shape.top/Inches(1)>7:
                if shape.left/Inches(1)>12:
                    replace(slide,shape.text,str(i))
                elif 10<shape.left/Inches(1)<12:
                    nav='모델 개요' if i==1 else ('01 평가 설계·context' if i<=6 else '02 SOTA 비교 후보')
                    replace(slide,shape.text,nav)
        for shape in slide.shapes:
            for e in shape._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):e.set('idx','0')
            sp=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
            if sp is not None:
                for e in list(sp):
                    if etree.QName(e).localname in ['effectLst','effectDag']:sp.remove(e)
                sp.append(OxmlElement('a:effectLst'))
    candidate=QA/'feedback_candidate.pptx';prs.save(candidate)
    with zipfile.ZipFile(candidate) as src,zipfile.ZipFile(OUT.with_suffix('.pptx'),'w',zipfile.ZIP_DEFLATED) as dst:
        for entry in src.infolist():
            data=src.read(entry.filename)
            if entry.filename.startswith('ppt/theme/') and entry.filename.endswith('.xml'):
                tree=etree.fromstring(data)
                for e in tree.xpath('//*[local-name()="outerShdw" or local-name()="innerShdw" or local-name()="prstShdw"]'):e.getparent().remove(e)
                data=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
            dst.writestr(entry,data)
    candidate.unlink()

def import_revised_scenes():
    """Render the actual edited PPTX, preserving its geometry and figure bytes."""
    original=json.loads((USER_BASE.parent/'scenes.json').read_text())
    keys={s['title']:s['source_keys'] for s in original}
    for si,slide in enumerate(Presentation(OUT.with_suffix('.pptx')).slides,1):
        items=[]
        for j,sh in enumerate(slide.shapes):
            coords={key:getattr(sh,attr)/914400 for key,attr in [('x','left'),('y','top'),('w','width'),('h','height')]}
            if sh.has_text_frame and sh.text:
                r=next(r for p in sh.text_frame.paragraphs for r in p.runs)
                assert r.font.size is not None
                align={PP_ALIGN.LEFT:'left',PP_ALIGN.RIGHT:'right',PP_ALIGN.CENTER:'center',None:'left'}[sh.text_frame.paragraphs[0].alignment]
                item=dict(kind='text',**coords,text=sh.text,size=r.font.size.pt,bold=bool(r.font.bold),color=str(r.font.color.rgb),align=align,url=r.hyperlink.address)
                items.append(item)
            elif sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
                assert not any([sh.crop_left,sh.crop_right,sh.crop_top,sh.crop_bottom]), 'Unexpected cropped picture'
                p=ASSETS/f'feedback_s{si}_{j}.{sh.image.ext}';p.write_bytes(sh.image.blob)
                items.append(dict(kind='image',**coords,path=str(p)))
            elif sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
                kind={MSO_SHAPE.RECTANGLE:'box',MSO_SHAPE.RIGHT_ARROW:'arrow',MSO_SHAPE.DOWN_ARROW:'down'}[sh.auto_shape_type]
                border=str(sh.line.color.rgb) if sh.line.fill.type and sh.line.color.type else None
                items.append(dict(kind=kind,**coords,fill=str(sh.fill.fore_color.rgb),border=border))
            else:raise ValueError((si,j,sh.shape_type))
        title=next(it['text'] for it in items if it['kind']=='text')
        SCENES.append(dict(title=title,source_keys=keys.get(title,['lotbonc','streams','icd','local','boost','dist','imbalance']),notes=slide.notes_slide.notes_text_frame.text,items=items))

def validate_scenes():
    issues=[]
    for page_no,scene in enumerate(SCENES,1):
        texts=[]
        for i,it in enumerate(scene['items']):
            x,y,w,h=[it[k] for k in ('x','y','w','h')]
            if min(x,y)<0 or x+w>W+.01 or y+h>H+.01:issues.append(dict(page=page_no,item=i,type='bounds',text=it.get('text','')))
            if it['kind']!='text':continue
            font=ImageFont.truetype(str(BOLD if it['bold'] else FONT),round(it['size']*4))
            width=max((font.getlength(s) for s in it['text'].splitlines()),default=0)/4/72
            if width>w-.015:issues.append(dict(page=page_no,item=i,type='width',width=width,available=w,text=it['text']))
            if re.search(r'합니다|입니다|습니다|—',it['text']):issues.append(dict(page=page_no,item=i,type='style',text=it['text']))
            if y<7.04 and y+h>6.98:issues.append(dict(page=page_no,item=i,type='footer_intrusion',text=it['text']))
            texts.append((i,it))
        for k,(ia,a) in enumerate(texts):
            for ib,b in texts[k+1:]:
                dx=min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
                dy=min(a['y']+a['h'],b['y']+b['h'])-max(a['y'],b['y'])
                if dx>.015 and dy>.015:issues.append(dict(page=page_no,type='text_overlap',items=[ia,ib],texts=[a['text'],b['text']]))
    result=dict(passed=not issues,issues=issues,slides=len(SCENES),intentional_overlaps=['Text and its background rectangle','Illustrations within figure areas'])
    (QA/'layout_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    if issues:raise RuntimeError(json.dumps(issues,ensure_ascii=False,indent=2))
    assert len(SCENES)==N_SLIDES

def pptx():
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    prs.core_properties.title='TabPFN 관련 연구 · 평가 시나리오 · Context 최적화'
    prs.core_properties.subject='20261021_references 보고서 1·3·4절 설명'
    for scene in SCENES:
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.notes_slide.notes_text_frame.text=scene['notes']
        for it in scene['items']:
            x,y,w,h=[Inches(it[k]) for k in ('x','y','w','h')]
            if it['kind']=='text':
                shape=slide.shapes.add_textbox(x,y,w,h);f=shape.text_frame;f.clear()
                f.word_wrap=False;f.auto_size=MSO_AUTO_SIZE.NONE;f.vertical_anchor=MSO_ANCHOR.TOP
                f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
                for i,line in enumerate(it['text'].splitlines()):
                    p=f.paragraphs[0] if i==0 else f.add_paragraph()
                    p.space_before=p.space_after=Pt(0);p.line_spacing=1.18
                    p.alignment={'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER}[it['align']]
                    r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(it['size']);r.font.bold=it['bold'];r.font.color.rgb=RGBColor.from_string(it['color'])
                    prop=r._r.get_or_add_rPr();prop.set('lang','ko-KR')
                    for tag in ['a:ea','a:cs']:
                        e=OxmlElement(tag);e.set('typeface','Pretendard');prop.append(e)
                    if it.get('url'):r.hyperlink.address=it['url']
            elif it['kind']=='image':slide.shapes.add_picture(it['path'],x,y,width=w,height=h)
            else:
                typ={'arrow':MSO_SHAPE.RIGHT_ARROW,'down':MSO_SHAPE.DOWN_ARROW}.get(it['kind'],MSO_SHAPE.RECTANGLE)
                s=slide.shapes.add_shape(typ,x,y,w,h);s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(it['fill']);s.line.fill.background()
                if it.get('border'):s.line.color.rgb=RGBColor.from_string(it['border']);s.line.width=Pt(.8)
            element=slide.shapes[-1]._element
            for effect in element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):effect.set('idx','0')
            sp=element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
            if sp is not None:
                for c in list(sp):
                    if etree.QName(c).localname in ['effectLst','effectDag']:sp.remove(c)
                sp.append(OxmlElement('a:effectLst'))
    candidate=QA/'candidate.pptx';prs.save(candidate)
    with zipfile.ZipFile(candidate) as src,zipfile.ZipFile(OUT.with_suffix('.pptx'),'w',zipfile.ZIP_DEFLATED) as dst:
        for entry in src.infolist():
            data=src.read(entry.filename)
            if entry.filename.startswith('ppt/theme/') and entry.filename.endswith('.xml'):
                tree=etree.fromstring(data)
                for sh in tree.xpath('//*[local-name()="outerShdw" or local-name()="innerShdw" or local-name()="prstShdw"]'):sh.getparent().remove(sh)
                data=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
            dst.writestr(entry,data)
    candidate.unlink()

def pdf():
    pages=[]
    for si,scene in enumerate(SCENES):
        tags=[]
        for i,it in enumerate(scene['items']):
            style=f'left:{it["x"]}in;top:{it["y"]}in;width:{it["w"]}in;height:{it["h"]}in;'
            if it['kind']=='text':
                style+=f'font-size:{it["size"]}pt;font-weight:{700 if it["bold"] else 400};color:#{it["color"]};text-align:{it["align"]};'
                content=html.escape(it['text'])
                if it.get('url'):content=f'<a href="{html.escape(it["url"],quote=True)}">{content}</a>'
                tags.append(f'<div id="s{si}i{i}" class="text" style="{style}">{content}</div>')
            elif it['kind']=='image':tags.append(f'<img class="figure" src="{Path(it["path"]).as_uri()}" style="{style}">')
            else:
                style+=f'background:#{it["fill"]};'
                if it.get('border'):style+=f'border:.8pt solid #{it["border"]};'
                if it['kind']=='arrow':style+='clip-path:polygon(0 25%,60% 25%,60% 0,100% 50%,60% 100%,60% 75%,0 75%);'
                if it['kind']=='down':style+='clip-path:polygon(25% 0,75% 0,75% 60%,100% 60%,50% 100%,0 60%,25% 60%);'
                tags.append(f'<div class="shape" style="{style}"></div>')
        pages.append('<section>'+''.join(tags)+'</section>')
    page=QA/'references_ko.html'
    page.write_text('<!doctype html><html lang="ko"><meta charset="utf-8"><style>'
       f'@font-face{{font-family:Pretendard;src:url("{FONT.as_uri()}");font-weight:400}}'
       f'@font-face{{font-family:Pretendard;src:url("{BOLD.as_uri()}");font-weight:700}}'
       f'@page{{size:{W}in {H}in;margin:0}}'
       '*{box-sizing:border-box}html,body{margin:0;padding:0}body{font-family:Pretendard,sans-serif;-webkit-print-color-adjust:exact;print-color-adjust:exact}'
       f'section{{position:relative;width:{W}in;height:{H}in;page-break-after:always;overflow:hidden}}'
       'section:last-child{page-break-after:auto}.text,.shape,.figure{position:absolute}.text{white-space:pre;line-height:1.18}a{color:inherit;text-decoration:none}'
       '</style>'+''.join(pages)+'</html>')
    candidate=QA/'rendered_ko.pdf'
    with (QA/'render.log').open('w') as log:
        subprocess.run(['/opt/google/chrome/chrome','--headless','--no-sandbox','--disable-dev-shm-usage','--disable-gpu','--no-pdf-header-footer','--allow-file-access-from-files','--run-all-compositor-stages-before-draw','--virtual-time-budget=5000',f'--print-to-pdf={candidate}',page.as_uri()],stdout=log,stderr=log,check=True,timeout=60)
    for p in QA.glob('page-*.png'):p.unlink()
    for cmd in [['pdftotext','-layout',str(candidate),str(QA/'rendered_text.txt')],['pdftoppm','-scale-to','1600','-png',str(candidate),str(QA/'page')]]:
        r=subprocess.run(cmd,check=True,capture_output=True,text=True);assert not r.stderr.strip(),r.stderr
    paths=sorted(QA.glob('page-*.png'));assert len(paths)==N_SLIDES
    text=(QA/'rendered_text.txt').read_text().split('\f');checks=[]
    for i,path in enumerate(paths):
        nonwhite=float(np.mean(np.any(np.asarray(Image.open(path).convert('RGB'))<240,axis=2)))
        assert nonwhite>.05 and len(text[i].strip())>150,(i,nonwhite)
        checks.append(dict(page=i+1,nonwhite_fraction=nonwhite,text_characters=len(text[i])))
    shutil.copyfile(candidate,OUT.with_suffix('.pdf'))
    for k in range(2):
        subset=paths[k*5:(k+1)*5];canvas=Image.new('RGB',(1500,3*445),'#DCE3EB')
        for j,p in enumerate(subset):
            im=Image.open(p);im.thumbnail((730,411));x=(j%2)*750+10;y=(j//2)*445+8
            canvas.paste(im,(x,y));ImageDraw.Draw(canvas).text((x,y+414),str(k*5+j+1),fill='black')
        canvas.save(QA/f'contact-{k+1}.png')
    return checks

def main():
    QA.mkdir(parents=True,exist_ok=True);ASSETS.mkdir(parents=True,exist_ok=True)
    if USER_BASE.exists():
        revise_user_pptx();import_revised_scenes();validate_scenes()
    else:
        assets();scenes();validate_scenes();pptx()
    checks=pdf()
    with zipfile.ZipFile(OUT.with_suffix('.pptx')) as z:
        shadows=0
        for name in z.namelist():
            if name.endswith('.xml'):
                tree=etree.fromstring(z.read(name));shadows+=len(tree.xpath('//*[local-name()="outerShdw" or local-name()="innerShdw" or local-name()="prstShdw"]'))
        assert shadows==0
        assert len([n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)])==N_SLIDES
    (QA/'scenes.json').write_text(json.dumps(SCENES,ensure_ascii=False,indent=2)+'\n')
    manifest=dict(date='2026-10-23',source_report=str(REPORT.relative_to(ROOT)),source_report_sha256=sha(REPORT),report_sections=[1,3,4],slides=N_SLIDES,sources=SOURCES,source_figure_manifest='docs/lablog/evidence/references/20261021/report_figure_sources_20261022.json',scope='Korean explanatory slides only; manuscript deferred by user',editable='All text and explanatory diagrams editable; original figures are attributed images')
    if USER_BASE.exists():
        manifest.update(source_user_pptx=str(USER_BASE.relative_to(ROOT)),source_user_pptx_sha256=sha(USER_BASE),slide_groups={'overview':[1],'agenda':[2],'lotbonc':[3,4],'streams':[5],'icd':[6],'sota_candidates':[7,8]},feedback_resolution=['Edited the uploaded PPTX directly and retained its first eight slides in user order','Converted page 2 memo to an agenda with separate purposes for LotBoNC, Streams, ICD and comparison methods','Preserved user explanations of prototype/prior and rare novel attacks; adjusted wrapping and footer wording','Removed uploaded pages 9–11 from presentation; feedback source preserved in archive','Renumbered pages and navigation; retained source links and speaker notes'])
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    (QA/'validation.json').write_text(json.dumps(dict(passed=True,pages=checks,pptx_shadow_count=shadows,speaker_notes=N_SLIDES,pdf_sha256=sha(OUT.with_suffix('.pdf')),pptx_sha256=sha(OUT.with_suffix('.pptx')),visual_review_pending=True),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(pptx=str(OUT.with_suffix('.pptx')),pdf=str(OUT.with_suffix('.pdf')),slides=N_SLIDES),ensure_ascii=False))

if __name__=='__main__':main()
