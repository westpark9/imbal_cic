"""Editable Korean PCAP research slides and matching PDF, with source/geometry QA."""
from pathlib import Path
import csv
import hashlib
import json
import sys

ROOT=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/experiment_paths.py').exists())
sys.path.insert(0,str(ROOT/'scripts/common'))
from experiment_paths import repo_root
import build_20261023_references as r

HERE=ROOT/'docs/slides'
QA=HERE/'_support/qa/20261029_pcap'
OUT=HERE/'kor/20261029_PCAP_연구방향'
EVIDENCE=ROOT/'docs/lablog/evidence/20261029_pcap'
RUN=ROOT/'tabpfn/results/v7/exp92/20261029_210839_cic2018_verified_pcap_4d62ceac'
SOURCE_DATA=json.loads((EVIDENCE/'sources.json').read_text())['sources']
r.QA=QA;r.ASSETS=HERE/'_support/assets/20261029_pcap';r.OUT=OUT;r.N_SLIDES=11;r.SCENES=[]
r.SOURCES={k:(v['title'],f"{v['venue']}, {v['year']}",v['url'],v['evidence']) for k,v in SOURCE_DATA.items()}
txt,box,arrow,block,page=r.txt,r.box,r.arrow,r.block,r.page
NAVY,BLUE,INK,GREY,PALE=r.NAVY,r.BLUE,r.INK,r.GREY,r.PALE
RED='B44948'

def card(x,y,w,h,title,body,size=21):
    return [box(x,y,w,h,PALE),box(x,y,w,.48,'E1ECF5'),txt(x+.16,y+.09,w-.32,title,21,True,NAVY),txt(x+.18,y+.72,w-.36,body,size)]

def table(x,y,widths,headers,rows,height=.55,size=20,colors=None):
    items=[]
    for ri,row in enumerate([headers]+rows):
        xx=x
        for ci,(w,value) in enumerate(zip(widths,row)):
            fill=NAVY if ri==0 else ('FFFFFF' if ri%2 else PALE)
            items.append(box(xx,y+ri*height,w,height,fill))
            color='FFFFFF' if ri==0 else (colors.get(ci,INK) if colors else INK)
            align='left' if ci==0 else 'right'
            items.append(txt(xx+.12,y+ri*height+.13,w-.24,str(value),size,ri==0,color,align))
            xx+=w
    return items

def scenes():
    # 1. User-requested left upper process, left lower feature examples, right text.
    a=[txt(.65,1.15,7.55,'실험망의 정상·공격 트래픽 수집',23,True,NAVY),
       txt(8.65,1.15,4.0,'CIC 데이터의 구성',24,True,NAVY)]
    for x,w,title,body in [(0.65,2.0,'실험 네트워크','사용자 활동\n공격 시나리오'),(3.28,2.0,'패킷 캡처','원시 패킷\nPCAP 파일'),(5.91,2.0,'Flow 구성','동일 통신의\n양방향 패킷')]:
        a+=block(x,1.84,w,1.38,title,body,body_size=19)
    a += [arrow(2.85,2.42,.28),arrow(5.48,2.42,.28),arrow(6.75,3.24,.25,.22,down=True),
          txt(.65,3.52,7.55,'CICFlowMeter의 통계 특징 추출',23,True,NAVY)]
    a+=table(.65,4.1,[3.35,4.0],['특징 예시','의미'],[
        ['Flow Duration','첫 패킷부터 마지막 패킷까지의 시간'],
        ['Total Fwd Packets','순방향 패킷 수'],
        ['Fwd Packet Length Mean','순방향 패킷 길이의 평균']],height=.53,size=16)
    a += [txt(.76,6.45,7.3,'특징 벡터 x = [지속시간, 패킷 수, 평균 길이, …] → CSV',17,True,BLUE),
          txt(8.65,1.86,4.05,'CIC2017·CIC2018\n정상 활동과 여러 공격을 포함한\n침입탐지 평가용 트래픽',21),
          txt(8.65,3.18,4.05,'PCAP + 특징 CSV + 라벨\n원시 패킷과 통계 기반 표현의\n추적 가능한 대응 관계',21),
          txt(8.65,4.51,4.05,'라벨 부여 근거\n공격 시간·호스트·포트·프로토콜\n실제 flow 내용과의 대응 검증 필요',20),
          txt(8.65,6.08,4.05,'Intrusion Detection System (IDS)\nPacket Capture (PCAP)',15,color=GREY)]
    page('CIC 데이터의 제작 과정',a,'원시 패킷을 통계 벡터로 요약하고 공격 시나리오의 라벨을 연결','01 데이터와 문제',
         'IDS 장비가 데이터셋을 직접 만든다고 표현하지 않는다. IDS 평가를 위한 실험망의 트래픽을 캡처했다. CIC2017 공식 설명은 미러 포트를 통한 캡처와 시간·주소·포트·프로토콜을 이용한 라벨을 설명했다. 왼쪽 아래 벡터는 형식 예시이며 측정값은 아니다. CSV는 저장 형식이고 모델의 정보량은 어떤 특징을 추출했는지에 좌우된다.', ['cic17','cic18'])

    with (ROOT/'docs/lablog/data/20261027_exp89_class_comparison.csv').open() as f:
        old=[x for x in csv.DictReader(f) if x['dataset']=='cic2018' and x['panel']=='1' and x['n'] in ('2','20')]
    assert all(x['train_version']=='original' and x['test_version']=='original' and x['benign_training']=='False' and x['test_scope']=='attack' for x in old)
    names=sorted({x['name'] for x in old},key=lambda c:-int(next(x['support'] for x in old if x['name']==c)))
    pretty={'ddos':'DDoS','brute_force':'Brute force','dos':'DoS','bot':'Bot','infiltration':'Infiltration','web_attacks':'Web attacks'}
    vals=[]
    for c in names:
        cells=[pretty[c],f"{int(next(x['support'] for x in old if x['name']==c)):,}"]
        for n,m in [('2','xgb'),('2','pfn'),('20','xgb'),('20','pfn')]:
            cells.append(f"{float(next(x['mean_percent'] for x in old if x['name']==c and x['n']==n and x['method']==m)):.2f}")
        vals.append(cells)
    a=[txt(.65,1.12,12.0,'기존 CIC2018 실험 · 정제 전·정상 미포함 · 클래스당 n개 · 5개 seed 평균 F1 (%)',19,color=GREY)]
    a+=table(.65,1.83,[2.25,1.45,2.08,2.08,2.08,2.08],['평가 클래스','평가 수','2개 XGB','2개 TabPFN','20개 XGB','20개 TabPFN'],vals,height=.57,size=18,colors={2:RED,3:BLUE,4:RED,5:BLUE})
    a += [txt(.65,6.10,12.0,'DDoS·Bot은 높은 점수에 조기 도달 · Web attacks에는 여전히 큰 분류 난도',22,True,NAVY),
          txt(.65,6.57,12.0,'평가 표본 수 내림차순 · XGBoost (XGB) · 중간 n = 3·5·10과 표준편차는 EXP89 상세 기록',15,color=GREY)]
    page('일부 공격의 높은 few-shot F1',a,'데이터 전체의 난도보다 클래스별 포화·오류 원인의 분리가 우선','01 데이터와 문제',
         'EXP89에서 가져온 과거 관측이며 PCAP 재구축 후 새로 측정한 점수가 아니다. 선행 SOTA 비교의 문제의식을 소수 라벨 baseline으로 진단한 결과다. 모든 클래스가 90%대라는 해석은 틀렸다. 현재 표는 2개와 20개 끝점이며 5개 seed 평균이고, 소표본 변동도 원 보고서에서 함께 확인해야 한다. TabPFN: Tabular Prior-data Fitted Network. DDoS: Distributed Denial of Service. DoS: Denial of Service. 전체 원시 행과 표준편차: '+json.dumps(old,ensure_ascii=False),[])

    a=[]
    papers=[('engelen','Engelen 외 · IEEE SPW / WTMC 2021','첫 FIN에서 조기 종료 → 종료 패킷만 남은 별도 flow\n방향 오류·RST 무시 → 잘못된 flow 경계와 라벨',1.22),
            ('liu','Liu 외 · IEEE CNS 2022','CIC2017·2018의 생성·특징·라벨 오류 재검토\nCIC2018 FTP 사례: 닫힌 포트에 대한 연결 시도',3.07),
            ('rosay','Rosay 외 · ICISSP 2022','특징 계산·세션 종료·라벨 문제 분석\n세션 종료 오류와 유사 특징·상이 라벨의 관계',4.92)]
    short_titles={
      'engelen':'Troubleshooting an Intrusion Detection Dataset:\nthe CICIDS2017 Case Study',
      'liu':'Error Prevalence in NIDS datasets: A Case Study on\nCIC-IDS-2017 and CSE-CIC-IDS-2018',
      'rosay':'Network Intrusion Detection: A Comprehensive\nAnalysis of CIC-IDS2017'}
    for k,head,body,y in papers:
        a += [box(.65,y,12.03,1.58,PALE),txt(.84,y+.12,5.30,head,21,True,NAVY,url=SOURCE_DATA[k]['url']),
              txt(.84,y+.62,5.28,short_titles[k],14,color=GREY),txt(6.56,y+.32,5.83,body,20)]
    page('CIC 결함을 확인한 선행연구',a,'단순한 시간 구간 설명을 넘어 flow 생성과 실제 공격 내용까지 확인한 근거','01 데이터와 문제',
         'Transmission Control Protocol (TCP), Finish (FIN), Reset (RST). Engelen은 첫 FIN 뒤 ACK·FIN 잔여 패킷이 별도 flow로 분리되는 문제를 분석했다. Liu의 공식 후속 설명은 FTP 포트가 닫혀 SYN에 RST/ACK만 응답한 경우를 Attempted로 구분했다. Rosay가 말한 것은 similar flows이며, 모든 입력 열의 exact equality를 증명했다는 의미가 아니다. 원인 근거는 2021·2022 연구가 가장 직접적이고, 2023 수정 라벨과 2024 일반화 연구를 함께 확인했다.', ['engelen','liu','rosay'])

    a=card(.65,1.3,5.78,2.10,'① Flow 경계·라벨의 오류','패킷 분리·결합 오류와 방향 오류\n공격 시간 안의 비공격 연결까지 포함',22)
    a+=card(6.83,1.3,5.78,2.10,'② 관측 정보의 부족','다른 패킷 내용이 같은 통계값으로 요약\n세부 공격 의도를 단일 flow에서 구별 곤란',22)
    a += [arrow(3.25,3.62,.35,.42,down=True),arrow(9.43,3.62,.35,.42,down=True)]
    a+=card(.65,4.27,5.78,1.65,'대응: 라벨과 세션 재검증','수정된 경계·시간·공격 근거의 적용\n불확실 사례는 근거와 함께 별도 보존',20)
    a+=card(6.83,4.27,5.78,1.65,'대응: 패킷 기반 표현의 비교','패킷 순서·시간·헤더·내용의 추가 활용\n동일 flow에서 기존 특징과 임베딩 비교',20)
    a += [txt(.72,6.35,11.9,'현재 충돌 사례의 원인을 모두 확정한 상태는 아님 · PCAP과 라벨 근거의 직접 대조가 다음 검증',18,True,NAVY)]
    page('라벨 오류와 표현 정보 부족의 구분',a,'표현 변경과 라벨 교정을 별도 요인으로 평가','02 재구축 설계',
         '동일한 x에 다른 y가 있으면 단일 flow만 보는 결정적 분류기는 둘을 동시에 맞힐 수 없다. 이것이 항상 잘못된 라벨이라는 뜻은 아니다. IP·시각을 제외하거나 통계로 요약하면 다른 통신이 같은 관측 벡터가 될 수 있다. 반대로 실제 공격 내용이 없는 연결에 공격 라벨을 주었을 수도 있다. 새 임베딩은 추가 정보를 담지만 잘못된 라벨을 자동으로 고치지는 않는다. 현재 기존 입력의 모든 충돌을 CICFlowMeter 버그 하나로 설명하지 않는다.',[])

    a=[txt(.65,1.18,12,'오프라인 PCAP 처리 · 동일 통신 단위의 추적 정보 보존',23,True,NAVY)]
    blocks=[(.65,2.0,'PCAP','시각·패킷\n원본 식별자'),(3.15,2.0,'Flow·라벨','세션 경계\n라벨 근거'),(5.65,2.0,'분할','학습·검증·시험\n통신 그룹 보존'),(8.15,2.0,'고정 인코더','netFound v2\n512차원'),(10.65,2.0,'PCA·벡터','학습 자료로 적합\n64차원 출력')]
    for x,w,title,body in blocks:a+=block(x,1.98,w,1.58,title,body,body_size=17)
    for x in [2.80,5.30,7.80,10.30]:a.append(arrow(x,2.61,.23))
    a+=card(.65,4.13,3.77,1.57,'라벨 근거 연결','시간·호스트·포트 + 내용\n경계 충돌·미매칭 별도 표시',18)
    a+=card(4.78,4.13,3.77,1.57,'통신 그룹별 분리','동일 통신의 분할 간 중복 방지\n인코더 가중치 갱신 없음',18)
    a+=card(8.91,4.13,3.77,1.57,'학습 자료만으로 PCA','검증·시험은 같은 PCA로 변환\nTabPFN·XGB에 같은 행 입력',18)
    a += [
          txt(.72,6.14,12,'Principal Component Analysis (PCA) · 차원 축소 기준은 학습 분할에서만 결정',19,color=GREY),
          txt(.72,6.58,12,'재전송 없이 패킷 파일에서 추출 · 주소·절대시각·라벨은 모델 입력에서 제외',18,True,BLUE)]
    page('PCAP 기반 학습 벡터 생성',a,'라벨·통신 식별자와 모델 입력을 분리해 재현 가능한 벡터 구성','02 재구축 설계',
         '이 도식은 표현 추출과 데이터 준비이며 아직 제안 분류 모델의 성능 개선 결과는 아니다. 인코더는 frozen이고 fine-tuning하지 않았다. 모델 입력과 별개로 PCAP SHA256, packet index, 5-tuple, UTC nanoseconds, label rule IDs를 보존한다. group hash로 먼저 분리하고 이후 표본을 선택한다. PCA는 학습 구간의 미라벨 자료까지 포함하는 train-pool 조건이다. 엄격한 n개 전체 입력 제약을 실험하려면 PCA 적합 자료도 별도로 제한해야 한다. tcpreplay는 패킷 재전송 도구이며 오프라인 임베딩 추출에는 필요하지 않다.', ['netfound'])

    a=[]
    for x,title,venue,body,status,key in [
        (.65,'netFound v2','arXiv v5 · 2026','헤더 필드·시간·방향\nBurst와 flow의 계층 표현\n공개 가중치는 payload 미사용','구현·실제 PCAP 검증 완료','netfound'),
        (4.78,'ET-BERT','WWW · 2022','Datagram의 바이트 표현\nTransformer 기반 사전학습\n패킷 내용 정보의 활용 후보','다음 인코더 비교 후보','etbert'),
        (8.91,'NetMamba','ICNP · 2024','헤더·payload 바이트 시퀀스\nMamba 기반 순차 처리\n속도·표현 효과의 비교 후보','다음 인코더 비교 후보','netmamba')]:
        a += [box(x,1.48,3.77,4.76,PALE),box(x,1.48,3.77,.70,NAVY),txt(x+.18,1.65,3.41,title,25,True,'FFFFFF'),
              txt(x+.18,2.43,3.41,venue,19,True,BLUE,url=SOURCE_DATA[key]['url']),
              txt(x+.18,3.11,3.41,body,20),txt(x+.18,5.60,3.41,status,16,True,NAVY)]
    a += [txt(.72,6.58,11.9,'내용 정보가 필요한 공격의 구분에는 payload 포함 인코더와의 후속 대조 필요',19,True,NAVY)]
    page('사전학습 인코더 후보',a,'netFound로 공통 파이프라인 검증 후 인코더의 관측 정보와 비용 비교','02 재구축 설계',
         'netFound는 원 논문 v1이 아니라 2026 v5 논문에 대응하는 v2 코드/Small checkpoint를 사용했다. 공개 가중치는 IP와 payload를 사용하지 않는다. ET-BERT는 Encrypted Traffic Bidirectional Encoder Representations from Transformer로, datagram 바이트를 사전학습했다. NetMamba는 ICNP 2024 모델을 가리키며 NetMamba+ 2026과 구분한다. 후보의 논문 성능을 현재 IDS 입력에서의 우위로 옮겨 쓰지 않는다. 세 모델을 함께 앙상블했다는 뜻도 아니다. ET-BERT·NetMamba 추론은 이번 구현 검증 범위에 포함하지 않았다.', ['etbert','netmamba'])

    a=[txt(.65,1.13,12.0,'PCAP 제공 여부와 재구축한 flow에 연결할 라벨 근거를 함께 검토',21,color=GREY)]
    rows=[['CIC2017 · 우선','PCAP + flow CSV + 공격 일정','교정 라벨 규칙과 패킷 경계 대조'],
          ['CIC2018 · 우선','호스트별 PCAP + 로그 + CSV','동일 통신의 중복 캡처·시간 확인'],
          ['IoT-23 · 다음 후보','PCAP + 라벨된 Zeek conn.log','개별 시나리오부터 PCAP·라벨 연결'],
          ['UNSW-NB15','PCAP + GT·공격 이벤트 표','시간·5-tuple로 ground truth 연결'],
          ['TON_IoT','네트워크 PCAP + Zeek·CSV','운영체제·센서 자료와 네트워크 구별'],
          ['CICIoT2023','공격별 PCAP + CSV·처리 코드','캡처명과 flow 라벨의 대응 검증']]
    # Text columns are left aligned; dataset names link to their official sources.
    xs=[.65,3.45,7.64];ws=[2.8,4.19,5.04]
    for ri,row in enumerate([['데이터','공개 원시 자료','선정·전처리 기준']]+rows):
        for ci,value in enumerate(row):
            y=1.82+ri*.62
            a += [box(xs[ci],y,ws[ci],.62,NAVY if ri==0 else PALE if ri%2==0 else 'FFFFFF'),
                  txt(xs[ci]+.12,y+.17,ws[ci]-.24,value,17,ri==0,'FFFFFF' if ri==0 else INK,
                      url=SOURCE_DATA[['cic17','cic18','iot23','unsw','ton','iot2023'][ri-1]]['url'] if ri and ci==0 else None)]
    a += [txt(.72,6.45,11.9,'CTU-13은 전체 정상·배경 PCAP 미공개 · 정상 포함 분류의 우선 비교에서는 후순위',18,True,NAVY)]
    page('PCAP·라벨을 함께 제공하는 데이터',a,'CIC2017·2018 우선 · 작은 시나리오 단위로 확장 가능한 IoT-23 후속 검토','02 재구축 설계',
         '추가 후보는 이름만 나열하지 않고 원시 패킷과 라벨을 다시 연결할 수 있는지로 선정했다. IoT-23 v1은 20 malware와 3 benign 시나리오이며 개별 다운로드를 제공했다. UNSW-NB15는 100GB PCAP과 GT/이벤트 표를 공개했다. TON_IoT는 네트워크 PCAP을 포함하지만 센서·OS 모달리티와 구분해야 한다. CICIoT2023은 33 공격·7종류 분류와 PCAP/CSV/처리 코드가 있으나 flow별 라벨을 일괄 추정하지 않는다. CTU-13 전체 혼합 PCAP은 공개되지 않았고 botnet-only PCAP만 제공했다. 이는 모든 후보의 로컬 다운로드와 라벨 join을 완료했다는 의미가 아니다.', ['iot23','unsw','ton','iot2023','ctu13'])

    a=[txt(.65,1.15,12,'동일 flow·동일 분할·동일 라벨 예산으로 표현과 라벨 효과 분리',22,True,NAVY),
       txt(2.95,1.92,4.23,'통계 특징',23,True,NAVY,'center'),txt(7.69,1.92,4.23,'사전학습 임베딩',23,True,NAVY,'center'),
       txt(.72,2.82,2.0,'기존 라벨',21,True,NAVY),txt(.72,4.13,2.0,'교정 라벨',21,True,NAVY)]
    a+=block(2.90,2.49,4.27,1.03,'A · 기준','기존 표현 + 기존 라벨',body_size=19)
    a+=block(7.62,2.49,4.27,1.03,'B · 표현 효과','새 표현 + 기존 라벨',body_size=19)
    a+=block(2.90,3.91,4.27,1.03,'C · 라벨 효과','기존 표현 + 교정 라벨',body_size=19)
    a+=block(7.62,3.91,4.27,1.03,'D · 결합 효과','새 표현 + 교정 라벨',body_size=19)
    a += [txt(.72,5.44,12,'정상 포함을 주 평가로 설정 · 공격만 평가하는 결과를 보조로 병기',22,True,NAVY),
          txt(.72,6.01,12,'n = 2·3·5·10·20 · 동일 seed와 test · 클래스별·macro F1 · 정상 오탐률',20),
          txt(.72,6.48,12,'희귀 공격·날짜/호스트 간 일반화 · 라벨 미매칭률 · 충돌률 · 추출·추론 비용',19,color=GREY)]
    page('효과를 분리하는 비교 실험',a,'목표는 더 어려운 데이터 제작이 아니라 소수 라벨 공격 분류의 개선 검증','03 검증 계획',
         '미실행 실험 설계. A–D 모두 PCAP에서 추적 가능한 공통 flow 집합과 공통 분할을 사용한다. 기존 보유 표 전체와 새 임베딩 일부를 점수만 비교하지 않는다. 동일 기준 라벨로 평가하거나 원라벨/교정라벨 두 평가를 함께 제공해 라벨 변경 자체의 점수 효과를 구분한다. 정상과 공격 각각 n개를 적용하는 기본 시나리오를 유지하고 PCA에 사용할 미라벨 train pool을 명시한다. n개 context와 미라벨 표현 적합 풀의 규모는 별개다. 캡처·호스트 분리는 추가 일반화 시험이며 단일 날짜 pilot은 이를 검증하지 못했다. Attempted는 감사용 주석으로 남기고 정상/공격 매핑은 목표에 따라 민감도 분석한다.', ['attempted'])

    prep=json.loads((RUN/'prepare.json').read_text());projection=json.loads((RUN/'projection.json').read_text())
    audit=json.loads((RUN/'audit.json').read_text());stat=next(iter(prep['raw_stats'].values()))
    a=card(.65,1.35,5.76,3.14,'CIC2018 실제 PCAP 파일럿',
           f"2018-02-14 · 공격 대상 호스트 1대\n패킷 {stat['raw_packets']:,}개 → flow {stat['flows']:,}개\n주석·분할별 표본 {projection['rows']:,}개 인코딩\n512차원 → PCA 64차원",22)
    a+=card(6.83,1.35,5.76,3.14,'검증한 항목',
            f"학습·검증·시험 flow 그룹 중복 0\nPCA 적합: 학습 {projection['pca_fit_rows']:,}행만 사용\n가중치 갱신 없이 공개 checkpoint 적용\n배치 길이 변경에도 같은 임베딩 확인",21)
    a += [txt(.72,4.94,11.9,'라벨 처리: Attempted 세부 유형은 메타데이터 · 이번 파일럿은 정상으로 매핑',21,True,NAVY),
          txt(.72,5.50,11.9,f"라벨 미확정 {audit['unresolved_rows']:,}행은 정상으로 추정하지 않고 별도 보존",21),
          txt(.72,6.04,11.9,'분류 성능·정상 오탐률은 후속 평가 · CIC2017과 다른 날짜·공격으로 확장',21),
          txt(.72,6.55,11.9,'초기 주석별 점검: 같은 임베딩의 서로 다른 Attempted 유형 확인 → 학습 라벨과 분리',16,color=GREY)]
    page('구현 결과와 다음 검증',a,'PCAP부터 학습용 벡터까지 실행 완료 · 분류 이점은 다음 비교 실험의 대상','03 검증 계획',
         f"수행 실행: {RUN.relative_to(ROOT)}. 데이터 크기 852,795,392 bytes. 원본 ZIP 39,913,353,098 bytes 전체 대신 호스트 member만 Range 다운로드했다. 원본 파일 끝의 불완전한 마지막 패킷 1개는 제외했고 관련 미완료 flow는 미라벨로 보존했다. 주석별·분할별 최대128개 reservoir의 구현 검증 표본이며 자연 분포 benchmark가 아니다. {audit['labeled_rows']}행은 학습 라벨을 가지고 나머지는 미매칭/경계 미확정으로 유지했다. 초기 감사에서는 두 Attempted 주석이 한 임베딩으로 겹쳤다. 교정 데이터 저자 지침에 따라 Attempted를 독립 ML 클래스로 쓰지 않고 BENIGN으로 매핑한 현재 산출물을 사용한다. 이 정상은 저자의 라벨 정책에 따른 것이며 충분한 일반 정상 활동을 대신한다고 주장하지 않는다. 현재 정상 오탐률·few-shot 우위를 평가하지 않았다.",['aws','rules'])

    a=[]
    a+=card(.65,1.35,12.0,2.10,'동일 특징·상이 라벨의 직접 관측 · Xie 외, Computers & Security 2022',
            'UNSW-NB15·CCCS-CIC-AndMal-2020에서 특징 중첩을 분석\n동일 관측에 여러 의미가 연결되는 현상의 근거 · CIC 원인의 직접 증거와 구별',21)
    a += [txt(.84,3.13,11.50,'Analysis and Detection against Network Attacks in the Overlapping Phenomenon of Behavior Attribute',14,color=GREY,url=SOURCE_DATA['xie']['url'])]
    a+=card(.65,3.95,12.0,2.10,'동일 데이터의 높은 점수와 일반화의 차이 · Cantone 외, IEEE Access 2024',
            'CIC2017·2018과 교정 자료에서 데이터 간 전이 성능을 비교\n동일 데이터의 높은 점수만으로 새로운 환경의 탐지 성능을 판단하기 어려움',21)
    a += [txt(.84,5.74,11.50,'Machine Learning in Network Intrusion Detection: A Cross-Dataset Generalization Study',15,color=GREY,url=SOURCE_DATA['cantone']['url']),
          txt(.72,6.52,11.9,'선정 문헌 5편: CIC 원인 분석 3편 + 동일 특징 현상 1편 + 일반화 검증 1편',18,True,NAVY)]
    page('보충 문헌과 근거의 적용 범위',a,'원인 근거·현상 근거·일반화 근거를 구별해 연구 주장에 연결','04 보충 근거',
         'Xie 논문은 arXiv 제출 연도 2023보다 앞선 Computers & Security 121 (2022) 102867 게재 논문이다. Cantone은 2024년 초 arXiv의 이전 제목 대신 최종 IEEE Access 제목과 DOI로 표기했다. 그 연구는 동일 데이터 내 높은 성능과 데이터 간 전이 실패를 구별했으며 우리 새 표현의 우위를 입증한 것은 아니다. 슬라이드3의 세 직접 근거와 합쳐 다섯 문헌을 선정했다.', ['xie','cantone'])

    # The historical SOTA comparison has different preparation and label budgets.
    # Keep it as an explicitly scoped appendix, not merged into the few-shot table.
    report=(ROOT/'docs/lablog/reports/20260904.tex').read_text()
    section=report.split(r'\subsection*{CIC2018 · 클래스별 F1 · test 3,042,473행}',1)[1].split(r'\bottomrule',1)[0]
    sota=[]
    for line in section.splitlines():
        if not line.startswith('{}'):continue
        fields=[v.strip() for v in line[2:].split('&')]
        sota.append([fields[0],fields[1],*[f'{float(fields[i])*100:.2f}' for i in [2,4,5,6]]])
    assert len(sota)==7 and sum(int(x[1].replace(',','')) for x in sota)==3042473
    assert [int(x[1].replace(',','')) for x in sota]==sorted([int(x[1].replace(',','')) for x in sota],reverse=True)
    a=[txt(.65,1.13,12.0,'CIC2018 · 정제 후·정상 포함 · 공통 학습 100,000행 · 전체 test 3,042,473행 · F1 (%)',18,color=GREY)]
    a+=table(.65,1.84,[2.1,1.6,2.09,2.08,2.08,2.08],['평가 클래스','평가 수','TabPFN v3','BoostPFN','LoCalPFN','DistPFN v3'],sota,height=.53,size=17)
    a += [txt(.72,6.27,11.9,'주요 공격의 높은 F1과 Infiltration·Web의 낮은 F1이 공존',22,True,NAVY),
          txt(.72,6.69,11.9,'기존 보고서 발췌 · BoostPFN·LoCalPFN은 v1 기반 · LoCalPFN은 validation을 이용한 FT 포함',14,color=GREY)]
    page('보충 결과: 기존 SOTA 비교',a,'소수 라벨 실험과 설정을 구별하고 클래스별 차이를 진단','04 보충 근거',
         '기존 HTML에서 통합된 docs/lablog/reports/20260904.tex의 공통100k·전체test 표를 발췌했다. 초기 18만 smoke 결과가 아니다. EXP59 Global C0와 같은 학습 ID를 사용한 비교이며, BoostPFN·LoCalPFN의 v1과 Global·DistPFN의 v3 backbone, LoCalPFN의 validation FT 차이가 남는다. 이는 정제 전·정상 미포함·n개·5seed인 본문 slide2와 다른 실험으로 수치를 직접 합치지 않는다. 모든 공격이 높은 성능이라는 주장을 피하기 위해 정상과 전체 공격 클래스를 유지했다. 추출 원문 구간:\n'+section,[])

def main():
    QA.mkdir(parents=True,exist_ok=True);r.ASSETS.mkdir(parents=True,exist_ok=True)
    scenes();r.validate_scenes();r.pptx()
    from pptx import Presentation
    prs=Presentation(OUT.with_suffix('.pptx'))
    prs.core_properties.title='CIC 데이터 결함 검토와 PCAP 기반 학습 벡터 재구축'
    prs.core_properties.subject='문헌 근거, 인코더·데이터 후보, EXP92 구현 검증'
    prs.save(OUT.with_suffix('.pptx'))
    checks=r.pdf()
    from PIL import Image,ImageDraw
    paths=sorted(QA.glob('page-*.png'))
    for start in range(0,len(paths),5):
        canvas=Image.new('RGB',(1500,1335),'#DCE3EB')
        for j,p in enumerate(paths[start:start+5]):
            im=Image.open(p);im.thumbnail((730,411));x=(j%2)*750+10;y=(j//2)*445+8
            canvas.paste(im,(x,y));ImageDraw.Draw(canvas).text((x,y+414),str(start+j+1),fill='black')
        canvas.save(QA/f'contact-{start//5+1}.png')
    (QA/'scenes.json').write_text(json.dumps(r.SCENES,ensure_ascii=False,indent=2)+'\n')
    visible='\n'.join(it.get('text','') for scene in r.SCENES for it in scene['items'])
    assert 'NF-v3' not in visible and 'NF-CIC' not in visible
    manifest=dict(date='2026-10-29',slides=len(r.SCENES),source_evidence=str(EVIDENCE.relative_to(ROOT)),
                  run=str(RUN.relative_to(ROOT)),old_results='EXP89 uncurated attack-only, five seeds; not a new PCAP benchmark',
                  validation='shared scenes for editable PPTX and Chrome PDF; geometry/font/text checks',
                  visual_review_pending=True,pages=checks,
                  pptx_sha256=hashlib.sha256(OUT.with_suffix('.pptx').read_bytes()).hexdigest(),
                  pdf_sha256=hashlib.sha256(OUT.with_suffix('.pdf').read_bytes()).hexdigest())
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(pptx=str(OUT.with_suffix('.pptx')),pdf=str(OUT.with_suffix('.pdf')),slides=len(r.SCENES)),ensure_ascii=False))

if __name__=='__main__':main()
