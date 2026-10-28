#!/usr/bin/env python3
"""Build an English four-slide extract from 20261021_Meeting_ko.

The source geometry and diagrams come directly from slides 1--4 of the Korean
deck generator. Only visible language and speaker notes are changed. Slide 2
states that the collection order is preserved while the raw traffic timeline
and exact time intervals are not reproduced.
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import zipfile
from pathlib import Path

from PIL import ImageFont
from lxml import etree
from pptx import Presentation


HERE = Path(__file__).resolve().parent
BASE_PATH = HERE / "build_20261021_meeting.py"
SPEC = importlib.util.spec_from_file_location("meeting_ko_base", BASE_PATH)
assert SPEC and SPEC.loader
base = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(base)

ROOT = base.ROOT
SLIDES = ROOT / "docs/slides"
QA = SLIDES / "_support/qa/20261021_meeting_en"
QA.mkdir(parents=True, exist_ok=True)
OUT = SLIDES / "eng/20261021_Meeting_en"


TRANSLATIONS = {
    # Slide 1
    "기존 IDS의 한계 · 새 공격과 제한된 라벨": "IDS update problem · New attacks with few labels",
    "침입 탐지 시스템 (IDS) · 통신 기록의 특징으로 정상·공격 유형을 구분하는 시스템": "Intrusion detection system (IDS) · Classifies network records as benign or attack types",
    "운영 시작 시 확보한 공격 라벨만으로\n이후 등장하는 공격까지 대응하기 어려운 문제": "Labels available at deployment are insufficient\nfor attack types observed later",
    "입력의 변화": "Changing inputs",
    "네트워크·서비스·공격의 다양성\n학습 시점과 다른 트래픽의 등장": "Variation in networks, services, and attacks\nTraffic differs from the training period",
    "제한된 공격 라벨": "Few attack labels",
    "새 공격 조사 후 확보하는 사례\n충분한 학습 자료 확보의 어려움": "Labelled examples after investigation\nLimited data for model updating",
    "오탐과 갱신 부담": "False alarms and update cost",
    "정상 트래픽에 대한 잘못된 경보\n새 사례 반영에 필요한 추가 계산": "Incorrect alarms on benign traffic\nAdditional computation for each update",
    "고정된 클래스의 한 번의 점수에 더해, 새 라벨을 반영한 뒤의 탐지 품질과 비용 확인": "Evaluate quality and cost after incorporating new labels, beyond one fixed closed-set score",
    "for Network Intrusion Detection, IEEE S&P 2010 · 학습 자료 부족·입력 변동성·오류 비용": "for Network Intrusion Detection, IEEE S&P 2010 · Data scarcity, input variability, and error costs",
    "연구 문제 · 소수의 신규 공격 사례를 기존 탐지기에 반영하는 방법": "Research question · Incorporating a few labelled new attacks into an existing detector",
    "01 문제": "01 Problem",
    # Slide 2
    "공격 수집 순서를 활용한 단계별 탐지기 갱신": "Stage-wise detector updates based on the attack collection order",
    "CSE-CIC-IDS2018 · AWS 시험망에서 정상 동작과 계획된 공격을 실행·수집": "CSE-CIC-IDS2018 · Benign activity and scheduled attacks collected in an AWS testbed",
    "CSE · UNB CIC, CSE-CIC-IDS2018 on AWS, 2018 · 공식 공격 수집 일정": "CSE · UNB CIC, CSE-CIC-IDS2018 on AWS, 2018 · Official attack collection schedule",
    "공격 수집 순서를 따라 신규 공격 유형을 단계별로 도입하고,\n같은 train에서 공격별 라벨 사례(n개)를 선택해 탐지기를 순차 갱신": "Follow the original order of attack types,\nthen select n labelled rows per attack from the same training split for each update",
    "비교 항목 · 갱신 후 탐지 품질 / 기존 클래스 성능 변화 / 갱신에 걸린 시간": "Compare · Post-update detection quality / retention of known classes / update time",
    "유형별 사례는 seed로 추출 · ToN-IoT는 자체 train과 별도 공격 도입 단계로 독립 평가": "Order retained; exact intervals and the complete traffic timeline are not replayed · ToN-IoT evaluated independently",
    "신규 공격의 라벨 사례를 소량씩 제공하고 갱신 후의 품질과 비용 비교": "Controlled label-release replay · Collection order retained without claiming temporal reproduction",
    # Slide 3
    "TabPFN · 표 데이터를 위한 사전학습 Transformer": "TabPFN · A pretrained Transformer for tabular data",
    "TabPFN · 합성 표 데이터 과제로 사전학습한 Transformer 기반 예측 모델": "TabPFN · Transformer predictor pretrained on synthetic tabular tasks",
    "사용 모델: v3 · Grinsztajn et al., TabPFN-3: Technical Report, arXiv 2026": "Model used: v3 · Grinsztajn et al., TabPFN-3: Technical Report, arXiv 2026",
    "Context · 라벨 사례": "Context · labelled examples",
    "기존 정상·공격 + 신규 공격\n특징과 정답을 함께 입력": "Known classes + new attacks\nFeatures and labels together",
    "Query · 분류할 통신 기록": "Query · network record to classify",
    "특징만 입력 · 정답은 평가에 사용": "Features only · label for evaluation",
    "사전학습 가중치 고정": "Pretrained weights frozen",
    "정상·공격별 확률": "Class probabilities",
    "context를 참조한 추론": "Prediction conditioned on context",
    "가중치 추가 학습 없이 · 주어진 라벨 사례를 참조해 query의 정답 예측": "Predict a query from labelled examples without additional weight optimization",
    "채택 가설": "Working hypothesis",
    "소량의 신규 사례로 탐지 품질을 유지하면서 재학습보다 낮은 갱신 비용": "Maintain detection quality with fewer new examples and lower update cost than retraining",
    "이번 비교: IDS 가중치 fine-tuning 미수행 · context 전처리와 캐시 준비에 걸린 시간은 측정": "This comparison: no IDS fine-tuning · context preprocessing and cache preparation are timed",
    "가중치를 고정한 in-context learning · 신규 라벨을 context로 제공해 현재 IDS 문제에 적용": "Frozen-weight in-context learning · New labels enter the context for the IDS task",
    "02 접근": "02 Approach",
    # Slide 4
    "실험 전체 흐름 · 소량 라벨로 탐지기 갱신": "Evaluation workflow · Updating a detector with few labels",
    "① 라벨 확보·보유 · CIC2018 예시": "① Label release and retention · CIC2018 example",
    "n = 5 · 10 · 20 · 30 · 50 · 100개 / 신규 공격 · 3 seeds": "n = 5 · 10 · 20 · 30 · 50 · 100 per new attack · 3 seeds",
    "초기 · 1단계 · 1,000개": "Initial · Stage 1 · 1,000 rows",
    "정상\n750개": "Benign\n750",
    "Brute force\n125개": "Brute force\n125",
    "DoS\n125개": "DoS\n125",
    "신규 · 2단계": "New · Stage 2",
    "DDoS\nn개": "DDoS\nn rows",
    "신규 · 3단계": "New · Stage 3",
    "Web attacks\nn개": "Web attacks\nn rows",
    "신규 · 4단계": "New · Stage 4",
    "Infiltration\nn개": "Infiltration\nn rows",
    "신규 · 5단계": "New · Stage 5",
    "Bot\nn개": "Bot\nn rows",
    "ToN은 자체 train·공격 단계로 별도 수행 · 상세 배분은 부록": "ToN uses its own training split and attack stages · allocations in the appendix",
    "누적": "Accum.",
    "최종\n1,000+4n": "Final\n1,000+4n",
    "지금까지 제공된 사례 모두 보유": "Retain every labelled row released so far",
    "버퍼": "Buffer",
    "1,000개": "1,000 rows",
    "기존 삭제 → 신규 n개 전량 편입": "Evict old rows → insert all n new rows",
    "삭제 정책 · FIFO / 무작위 / 유사도": "Eviction · FIFO / random / similarity",
    "② 매 단계 갱신": "② Update at every stage",
    "XGB\n모델 재학습": "XGB\nrefit model",
    "TabPFN\ncontext·캐시 갱신": "TabPFN\nrefresh context and cache",
    "③ 같은 test 평가": "③ Evaluate on the same test set",
    "최종 모델 · 전체 test": "Final model · full test set",
    "3,042,473개": "3,042,473 rows",
    "2,271,723개": "2,271,723 rows",
    "각 벤치마크 안에서 독립 평가": "Separate per benchmark",
    "④ 비교 항목": "④ Outcomes",
    "탐지 품질": "Detection quality",
    "전체 test · 전체·신규 공격 F1": "Full test · overall and new-attack F1",
    "기존 클래스 성능저하": "Retention of known classes",
    "고정 패널 · 최종 − 초기 F1": "Fixed panel · final − initial F1",
    "갱신 비용": "Update cost",
    "2단계 이후 합계 · 추론 제외": "Sum after Stage 1 · inference excluded",
    "같은 보유 정책·같은 사례로 XGB/PFN 비교 · 최종 품질과 기존 성능 변화·갱신 비용 측정": "Matched rows and retention policies · Quality, retention, and update cost",
    "03 평가": "03 Evaluation",
}


NOTES = [
    "The setting begins with attack labels available at deployment and later receives a few labelled examples of additional attacks. The question is whether the detector can incorporate them while retaining benign and previously labelled attack discrimination. The cited study motivates data scarcity, changing inputs, and error costs; it is not a direct XGBoost–TabPFN comparison.",
    "CIC2018 scheduled different attack types on different collection dates. We retain the order in which those attack types appeared and reveal labelled training rows in that order. We do not sort every row by timestamp, preserve the original time gaps, or replay the complete benign and attack traffic timeline. This is therefore a controlled label-release replay rather than temporal reproduction. ToN-IoT is evaluated independently using its own training split and attack-introduction order.",
    "TabPFN is a Transformer-based predictor pretrained on synthetic tabular tasks. The labelled context and an unlabelled query are supplied at inference. The backbone remains frozen in this comparison, while context preprocessing and cache preparation are still measured. Lower update cost and useful quality with few labels are hypotheses evaluated by the experiment.",
    "CIC2018 begins with 750 benign, 125 Brute force, and 125 DoS rows. DDoS, Web attacks, Infiltration, and Bot are then released with n labelled rows per class. Cumulative retention keeps all released rows, while the fixed buffer removes existing rows and inserts every new row. XGBoost is refitted and TabPFN refreshes its context and cache using the same retained row IDs. Final quality is measured on the same fixed test set, retention on a fixed known-class panel, and update cost after Stage 1 with inference excluded.",
]


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _fit_text_widths() -> None:
    """Keep the source geometry and reduce English text only when needed."""
    for scene in base.SCENES:
        for item in scene["items"]:
            if item["kind"] != "text":
                continue
            font_path = base.FONT_BOLD if item["bold"] else base.FONT
            font = ImageFont.truetype(str(font_path), round(item["size"] * 4))
            actual = max((font.getlength(line) for line in item["text"].splitlines()), default=0) / 4 / 72
            available = item["w"] - 0.015
            if actual > available:
                item["size"] = round(item["size"] * available / actual * 0.98, 2)
                assert item["size"] >= 11, (scene["title"], item["text"], item["size"])


def _patch_pptx_language_and_notes() -> None:
    prs = Presentation(OUT.with_suffix(".pptx"))
    assert len(prs.slides) == 4
    for slide, note in zip(prs.slides, NOTES):
        slide.notes_slide.notes_text_frame.text = note
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run._r.get_or_add_rPr().set("lang", "en-US")
    prs.save(OUT.with_suffix(".pptx"))


def _assert_no_korean() -> None:
    prs = Presentation(OUT.with_suffix(".pptx"))
    visible = "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in prs.slides)
    assert not re.search(r"[가-힣]", visible + notes), "Korean text remains in English deck"


def _assert_no_shadows() -> int:
    with zipfile.ZipFile(OUT.with_suffix(".pptx")) as archive:
        count = 0
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            tree = etree.fromstring(archive.read(name))
            count += len(tree.xpath('//*[local-name()="outerShdw" or local-name()="innerShdw" or local-name()="prstShdw"]'))
    assert count == 0
    return count


def main() -> None:
    base.SCENES.clear()
    base.build_scenes()
    base.SCENES[:] = base.SCENES[:4]
    for scene in base.SCENES:
        for item in scene["items"]:
            if item["kind"] == "text" and item["text"] in TRANSLATIONS:
                item["text"] = TRANSLATIONS[item["text"]]
        scene["title"] = TRANSLATIONS.get(scene["title"], scene["title"])
    remaining = sorted({
        item["text"]
        for scene in base.SCENES
        for item in scene["items"]
        if item["kind"] == "text" and re.search(r"[가-힣]", item["text"])
    })
    assert not remaining, remaining
    for scene, note in zip(base.SCENES, NOTES):
        scene["notes"] = note
    _fit_text_widths()

    base.QA = QA
    base.OUT = OUT
    base.render()

    # base.render uses Korean internal support filenames; rename only the QA copies.
    for old_name, new_name in [("meeting_ko.html", "meeting_en.html"), ("rendered_ko.pdf", "rendered_en.pdf")]:
        old = QA / old_name
        if old.exists():
            old.replace(QA / new_name)

    for scene, note in zip(base.SCENES, NOTES):
        scene["notes"] = note
    (QA / "scenes.json").write_text(json.dumps(base.SCENES, ensure_ascii=False, indent=2) + "\n")
    _patch_pptx_language_and_notes()
    _assert_no_korean()
    shadow_count = _assert_no_shadows()

    manifest = {
        "date": "2026-10-21",
        "purpose": "English translation of slides 1–4 from 20261021_Meeting_ko for evaluation-scenario review",
        "source_deck": "docs/slides/kor/20261021_Meeting_ko.pptx",
        "source_slides": [1, 2, 3, 4],
        "output": str(OUT.relative_to(ROOT)),
        "slides": 4,
        "structural_changes": [],
        "slide_2_clarification": "Attack-type order is retained; exact time intervals and the complete raw traffic timeline are not replayed, so this is not temporal reproduction.",
        "files": {
            str(OUT.with_suffix(".pptx").relative_to(ROOT)): _sha(OUT.with_suffix(".pptx")),
            str(OUT.with_suffix(".pdf").relative_to(ROOT)): _sha(OUT.with_suffix(".pdf")),
        },
    }
    (QA / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    validation = json.loads((QA / "validation.json").read_text())
    validation.update({
        "passed": True,
        "slides": 4,
        "speaker_notes": 4,
        "visible_korean_characters": 0,
        "pptx_shadow_count": shadow_count,
        "pptx_sha256": _sha(OUT.with_suffix(".pptx")),
        "pdf_sha256": _sha(OUT.with_suffix(".pdf")),
        "visual_review_pending": True,
    })
    for stale in ["source_conditions", "new_xgb_conditions", "seeds", "cumulative_color", "cumulative_drawn_on_top", "matched_panel_axes", "revision"]:
        validation.pop(stale, None)
    (QA / "validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"slides": 4, "pptx": str(OUT.with_suffix('.pptx')), "pdf": str(OUT.with_suffix('.pdf'))}, ensure_ascii=False))


if __name__ == "__main__":
    main()
