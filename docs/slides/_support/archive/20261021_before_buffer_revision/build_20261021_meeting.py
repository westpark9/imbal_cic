#!/usr/bin/env python3
"""Problem -> ICL -> explicit evaluation construction -> verified EXP84 results.

Korean editable PPTX and PDF share one scene description. Charts use recorded
EXP84 values; no new experiments or edits to the English manuscript.
"""
from pathlib import Path
import csv
import hashlib
import html
import json
import math
import re
import subprocess
import sys

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = next(p for p in Path(__file__).resolve().parents if (p/'AGENTS.md').exists())
SLIDES = ROOT/'docs/slides'
QA = SLIDES/'_support/qa/20261021_meeting'
ASSETS = QA/'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
OUT = SLIDES/'kor/20261021_Meeting_ko'
RUN = ROOT/'tabpfn/results/v7/exp84/20261020_160101_small_labels_5_100_buffer_retention_s43_3204a372'
DATA = ROOT/'docs/lablog/data'
FONT = Path('/home/user/.local/share/fonts/Pretendard/Pretendard-Regular.otf')
FONT_BOLD = FONT.with_name('Pretendard-Bold.otf')
W,H = 13.333333,7.5
NAVY,BLUE,INK,GREY,PALE,TEAL,AMBER = '00306C','007BC6','20262E','617080','F1F5F9','2F8C86','A57A30'
METHODS=['xgb','pfn_accumulate','pfn_fifo','pfn_random','pfn_similar']
LABELS=['XGB 누적 재학습','PFN 누적','PFN FIFO','PFN 무작위','PFN 유사도']
COLORS=['617080',NAVY,BLUE,TEAL,AMBER]
N=[5,10,20,30,50,100]
DS=['cic2018','toniot']
DS_LABEL={'cic2018':'CIC2018','toniot':'ToN-IoT'}
CLASS_LABEL={'benign':'Benign','bot':'Bot','brute_force':'Brute force','ddos':'DDoS','dos':'DoS','infiltration':'Infiltration','web_attacks':'Web attacks','backdoor':'Backdoor','injection':'Injection','mitm':'MITM','password':'Password','ransomware':'Ransomware','scanning':'Scanning','xss':'XSS'}
URLS={
 'ids':'https://www.icir.org/robin/papers/oakland10-ml.pdf',
 'cic':'https://www.unb.ca/cic/datasets/ids-2018.html',
 'ton':'https://research.unsw.edu.au/projects/toniot-datasets',
 'pfn':'https://arxiv.org/abs/2605.13986',
 'arch':'https://arxiv.org/html/2605.13986v1#S2.SS1',
}
SOURCES=[
 dict(id='ids',authors='Robin Sommer; Vern Paxson',title='Outside the Closed World: On Using Machine Learning for Network Intrusion Detection',venue='IEEE Symposium on Security and Privacy',year=2010,url=URLS['ids'],claims=['Training-data scarcity, variability of network traffic, and error costs','Operational evaluation must be specific to its setting'],scope='A critique focused on anomaly-based network IDS; not a direct evaluation of XGBoost or TabPFN'),
 dict(id='cic',authors='CSE; Canadian Institute for Cybersecurity, UNB',title='CSE-CIC-IDS2018 on AWS',year=2018,url=URLS['cic'],claims=['Controlled testbed traffic and scheduled attacks','Table 2 dates and section 3 flow feature extraction'],scope='Official collection schedule illustrates attack variety; it does not establish real-time performance or the experimental label budgets'),
 dict(id='ton',authors='UNSW Canberra',title='The TON_IoT Datasets',year=None,url=URLS['ton'],claims=['Traffic collected in a designed Cyber Range and IoT Lab','Network data used independently of CIC'],scope='Controlled research testbed, not an untouched production deployment'),
 dict(id='pfn',authors='Léo Grinsztajn et al.',title='TabPFN-3: Technical Report',year=2026,url=URLS['pfn'],claims=['Synthetic-task pretraining','In-context prediction with fixed weights','Many-class decoder over context labels','Context/cache computation remains necessary'],sections=['2.1','2.2','2.4.2','2.5'],scope='Model mechanism, not evidence of IDS accuracy or low-cost updating in this experiment'),
]

def read(p):return json.loads(Path(p).read_text())
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
ROWS=list(csv.DictReader((DATA/'exp84_budget_values.csv').open()))
CLASSES=list(csv.DictReader((DATA/'exp84_per_class.csv').open()))

def metric(ds,n,method,key):
    r=next(r for r in ROWS if r['dataset']==ds and int(r['n'])==n and r['method']==method)
    return float(r[key])

def cls(ds,n,method,name):return next(r for r in CLASSES if r['dataset']==ds and int(r['n'])==n and r['method']==method and r['name']==name)

def txt(x,y,w,h,s,size=21,bold=False,color=INK,align='left',url=None):
    h=max(h,len(s.splitlines())*size*1.16/72*1.05)
    return dict(kind='text',x=x,y=y,w=w,h=h,text=s,size=size,bold=bold,color=color,align=align,url=url)
def box(x,y,w,h,fill=PALE):return dict(kind='box',x=x,y=y,w=w,h=h,fill=fill)
def arr(x,y,w=.38,h=.26,down=False):return dict(kind='down' if down else 'arrow',x=x,y=y,w=w,h=h,fill=BLUE)
def img(path,x,y,w,h,crop=None):return dict(kind='image',x=x,y=y,w=w,h=h,path=str(path),crop=crop or [0,0,0,0])
def rule(x,y,w,fill='D7E0EA'):return box(x,y,w,.012,fill)
def label(x,y,w,s,size=16,url=None):return txt(x,y,w,.29,s,size,False,GREY,url=url)

SCENES=[]
def page(title,items,message='',section='01 문제',notes=''):
    n=len(SCENES)+1
    footer=[box(0,7.05,W,.45,NAVY),txt(.58,7.15,9.6,.22,message,14,color='FFFFFF'),
            txt(10.3,7.15,1.85,.22,section,12,color='C9D8EB',align='right'),txt(12.35,7.14,.40,.23,str(n),13,color='FFFFFF',align='right')]
    SCENES.append(dict(title=title,notes=notes,items=[txt(.58,.32,12.18,.6,title,30,True,NAVY)]+items+footer))

def simple_table(x,y,widths,headers,rows,row_heights=None,size=17):
    out=[];xx=x
    for j,(head,w) in enumerate(zip(headers,widths)):
        out += [box(xx,y,w,.49,NAVY),txt(xx+.10,y+.10,w-.20,.29,head,size,True,'FFFFFF',align='left' if j<2 else 'right')];xx+=w
    yy=y+.49
    for i,row in enumerate(rows):
        h=row_heights[i] if row_heights else .46
        out.append(box(x,yy,sum(widths),h,PALE if i%2==0 else 'FFFFFF'))
        xx=x
        for j,(s,w) in enumerate(zip(row,widths)):
            align='left' if j==0 or (j==1 and headers[1]!='test 수') else 'right'
            out.append(txt(xx+.10,yy+.09,w-.20,h-.14,str(s),size,align=align));xx+=w
        yy+=h
    return out

def setup_plot():
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt,font_manager
    font_manager.fontManager.addfont(str(FONT));font_manager.fontManager.addfont(str(FONT_BOLD))
    plt.rcParams.update({'font.family':'Pretendard','font.size':16,'axes.unicode_minus':False,'axes.spines.top':False,'axes.spines.right':False,'axes.edgecolor':'#C8D2DE','axes.labelcolor':'#'+INK,'xtick.color':'#'+GREY,'ytick.color':'#'+GREY,'svg.fonttype':'none','savefig.facecolor':'white'})
    return plt

def plot_curves(name,key,ylabel,ylim,ticks=None):
    plt=setup_plot();fig,axes=plt.subplots(1,2,figsize=(12.2,3.65))
    for ax,ds in zip(axes,DS):
        for i,m in enumerate(METHODS):
            ax.plot(range(6),[metric(ds,n,m,key) for n in N],label=LABELS[i],color='#'+COLORS[i],marker=['s','o','^','D','v'][i],linewidth=2,markersize=5,linestyle=['-','-','--','-.',':'][i])
        ax.set_title(DS_LABEL[ds],fontsize=20,fontweight='bold',color='#'+NAVY,pad=11)
        ax.set_xticks(range(6),N);ax.set_xlabel('신규 공격당 제공 라벨 수 n (개)',fontsize=15,labelpad=6)
        ax.set_ylim(*ylim);ax.grid(axis='y',color='#E1E7EF',linewidth=.7)
        if ticks is not None:ax.set_yticks(ticks)
        if 'change' in key:ax.axhline(0,color='#708090',lw=1)
        ax.set_ylabel(ylabel,fontsize=15)
    fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',bbox_to_anchor=(.5,-.005),ncol=5,frameon=False,fontsize=13.5,handlelength=2,columnspacing=1.2)
    fig.subplots_adjust(left=.065,right=.985,top=.88,bottom=.26,wspace=.26)
    for ext in ['png','svg']:fig.savefig(ASSETS/f'{name}.{ext}',dpi=190)
    plt.close(fig)

def plot_tail():
    plt=setup_plot();fig,axes=plt.subplots(1,2,figsize=(12.2,3.2))
    groups=[('cic2018',['infiltration','web_attacks']),('toniot',['mitm','ransomware'])]
    for ax,(ds,names) in zip(axes,groups):
        for j,m in enumerate(METHODS):
            vals=[100*float(cls(ds,100,m,c)['f1']) for c in names]
            bars=ax.bar(np.arange(2)+(j-2)*.15,vals,width=.14,color='#'+COLORS[j],label=LABELS[j])
            ax.bar_label(bars,labels=[f'{v:.2f}' for v in vals],fontsize=11.5,padding=2,rotation=0)
        ax.set_xticks(range(2),[CLASS_LABEL[c]+'\n'+f"test {int(cls(ds,100,'xgb',c)['support']):,}개" for c in names],fontsize=15)
        ax.set_ylim(0,18);ax.set_yticks([0,5,10,15]);ax.set_ylabel('클래스별 F1 (%)',fontsize=15)
        ax.set_title(DS_LABEL[ds],fontsize=19,fontweight='bold',color='#'+NAVY)
        ax.grid(axis='y',color='#E1E7EF',lw=.6);ax.set_axisbelow(True)
    fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',bbox_to_anchor=(.5,-.01),ncol=5,frameon=False,fontsize=13.5)
    fig.subplots_adjust(left=.065,right=.985,top=.86,bottom=.30,wspace=.25)
    for ext in ['png','svg']:fig.savefig(ASSETS/f'tail.{ext}',dpi=190)
    plt.close(fig)

def endpoint_row(key,unit='%',digits=2,y=5.58):
    # Compact numeric companion for the rightmost point; all sweep values remain in curves/CSV.
    out=[label(.65,y,1.35,'n = 100 값',16)]
    for i,m in enumerate(METHODS):
        x=2.13+i*2.13
        out += [txt(x,y,2.0,.25,LABELS[i],14.5,True,COLORS[i]),
                txt(x,y+.36,.93,.24,'CIC',13,color=GREY,align='right'),txt(x+1.05,y+.36,.93,.24,'ToN',13,color=GREY,align='right'),
                txt(x,y+.70,.93,.28,f'{metric("cic2018",100,m,key):.{digits}f}',17,align='right'),
                txt(x+1.05,y+.70,.93,.28,f'{metric("toniot",100,m,key):.{digits}f}',17,align='right')]
    out += [label(.65,y+.70,1.35,'단위 '+unit,14)]
    return out

def build_scenes():
    # 1. Problem first; the literature is evidence, not a claim of universal failure.
    items=[txt(.62,1.10,12,.38,'침입 탐지 시스템 (IDS) · 통신 기록의 특징으로 정상·공격 유형을 구분하는 시스템',20),
           txt(.62,1.84,12,1.05,'운영 시작 시 확보한 공격 라벨만으로\n이후 등장하는 공격까지 대응하기 어려운 문제',29,True,NAVY)]
    for i,(title,body) in enumerate([
        ('입력의 변화','네트워크·서비스·공격의 다양성\n학습 시점과 다른 트래픽의 등장'),
        ('제한된 공격 라벨','새 공격 조사 후 확보하는 사례\n충분한 학습 자료 확보의 어려움'),
        ('오탐과 갱신 부담','정상 트래픽에 대한 잘못된 경보\n새 사례 반영에 필요한 추가 계산')]):
        x=.62+4.23*i
        items += [box(x,3.23,3.64,1.66),box(x,3.23,3.64,.48,NAVY),txt(x+.16,3.33,3.3,.3,title,21,True,'FFFFFF'),txt(x+.16,3.94,3.30,.75,body,18)]
    items += [txt(.66,5.25,12,.40,'고정된 클래스의 한 번의 점수에 더해, 새 라벨을 반영한 뒤의 탐지 품질과 비용 확인',21,True,BLUE),
              label(.66,6.02,12,'Sommer · Paxson, Outside the Closed World: On Using Machine Learning',15,url=URLS['ids']),
              label(.66,6.34,12,'for Network Intrusion Detection, IEEE S&P 2010 · 학습 자료 부족·입력 변동성·오류 비용',15,url=URLS['ids'])]
    page('기존 IDS의 한계 · 새 공격과 제한된 라벨',items,'연구 문제 · 소수의 신규 공격 사례를 기존 탐지기에 반영하는 방법',notes='IDS를 먼저 정의한다. 선행연구는 주로 anomaly-based IDS의 운영상 어려움을 논의했으며 XGB나 TabPFN의 직접 비교가 아니다. 새 공격 라벨 확보 이후의 갱신을 본 연구 문제로 구체화한다. 신규 공격을 라벨 확보 전에 탐지했다는 주장은 하지 않는다. '+URLS['ids'])

    # 2. Official collection evidence appears after the operational problem.
    items=[txt(.62,1.12,12,.38,'CSE-CIC-IDS2018 · AWS 시험망에서 정상 동작과 계획된 공격을 실행·수집',22,True,NAVY),
           label(.62,1.70,12,'CSE · UNB CIC, CSE-CIC-IDS2018 on AWS, 2018 · 공식 문서의 공격 일정 Table 2',16,url=URLS['cic'])]
    labels=[('2/14','Brute force'),('2/15·16','DoS'),('2/20·21','DDoS'),('2/22·23','Web attacks'),('2/28·3/1','Infiltration'),('3/2','Bot')]
    for i,(date,name) in enumerate(labels):
        x=.62+i*2.10
        items += [box(x,2.37,1.81,1.03),txt(x+.10,2.52,1.61,.3,date,21,True,BLUE,align='center'),txt(x+.10,3.01,1.61,.27,name,15.5,align='center')]
        if i<5:items.append(arr(x+1.87,2.77,.17,.20))
    items += [txt(.65,3.88,12,.75,'수집 기간에 따라 달라지는 공격 유형\n패킷에서 추출한 통신량·지속시간 등의 표 특징과 공격 라벨',22),
              rule(.65,4.94,12),
              txt(.65,5.19,12,.75,'이번 실험에서 설계한 부분 · 확보되는 공격 라벨의 순서와 개수\n원본 날짜별 전체 트래픽 재생과 구별되는 단계적 라벨 제공',21,True,NAVY),
              label(.65,6.30,12,'ToN-IoT도 별도 시험망에서 수집된 네트워크 자료 사용 · UNSW Canberra, The TON_IoT Datasets',15.5,url=URLS['ton'])]
    page('공격 변화의 수집 사례 · 공식 벤치마크',items,'수집 기록을 활용한 통제 실험 · 공격 라벨이 추가되는 상황의 구성',notes='2018년 날짜이며 모든 이벤트를 나열한 표는 아니다. 공식 문서의 최초 등장 순서를 요약한다. 날짜가 달라졌다는 이유만으로 재학습이 반드시 필요하다고 단정하지 않는다. 두 벤치마크 모두 통제된 시험망에서 생성·수집된 트래픽이며 실제 운영망의 그대로의 관측이라고 표현하지 않는다. 본 연구는 특징 행을 새로 합성하거나 두 벤치마크를 결합하지 않는다. '+URLS['cic']+' '+URLS['ton'])

    # 3. ICL and FT distinction at the actual input-output boundary.
    items=[txt(.62,1.08,12,.35,'TabPFN · 표 데이터의 사례와 정답으로부터 예측하는 방법을 사전학습한 모델',22,True,NAVY),
           label(.62,1.62,12,'Grinsztajn et al., TabPFN-3: Technical Report, arXiv 2026 · 합성 표 과제로 사전학습',16,url=URLS['pfn']),
           box(.62,2.20,3.43,1.25),txt(.82,2.36,3.02,.3,'Context · 라벨 사례',20,True,BLUE),txt(.82,2.90,3.02,.4,'기존 정상·공격 + 신규 공격\n특징과 정답을 함께 입력',17),
           box(.62,3.73,3.43,.94),txt(.82,3.87,3.02,.30,'Query · 분류할 통신 기록',20,True,BLUE),txt(.82,4.29,3.02,.26,'특징만 입력 · 정답은 평가에 사용',16.5),
           arr(4.25,2.79,.47,.31),arr(4.25,4.04,.47,.31),box(4.94,2.64,3.46,1.66,NAVY),txt(5.18,2.95,2.98,.43,'TabPFN v3',27,True,'FFFFFF'),txt(5.18,3.68,2.98,.35,'사전학습 가중치 고정',20,color='FFFFFF'),
           arr(8.61,3.30,.46,.31),box(9.28,2.64,3.43,1.66),txt(9.50,2.97,2.99,.38,'정상·공격별 확률',22,True,BLUE),txt(9.50,3.67,2.99,.31,'context를 참조한 추론',18),
           txt(.65,5.06,2.52,.35,'In-context learning',20,True,BLUE),txt(3.45,5.06,9.20,.68,'입력 사례를 바꾸어 현재 분류 문제에 적응\n가중치에 사례를 저장하거나 역전파하는 과정 없음',20),
           txt(.65,6.12,2.52,.35,'Fine-tuning',20,True,GREY),txt(3.45,6.12,9.20,.35,'오차를 역전파해 가중치를 바꾸는 별도 절차 · 이번 실험은 미수행',18)]
    page('TabPFN의 in-context learning',items,'갱신 대상 · 사전학습 가중치 외부의 context와 계산 캐시',section='02 접근',notes='Context는 가중치 안에 재구성하는 데이터가 아니다. 현재 IDS 라벨 의미를 지정하는 외부 입력이다. v3는 context의 라벨을 참조하는 decoder를 사용하므로 기존 클래스 사례도 필요하다. 사전학습은 합성 과제에 관한 것이며 평가용 IDS 행을 합성했다는 뜻이 아니다. 고정 가중치 모델도 전처리와 캐시 준비 비용을 지불한다. '+URLS['pfn'])

    # 4. Purpose must precede the dataset schedule.
    items=[txt(.65,1.16,12,1.02,'새 공격의 라벨 사례가 소량 확보될 때\nTabPFN의 context 갱신으로 어느 품질과 비용까지 가능한지 평가',27,True,NAVY)]
    for i,(h,b) in enumerate([('출발 조건','정상과 일부 공격의 라벨 확보\n초기 탐지기 구성'),('변화 조건','새 공격의 라벨 n개 추가 확보\n동일 사례를 두 모델에 제공'),('비교 대상','XGB의 누적 재학습\nPFN의 누적·버퍼 context 갱신')]):
        x=.65+4.22*i;items += [box(x,2.71,3.61,1.40),txt(x+.17,2.87,3.24,.32,h,21,True,BLUE),txt(x+.17,3.43,3.24,.60,b,18)]
        if i<2:items.append(arr(x+3.80,3.27,.24,.25))
    items += [txt(.65,4.54,3.70,.35,'적은 신규 라벨',22,True,BLUE),txt(4.87,4.54,3.70,.35,'기존 클래스 유지',22,True,BLUE),txt(9.09,4.54,3.61,.35,'갱신 비용',22,True,BLUE),
              txt(.65,5.12,3.70,.67,'신규 공격별 F1·재현율\n제공 라벨 수에 따른 변화',19),txt(4.87,5.12,3.70,.67,'같은 평가 사례에서\n갱신 전후 성능 변화',19),txt(9.09,5.12,3.61,.67,'추가 라벨 반영에 걸린 시간\n초기 준비·추론과 분리',19),
              label(.65,6.30,12,'평가 범위 · 라벨 확보 후의 공격 유형 분류 / 미지 공격의 사전 탐지·다른 네트워크 일반화는 별도 과제',17)]
    page('평가 목적 · 소수 라벨로 가능한 갱신',items,'비교 가설 · 재학습 없는 적응의 품질과 비용을 함께 확인',section='03 평가',notes='데이터 구성보다 연구 질문을 먼저 설명한다. 모든 신규 공격이 처음부터 학습에 포함된 단일 평가가 아니라 공격 라벨을 단계적으로 제공한다. 최종 예측은 다중 클래스 분류이고 미지 공격 reject 검출은 수행하지 않는다. 세 결과 축을 뒤의 신규 F1, 초기 클래스 변화, 후속 갱신 시간에 대응시킨다.')

    # 5. A/B ambiguity resolved with separate source -> train/test -> provided data lanes.
    items=[txt(.62,1.10,12,.40,'각 벤치마크 안에서 train의 행을 선택 · CIC와 ToN 사이의 행 결합 없음',23,True,NAVY)]
    for i,(ds,test) in enumerate([('CIC2018','3,042,473'),('ToN-IoT','2,271,723')]):
        y=1.93+2.08*i
        items += [box(.62,y,1.68,1.44),txt(.77,y+.29,1.38,.40,ds,22,True,NAVY),txt(.77,y+.88,1.38,.30,'수집 벤치마크',15),arr(2.43,y+.57,.31,.27),
                  box(2.91,y,3.27,1.44),txt(3.08,y+.14,2.92,.34,'정제 후 train / test 분리',19,True,BLUE),txt(3.08,y+.72,2.92,.55,'상충 라벨 제거\n학습·평가 행 분리',16.5),
                  arr(6.33,y+.33,.28,.23),arr(6.33,y+1.00,.28,.23),
                  box(6.79,y,5.92,.67),txt(6.96,y+.11,5.55,.44,'Train → 초기 1,000개 + 신규 공격당 n개 선택',19,True,NAVY),
                  box(6.79,y+.80,5.92,.64),txt(6.96,y+.94,5.55,.32,f'Test → 고정 {test}개 · 학습/context에 미사용',18)]
    items += [txt(.66,5.82,12,.36,'학습 후보 중 일부를 단계별로 제공 → 신규 라벨을 소량 확보한 상황의 재현',20,True,NAVY),
              txt(.66,6.38,12,.36,'“현재 데이터셋”에 해당하는 것 · 그 단계까지 제공한 라벨 사례 집합',20,True,BLUE)]
    page('평가 방법 · 데이터 출처와 사례 선택',items,'수집된 특징 행을 선택한 라벨 제공 실험 · 새로운 특징 행의 합성 없음',section='03 평가',notes='청중이 말한 원본 A는 각 출처 벤치마크이고 현재 B는 그 train에서 해당 단계까지 제공한 사례 집합이다. 별도의 출처 B를 생성하거나 A의 공격을 다른 데이터셋에 붙이지 않는다. CIC에서 CIC train/test를, ToN에서 ToN train/test를 각각 사용한다. 두 벤치마크는 시험망 수집 자료다. 행 ID는 분리되어 있으나 동일 특징 벡터가 남을 수 있으므로 완전한 특징 중복 제거를 주장하지 않는다. 기존 중복 특징 제외 진단은 별도 결과에 저장되어 있다. '+URLS['cic']+' '+URLS['ton'])

    # 6. Actual stage allocations, all budgets, and one worked example.
    items=[txt(.62,1.09,12,.39,'n = 5 · 10 · 20 · 30 · 50 · 100개 / 신규 공격 클래스 · seed 43 한 번',23,True,NAVY),
           label(.62,1.60,12,'n은 확보 가능한 라벨 수의 실험 조건 · 하루 공격량이나 최적 표본 수를 뜻하지 않는 값',17)]
    cic=[['1','Benign / Brute force / DoS','750 / 125 / 125','1,000'],['2','DDoS','n','1,000+n'],['3','Web attacks','n','1,000+2n'],['4','Infiltration','n','1,000+3n'],['5','Bot','n','1,000+4n']]
    ton=[['1','Benign / DoS / Scanning','750 / 125 / 125','1,000'],['2','DDoS / Injection','각 n','1,000+2n'],['3','Password','n','1,000+3n'],['4','XSS','n','1,000+4n'],['5','Backdoor / Ransomware','각 n','1,000+6n'],['6','MITM','n','1,000+7n']]
    # Multiline stage 1 keeps the class-specific allocation explicit and aligned.
    cic[0]=['1','Benign\nBrute force\nDoS','750\n125\n125','1,000']
    ton[0]=['1','Benign\nDoS\nScanning','750\n125\n125','1,000']
    ton[1][1]='DDoS\nInjection';ton[1][2]='n\nn';ton[4][1]='Backdoor\nRansomware';ton[4][2]='n\nn'
    items += [txt(.62,1.99,5.85,.36,'CIC2018',22,True,BLUE),txt(6.87,1.99,5.85,.36,'ToN-IoT',22,True,BLUE)]
    widths=[.72,2.27,1.20,1.67]
    items += simple_table(.62,2.45,widths,['단계','제공 클래스','클래스당','누적 라벨'],cic,[.90,.48,.48,.48,.48],15)
    items += simple_table(6.87,2.45,widths,['단계','제공 클래스','클래스당','누적 라벨'],ton,[.90,.60,.38,.38,.60,.38],15)
    items += [label(.65,6.54,12,'공격 도입 단계순 · 누적 라벨 = 지금까지 제공한 총 사례 / 버퍼 방식의 실제 보유량은 1,000개',15.5)]
    page('평가 방법 · 단계별 클래스와 제공 사례 수',items,'신규 사례 전량 반영 · 초기 1,000개와 동일한 신규 행으로 방법 간 비교',section='03 평가',notes='초기 benign 750, 공격 2종 각125는 통제 조건이며 운영 분포나 최적값이라고 주장하지 않는다. 단계는 원본의 공격 등장 순서를 참고한 라벨 공개 순서이고 raw timestamp replay가 아니다. n=100은 6개 조건 중 상한이며 대표 최적값이 아니다. n=20의 CIC 단계2 예시는 기존1000 + DDoS20 = 누적1020, 버퍼는 기존20개삭제 + 신규20개 =1000이다. 클래스별 표본 동일 조건에서도 단계별 새 클래스 수에 따라 신규 총량이 달라진다.')

    # 7. Explicit update operation, including cost and eval population differences.
    items=[txt(.62,1.12,12,.37,'CIC 2단계 · n = 20 예시 · 같은 train에서 선택한 DDoS 20개를 전량 제공',22,True,NAVY)]
    policies=[('XGB 누적 재학습','기존 1,000 + 신규 20 = 1,020개','누적 사례 전체로 새 모델 학습'),('PFN 누적 context','기존 1,000 + 신규 20 = 1,020개','가중치 고정 · context/cache 재준비'),('PFN 1,000개 버퍼','기존 20개 삭제 + 신규 20개 편입','FIFO / 무작위 / 유사도 삭제')]
    for i,(h,b,c) in enumerate(policies):
        x=.62+4.23*i
        items += [box(x,1.93,3.65,1.75),box(x,1.93,3.65,.47,NAVY),txt(x+.15,2.04,3.35,.28,h,19,True,'FFFFFF'),txt(x+.15,2.64,3.35,.43,b,17),txt(x+.15,3.20,3.35,.34,c,16.5)]
    items += [txt(.65,4.00,12,.66,'FIFO · 먼저 들어온 사례 삭제     무작위 · 기존 사례에서 무작위 삭제\n유사도 · 새 사례와 가까운 기존 사례 삭제 / 기존 클래스별 최소 1개 유지',19),
              rule(.65,4.92,12),
              txt(.65,5.18,2.06,.34,'최종 품질',20,True,BLUE),txt(2.92,5.18,9.65,.35,'마지막 모델로 전체 test 평가 · 신규 공격별 F1을 동일 비중으로 평균',18),
              txt(.65,5.72,2.06,.34,'기존 성능 유지',20,True,BLUE),txt(2.92,5.72,9.65,.35,'초기 3개 클래스의 고정 평가 표본에서 최종 − 초기 Macro-F1',18),
              txt(.65,6.26,2.06,.34,'갱신 비용',20,True,BLUE),txt(2.92,6.26,9.65,.35,'2단계 이후 선택·삭제·모델 준비·fit/cache 시간 합계 · 추론 제외',18)]
    page('평가 방법 · 동일 사례와 서로 다른 갱신',items,'가중치 재학습 없음 ≠ 갱신 비용 없음 · context 처리 비용까지 실측',section='03 평가',notes='XGB cumulative refit는 이번 비교에서 선택한 baseline이며 모든 XGBoost 사용법이 항상 full retraining을 요구한다고 주장하지 않는다. FIFO는 실제 outdated 판정이 아니라 편입 순서다. Similar는 OLD context로 StandardScaler를 맞춘 뒤 신규 행까지의 최소 Euclidean distance가 작은 기존 행부터 제거한다. 기존 클래스당 최소1개만 보존하며 신규 편입상한이나 새 클래스별 quota는 없다. 유지 패널은 클래스당 최대2000개의 같은 test행과 같은 query 구성을 사용하고, 전체test F1과 다른 평가 모집단이다. 초기준비, 모든추론, 체크포인트저장, 검증은 주 갱신시간에서 제외한다.')

    # 8. New-class score as the primary outcome, with an unambiguous definition.
    items=[txt(.62,1.09,12,.38,'최종 모델 · 전체 test에서 계산한 신규 공격 클래스별 F1의 산술평균',21,True,NAVY),
           label(.62,1.62,12,'CIC 4개 · ToN 7개 신규 공격 / 정상·초기 공격은 평균에서 제외, 이들에서 발생한 오탐은 F1에 포함',16.5),
           img(ASSETS/'new_attacks.png',.55,2.00,12.24,3.65)]
    items += endpoint_row('new_attack_f1_pct',y=5.72)
    page('결과 · 신규 공격 라벨 수에 따른 탐지 품질',items,'평균 성능의 유사성과 희소 공격 탐지의 충분성을 구분',section='04 결과',notes='EXP84 60조건 검증완료. n값끼리 훈련행은 nested이고 모든 방법에 동일한 행을 제공한다. 신규 평균은 각 조건의 마지막 모델 전체test 혼동행렬에서 post-initial 클래스F1을 골라 평균한다. 단계별 예측 합산이나 단계별F1평균이 아니다. 하단 수치는 n=100의 설명용 점이며 최적 조건으로 선택한 것이 아니다. CIC 신규={ddos,web_attacks,infiltration,bot}; ToN 신규={ddos,injection,password,xss,backdoor,ransomware,mitm}. 신규 그룹과 test 희소클래스 그룹은 다르다.')

    # 9. Retention is a paired panel diagnostic, not the full-test outcome.
    items=[txt(.62,1.09,12,.38,'초기 정상·공격 2종의 같은 평가 표본 · 최종 − 초기 Macro-F1 (%p)',21,True,NAVY),
           label(.62,1.62,12,'클래스당 최대 2,000개 고정 test 표본 / 같은 query 구성 / 0 아래는 초기보다 성능 하락',17),
           img(ASSETS/'retention.png',.55,2.00,12.24,3.65)]
    items += endpoint_row('initial_retention_change_pp',unit='%p',y=5.72)
    page('결과 · 갱신 이후 초기 클래스 성능 변화',items,'ToN의 큰 교체 조건에서 FIFO·유사도 삭제의 유지 성능 하락',section='04 결과',notes='전체 test F1과 다른 고정 패널 진단이다. 초기3개 클래스 행만의 F1이므로 이후공격에서초기클래스로향하는FP를 포함하지 않는다. 전체 현재클래스의 joint오탐과 각단계 직전클래스 유지진단은 원시 결과에 별도로 있다. 따라서 전체testF1 상승과 이지표 하락이 동시에 가능하다. 미래클래스 query특징이 패널에 포함될수 있는 오프라인 분석이고 정답은학습에입력하지않는다. 단일seed에서정책우열을확정하지않는다.')

    # 10. Cost is limited to measured post-initial work.
    items=[txt(.62,1.09,12,.38,'2단계 이후 갱신 시간 합계 · 초기 준비와 모든 추론 제외',23,True,NAVY),
           label(.62,1.62,12,'XGB 누적 재학습 / PFN 가중치 고정 후 context 선택·전처리·캐시 준비 · 동일 실행 환경에서 측정',17),
           img(ASSETS/'cost.png',.55,2.00,12.24,3.65)]
    items += endpoint_row('update_seconds',unit='초',digits=2,y=5.72)
    page('결과 · 소규모 context의 후속 갱신 비용',items,'이번 1,000~1,700개 조건의 비용 이점 · 대규모·전체 추론 비용으로 일반화 불가',section='04 결과',notes='GPU 동기화후실측. dataset별단계수가다르므로CIC와ToN시간차를직접model우열로비교하지않는다. PFNcache가free거나incremental하다는주장이아니다. 고정1000buffer가PFN누적보다추가로시간을줄이는효과는이범위에서뚜렷하지않다. 작은context범위의후속갱신시간만XGB보다낮았다. wallclock초기첫모델로드와추론은별도저장되며이그림에합치지않는다.')

    # 11. The unresolved Tail issue closes the result narrative honestly.
    items=[txt(.62,1.09,12,.38,'n = 100 · 각 데이터셋 test에서 가장 희소한 공격 2종 · test 수 내림차순',21,True,NAVY),
           img(ASSETS/'tail.png',.55,1.68,12.24,3.20),
           txt(.65,5.03,12,.39,'ToN Ransomware · PFN 무작위의 재현율 95.20% / 정밀도 1.30%',22,True,BLUE),
           txt(.65,5.63,12,.73,'삭제 없는 PFN 누적에서도 낮은 Tail F1 → 버퍼만으로 설명되지 않는 한계\n다음 확인 · 기존 대규모 결과와 클래스별 비교 후 가중치 적응·오탐 억제 효과 분리',20)]
    page('결과 해석 · 낮은 갱신 비용과 남은 Tail 문제',items,'확인된 성과 · 소규모 갱신 비용 절감 / 개선 과제 · 희소 공격의 오탐과 분류 품질',section='04 결과',notes='Tail은전체test에서희소한공격이며새클래스에제공한n과다른개념. n100은스윕상한의진단예시다. PFNrandomRansomwareTP675,FP51195,FN34;FP중benign50596. 재현율95.20precision1.30F12.57. PFN누적도2.68이므로삭제만의인과로단정하지않는다. 사용자가언급한전체training결과와현재설정의정합확인은후속분석이며이번슬라이드에서새롭게검증한사실로인용하지않는다. 모델자체분리능력과표본구성/확률보정가능성을분리해야한다. FT/보정효과는아직검증하지않았다. 한seed탐색결과로최소필요표본수또는일반적우위를확정하지않는다.')

    # Appendices retain secondary aggregates and architecture without interrupting the scenario.
    items=[txt(.62,1.09,12,.38,'마지막 모델로 전체 고정 test 평가 · 모든 클래스 F1의 산술평균',22,True,NAVY),
           label(.62,1.62,12,'CIC 3,042,473개 / ToN 2,271,723개 · 단계별 예측 합산 없음 · seed 43 한 번',17),
           img(ASSETS/'overall.png',.55,2.00,12.24,3.65)]+endpoint_row('macro_f1_pct',y=5.72)
    page('부록 · 전체 클래스 Macro-F1',items,'전체 평균의 상승이 모든 Tail 클래스의 개선을 뜻하지 않는 결과',section='부록',notes='모든클래스동일비중MacroF1. 신규공격평균,초기클래스고정패널유지지표와범위가다르다. 전체60조건data/exp84_budget_values.csv와연결.')
    for ds in DS:
        names=sorted({r['name'] for r in CLASSES if r['dataset']==ds},key=lambda c:-int(cls(ds,100,'xgb',c)['support']))
        rows=[[CLASS_LABEL[c],f"{int(cls(ds,100,'xgb',c)['support']):,}"]+[f"{100*float(cls(ds,100,m,c)['f1']):.2f}" for m in METHODS] for c in names]
        items=[label(.62,1.12,12,'n = 100 · 마지막 모델 전체 test · F1 (%) · 평가 표본 수 내림차순',19)]
        items += simple_table(.62,1.77,[2.27,1.62,1.66,1.66,1.66,1.66,1.56],['클래스','test 수','XGB','PFN 누적','FIFO','무작위','유사도'],rows,[.43]*len(rows),17)
        items += [label(.65,6.60,12,'각 방법은 동일한 초기·신규 사례 사용 · 전체 n의 precision·recall·F1은 실험 보고서에 수록',16)]
        page('부록 · '+DS_LABEL[ds]+' 클래스별 성능',items,section='부록',notes='클래스별F1과support는검증된exp84_per_class.csv에서직접생성. n100은설명용상한조건이며최적값아님. 전체n별결과를메인곡선과실험보고서에보존.')

    arch_path=SLIDES/'_support/qa/kor/1015_labmeeting_qa/v3_architecture_page.png'
    first=[img(arch_path,.62,1.09,3.95,5.88,[.20,.082,.19,.285])]
    for i,(h,b) in enumerate([('1  입력값 표현','통신 기록의 각 값을 벡터로 변환\n결측 표시와 context 라벨 정보 반영'),('2  열의 분포 요약','여러 기록에서 같은 특징을 살펴\n그 특징의 분포를 표현'),('3  한 기록의 특징 통합','여러 열을 행 단위 표현으로 압축\n다음 단계에서 사례끼리 비교')]):
        y=1.39+i*1.72;first += [txt(4.98,y,7.72,.40,h,23,True,BLUE),txt(5.25,y+.67,7.18,.78,b,21)]
    page('부록 · TabPFN v3의 기록 표현',first,section='부록',notes='Grinsztajn etal2026 Figure5 원도식. 그림캡션은사용자요청에따라표시하지않는다. 입력연산단계이며IDS가중치학습단계를뜻하지않는다. '+URLS['arch'])
    second=[img(arch_path,.62,1.09,3.95,5.88,[.20,.082,.19,.285]),
            txt(4.98,1.40,7.72,.40,'4  Context를 참조한 예측',23,True,BLUE),txt(5.25,2.03,7.18,.78,'query가 라벨이 있는 기록을 참조\nattention으로 참고 비중 계산',21),
            txt(4.98,3.12,7.72,.40,'5  라벨을 이용한 확률 출력',23,True,BLUE),txt(5.25,3.75,7.18,.78,'context 라벨을 가중합하는 decoder\n기존·신규 클래스 사례 모두 필요',21),
            box(4.98,5.10,7.72,1.57),txt(5.20,5.29,7.25,.35,'별도의 가중치 fine-tuning',22,True,NAVY),txt(5.20,5.86,7.25,.65,'라벨 사례를 context와 query로 나눠 query 오차 역전파\n이번 실험은 미수행 · context/cache 갱신만 사용',19)]
    page('부록 · TabPFN v3의 추론과 fine-tuning',second,section='부록',notes='원문2.2manyclassdecoder는onehotcontextlabel을value로사용한다. context밖의class에대한지도학습다중클래스확률을만들기위해관련label사례가필요하다. FT는별도가중치최적화. 이번실험에서는weightdigest동일성을확인했고FT시간0. '+URLS['arch'])

def validate_scenes():
    issues=[]
    for page_no,scene in enumerate(SCENES,1):
        texts=[]
        for i,item in enumerate(scene['items']):
            x,y,w,h=[item[k] for k in ('x','y','w','h')]
            if min(x,y)<0 or x+w>W+.01 or y+h>H+.01:
                issues.append(dict(page=page_no,item=i,kind='outside_slide',text=item.get('text','')))
            if item['kind']!='text':continue
            font=ImageFont.truetype(str(FONT_BOLD if item['bold'] else FONT),round(item['size']*4))
            # Font widths measured at 4x points, 1 inch = 72 points.
            max_width=max((font.getlength(s) for s in item['text'].splitlines()),default=0)/4/72
            if max_width>w-.005:
                issues.append(dict(page=page_no,item=i,kind='text_width',actual=round(max_width,3),width=w,text=item['text']))
            if re.search(r'합니다|입니다|습니다|—',item['text']):
                issues.append(dict(page=page_no,item=i,kind='style',text=item['text']))
            texts.append((i,item))
        for j,(ia,a) in enumerate(texts):
            for ib,b in texts[j+1:]:
                dx=min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
                dy=min(a['y']+a['h'],b['y']+b['h'])-max(a['y'],b['y'])
                if dx>.015 and dy>.015:
                    issues.append(dict(page=page_no,kind='text_overlap',items=[ia,ib],texts=[a['text'],b['text']],overlap=[dx,dy]))
    (QA/'layout_validation.json').write_text(json.dumps(dict(passed=not issues,issues=issues,slides=len(SCENES),intentional_overlaps=['Text over its background rectangles','Picture content within image rectangles']),ensure_ascii=False,indent=2)+'\n')
    if issues:
        print(json.dumps(issues,ensure_ascii=False,indent=2))
        raise RuntimeError(f'{len(issues)} layout issues')

def render():
    validate_scenes()
    narrations=[
        '오늘 다루는 문제는 운영을 시작할 때 없었던 공격의 라벨 사례가 나중에 소량 확보되는 상황입니다. 기존 정상과 공격을 계속 구분하면서 그 새 사례를 반영해야 합니다. 선행연구가 지적한 자료 부족과 오탐 문제를 이 갱신 상황에서 살펴보겠습니다.',
        '공식 CIC 자료는 서로 다른 날짜에 여러 공격을 실행한 시험망 수집 자료입니다. 이는 공격 종류가 달라지는 구체적인 사례입니다. 우리는 날짜별 전체 트래픽을 재생하는 대신, 공격 라벨을 단계적으로 확보하는 부분을 통제해서 평가합니다.',
        'TabPFN이 사전학습한 것은 특정 IDS의 공격 이름이 아니라 라벨 사례를 참고해 표 데이터를 예측하는 방법입니다. 그래서 현재 문제의 정상과 공격 사례를 context로 함께 넣습니다. 이번 실험에서는 가중치를 바꾸는 fine-tuning 없이 context를 바꾸고 그 준비 비용을 측정합니다.',
        '이제 검증할 질문을 세 가지로 나눕니다. 새 공격 사례가 적을 때 얼마나 잘 분류하는지, 기존 클래스의 성능이 얼마나 유지되는지, 그 상태로 갱신하는 데 시간이 얼마나 걸리는지입니다. 다음 장의 데이터 선택은 이 질문을 통제된 조건에서 확인하기 위한 것입니다.',
        'CIC와 ToN은 각각 독립적으로 평가합니다. 각 자료를 정제하고 train과 test로 나눈 뒤, 같은 train에서 초기 사례와 새 공격 사례를 선택합니다. 여기서 현재 데이터셋이라는 표현은 다른 출처의 자료가 아니라 그때까지 모델에 제공한 사례 집합을 뜻합니다. 일부만 제공하는 이유는 새 공격의 라벨을 조금씩 확보하는 상황을 만들기 위해서입니다.',
        '초기에는 정상 750개와 공격 두 종류를 각각 125개씩 제공합니다. 이후 단계마다 새 공격 클래스당 n개를 제공합니다. n은 하루 공격량이 아니라 실험에서 통제하는 라벨 수이며, 다섯 개부터 백 개까지 여섯 조건을 실행했습니다.',
        '예를 들어 CIC에 DDoS 20개를 제공하면 XGB는 누적 1,020개로 다시 학습하고, PFN 누적은 같은 1,020개를 context로 사용합니다. 버퍼 방식은 기존 사례 20개를 빼고 새 20개를 모두 넣어 1,000개를 유지합니다. 이후 결과에서 최종 품질, 기존 클래스 변화, 갱신 시간을 구분해 보겠습니다.',
        '이 그림은 모든 갱신이 끝난 모델의 전체 test 결과입니다. 그중 초기 이후 도입된 공격의 클래스별 F1만 골라 같은 비중으로 평균했습니다. 신규 공격 평균이 비슷하더라도 특정 희소 공격의 성능까지 충분하다는 뜻은 아닙니다.',
        '기존 성능은 초기의 정상과 두 공격에 해당하는 같은 평가 표본으로 전후를 비교했습니다. 전체 test 점수와는 다른 진단입니다. 특히 ToN에서 많은 사례를 교체하면 FIFO와 유사도 삭제의 성능 유지가 나빠지는 조건이 관찰됩니다.',
        '초기 준비와 추론을 제외하고 두 번째 단계부터 갱신하는 데 든 시간을 합했습니다. 이번 작은 context에서는 PFN의 갱신이 XGB 누적 재학습보다 빨랐습니다. 다만 버퍼 방식이 PFN 누적보다 추가로 빨라진 것은 아니며, 이 결과를 큰 데이터나 전체 추론 비용에 적용할 수는 없습니다.',
        '마지막으로 test에서 가장 희소한 공격을 보면 해결되지 않은 문제가 남습니다. PFN은 공격을 찾아내면서도 정상 트래픽을 해당 공격으로 너무 많이 분류할 수 있습니다. 삭제하지 않는 누적 방식에서도 낮은 F1이 나타나므로, 버퍼만 바꾸는 실험보다 기존 대규모 결과와 함께 모델의 구분 능력과 오탐 원인을 먼저 확인해야 합니다.',
    ]
    for scene,narration in zip(SCENES,narrations):
        scene['notes']='발표 설명\n'+narration+'\n\n근거와 해석 범위\n'+scene['notes']
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    for scene in SCENES:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text=scene['notes']
        for it in scene['items']:
            x,y,w,h=[Inches(it[k]) for k in ('x','y','w','h')]
            if it['kind']=='text':
                shape=slide.shapes.add_textbox(x,y,w,h);f=shape.text_frame;f.clear()
                f.word_wrap=False;f.auto_size=MSO_AUTO_SIZE.NONE;f.vertical_anchor=MSO_ANCHOR.TOP
                f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
                for n,line in enumerate(it['text'].splitlines()):
                    p=f.paragraphs[0] if n==0 else f.add_paragraph()
                    p.space_before=p.space_after=Pt(0);p.line_spacing=1.16
                    p.alignment={'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER}[it['align']]
                    r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(it['size']);r.font.bold=it['bold'];r.font.color.rgb=RGBColor.from_string(it['color'])
                    prop=r._r.get_or_add_rPr();prop.set('lang','ko-KR')
                    for tag in ['a:ea','a:cs']:
                        e=OxmlElement(tag);e.set('typeface','Pretendard');prop.append(e)
                    if it.get('url'):r.hyperlink.address=it['url']
            elif it['kind']=='image':
                s=slide.shapes.add_picture(it['path'],x,y,width=w,height=h)
                s.crop_left,s.crop_top,s.crop_right,s.crop_bottom=it['crop']
            else:
                typ={'arrow':MSO_SHAPE.RIGHT_ARROW,'down':MSO_SHAPE.DOWN_ARROW}.get(it['kind'],MSO_SHAPE.RECTANGLE)
                s=slide.shapes.add_shape(typ,x,y,w,h);s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(it['fill']);s.line.fill.background()
    prs.save(OUT.with_suffix('.pptx'))
    pages=[]
    for si,scene in enumerate(SCENES):
        tags=[]
        for i,it in enumerate(scene['items']):
            style=f'left:{it["x"]}in;top:{it["y"]}in;width:{it["w"]}in;height:{it["h"]}in;'
            if it['kind']=='text':
                style+=f'font-size:{it["size"]}pt;font-weight:{700 if it["bold"] else 400};color:#{it["color"]};text-align:{it["align"]};'
                content=html.escape(it['text'])
                if it.get('url'):content=f'<a href="{html.escape(it["url"],quote=True)}">{content}</a>'
                tags.append(f'<div id="s{si}i{i}" class="text" style="{style}">{content}</div>')
            elif it['kind']=='image':
                l,t,r,b=it['crop'];iw=it['w']/(1-l-r);ih=it['h']/(1-t-b)
                tags.append(f'<div class="clip" style="{style}"><img src="{Path(it["path"]).as_uri()}" style="width:{iw}in;height:{ih}in;left:{-l*iw}in;top:{-t*ih}in"></div>')
            else:
                style+=f'background:#{it["fill"]};'
                if it['kind']=='arrow':style+='clip-path:polygon(0 25%,60% 25%,60% 0,100% 50%,60% 100%,60% 75%,0 75%);'
                tags.append(f'<div class="shape" style="{style}"></div>')
        pages.append('<section>'+''.join(tags)+'</section>')
    output_html=QA/'meeting_ko.html'
    output_html.write_text('<!doctype html><html lang="ko"><meta charset="utf-8"><style>'
       f'@page{{size:{W}in {H}in;margin:0}}'
       '*{box-sizing:border-box}html,body{margin:0;padding:0}body{font-family:Pretendard,sans-serif;-webkit-print-color-adjust:exact;print-color-adjust:exact}'
       f'section{{position:relative;width:{W}in;height:{H}in;page-break-after:always;overflow:hidden}}'
       'section:last-child{page-break-after:auto}.text,.shape,.clip{position:absolute}.text{white-space:pre;line-height:1.16}.clip{overflow:hidden}.clip img{position:absolute}a{color:inherit;text-decoration:none}'
       '</style>'+''.join(pages)+'</html>')
    command=['/opt/google/chrome/chrome','--headless','--no-sandbox','--disable-dev-shm-usage','--disable-gpu','--no-pdf-header-footer','--allow-file-access-from-files','--run-all-compositor-stages-before-draw','--virtual-time-budget=5000',f'--print-to-pdf={OUT.with_suffix(".pdf")}',output_html.as_uri()]
    with (QA/'render.log').open('w') as log:subprocess.run(command,stdout=log,stderr=log,check=True,timeout=60)
    subprocess.run(['pdftotext','-layout',str(OUT.with_suffix('.pdf')),str(QA/'rendered_text.txt')],check=True)
    subprocess.run(['pdftoppm','-scale-to','1400','-png',str(OUT.with_suffix('.pdf')),str(QA/'page')],check=True,stdout=subprocess.DEVNULL)
    paths=sorted(QA.glob('page-*.png'))
    assert len(paths)==len(SCENES),(len(paths),len(SCENES))
    for part in range(math.ceil(len(paths)/6)):
        subset=paths[part*6:(part+1)*6];canvas=Image.new('RGB',(1440,3*425),'#D1D9E3')
        for j,p in enumerate(subset):
            im=Image.open(p);im.thumbnail((700,394));canvas.paste(im,((j%2)*720+10,(j//2)*425+10))
            ImageDraw.Draw(canvas).text(((j%2)*720+12,(j//2)*425+405),f'{part*6+j+1}',fill='black')
        canvas.save(QA/f'contact-{part+1}.png')
    manifest=dict(date='2026-10-21',revision=1,source_deck='docs/slides/kor/20261020_Meeting_intro_ko.pdf',output=str(OUT.relative_to(ROOT)),main_slides=11,appendix_slides=len(SCENES)-11,sources=SOURCES,experiment_run=str(RUN.relative_to(ROOT)),verified_jobs=60,seed=43,no_ids_finetuning=True,
      measurements={'quality':'Final model on entire fixed test','new_attack_mean':'Arithmetic mean of post-initial class F1 from full-test confusion matrix','retention':'Final minus initial Macro-F1 on the same initial-class panel; distinct from full-test metrics','update_cost':'Sum after initial stage; selection/eviction/construction/fit/cache; excludes all inference'},
      data_sha256={str(p.relative_to(ROOT)):sha(p) for p in [DATA/'exp84_budget_values.csv',DATA/'exp84_per_class.csv',RUN/'validation.json',ROOT/'configs/experiments/exp84_protocol.json']},
      files={str(p.relative_to(ROOT)):sha(p) for p in [OUT.with_suffix('.pptx'),OUT.with_suffix('.pdf')]},
      feedback_mapping={'problem_first':[1,2],'in_context_learning':[3],'evaluation_purpose':[4],'A_B_origin_and_sampling':[5],'class_stage_counts_and_why_n':[6],'update_operations_and_metrics':[7],'new_class_quality':[8],'old_class_retention':[9],'measured_update_cost':[10],'unresolved_tail_issue':[11],'overall_and_per_class_appendix':[12,13,14],'v3_architecture_appendix':[15,16]},
      limitations=['Single seed exploratory experiment','Previously inspected fixed test','Row IDs disjoint; identical feature vectors can remain','Cost excludes initial preparation and all inference','Retention panel differs from full test; no guarantee for all old classes','Architecture and charts are embedded images; all other slide text/diagrams editable'])
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    (QA/'scenes.json').write_text(json.dumps(SCENES,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(slides=len(SCENES),pptx=str(OUT.with_suffix('.pptx')),pdf=str(OUT.with_suffix('.pdf'))),ensure_ascii=False))

def main():
    assert len(ROWS)==60 and read(RUN/'validation.json')['passed']
    assert read(RUN/'results.json')['state']=='complete'
    plot_curves('new_attacks','new_attack_f1_pct','신규 공격 평균 F1 (%)',(40,75),[40,50,60,70])
    plot_curves('retention','initial_retention_change_pp','초기 대비 변화 (%p)',(-32,3),[-30,-20,-10,0])
    plot_curves('cost','update_seconds','후속 갱신 시간 합계 (초)',(0,6.5),[0,2,4,6])
    plot_curves('overall','macro_f1_pct','전체 Macro-F1 (%)',(50,80),[50,60,70,80])
    plot_tail();build_scenes();render()

if __name__=='__main__':main()
