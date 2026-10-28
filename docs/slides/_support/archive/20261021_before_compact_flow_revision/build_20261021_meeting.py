#!/usr/bin/env python3
"""Problem -> ICL -> matched examples -> verified EXP84/85 results.

Korean editable PPTX and PDF share one scene description. Charts use recorded
EXP84/85 values; no edits to the English manuscript.
"""
from pathlib import Path
import csv
import hashlib
import html
import json
import math
import re
import shutil
import subprocess
import sys
import zipfile

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree

ROOT = next(p for p in Path(__file__).resolve().parents if (p/'AGENTS.md').exists())
SLIDES = ROOT/'docs/slides'
QA = SLIDES/'_support/qa/20261021_meeting'
ASSETS = QA/'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
OUT = SLIDES/'kor/20261021_Meeting_ko'
RUN = ROOT/'tabpfn/results/v7/exp84/20261020_160101_small_labels_5_100_buffer_retention_s43_3204a372'
MATCHED = ROOT/'tabpfn/results/v7/exp85/20261021_142357_matched_xgb_buffers_s43_s44_s45_7b56fee4'
DATA = ROOT/'docs/lablog/data'
FONT = Path('/home/user/.local/share/fonts/Pretendard/Pretendard-Regular.otf')
FONT_BOLD = FONT.with_name('Pretendard-Bold.otf')
W,H = 13.333333,7.5
NAVY,BLUE,INK,GREY,PALE,TEAL,AMBER = '00306C','007BC6','20262E','617080','F1F5F9','2F8C86','A57A30'
POLICIES=['accumulate','fifo','random','similar']
POLICY_LABELS=['누적','FIFO','무작위','유사도']
COLORS=['111111','D55E00','009E73','B45A9B']
def method(model,policy):return 'xgb' if model=='xgb' and policy=='accumulate' else model+'_'+policy
METHODS=[method(model,policy) for model in ['xgb','pfn'] for policy in POLICIES]
N=[5,10,20,30,50,100]
DS=['cic2018','toniot']
DS_LABEL={'cic2018':'CIC2018','toniot':'ToN-IoT'}
CLASS_LABEL={'benign':'Benign','bot':'Bot','brute_force':'Brute force','ddos':'DDoS','dos':'DoS','infiltration':'Infiltration','web_attacks':'Web attacks','backdoor':'Backdoor','injection':'Injection','mitm':'MITM','password':'Password','ransomware':'Ransomware','scanning':'Scanning','xss':'XSS'}
URLS={
 'ids':'https://www.icir.org/robin/papers/oakland10-ml.pdf',
 'cic':'https://www.unb.ca/cic/datasets/ids-2018.html',
 'ton':'https://research.unsw.edu.au/projects/toniot-datasets',
 'pfn':'https://arxiv.org/abs/2605.13986',
 'arch':'https://arxiv.org/html/2605.13986v1#S2.SS1',
}
SOURCES=[
 dict(id='ids',authors='Robin Sommer; Vern Paxson',title='Outside the Closed World: On Using Machine Learning for Network Intrusion Detection',venue='IEEE Symposium on Security and Privacy',year=2010,url=URLS['ids'],claims=['Training-data scarcity, variability of network traffic, and error costs','Operational evaluation must be specific to its setting'],scope='A critique focused on anomaly-based network IDS; not a direct evaluation of XGBoost or TabPFN'),
 dict(id='cic',authors='CSE; Canadian Institute for Cybersecurity, UNB',title='CSE-CIC-IDS2018 on AWS',year=2018,url=URLS['cic'],claims=['Controlled testbed traffic and scheduled attacks','Table 2 dates and section 3 flow feature extraction'],scope='Official collection schedule illustrates attack variety; it does not establish real-time performance or the experimental label budgets'),
 dict(id='ton',authors='UNSW Canberra',title='The TON_IoT Datasets',year=None,url=URLS['ton'],claims=['Traffic collected in a designed Cyber Range and IoT Lab','Network data used independently of CIC'],scope='Controlled research testbed, not an untouched production deployment'),
 dict(id='pfn',authors='Léo Grinsztajn et al.',title='TabPFN-3: Technical Report',year=2026,url=URLS['pfn'],claims=['Synthetic-task pretraining','In-context prediction with fixed weights','Many-class decoder over context labels','Context/cache computation remains necessary'],sections=['2.1','2.2','2.4.2','2.5'],scope='Model mechanism, not evidence of IDS accuracy or low-cost updating in this experiment'),
]

def read(p):return json.loads(Path(p).read_text())
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
ROWS=list(csv.DictReader((DATA/'exp85_summary.csv').open()))
CLASSES=list(csv.DictReader((DATA/'exp85_per_class_summary.csv').open()))

def metric(ds,n,method,key,stat='mean'):
    r=next(r for r in ROWS if r['dataset']==ds and int(r['n'])==n and r['method']==method)
    return float(r[key+'_'+stat])

def cls(ds,n,method,name):return next(r for r in CLASSES if r['dataset']==ds and int(r['n'])==n and r['method']==method and r['name']==name)

def txt(x,y,w,h,s,size=21,bold=False,color=INK,align='left',url=None):
    h=max(h,len(s.splitlines())*size*1.16/72*1.05)
    return dict(kind='text',x=x,y=y,w=w,h=h,text=s,size=size,bold=bold,color=color,align=align,url=url)
def box(x,y,w,h,fill=PALE):return dict(kind='box',x=x,y=y,w=w,h=h,fill=fill)
def outline(x,y,w,h,color=BLUE,dashed=False):return dict(kind='outline',x=x,y=y,w=w,h=h,fill='FFFFFF',border=color,dashed=dashed)
def arr(x,y,w=.38,h=.26,down=False):return dict(kind='down' if down else 'arrow',x=x,y=y,w=w,h=h,fill=BLUE)
def img(path,x,y,w,h,crop=None):return dict(kind='image',x=x,y=y,w=w,h=h,path=str(path),crop=crop or [0,0,0,0])
def rule(x,y,w,fill='D7E0EA'):return box(x,y,w,.012,fill)
def label(x,y,w,s,size=16,url=None):return txt(x,y,w,.29,s,size,False,GREY,url=url)

DATA_COLORS=['9AA7B4','0072B2','E69F00','B45A9B']

def sample_bar(x,y,w,h,counts,max_total=None,show_counts=True,size=14):
    """Editable blocks; block lengths encode actual row counts."""
    total=sum(counts);scale=w/(max_total or total);items=[];xx=x
    for count,color in zip(counts,DATA_COLORS):
        if not count:continue
        width=count*scale
        items.append(box(xx,y,width,h,color))
        if show_counts and width>.33:
            items.append(txt(xx+.025,y+(h-.24)/2,max(.01,width-.05),.24,f'{count:,}',size,True,'FFFFFF',align='center'))
        xx+=width
    return items

def reference_counts(ds,n,policy,stage=-1):
    result=read(RUN/f'new_{n}'/'jobs'/f'{ds}_pfn_{policy}'/'results.json')
    names=read(RUN/'data'/ds/'manifest.json')['class_names']
    return dict(zip(names,result['fits'][stage]['context_class_counts']))

SCENES=[]
def page(title,items,message='',section='01 문제',notes=''):
    n=len(SCENES)+1
    footer=[box(0,7.05,W,.45,NAVY),txt(.58,7.15,9.6,.22,message,14,color='FFFFFF'),
            txt(10.3,7.15,1.85,.22,section,12,color='C9D8EB',align='right'),txt(12.35,7.14,.40,.23,str(n),13,color='FFFFFF',align='right')]
    SCENES.append(dict(title=title,notes=notes,items=[txt(.58,.32,12.18,.6,title,30,True,NAVY)]+items+footer))

def simple_table(x,y,widths,headers,rows,row_heights=None,size=17,emphasis_cols=()):
    out=[];xx=x
    for j,(head,w) in enumerate(zip(headers,widths)):
        out += [box(xx,y,w,.49,TEAL if j in emphasis_cols else NAVY),txt(xx+.10,y+.10,w-.20,.29,head,size,True,'FFFFFF',align='left' if j<2 else 'right')];xx+=w
    yy=y+.49
    for i,row in enumerate(rows):
        h=row_heights[i] if row_heights else .46
        out.append(box(x,yy,sum(widths),h,PALE if i%2==0 else 'FFFFFF'))
        xx=x
        for j,(s,w) in enumerate(zip(row,widths)):
            align='left' if j==0 or (j==1 and headers[1]!='test 수') else 'right'
            out.append(txt(xx+.10,yy+.09,w-.20,h-.14,str(s),size,bold=j in emphasis_cols,color=TEAL if j in emphasis_cols else INK,align=align));xx+=w
        yy+=h
    return out

def setup_plot():
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt,font_manager
    font_manager.fontManager.addfont(str(FONT));font_manager.fontManager.addfont(str(FONT_BOLD))
    plt.rcParams.update({'font.family':'Pretendard','font.size':16,'axes.unicode_minus':False,'axes.spines.top':False,'axes.spines.right':False,'axes.edgecolor':'#C8D2DE','axes.labelcolor':'#'+INK,'xtick.color':'#'+GREY,'ytick.color':'#'+GREY,'svg.fonttype':'none','savefig.facecolor':'white'})
    return plt

def plot_curves(name,key,ylabel):
    plt=setup_plot();fig,axes=plt.subplots(2,2,figsize=(12.2,4.72))
    for row,ds in enumerate(DS):
        all_low=[metric(ds,n,m,key)-metric(ds,n,m,key,'sd') for n in N for m in METHODS]
        all_high=[metric(ds,n,m,key)+metric(ds,n,m,key,'sd') for n in N for m in METHODS]
        lo,hi=min(all_low),max(all_high);margin=max((hi-lo)*.15,.08)
        if 'seconds' in key:lo=0
        if 'change' in key:hi=max(hi,0)
        limits=(lo-margin,hi+margin)
        if 'seconds' in key:limits=(0,hi+margin)
        for col,model in enumerate(['xgb','pfn']):
            ax=axes[row,col];last=[]
            for j,policy in enumerate(POLICIES):
                m=method(model,policy)
                values=np.array([metric(ds,n,m,key) for n in N]);sd=np.array([metric(ds,n,m,key,'sd') for n in N])
                color='#'+COLORS[j]
                ax.errorbar(range(6),values,yerr=sd,label=POLICY_LABELS[j],color=color,
                    marker=['s','^','D','v'][j],linestyle=['-','--','-.',':'][j],
                    linewidth=2.3 if j==0 else 1.8,markersize=5.5,
                    markerfacecolor='white' if j==0 else color,markeredgewidth=1.3,
                    elinewidth=.8,capsize=2,zorder=10 if j==0 else 3)
                last.append((values[-1],j))
            # Labels give all four n=100 values without covering the curves.
            last.sort();gap=(limits[1]-limits[0])*.115
            positions=[]
            for value,j in last:positions.append(max(value,positions[-1]+gap) if positions else value)
            overflow=max(0,positions[-1]-(limits[1]-.06*(limits[1]-limits[0])))
            positions=[v-overflow for v in positions]
            for (value,j),position in zip(last,positions):
                ax.annotate(f'{value:.2f}' if 'seconds' in key else f'{value:.1f}',xy=(5,value),xytext=(5.48,position),
                    fontsize=11.5,color='#'+COLORS[j],va='center',
                    arrowprops=dict(arrowstyle='-',lw=.65,color='#'+COLORS[j]))
            model_title='XGB 재학습' if model=='xgb' else 'PFN context 갱신'
            ax.set_title(DS_LABEL[ds]+' · '+model_title,fontsize=17,fontweight='bold',color='#'+NAVY,pad=8)
            ax.set_xlim(-.18,6.25);ax.set_ylim(*limits)
            ax.set_xticks(range(6),N,fontsize=12);ax.tick_params(axis='y',labelsize=12)
            ax.grid(axis='y',color='#E1E7EF',linewidth=.65)
            if col==0:ax.set_ylabel(ylabel,fontsize=12.5,labelpad=4)
            if row==1:ax.set_xlabel('신규 공격당 라벨 수 n (개)',fontsize=12.5,labelpad=4)
            if 'change' in key:ax.axhline(0,color='#708090',lw=.8)
    fig.legend(*axes[0,0].get_legend_handles_labels(),loc='lower center',bbox_to_anchor=(.5,-.005),ncol=4,frameon=False,fontsize=14,handlelength=2.2,columnspacing=2)
    fig.subplots_adjust(left=.068,right=.985,top=.92,bottom=.215,wspace=.21,hspace=.57)
    for ext in ['png','svg']:fig.savefig(ASSETS/f'{name}.{ext}',dpi=190)
    plt.close(fig)

def plot_tail():
    plt=setup_plot();fig,axes=plt.subplots(2,2,figsize=(12.2,4.65))
    groups=[('cic2018',['infiltration','web_attacks']),('toniot',['mitm','ransomware'])]
    for row,(ds,names) in enumerate(groups):
        for col,c in enumerate(names):
            ax=axes[row,col];maxv=0
            for j,model in enumerate(['xgb','pfn']):
                values=[float(cls(ds,100,method(model,p),c)['f1_pct_mean']) for p in POLICIES]
                sd=[float(cls(ds,100,method(model,p),c)['f1_pct_sd']) for p in POLICIES]
                maxv=max(maxv,max(np.array(values)+np.array(sd)))
                bars=ax.bar(np.arange(4)+(j-.5)*.34,values,width=.31,yerr=sd,capsize=2,
                    color='#111111' if j==0 else '#0072B2',label='XGB' if j==0 else 'PFN',error_kw={'elinewidth':.8})
                for b,v,e in zip(bars,values,sd):
                    ax.text(b.get_x()+b.get_width()/2,v+e+max(.15,maxv*.035),f'{v:.1f}',ha='center',fontsize=11.5)
            ax.set_xticks(range(4),POLICY_LABELS,fontsize=12)
            ax.set_ylim(0,max(2,maxv*1.34));ax.set_ylabel('F1 (%)',fontsize=12.5)
            ax.tick_params(axis='y',labelsize=12)
            support=int(cls(ds,100,'xgb',c)['support'])
            ax.set_title(f'{DS_LABEL[ds]} · {CLASS_LABEL[c]} · test {support:,}개',fontsize=15,fontweight='bold',color='#'+NAVY,pad=7)
            ax.grid(axis='y',color='#E1E7EF',lw=.6);ax.set_axisbelow(True)
    fig.legend(*axes[0,0].get_legend_handles_labels(),loc='lower center',bbox_to_anchor=(.5,-.005),ncol=2,frameon=False,fontsize=14)
    fig.subplots_adjust(left=.068,right=.985,top=.92,bottom=.17,wspace=.21,hspace=.58)
    for ext in ['png','svg']:fig.savefig(ASSETS/f'tail.{ext}',dpi=190)
    plt.close(fig)

def build_scenes():
    # 1. Problem first; the literature is evidence, not a claim of universal failure.
    items=[txt(.62,1.10,12,.38,'침입 탐지 시스템 (IDS) · 통신 기록의 특징으로 정상·공격 유형을 구분하는 시스템',20),
           txt(.62,1.84,12,1.05,'운영 시작 시 확보한 공격 라벨만으로\n이후 등장하는 공격까지 대응하기 어려운 문제',29,True,NAVY)]
    for i,(title,body) in enumerate([
        ('입력의 변화','네트워크·서비스·공격의 다양성\n학습 시점과 다른 트래픽의 등장'),
        ('제한된 공격 라벨','새 공격 조사 후 확보하는 사례\n충분한 학습 자료 확보의 어려움'),
        ('오탐과 갱신 부담','정상 트래픽에 대한 잘못된 경보\n새 사례 반영에 필요한 추가 계산')]):
        x=.62+4.23*i
        items += [box(x,3.23,3.64,1.66),box(x,3.23,3.64,.48,NAVY),txt(x+.16,3.33,3.3,.3,title,21,True,'FFFFFF'),txt(x+.16,3.94,3.30,.75,body,18)]
    items += [txt(.66,5.25,12,.40,'고정된 클래스의 한 번의 점수에 더해, 새 라벨을 반영한 뒤의 탐지 품질과 비용 확인',21,True,BLUE),
              label(.66,6.02,12,'Sommer · Paxson, Outside the Closed World: On Using Machine Learning',15,url=URLS['ids']),
              label(.66,6.34,12,'for Network Intrusion Detection, IEEE S&P 2010 · 학습 자료 부족·입력 변동성·오류 비용',15,url=URLS['ids'])]
    page('기존 IDS의 한계 · 새 공격과 제한된 라벨',items,'연구 문제 · 소수의 신규 공격 사례를 기존 탐지기에 반영하는 방법',notes='IDS를 먼저 정의한다. 선행연구는 주로 anomaly-based IDS의 운영상 어려움을 논의했으며 XGB나 TabPFN의 직접 비교가 아니다. 새 공격 라벨 확보 이후의 갱신을 본 연구 문제로 구체화한다. 신규 공격을 라벨 확보 전에 탐지했다는 주장은 하지 않는다. '+URLS['ids'])

    # 2. Official collection evidence appears after the operational problem.
    items=[txt(.62,1.12,12,.38,'CSE-CIC-IDS2018 · AWS 시험망에서 정상 동작과 계획된 공격을 실행·수집',22,True,NAVY),
           label(.62,1.70,12,'CSE · UNB CIC, CSE-CIC-IDS2018 on AWS, 2018 · 공식 문서의 공격 일정 Table 2',16,url=URLS['cic'])]
    labels=[('2/14','Brute force'),('2/15·16','DoS'),('2/20·21','DDoS'),('2/22·23','Web attacks'),('2/28·3/1','Infiltration'),('3/2','Bot')]
    for i,(date,name) in enumerate(labels):
        x=.62+i*2.10
        items += [box(x,2.37,1.81,1.03),txt(x+.10,2.52,1.61,.3,date,21,True,BLUE,align='center'),txt(x+.10,3.01,1.61,.27,name,15.5,align='center')]
        if i<5:items.append(arr(x+1.87,2.77,.17,.20))
    items += [txt(.65,3.83,12,.38,'수집된 공격 유형을 라벨 확보 단계로 구성 · 같은 train에서 단계별 사례 선택',22,True,NAVY)]
    for x,title,body in [(.65,'초기 사례 1,000개','정상 750개\nBrute force·DoS 각 125개'),
                         (4.90,'신규 공격당 n개','DDoS · Web · Infiltration · Bot\n각 단계의 신규 유형에서 n행 선택'),
                         (9.15,'단계별 탐지기 갱신','같은 사례로 XGB/PFN 갱신\n고정 test에서 품질·비용 비교')]:
        items += [box(x,4.52,3.55,1.38),txt(x+.15,4.67,3.25,.32,title,21,True,BLUE),txt(x+.15,5.22,3.25,.52,body,16.5)]
        if x<9:items.append(arr(x+3.76,5.08,.27,.24))
    items += [label(.65,6.35,12,'ToN-IoT도 자체 train에서 같은 절차로 독립 평가 · UNSW Canberra, The TON_IoT Datasets',16,url=URLS['ton'])]
    page('공격 수집 기록을 단계별 라벨 공급으로 활용',items,'평가 목적 · 새로운 공격 사례를 조금씩 확보할 때의 갱신 품질과 비용',notes='CIC의 수집 일정에서 나타난 공격 유형을 초기/후속 라벨 공급 단계로 사용한다. 날짜별 모든 행을 재생하는 것이 아니라 같은 정제 train에서 초기1000개와 후속유형당n개를 선택한다. ToN도 자체train에서독립적으로진행한다. 수집기록을어떻게사용하는지긍정적으로설명한다. 두벤치마크는시험망수집자료이고새로운특징행을합성하지않는다. '+URLS['cic']+' '+URLS['ton'])

    # 3. ICL and FT distinction at the actual input-output boundary.
    items=[txt(.62,1.08,12,.35,'새 공격 반영 시 가중치 재학습을 생략할 수 있는 in-context learning 활용',22,True,NAVY),
           label(.62,1.62,12,'Grinsztajn et al., TabPFN-3: Technical Report, arXiv 2026 · 합성 표 과제로 사전학습',16,url=URLS['pfn']),
           box(.62,2.20,3.43,1.25),txt(.82,2.36,3.02,.3,'Context · 라벨 사례',20,True,BLUE),txt(.82,2.90,3.02,.4,'기존 정상·공격 + 신규 공격\n특징과 정답을 함께 입력',17),
           box(.62,3.73,3.43,.94),txt(.82,3.87,3.02,.30,'Query · 분류할 통신 기록',20,True,BLUE),txt(.82,4.29,3.02,.26,'특징만 입력 · 정답은 평가에 사용',16.5),
           arr(4.25,2.79,.47,.31),arr(4.25,4.04,.47,.31),box(4.94,2.64,3.46,1.66,NAVY),txt(5.18,2.95,2.98,.43,'TabPFN v3',27,True,'FFFFFF'),txt(5.18,3.68,2.98,.35,'사전학습 가중치 고정',20,color='FFFFFF'),
           arr(8.61,3.30,.46,.31),box(9.28,2.64,3.43,1.66),txt(9.50,2.97,2.99,.38,'정상·공격별 확률',22,True,BLUE),txt(9.50,3.67,2.99,.31,'context를 참조한 추론',18),
           txt(.65,5.07,2.20,.35,'채택 이유',21,True,BLUE),txt(3.10,5.07,9.55,.38,'사전학습된 예측 규칙을 활용 · 라벨 사례 입력으로 현재 IDS 문제에 적응',19.5),
           txt(.65,5.70,2.20,.35,'검증할 이점',21,True,BLUE),txt(3.10,5.70,9.55,.38,'소량 라벨·제한된 보유량에서 품질을 유지하며 갱신 비용을 줄일 가능성',19.5),
           label(.65,6.40,12,'가중치 fine-tuning은 별도 최적화 절차 · 이번 실험은 미수행 / context 전처리·캐시 준비 시간은 측정',16.5)]
    page('TabPFN을 검토하는 이유 · 가중치 재학습 없는 적응',items,'구조적 특징 · 가중치 고정 / 실험으로 확인할 가설 · 소량 라벨에서의 품질과 갱신 비용',section='02 접근',notes='표 예측 규칙을 사전학습한 모델이므로 개별 IDS 문제에서 가중치를 다시 최적화하지 않고context와query로예측할수있다는것이채택이유다. 소량데이터품질과낮은갱신비용은확인할가설이다. Context는가중치외부의입력이고새공격라벨의의미를제공한다. 기존클래스사례도필요하다. 캐시준비비용을측정하며추론비용이낮다는주장은하지않는다. '+URLS['pfn'])

    # 4. Purpose must precede the dataset schedule.
    items=[txt(.65,1.14,12,.90,'새 공격의 라벨 사례가 소량 확보될 때\n누적 재학습과 context 갱신의 탐지 품질·계산 비용 비교',26,True,NAVY),
           box(.65,2.24,12.02,.54),txt(.84,2.35,11.65,.31,'공통 입력 · 초기 1,000개 + 단계별 신규 공격 라벨 / 모든 방법에 동일 사례 제공',20,True,NAVY)]
    items += simple_table(.65,3.03,[4.0,4.0,4.02],['실제 사용하는 사례','XGB','TabPFN'],
        [['누적 · 지금까지 제공한 전부','누적 사례로 재학습','누적 사례로 context 갱신'],
         ['1,000개 · FIFO / 무작위 / 유사도','같은 버퍼 행으로 재학습','같은 버퍼 행으로 context 갱신']],
        [.60,.60],18)
    for x,h,b in [(.65,'탐지 품질','전체·신규 공격별 F1'),(4.87,'기존 클래스 유지','동일 평가 사례의 갱신 전후 변화'),(9.09,'갱신 비용','초기 준비·추론과 분리한 시간')]:
        items += [txt(x,5.12,3.61,.35,h,22,True,BLUE),txt(x,5.69,3.61,.38,b,18)]
    items += [label(.65,6.44,12,'평가 범위 · 신규 공격의 라벨 확보 후 분류 / 라벨 없는 미지 공격 탐지와 네트워크 간 전이는 별도 과제',16.5)]
    page('평가 목적 · 같은 사례로 비교하는 모델과 보유 방식',items,'같은 정책·같은 seed에서 행 ID까지 일치 · 모델 차이와 사례 구성 효과의 분리',section='03 평가',notes='누적은XGB/PFN같은사례를사용한다. 버퍼역시FIFO,무작위,유사도각각두모델에정확히같은행ID와순서를제공한다. 따라서모델효과와사례선택효과를구분할수있다. 신규라벨수와보유량은별개조건이다. 최종성능뿐아니라후속갱신시간합계를측정하므로XGB는매단계재학습한다.')

    # 5. A/B ambiguity resolved with separate source -> train/test -> provided data lanes.
    items=[txt(.62,1.10,12,.40,'각 벤치마크 안에서 train의 행을 선택 · CIC와 ToN 사이의 행 결합 없음',23,True,NAVY)]
    for i,(ds,test) in enumerate([('CIC2018','3,042,473'),('ToN-IoT','2,271,723')]):
        y=1.93+2.08*i
        items += [box(.62,y,1.68,1.44),txt(.77,y+.29,1.38,.40,ds,22,True,NAVY),txt(.77,y+.88,1.38,.30,'수집 벤치마크',15),arr(2.43,y+.57,.31,.27),
                  box(2.91,y,3.27,1.44),txt(3.08,y+.14,2.92,.34,'정제 후 train / test 분리',19,True,BLUE),txt(3.08,y+.72,2.92,.55,'상충 라벨 제거\n학습·평가 행 분리',16.5),
                  arr(6.33,y+.33,.28,.23),arr(6.33,y+1.00,.28,.23),
                  box(6.79,y,5.92,.67),txt(6.96,y+.11,5.55,.44,'Train → 초기 1,000개 + 신규 공격당 n개 선택',19,True,NAVY),
                  box(6.79,y+.80,5.92,.64),txt(6.96,y+.94,5.55,.32,f'Test → 고정 {test}개 · 학습/context에 미사용',18)]
    items += [txt(.66,5.82,12,.36,'제공 사례 · 같은 train에서 단계별로 선택한 라벨 사례 / 별도 출처의 데이터셋 생성 없음',19.5,True,NAVY),
              txt(.66,6.38,12,.36,'실제 사용 사례 · 누적은 지금까지 제공한 전부 / 버퍼는 삭제 후 남긴 사례와 신규 사례 1,000개',19.5,True,TEAL)]
    page('평가 방법 · 데이터 출처와 사례 선택',items,'수집된 특징 행을 선택한 라벨 제공 실험 · 새로운 특징 행의 합성 없음',section='03 평가',notes='청중이 말한 원본 A는 각 출처 벤치마크이고 현재 B는 그 train에서 해당 단계까지 제공한 사례 집합이다. 별도의 출처 B를 생성하거나 A의 공격을 다른 데이터셋에 붙이지 않는다. CIC에서 CIC train/test를, ToN에서 ToN train/test를 각각 사용한다. 두 벤치마크는 시험망 수집 자료다. 행 ID는 분리되어 있으나 동일 특징 벡터가 남을 수 있으므로 완전한 특징 중복 제거를 주장하지 않는다. 기존 중복 특징 제외 진단은 별도 결과에 저장되어 있다. '+URLS['cic']+' '+URLS['ton'])

    # 6. Actual stage allocations, all budgets, and one worked example.
    items=[txt(.62,1.09,12,.39,'n = 5 · 10 · 20 · 30 · 50 · 100개 / 신규 공격 클래스 · seed 43·44·45',23,True,NAVY),
           label(.62,1.60,12,'n은 확보 가능한 라벨 수의 실험 조건 · 하루 공격량이나 최적 표본 수를 뜻하지 않는 값',17)]
    cic=[['1','Benign / Brute force / DoS','750 / 125 / 125','1,000'],['2','DDoS','n','1,000+n'],['3','Web attacks','n','1,000+2n'],['4','Infiltration','n','1,000+3n'],['5','Bot','n','1,000+4n']]
    ton=[['1','Benign / DoS / Scanning','750 / 125 / 125','1,000'],['2','DDoS / Injection','각 n','1,000+2n'],['3','Password','n','1,000+3n'],['4','XSS','n','1,000+4n'],['5','Backdoor / Ransomware','각 n','1,000+6n'],['6','MITM','n','1,000+7n']]
    # Multiline stage 1 keeps the class-specific allocation explicit and aligned.
    cic[0]=['1','Benign\nBrute force\nDoS','750\n125\n125','1,000']
    ton[0]=['1','Benign\nDoS\nScanning','750\n125\n125','1,000']
    ton[1][1]='DDoS\nInjection';ton[1][2]='n\nn';ton[4][1]='Backdoor\nRansomware';ton[4][2]='n\nn'
    items += [txt(.62,1.99,5.85,.36,'CIC2018',22,True,BLUE),txt(6.87,1.99,5.85,.36,'ToN-IoT',22,True,BLUE)]
    widths=[.60,1.88,.90,1.42,1.06]
    cic=[r+['1,000'] for r in cic];ton=[r+['1,000'] for r in ton]
    items += simple_table(.62,2.45,widths,['단계','제공 클래스','제공 수','누적 보유','버퍼'],cic,[.90,.48,.48,.48,.48],14,emphasis_cols=(4,))
    items += simple_table(6.87,2.45,widths,['단계','제공 클래스','제공 수','누적 보유','버퍼'],ton,[.90,.60,.38,.38,.60,.38],14,emphasis_cols=(4,))
    items += [label(.65,6.54,12,'제공 수는 클래스당 사례 수 · 누적·버퍼 모두 XGB/PFN 공통 · 표의 보유량은 매 단계 실제 사용 행 수',15.5)]
    page('평가 방법 · 단계별 공급량과 실제 보유량',items,'같은 신규 사례 전량 반영 · 누적 보유량 증가와 1,000개 버퍼 유지 비교',section='03 평가',notes='초기 benign 750, 공격 2종 각125는 통제 조건이며 운영 분포나 최적값이라고 주장하지 않는다. 단계는 원본의 공격 등장 순서를 참고한 라벨 공개 순서이고 raw timestamp replay가 아니다. n=100은 6개 조건 중 상한이며 대표 최적값이 아니다. 각 단계에서 누적 보유량과 버퍼 실제 크기를 별도 열로 표시한다. 클래스별 표본 동일 조건에서도 단계별 새 클래스 수에 따라 신규 총량이 달라진다.')

    # 7. Actual data blocks -> retention choice -> model work -> same test.
    items=[txt(.62,1.09,12,.36,'예시 · CIC 3단계, 신규 공격당 100개 · 같은 사례를 서로 다른 방식으로 갱신',21,True,NAVY),
           label(.65,1.73,1.65,'라벨 확보 순서',17),
           txt(2.35,1.65,5.20,.30,'초기 정상 750개 · 공격 250개',17,True,NAVY),
           txt(8.65,1.65,.72,.30,'DDoS',15,True,DATA_COLORS[2],align='center'),
           txt(9.36,1.65,.62,.30,'Web',15,True,DATA_COLORS[3],align='center')]
    items += sample_bar(2.35,2.12,7.62,.43,[750,250,100,100],show_counts=True,size=15)
    items += [arr(4.26,2.70,.35,.28,down=True),
              outline(.65,3.07,4.80,2.94,'9AA7B4'),txt(.85,3.23,4.40,.31,'보유 방식 선택 · XGB/PFN 공통',19,True,NAVY)]
    c=reference_counts('cic2018',100,'accumulate',2)
    f=reference_counts('cic2018',100,'random',2)
    grouping=lambda a:[a['benign'],a['brute_force']+a['dos'],a['ddos'],a['web_attacks']]
    items += [txt(.85,3.84,1.05,.30,'누적',18,True,BLUE)]
    items += sample_bar(2.04,3.82,3.16,.42,grouping(c),max_total=1200,show_counts=False)
    items += [label(2.04,4.32,3.16,'1,200개 · 기존 사례 모두 보유',15)]
    items += [txt(.85,4.92,1.05,.30,'버퍼',18,True,TEAL)]
    items += sample_bar(2.04,4.90,3.16,.42,grouping(f),max_total=1200,show_counts=False)
    items += [label(2.04,5.40,3.16,'1,000개 · 기존 삭제 + 신규 전량',15),
              label(.85,6.19,4.5,'버퍼 예시: 무작위, seed 43 · 막대 길이 = 행 수',14.5)]
    # A shared fork means either retention choice is supplied to BOTH models.
    items += [rule(5.45,4.52,.25,BLUE),box(5.69,3.80,.018,1.74,BLUE),
              arr(5.69,3.67,.40,.26),arr(5.69,5.39,.40,.26),
              label(6.17,2.80,4.12,'XGB · 매 단계 모델 재학습',19),
              outline(6.16,3.25,4.15,1.05,'111111',dashed=True),
              txt(6.37,3.44,3.74,.32,'트리 구조·리프 출력 학습',22,True,'111111'),
              txt(6.37,3.95,3.74,.25,'해당 단계의 보유 사례로 새 모델 구성',16),
              label(6.17,4.58,4.12,'PFN · 사전학습 가중치 유지',19),
              outline(6.16,5.04,1.72,1.05,TEAL,dashed=True),
              txt(6.29,5.20,1.46,.62,'context\n캐시 갱신',18,True,TEAL,align='center'),
              arr(8.00,5.44,.28,.24),box(8.42,5.04,1.89,1.05,NAVY),
              txt(8.55,5.20,1.63,.62,'TabPFN\n가중치 고정',18,True,'FFFFFF',align='center'),
              arr(10.42,3.67,.30,.24),arr(10.42,5.43,.30,.24),
              box(10.84,3.45,1.65,.69),txt(10.96,3.63,1.41,.29,'예측·평가',18,True,BLUE,align='center'),
              box(10.84,5.21,1.65,.69),txt(10.96,5.39,1.41,.29,'예측·평가',18,True,BLUE,align='center'),
              outline(10.84,1.93,1.86,.83,BLUE),txt(10.97,2.10,1.60,.45,'같은 고정 test',17,True,BLUE,align='center'),
              box(12.75,2.34,.016,3.22,BLUE),rule(12.70,2.34,.065,BLUE),
              dict(kind='left',x=12.49,y=3.67,w=.28,h=.24,fill=BLUE),
              dict(kind='left',x=12.49,y=5.43,w=.28,h=.24,fill=BLUE),
              label(6.17,6.42,6.50,'점선: 매 단계 갱신 영역 · 전체 test는 마지막 모델에만 평가',15)]
    page('평가 흐름 · 사례 보유에서 모델 갱신과 test까지',items,'각 단계: 갱신 시간·고정 패널 확인 / 마지막 단계: 전체 test 품질 평가',section='03 평가',notes='블록은CIC의초기정상750,초기공격250,DDoS100,Web100을나타내며길이는실제행수에비례한다. 현재버퍼예시는seed43무작위정책의실제3단계구성이다. 누적과버퍼중선택한동일사례를XGB와PFN모두에게제공한다. XGB는각단계새트리모델을학습하고PFN은context와캐시를갱신하며가중치를고정한다. 고정test는학습에입력되지않고모델구성후예측평가에사용한다. 최종품질만필요하면XGB마지막학습한번으로가능하지만갱신시간합계와유지를위해매단계학습한다. 전체test는마지막에만평가하며각단계는고정패널진단이다.')

    # 8. Retention changes class composition as well as the number of examples.
    items=[txt(.62,1.10,12,.37,'예시 · ToN 최종 단계, 신규 공격당 100개, seed 43 · 각 막대는 실제 보유 행 수',21,True,NAVY)]
    for x,color,name in [(2.60,DATA_COLORS[0],'정상'),(4.75,DATA_COLORS[1],'초기 공격 2종'),(7.95,DATA_COLORS[2],'이후 신규 공격 7종')]:
        items += [box(x,1.74,.22,.20,color),txt(x+.34,1.70,1.70 if name=='정상' else 2.70,.30,name,17)]
    compositions=[('초기','initial',{'benign':750,'dos':125,'scanning':125})]
    compositions += [(name,policy,reference_counts('toniot',100,policy)) for name,policy in [('누적','accumulate'),('FIFO','fifo'),('무작위','random'),('유사도','similar')]]
    for i,(name,policy,c) in enumerate(compositions):
        y=2.30+i*.70
        normal=c['benign'];old=c['dos']+c['scanning'];new=sum(c.values())-normal-old;total=sum(c.values())
        items += [txt(.65,y+.04,1.48,.30,name,20,True,NAVY)]
        items += sample_bar(2.38,y,8.00,.43,[normal,old,new],max_total=1700,show_counts=True,size=15)
        items += [txt(10.63,y+.04,2.05,.30,f'{total:,}개',19,True,NAVY,align='right')]
    items += [txt(.65,6.11,12,.36,'정상 비율 · 초기 75.0% → 누적 44.1% / FIFO 22.4% / 무작위 35.1% / 유사도 37.3%',20,True,BLUE),
              label(.65,6.59,12,'삭제 기준 · FIFO: 먼저 확보한 행 / 무작위: 임의의 행 / 유사도: 신규와 가까운 행 · 클래스별 최소 1개 보존',15.5)]
    page('버퍼 교체에 따른 클래스 구성 변화',items,'버퍼는 표본 수와 클래스 비율을 함께 변경 · 같은 구성의 XGB 대조군으로 효과 분리',section='03 평가',notes='성능추정치는3seed평균을사용하지만이페이지는실제행구성을이해하기위한seed43예시다. 정상/초기두공격/후속일곱공격으로집계했으며새공격개별클래스수가모두같다는뜻은아니다. ToN FIFO최종은정상224,초기76,신규700;무작위351/112/537;유사도373/148/479. FIFO는오래된네트워크특성을판정하지않고편입순서로삭제한다. 실제운영에서는기존정상공격도계속들어올수있지만본실험은신규공격라벨을추가하는통제조건이다.')

    # 9. Overall quality first; all eight model/policy combinations, three seeds.
    items=[txt(.62,1.09,12,.38,'최종 모델 · 전체 고정 test · 모든 클래스 F1의 산술평균',22,True,NAVY),
           label(.62,1.60,12,'평균 ± 표준편차 · seed 43·44·45 · 같은 데이터셋의 두 패널은 같은 축 · 우측 숫자: n=100',16.5),
           img(ASSETS/'overall.png',.55,1.97,12.24,4.72)]
    page('결과 · 전체 클래스의 최종 분류 성능',items,'같은 사례에서 모델 차이 비교 / 같은 모델에서 누적·버퍼 정책 차이 비교',section='04 결과',notes='CIC7개와ToN10개클래스를동일비중으로평균한다. 전체test는CIC3042473,ToN2271723이며단계별예측을합치지않는다. XGB와PFN을분리한패널이지만데이터셋별y축은일치한다. 오차막대는고정분할에서3seed표본표준편차이며신뢰구간이나동등성검정이아니다. 우측은모든정책의n100평균이며최적조건으로선택한값이아니다.')

    # 10. New-class quality from the same final confusion matrices.
    items=[txt(.62,1.09,12,.38,'최종 모델 · 전체 test에서 계산한 신규 공격 클래스별 F1의 산술평균',21,True,NAVY),
           label(.62,1.60,12,'CIC 4개 · ToN 7개 신규 공격 · 정상·초기 공격에서 발생한 오탐도 F1에 포함 · 3 seed 평균 ± 표준편차',16.5),
           img(ASSETS/'new_attacks.png',.55,1.97,12.24,4.72)]
    page('결과 · 신규 공격 라벨 수에 따른 탐지 품질',items,'적은 라벨로 반영되는 정도와 클래스별 실패를 함께 확인',section='04 결과',notes='같은최종전체test혼동행렬에서초기이후공격F1만평균한다. 단계평균아님. CIC신규는ddos,web_attacks,infiltration,bot이고ToN은ddos,injection,password,xss,backdoor,ransomware,mitm이다. 초기정상공격은평균항목에서제외하지만그행에서신규로오분류한FP는해당F1에포함한다. 신규그룹은test희소클래스그룹과다르다. XGB고정버퍼도같은행을사용한다.')

    # 11. Retention on a fixed initial-class panel, not the entire test.
    items=[txt(.62,1.09,12,.38,'초기 정상·공격 2종의 같은 평가 표본 · 최종 − 초기 Macro-F1 (%p)',21,True,NAVY),
           label(.62,1.60,12,'클래스당 최대 2,000개 고정 test 표본 · 0 아래는 초기보다 성능 하락 · 3 seed 평균 ± 표준편차',16.5),
           img(ASSETS/'retention.png',.55,1.97,12.24,4.72)]
    page('결과 · 갱신 이후 초기 클래스 성능 변화',items,'신규 클래스 평균과 별도로 확인하는 기존 클래스 유지 · 전체 test와 구분된 진단',section='04 결과',notes='초기3개클래스행만의고정패널F1이다. 이후공격에서초기클래스로향하는FP는포함하지않으므로전체testF1과다른모집단이다. 전체현재클래스joint오탐및직전클래스유지진단은원시결과에별도보존했다. 전체testF1상승과이진단하락은동시에가능하다. 미래클래스query특징을포함할수있는고정query구성으로오프라인진단하고정답은학습하지않는다.')

    # 12. Match update timing boundaries; disclose sessions without shifting the focus.
    items=[txt(.62,1.09,12,.38,'2단계 이후 갱신 시간 합계 · 초기 준비와 모든 추론 제외',23,True,NAVY),
           label(.62,1.60,12,'사례 선택·삭제 + XGB 재학습 또는 PFN 전처리·캐시 준비 · 같은 장비, 별도 실행 · 3 seed 평균 ± 표준편차',16),
           img(ASSETS/'cost.png',.55,1.97,12.24,4.72)]
    page('결과 · 같은 사례를 반영하는 후속 갱신 비용',items,'가중치 학습 생략 효과를 실제 준비 시간으로 평가 · 작은 context 범위의 측정값',section='04 결과',notes='기존EXP84의PFN및누적XGB시간과새EXP85버퍼XGB시간을비교하며실행세션이다르다. 모두같은장비설정및동일시간범위이며GPU동기화로측정했다. PFN은재실행하지않았고XGB는정책별삭제계산을실행해시간을다시측정했다. 초기준비,모든추론,체크포인트와검증은제외한다. 전체test는마지막에만평가한다. 서로다른데이터셋은단계수도다르다. 이그림의비용차이를전체운영비용으로해석하지않는다.')

    # 13. Tail is defined by test support, not by the label budget n.
    r=cls('toniot',100,'pfn_random','ransomware')
    items=[txt(.62,1.09,12,.38,'n = 100 · test에서 가장 희소한 공격 2종씩 · 3 seed 평균 ± 표준편차',21,True,NAVY),
           img(ASSETS/'tail.png',.55,1.65,12.24,4.65),
           label(.65,6.43,12,f"ToN Ransomware · PFN 무작위 평균 재현율 {float(r['recall_pct_mean']):.1f}% / 정밀도 {float(r['precision_pct_mean']):.1f}% · 다음 분석: 오탐의 출처와 보정 효과",16.5)]
    page('결과 해석 · 같은 버퍼에서도 남는 희소 공격 오류',items,'낮은 F1의 원인 분리 · 사례 삭제 효과와 모델의 오탐 문제를 구분해 개선',section='04 결과',notes='각데이터셋test에서가장희소한공격2종을support내림차순으로제시한다. n100은스윕상한진단예시이며모든n의클래스별결과는보고서에보존한다. XGB/PFN같은버퍼조건으로비교한다. PFN누적에서도낮은tailF1이남으므로버퍼삭제만으로원인을단정하지않는다. 높은재현율과낮은정밀도를구분하여정상에서발생한오탐,클래스비율변화및모델분리능력을분석한다. 가중치적응이나보정효과는아직입증하지않았다.')

    # Class tables retain every model/policy pair, means from all three seeds.
    for ds in DS:
        names=sorted({r['name'] for r in CLASSES if r['dataset']==ds},key=lambda c:-int(cls(ds,100,'xgb',c)['support']))
        rows=[[CLASS_LABEL[c],f"{int(cls(ds,100,'xgb',c)['support']):,}"]+[f"{float(cls(ds,100,m,c)['f1_pct_mean']):.2f}" for m in METHODS] for c in names]
        items=[label(.62,1.12,12,'n = 100 · 마지막 모델 전체 test · F1 (%)의 3 seed 평균 · 평가 표본 수 내림차순',18),
               txt(4.06,1.65,4.24,.30,'XGB 재학습',20,True,'111111',align='center'),
               txt(8.30,1.65,4.24,.30,'PFN context 갱신',20,True,BLUE,align='center')]
        items += simple_table(.62,2.15,[2.04,1.40]+[1.06]*8,['클래스','test 수']+POLICY_LABELS*2,rows,[.40]*len(rows),14)
        items += [label(.65,6.68,12,'같은 정책·같은 seed에서 XGB/PFN의 행 ID 일치 · 전체 n의 수치와 표준편차는 실험 보고서에 수록',15)]
        page('부록 · '+DS_LABEL[ds]+' 클래스별 성능',items,section='부록',notes='검증된exp85_per_class_summary.csv의3seed평균F1과support를사용한다. n100은설명용상한조건이며최적값이아니다. 왼쪽네정책은XGB,오른쪽네정책은PFN이다. 표준편차와precision/recall은보고서에보존했다.')

    arch_path=SLIDES/'_support/qa/kor/1015_labmeeting_qa/v3_architecture_page.png'
    first=[img(arch_path,.62,1.09,3.95,5.88,[.20,.082,.19,.285])]
    for i,(h,b) in enumerate([('1  입력값 표현','통신 기록의 각 값을 벡터로 변환\n결측 표시와 context 라벨 정보 반영'),('2  열의 분포 요약','여러 기록에서 같은 특징을 살펴\n그 특징의 분포를 표현'),('3  한 기록의 특징 통합','여러 열을 행 단위 표현으로 압축\n다음 단계에서 사례끼리 비교')]):
        y=1.39+i*1.72;first += [txt(4.98,y,7.72,.40,h,23,True,BLUE),txt(5.25,y+.67,7.18,.78,b,21)]
    page('부록 · TabPFN v3의 기록 표현',first,section='부록',notes='Grinsztajn etal2026 Figure5 원도식. 그림캡션은사용자요청에따라표시하지않는다. 입력연산단계이며IDS가중치학습단계를뜻하지않는다. '+URLS['arch'])
    second=[img(arch_path,.62,1.09,3.95,5.88,[.20,.082,.19,.285]),
            txt(4.98,1.40,7.72,.40,'4  Context를 참조한 예측',23,True,BLUE),txt(5.25,2.03,7.18,.78,'query가 라벨이 있는 기록을 참조\nattention으로 참고 비중 계산',21),
            txt(4.98,3.12,7.72,.40,'5  라벨을 이용한 확률 출력',23,True,BLUE),txt(5.25,3.75,7.18,.78,'context 라벨을 가중합하는 decoder\n기존·신규 클래스 사례 모두 필요',21),
            box(4.98,5.10,7.72,1.57),txt(5.20,5.29,7.25,.35,'별도의 가중치 fine-tuning',22,True,NAVY),txt(5.20,5.86,7.25,.65,'라벨 사례를 context와 query로 나눠 query 오차 역전파\n이번 실험은 미수행 · context/cache 갱신만 사용',19)]
    page('부록 · TabPFN v3의 추론과 fine-tuning',second,section='부록',notes='원문2.2manyclassdecoder는onehotcontextlabel을value로사용한다. context밖의class에대한지도학습다중클래스확률을만들기위해관련label사례가필요하다. FT는별도가중치최적화. 이번실험에서는weightdigest동일성을확인했고FT시간0. '+URLS['arch'])

def validate_scenes():
    issues=[]
    for page_no,scene in enumerate(SCENES,1):
        texts=[]
        for i,item in enumerate(scene['items']):
            x,y,w,h=[item[k] for k in ('x','y','w','h')]
            if min(x,y)<0 or x+w>W+.01 or y+h>H+.01:
                issues.append(dict(page=page_no,item=i,kind='outside_slide',text=item.get('text','')))
            if item['kind']!='text':continue
            font=ImageFont.truetype(str(FONT_BOLD if item['bold'] else FONT),round(item['size']*4))
            # Font widths measured at 4x points, 1 inch = 72 points.
            max_width=max((font.getlength(s) for s in item['text'].splitlines()),default=0)/4/72
            if max_width>w-.005:
                issues.append(dict(page=page_no,item=i,kind='text_width',actual=round(max_width,3),width=w,text=item['text']))
            if re.search(r'합니다|입니다|습니다|—',item['text']):
                issues.append(dict(page=page_no,item=i,kind='style',text=item['text']))
            texts.append((i,item))
        for j,(ia,a) in enumerate(texts):
            for ib,b in texts[j+1:]:
                dx=min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
                dy=min(a['y']+a['h'],b['y']+b['h'])-max(a['y'],b['y'])
                if dx>.015 and dy>.015:
                    issues.append(dict(page=page_no,kind='text_overlap',items=[ia,ib],texts=[a['text'],b['text']],overlap=[dx,dy]))
    (QA/'layout_validation.json').write_text(json.dumps(dict(passed=not issues,issues=issues,slides=len(SCENES),intentional_overlaps=['Text over its background rectangles','Picture content within image rectangles']),ensure_ascii=False,indent=2)+'\n')
    if issues:
        print(json.dumps(issues,ensure_ascii=False,indent=2))
        raise RuntimeError(f'{len(issues)} layout issues')

def render():
    validate_scenes()
    narrations=['운영을 시작할 때 없었던 공격의 라벨 사례가 나중에 소량 확보되는 상황을 다룹니다. 기존 정상과 공격을 계속 구분하면서 새 사례를 반영해야 합니다. 자료 부족과 입력 변화를 이 갱신 상황에서 살펴보겠습니다.', 'CIC는 시험망에서 날짜별로 서로 다른 공격을 실행해 수집했습니다. 우리는 이 공격 유형을 초기와 후속 단계로 나누고, 같은 train에서 초기 천 개와 신규 유형당 n개를 선택합니다. 이 사례를 두 모델에 똑같이 제공하고 별도로 둔 test로 평가합니다. ToN에도 자체 자료 안에서 같은 절차를 적용합니다.', 'TabPFN은 개별 IDS에 맞춰 가중치를 다시 학습하지 않고도 라벨 사례를 입력으로 받아 예측할 수 있다는 점 때문에 검토합니다. 사전학습한 것은 특정 공격 이름이 아니라 표 데이터의 예측 규칙입니다. Context가 현재 IDS의 특징과 정답 관계를 제공합니다. 소량 라벨에서도 품질을 유지하고 갱신 비용을 줄일 수 있는지는 다음 실험에서 확인합니다.', '모델과 사례 보유 방식을 분리해서 비교합니다. 누적뿐 아니라 FIFO, 무작위, 유사도 버퍼도 XGB와 PFN에 정확히 같은 행을 제공합니다. 이 대조군이 있어야 버퍼 자체의 효과를 PFN의 효과로 오해하지 않습니다. 최종 품질과 기존 성능 유지, 후속 갱신 비용을 확인합니다.', 'CIC와 ToN은 각각 독립적으로 평가합니다. 정제 train에서 초기와 새 공격 사례를 선택합니다. 현재 데이터셋이라는 표현은 다른 출처의 자료가 아니라 단계별로 제공된 라벨 사례 집합을 뜻합니다. 일부만 선택하는 이유는 조사로 새 공격 라벨을 조금씩 확보하는 상황을 통제하기 위해서입니다.', '초기는 정상 750개, 공격 두 종류 각 125개입니다. 이후 새 클래스당 다섯 개부터 백 개까지 여섯 라벨 수를 비교했습니다. 누적 보유량은 계속 증가하고 버퍼는 천 개로 유지됩니다. 두 방식 모두 XGB와 PFN에 적용하며 세 seed를 사용했습니다.', '그림 위쪽은 라벨 확보 순서입니다. CIC 세 번째 단계에서 초기 천 개에 DDoS 백 개와 Web 백 개가 들어옵니다. 누적은 천이백 개를 쓰고 버퍼는 기존 행을 삭제해 천 개를 유지합니다. 선택한 같은 사례가 오른쪽 두 모델에 들어갑니다. 점선 안에서 XGB는 모델을 다시 학습하고 PFN은 context와 캐시를 준비합니다. 오른쪽의 같은 test는 학습에 쓰이지 않고 예측 평가에 사용됩니다.', '버퍼 크기만 같은 것이 충분하지는 않습니다. ToN의 예시를 보면 삭제 정책에 따라 정상과 초기 공격이 남는 비율도 달라집니다. 이 그림은 seed 43의 실제 행 수이며 결과 평균과 구분한 예시입니다. 그래서 같은 정책에서는 XGB와 PFN에 동일한 행을 줘야 모델 차이를 판단할 수 있습니다.', '먼저 전체 test에서 마지막 모델의 Macro-F1을 확인합니다. XGB가 왼쪽, PFN이 오른쪽이고 같은 데이터셋의 축을 맞췄습니다. 선 색은 보유 정책이며 오차막대는 세 seed의 표준편차입니다. 우측 숫자는 모든 정책의 백 개 조건 평균입니다.', '같은 최종 예측에서 신규 공격 클래스의 F1만 골라 평균합니다. 정상과 초기 공격은 평균 항목에서 빠지지만 그 행에서 새 공격으로 오분류한 경우는 해당 공격의 F1에 반영됩니다. 단계 점수를 평균한 값이 아닙니다.', '유지는 초기 세 클래스에 해당하는 고정 평가 행에서 최종과 초기 점수의 차이로 확인합니다. 전체 test F1과 다른 진단입니다. 최종 전체 점수만 보면 드러나지 않는 기존 클래스 변화를 살펴봅니다.', '초기 준비와 모든 추론을 제외하고, 두 번째 단계부터 사례 선택과 삭제 및 모델 갱신 시간을 합했습니다. XGB 버퍼도 모든 단계에서 재학습했습니다. 기존 PFN과 새 XGB는 같은 장비에서 다른 세션에 측정한 값입니다.', '마지막으로 test에서 가장 희소한 공격 두 종류씩을 확인합니다. 같은 버퍼에서도 모델별 오탐과 탐지의 차이가 남습니다. 삭제하지 않은 PFN 누적의 낮은 F1까지 버퍼 탓으로 돌릴 수는 없습니다. 다음 개선은 오탐의 출처와 클래스 구성, 모델의 분리 능력을 구분하는 분석에 연결합니다.']
    for scene,narration in zip(SCENES,narrations):
        scene['notes']='발표 설명\n'+narration+'\n\n근거와 해석 범위\n'+scene['notes']
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    for scene in SCENES:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text=scene['notes']
        for it in scene['items']:
            x,y,w,h=[Inches(it[k]) for k in ('x','y','w','h')]
            if it['kind']=='text':
                shape=slide.shapes.add_textbox(x,y,w,h);f=shape.text_frame;f.clear()
                f.word_wrap=False;f.auto_size=MSO_AUTO_SIZE.NONE;f.vertical_anchor=MSO_ANCHOR.TOP
                f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
                for n,line in enumerate(it['text'].splitlines()):
                    p=f.paragraphs[0] if n==0 else f.add_paragraph()
                    p.space_before=p.space_after=Pt(0);p.line_spacing=1.16
                    p.alignment={'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER}[it['align']]
                    r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(it['size']);r.font.bold=it['bold'];r.font.color.rgb=RGBColor.from_string(it['color'])
                    prop=r._r.get_or_add_rPr();prop.set('lang','ko-KR')
                    for tag in ['a:ea','a:cs']:
                        e=OxmlElement(tag);e.set('typeface','Pretendard');prop.append(e)
                    if it.get('url'):r.hyperlink.address=it['url']
            elif it['kind']=='image':
                s=slide.shapes.add_picture(it['path'],x,y,width=w,height=h)
                s.crop_left,s.crop_top,s.crop_right,s.crop_bottom=it['crop']
            else:
                typ={'arrow':MSO_SHAPE.RIGHT_ARROW,'left':MSO_SHAPE.LEFT_ARROW,'down':MSO_SHAPE.DOWN_ARROW}.get(it['kind'],MSO_SHAPE.RECTANGLE)
                s=slide.shapes.add_shape(typ,x,y,w,h);s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(it['fill']);s.line.fill.background()
                if it['kind']=='outline':
                    s.line.color.rgb=RGBColor.from_string(it['border']);s.line.width=Pt(1.3)
                    if it.get('dashed'):s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
            # PowerPoint's default rectangle theme otherwise references a shadow.
            element=slide.shapes[-1]._element
            for effect in element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):
                effect.set('idx','0')
            sp=element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
            if sp is not None:
                for child in list(sp):
                    if etree.QName(child).localname in ['effectLst','effectDag']:sp.remove(child)
                sp.append(OxmlElement('a:effectLst'))
    candidate_pptx=QA/'meeting_no_shadow.pptx'
    prs.save(candidate_pptx)
    with zipfile.ZipFile(candidate_pptx) as src,zipfile.ZipFile(OUT.with_suffix('.pptx'),'w',zipfile.ZIP_DEFLATED) as dst:
        for entry in src.infolist():
            data=src.read(entry.filename)
            if entry.filename.startswith('ppt/theme/') and entry.filename.endswith('.xml'):
                tree=etree.fromstring(data)
                for shadow in tree.xpath('//*[local-name()="outerShdw" or local-name()="innerShdw" or local-name()="prstShdw"]'):
                    shadow.getparent().remove(shadow)
                data=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
            dst.writestr(entry,data)
    candidate_pptx.unlink()
    pages=[]
    for si,scene in enumerate(SCENES):
        tags=[]
        for i,it in enumerate(scene['items']):
            style=f'left:{it["x"]}in;top:{it["y"]}in;width:{it["w"]}in;height:{it["h"]}in;'
            if it['kind']=='text':
                style+=f'font-size:{it["size"]}pt;font-weight:{700 if it["bold"] else 400};color:#{it["color"]};text-align:{it["align"]};'
                content=html.escape(it['text'])
                if it.get('url'):content=f'<a href="{html.escape(it["url"],quote=True)}">{content}</a>'
                tags.append(f'<div id="s{si}i{i}" class="text" style="{style}">{content}</div>')
            elif it['kind']=='image':
                l,t,r,b=it['crop'];iw=it['w']/(1-l-r);ih=it['h']/(1-t-b)
                tags.append(f'<div class="clip" style="{style}"><img src="{Path(it["path"]).as_uri()}" style="width:{iw}in;height:{ih}in;left:{-l*iw}in;top:{-t*ih}in"></div>')
            else:
                style+=f'background:#{it["fill"]};'
                if it['kind']=='outline':style+=f'border:1.3pt {"dashed" if it.get("dashed") else "solid"} #{it["border"]};'
                if it['kind']=='arrow':style+='clip-path:polygon(0 25%,60% 25%,60% 0,100% 50%,60% 100%,60% 75%,0 75%);'
                if it['kind']=='left':style+='clip-path:polygon(100% 25%,40% 25%,40% 0,0 50%,40% 100%,40% 75%,100% 75%);'
                if it['kind']=='down':style+='clip-path:polygon(25% 0,75% 0,75% 60%,100% 60%,50% 100%,0 60%,25% 60%);'
                tags.append(f'<div class="shape" style="{style}"></div>')
        pages.append('<section>'+''.join(tags)+'</section>')
    output_html=QA/'meeting_ko.html'
    output_html.write_text('<!doctype html><html lang="ko"><meta charset="utf-8"><style>'
       f'@page{{size:{W}in {H}in;margin:0}}'
       '*{box-sizing:border-box}html,body{margin:0;padding:0}body{font-family:Pretendard,sans-serif;-webkit-print-color-adjust:exact;print-color-adjust:exact}'
       f'section{{position:relative;width:{W}in;height:{H}in;page-break-after:always;overflow:hidden}}'
       'section:last-child{page-break-after:auto}.text,.shape,.clip{position:absolute}.text{white-space:pre;line-height:1.16}.clip{overflow:hidden}.clip img{position:absolute}a{color:inherit;text-decoration:none}'
       '</style>'+''.join(pages)+'</html>')
    candidate_pdf=QA/'rendered_ko.pdf'
    command=['/opt/google/chrome/chrome','--headless','--no-sandbox','--disable-dev-shm-usage','--disable-gpu','--no-pdf-header-footer','--allow-file-access-from-files','--run-all-compositor-stages-before-draw','--virtual-time-budget=5000',f'--print-to-pdf={candidate_pdf}',output_html.as_uri()]
    with (QA/'render.log').open('w') as log:subprocess.run(command,stdout=log,stderr=log,check=True,timeout=60)
    for command in [
        ['pdfinfo',str(candidate_pdf)],
        ['pdftotext','-layout',str(candidate_pdf),str(QA/'rendered_text.txt')],
        ['pdftoppm','-scale-to','1400','-png',str(candidate_pdf),str(QA/'page')]]:
        result=subprocess.run(command,check=True,capture_output=True,text=True)
        assert not result.stderr.strip(),result.stderr
    paths=sorted(QA.glob('page-*.png'))
    assert len(paths)==len(SCENES),(len(paths),len(SCENES))
    extracted=(QA/'rendered_text.txt').read_text().split('\f')
    assert all(len(text.strip())>150 for text in extracted[:len(SCENES)])
    render_checks=[]
    for i,path in enumerate(paths):
        pixels=np.asarray(Image.open(path).convert('RGB'))
        nonwhite=float(np.mean(np.any(pixels<240,axis=2)))
        assert nonwhite>.05,(i+1,'blank page',nonwhite)
        render_checks.append(dict(page=i+1,nonwhite_fraction=nonwhite,text_characters=len(extracted[i])))
    # Publish only after validating the binary PDF; copying must remain binary-safe.
    shutil.copyfile(candidate_pdf,OUT.with_suffix('.pdf'))
    assert sha(candidate_pdf)==sha(OUT.with_suffix('.pdf'))
    for part in range(math.ceil(len(paths)/6)):
        subset=paths[part*6:(part+1)*6];canvas=Image.new('RGB',(1440,3*425),'#D1D9E3')
        for j,p in enumerate(subset):
            im=Image.open(p);im.thumbnail((700,394));canvas.paste(im,((j%2)*720+10,(j//2)*425+10))
            ImageDraw.Draw(canvas).text(((j%2)*720+12,(j//2)*425+405),f'{part*6+j+1}',fill='black')
        canvas.save(QA/f'contact-{part+1}.png')
    manifest=dict(date='2026-10-21',revision=3,source_deck='docs/slides/kor/20261021_Meeting_ko.pdf (revision 2)',output=str(OUT.relative_to(ROOT)),main_slides=13,appendix_slides=len(SCENES)-13,sources=SOURCES,experiment_run=str(MATCHED.relative_to(ROOT)),reference_exp84_run=str(RUN.relative_to(ROOT)),verified_jobs=288,new_xgb_jobs=108,seeds=[43,44,45],no_ids_finetuning=True,
      measurements={'quality':'Final model on entire fixed test','new_attack_mean':'Arithmetic mean of post-initial class F1 from full-test confusion matrix','retention':'Final minus initial Macro-F1 on the same initial-class panel; distinct from full-test metrics','update_cost':'Sum after initial stage; selection/eviction/construction/fit/cache; excludes all inference','error_bars':'Sample standard deviation (ddof=1) over seeds 43,44,45','composition_example':'Actual seed 43 rows; not a seed-averaged composition'},
      data_sha256={str(p.relative_to(ROOT)):sha(p) for p in [DATA/'exp85_summary.csv',DATA/'exp85_per_class_summary.csv',DATA/'exp85_values.csv',MATCHED/'aggregate_validation.json',ROOT/'configs/experiments/exp85_protocol.json']},
      files={str(p.relative_to(ROOT)):sha(p) for p in [OUT.with_suffix('.pptx'),OUT.with_suffix('.pdf')]},
      feedback_mapping={'problem_first':[1,2],'why_TabPFN_and_ICL':[3],'model_by_retention_controls':[4],'A_B_origin_and_sampling':[5],'cumulative_and_buffer_counts':[6],'acquisition_retention_model_test_pipeline':[7],'changing_class_proportions':[8],'overall_quality':[9],'new_class_quality':[10],'old_class_retention':[11],'measured_update_cost':[12],'unresolved_tail_issue':[13],'per_class_appendix':[14,15],'v3_architecture_appendix':[16,17]},
      limitations=['Three seeds on a fixed, previously inspected train/test split','Row IDs disjoint; identical feature vectors can remain','Cost excludes initial preparation and all inference; XGB buffers measured in a separate session on the same hardware','Retention panel differs from full test; no guarantee for all old classes','Architecture and charts are embedded images; other slide text/diagrams editable'])
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    (QA/'scenes.json').write_text(json.dumps(SCENES,ensure_ascii=False,indent=2)+'\n')
    with zipfile.ZipFile(OUT.with_suffix('.pptx')) as z:
        shadow_count=0;effect_refs=[]
        for name in z.namelist():
            if name.endswith('.xml'):
                tree=etree.fromstring(z.read(name))
                shadow_count+=len(tree.xpath('//*[local-name()="outerShdw" or local-name()="innerShdw" or local-name()="prstShdw"]'))
                if name.startswith('ppt/slides/'):
                    effect_refs += tree.xpath('//*[local-name()="effectRef"]/@idx')
        assert shadow_count==0 and all(idx=='0' for idx in effect_refs)
    (QA/'validation.json').write_text(json.dumps(dict(passed=True,revision=3,pages=render_checks,
        pdf_parser_errors=0,pptx_shadow_count=shadow_count,pptx_effect_refs_disabled=True,
        source_conditions=288,new_xgb_conditions=108,seeds=[43,44,45],cumulative_color='111111',cumulative_drawn_on_top=True,matched_panel_axes=True,
        pdf_sha256=sha(OUT.with_suffix('.pdf')),pptx_sha256=sha(OUT.with_suffix('.pptx'))),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(slides=len(SCENES),pptx=str(OUT.with_suffix('.pptx')),pdf=str(OUT.with_suffix('.pdf'))),ensure_ascii=False))

def main():
    assert len(ROWS)==96 and all(int(r['seed_count'])==3 for r in ROWS)
    validation=read(MATCHED/'aggregate_validation.json')
    assert validation['passed'] and validation['matched_conditions']==108
    assert read(MATCHED/'results.json')['state']=='complete'
    assert len(list(csv.DictReader((DATA/'exp85_values.csv').open())))==288
    plot_curves('new_attacks','new_attack_f1_pct','신규 공격 평균 F1 (%)')
    plot_curves('retention','initial_retention_change_pp','초기 대비 변화 (%p)')
    plot_curves('cost','update_seconds','후속 갱신 합계 (초)')
    plot_curves('overall','macro_f1_pct','전체 Macro-F1 (%)')
    plot_tail();build_scenes();render()

if __name__=='__main__':main()
