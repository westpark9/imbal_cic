#!/usr/bin/env python3
"""Editable Korean lab meeting deck; original 1002 deck stays unchanged."""
from pathlib import Path
import sys
_root=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/experiment_paths.py').is_file())
sys.path.insert(0,str(_root/'scripts/common'))
from experiment_paths import repo_root
ROOT=repo_root(__file__)
import hashlib,json,re
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE,MSO_ANCHOR,PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

HERE=ROOT/'docs/slides';OUT=HERE/'kor/1015_labmeeting_ko.pptx';QA=HERE/'kor/1015_labmeeting_qa'
QA.mkdir(exist_ok=True)
BASE=HERE/'kor/1002_labmeeting_ko_draft.pptx'
prs=Presentation(BASE)
for s in list(prs.slides._sldIdLst):
    prs.part.drop_rel(s.rId);prs.slides._sldIdLst.remove(s)
W,H=prs.slide_width/Inches(1),prs.slide_height/Inches(1)
NAVY,BLUE,CYAN,INK,GREY,PALE='00306C','007BC6','00B8EE','21262E','5D6773','F0F5FA'
FONT='Pretendard';MONO='DejaVu Sans Mono'
records=[];text_boxes=[];numeric_cells=[]
PAPER='https://arxiv.org/pdf/2605.13986'
ICD='https://arxiv.org/pdf/2402.06971'
FT='https://github.com/PriorLabs/TabPFN/blob/main/src/tabpfn/finetuning/finetuned_classifier.py'

def rgb(c):return RGBColor.from_string(c)
def frame(tf,value,size=20,bold=False,color=INK,align=PP_ALIGN.LEFT,font=FONT):
    tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=Inches(.03);tf.margin_top=tf.margin_bottom=Inches(.03)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(3);p.line_spacing=1.08
        r=p.add_run();r.text=line;r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
        ea=OxmlElement('a:ea');ea.set('typeface',font);r._r.get_or_add_rPr().append(ea)
def text(s,x,y,w,h,value,**kw):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame(sh.text_frame,value,**kw)
    text_boxes.append(dict(page=len(prs.slides),x=x,y=y,w=w,h=h,text=str(value),font_size=kw.get('size',20)))
    return sh
def box(s,x,y,w,h,color=PALE):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.adjustments[0]=.035;sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();return sh
def arrow(s,x,y,w=.35):
    return text(s,x,y,w,.4,'→',size=23,color=BLUE,align=PP_ALIGN.CENTER)
def slide(title,section,band='',note=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    # Hide inherited template artwork so every content shape has a known geometry.
    s._element.set('showMasterSp','0')
    bg=s.background.fill;bg.solid();bg.fore_color.rgb=rgb('FFFFFF')
    text(s,.6,.32,12.1,.65,title,size=28,bold=True,color=NAVY)
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(7.02),prs.slide_width,Inches(.48));sh.fill.solid();sh.fill.fore_color.rgb=rgb(NAVY);sh.line.fill.background()
    if band:text(s,.62,7.07,9.1,.31,band,size=13,bold=True,color='FFFFFF')
    text(s,10.05,7.09,2.35,.24,section,size=10,color='D6E4F3',align=PP_ALIGN.RIGHT)
    text(s,12.55,7.08,.32,.25,str(len(prs.slides)),size=10,color='FFFFFF',align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=note
    records.append(dict(page=len(prs.slides),title=title,section=section,band=band,notes=note))
    return s
def subtitle(s,value):text(s,.62,1.04,12.05,.4,value,size=14,color=GREY)
def card(s,x,y,w,h,title,body,size=19):
    box(s,x,y,w,h);box(s,x,y,w,.6,'E0ECF6')
    text(s,x+.16,y+.08,w-.32,.42,title,size=20,bold=True,color=BLUE)
    text(s,x+.17,y+.8,w-.34,h-.95,body,size=size)
def table(s,x,y,widths,headers,rows,row_h=.5,size=16,numeric=()):
    sh=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(sum(widths)),Inches(row_h*(len(rows)+1)))
    t=sh.table
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for r,values in enumerate([headers]+rows):
        t.rows[r].height=Inches(row_h)
        for c,v in enumerate(values):
            cell=t.cell(r,c);cell.fill.solid();cell.fill.fore_color.rgb=rgb(NAVY if r==0 else PALE if r%2 else 'FFFFFF')
            isnum=r>0 and c in numeric
            frame(cell.text_frame,v,size=size,bold=r==0,color='FFFFFF' if r==0 else INK,
                  align=PP_ALIGN.RIGHT if isnum else PP_ALIGN.LEFT,font=MONO if isnum else FONT)
            if isnum:numeric_cells.append(dict(page=len(prs.slides),row=r,column=c,text=str(v),right_aligned=True))
    return sh
def cite(s,value):text(s,.64,6.57,12,.3,value,size=11,color=GREY)
def steps(s,items,y=1.8,h=2.3):
    w=(12.05-.4*(len(items)-1))/len(items)
    for i,(title,body) in enumerate(items):
        x=.62+i*(w+.4);card(s,x,y,w,h,title,body,size=17)
        if i<len(items)-1:arrow(s,x+w+.02,y+h/2-.2)

def architecture_panel(s):
    sh=s.shapes.add_picture(str(QA/'v3_architecture_page.png'),Inches(.57),Inches(1.15),width=Inches(4.65),height=Inches(5.55))
    sh.crop_left=.20;sh.crop_right=.19;sh.crop_top=.082;sh.crop_bottom=.285
    text(s,.65,6.72,4.55,.2,'TabPFN-3 · Figure 5 · 원문 PDF 6쪽',size=10,color=GREY)

def detail(s,y,heading,body,h=1.55):
    text(s,5.62,y,7.05,.43,heading,size=21,bold=True,color=BLUE)
    text(s,5.64,y+.50,7.0,h-.5,body,size=18)

# Research motivation precedes the model explanation, per the latest correction.
s=slide('TabPFN v3의 작동과 IDS context 갱신','랩미팅')
text(s,.68,1.65,11.7,1.4,'사전학습 모델의 활용과\n적은 사례로 새 공격을 반영하는 방식',size=36,bold=True,color=NAVY)
text(s,.7,3.55,11.8,1.1,'TabPFN과 v3 작동 원리 · 가중치 fine-tuning\n누적 context · 고정 버퍼 · 갱신 없는 추론',size=23,color=BLUE)
text(s,.7,5.55,10,.5,'랩미팅   2026년 10월 15일',size=19,color=GREY)

s=slide('0. 연구 관점 · 고정된 F1 경쟁에서 변화에 대응하는 IDS로','00 연구 방향',
        '새 사례 반영 품질·사용 데이터·갱신 비용으로 모델의 유용성 평가',
        'User item 0 precedes the TabPFN explanation. Network changes are the operational motivation, not a new cross-dataset experiment or an established robustness result. ICD Introduction motivates pretraining and transfer in tabular prediction; its §4 tuned-XGB comparison is historical benchmark evidence, not a universal conclusion about v3. '+ICD)
for y,head,body in [
    (1.5,'기존 관점','고정된 공격 클래스의 closed-set F1 향상에 집중\nICD의 정적 벤치마크에서도 tuned XGB의 강세\n수치 경쟁에서 실제 갱신 능력의 평가로 전환 필요'),
    (3.0,'운영의 요구','네트워크 환경·공격 양상이 바뀌는 상황\n기존 분포에서 높은 F1만으로 갱신 후 탐지 능력 설명 곤란'),
    (4.5,'연구의 전환','TabPFN을 기본 모델로 채택하고 새 사례를 context로 반영\n적은 라벨과 낮은 갱신 비용에서 필요한 탐지 품질 확보')]:
    text(s,.7,y,2.25,.5,head,size=23,bold=True,color=BLUE)
    text(s,3.15,y,9.4,1.28,body,size=21)
cite(s,'Ma et al., In-Context Data Distillation with TabPFN, ICLR 2024 ME-FoMo Workshop')

s=slide('TabPFN · 표 데이터의 예측 규칙을 사전학습한 모델','01 작동 원리',
        '라벨 사례를 context로 제공하고, 정답이 없는 query의 확률 예측',
        'PFN = Prior-data Fitted Network. Context is labelled task data, not a natural-language prompt. PriorLabs TabPFN-3 report §2; local architecture and classifier implementation. '+PAPER)
steps(s,[('사전학습','다양한 합성 표 데이터\n일부 라벨을 가린 예측 과제\n공통 가중치 θ 학습'),('새 IDS 데이터','알려진 특징 Xc와 라벨 yc\n예측 대상의 특징 xq\n정답 yq는 입력에서 제외'),('In-context learning','사례의 관계를 이용한 추론\n현재 데이터에 따른 예측\n기본 동작에서 θ 고정')],h=2.5)
text(s,.85,5.02,11.6,.7,'pθ(yq | xq, Xc, yc)',size=28,bold=True,color=NAVY,align=PP_ALIGN.CENTER)
cite(s,'Prior Labs, TabPFN-3: Technical Report, 2026 · §2')

s=slide('Context와 query · 모델에 전달되는 두 종류의 행','01 작동 원리',
        'Context에는 라벨 포함 · query의 정답은 평가에만 사용',
        'Illustrative schema, no fabricated performance values. Context rows and held-out query rows remain distinct. Classification output labels are those represented by the fitted context.')
table(s,.65,1.7,[2.05,2.25,2.25,2.1,3.3],['구분','Flow 특징 1','Flow 특징 2','공격 라벨','역할'],[
 ['Context','관측값','관측값','Benign','정상 사례 제공'],['Context','관측값','관측값','공격 A','기존 공격 사례 제공'],
 ['Context','관측값','관측값','공격 B','신규 라벨 반영 시 추가'],['Query','관측값','관측값','모델에 미제공','정상·공격 확률 예측']],row_h=.66,size=17)
text(s,.8,5.35,11.7,.75,'같은 query도 context의 구성·규모에 따라 예측 변화\n라벨 없는 query를 추가하는 것과 라벨된 context를 추가하는 것의 구분',size=21,color=BLUE)

s=slide('현재 실험의 TabPFN v3 설정','01 작동 원리',
        '1014 실행과 동일한 checkpoint·추론 설정을 기준으로 설명',
        'Actual saved checkpoint config inspected with torch.load. Parameters counted and hashed in docs/research/20261015/tabpfn_fit_mechanism_check.json. Version 8.2.0 is the installed library, not the model generation. Local checkpoint: tabpfn/tabpfn-v3-classifier-v3_20260417_multiclass.ckpt. Model report Appendix C; support ceilings are not measured IDS performance.')
table(s,.68,1.65,[3.1,8.85],['항목','실제 사용 설정'],[
 ['모델 / 라이브러리','TabPFN-v3 multiclass / tabpfn 8.2.0'],['입력','두 데이터셋의 정제 Flow 특징 46개'],
 ['모델 파라미터','53,153,144개 · 공통 사전학습 가중치'],['앙상블','서로 다른 전처리·순열의 estimator 4개'],
 ['Context 처리','fit_with_cache · 캐시 CPU 저장'],['체크포인트 설정','ICL block 24개 · 특징 통합 CLS token 4개']],row_h=.64,size=18)
cite(s,'실제 checkpoint config와 실행 코드 확인 · v3 원문 Appendix C와 대조')

s=slide('v3 구조 · 열의 분포에서 행 단위 표현으로','01 작동 원리',
        '열의 정보를 먼저 통합한 뒤 행 사이의 관계를 추론',
        'Redrawn from local architectures/tabpfn_v3.py TabPFNV3 docstring and forward path, with report Figure 5 as architecture cross-check. Local checkpoint: feature_group_size=3, dist_embed_num_blocks=3, feat_agg_num_blocks=3, feat_agg_num_cls_tokens=4. '+PAPER+'#page=6')
architecture_panel(s)
detail(s,1.37,'① 입력 표현 · Feature grouping & cell embedding',
       '특징 3개와 결측 표시를 묶어 128차원 표현으로 변환\nContext 행에 알려진 라벨의 embedding 추가\nQuery에는 정답 라벨을 넣지 않고 특징만 제공',h=1.55)
detail(s,3.10,'② 열별 분포 · Feature embedding',
       '특징 하나의 값들을 128개 inducing point로 요약\n분포를 요약하는 학습된 보조 표현과 attention 수행\n3개 block을 통해 현재 context의 분포 정보 반영',h=1.55)
detail(s,4.84,'③ 행별 통합 · Feature aggregation',
       '한 행 안의 특징들이 attention으로 정보 교환\n행을 요약하는 CLS token 4개에 특징 관계 통합\n4 × 128차원을 이어 행별 512차원 표현 생성',h=1.55)

s=slide('v3 구조 · context 관계와 클래스 확률의 계산','01 작동 원리',
        '새 클래스의 라벨 사례를 context에 넣어 출력 클래스에 반영',
        'Local ICLTransformerBlock uses train-only keys; ManyClassDecoder uses train representations and one-hot labels. Report §§2.1–2.2, Figure 5. The actual checkpoint has max_num_classes=160; this is a checkpoint constraint even though decoder weights are independent of class count. No unlabeled novel-class naming is implied.')
architecture_panel(s)
detail(s,1.5,'④ 사례 사이의 추론 · In-context learning',
       'Context의 행 표현에 라벨 표현을 추가\nContext ↔ context 관계와 query → context 참조\n24개 block에서 조건부 예측에 필요한 표현 계산\nQuery끼리의 정답 공유나 가중치 최적화 없음',h=1.95)
detail(s,3.65,'⑤ 라벨별 확률 · Many-class decoder',
       'Query 표현과 context 표현의 관련도를 attention으로 계산\n각 context 라벨의 one-hot 벡터를 관련도에 따라 가중합\n같은 클래스의 근거를 모아 클래스별 확률 산출\n새 클래스 사례 추가 시 출력 라벨에 포함',h=1.95)
text(s,5.65,5.88,7.02,.65,'현재 checkpoint의 클래스 상한 160개\n신규 라벨이 없는 context에서 새 공격 이름의 출력 불가',size=17,color=GREY)

s=slide('fit과 predict · context 캐시의 생성과 재사용','01 작동 원리',
        '가중치가 고정되어도 context 준비와 query 계산 비용 발생',
        '1014: scripts/v7/exp72/replay_models.py fit_stage calls TabPFNClassifier.fit, using get_model in exp71_bounded_updates.py. fit_with_cache builds context-dependent state and predict reuses it. Actual parameter SHA-256 unchanged after 1000→1020 fit: see mechanism check. The current library does not incrementally append 20 rows to the old cached state in this experiment. Report §2.4.2 distinguishes cache build vs cached predict.')
steps(s,[('fit(Xc, yc)','전처리와 라벨 대응 구성\nContext 측 표현·KV 계산\n추론용 상태와 캐시 저장'),('predict(Xq)','Query 전처리\n저장된 context 상태 참조\n확률 계산·앙상블 결합'),('새 context의 fit','변경된 사례로 상태 재구성\n1014는 누적 전체 재처리\n기존 사전학습 가중치 유지')],y=1.7,h=3)
text(s,.85,5.2,11.6,.8,'실제 확인: 1,000개 → 1,020개 fit 후\n53,153,144개 파라미터의 SHA-256 동일',size=23,bold=True,color=BLUE,align=PP_ALIGN.CENTER)

s=slide('데이터 적응의 세 방식과 학습 대상','02 미세조정',
        'Context 갱신과 가중치 fine-tuning은 서로 다른 적응 연산',
        'TabPFNClassifier versus FinetunedTabPFNClassifier in installed official 8.2.0 source. ICD optimizes input context features through a frozen model, as Ma et al. 2024. 1014 ordinary PFN used row 1; EXP75 used context-feature optimization; EXP74 trained external classification heads and is not the user-referenced experiment.')
table(s,.65,1.75,[2.65,2.95,2.45,4.15],['방식','변경 대상','역전파','추론 시 필요한 것'],[
 ['Context 갱신','라벨된 사례의 구성','없음','사전학습 θ + 현재 context'],['Context 최적화','Context의 특징값','입력 특징으로 전파','고정 θ + 최적화된 context'],
 ['가중치 fine-tuning','모델 파라미터 θ','모델 파라미터로 전파','갱신된 θ + 현재 context']],row_h=.88,size=18)
text(s,.8,5.7,11.8,.65,'1014 단일 TabPFN: 첫 번째 방식 · LoCalPFN: 별도의 미세조정 비교군',size=21,color=BLUE)

s=slide('v3 가중치 fine-tuning · 학습 절차','02 미세조정',
        'Query 정답의 손실로 θ 갱신 · 미세조정 이후에도 context를 사용',
        'The v3 report describes architecture and variants but is not a detailed task-specific fine-tuning tutorial. Procedure here is grounded in the installed official finetuning/finetuned_base.py and finetuned_classifier.py. Cross entropy over query labels; AdamW; explicit independent validation; final inference classifier fits X_, y_. '+FT)
architecture_panel(s)
conn=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(5.3),Inches(6.35),Inches(5.3),Inches(1.5))
conn.line.color.rgb=rgb('B54040');conn.line.width=Pt(2)
end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','med');end.set('len','med');conn.line._get_or_add_ln().append(end)
detail(s,1.35,'① 학습용 context/query 과제 구성',
       'Train을 context와 query로 분할, 예시 비율 80% / 20%\nForward 입력: context 특징·라벨 + query 특징\nQuery의 정답은 숨겼다가 loss 계산에만 사용',h=1.55)
detail(s,3.10,'② Query loss의 역전파와 가중치 갱신',
       'L = −Σq log pθ(yq | xq, context)\n출력부터 decoder·ICL·embedding의 θ로 gradient 전파\nloss.backward()와 AdamW.step()으로 θ 갱신',h=1.55)
detail(s,4.85,'③ Validation 선택과 갱신된 모델의 추론',
       '별도 validation의 log loss 등으로 checkpoint 선택\n갱신된 θ에 라벨 context를 다시 제공한 뒤 추론\n새 query에는 정답 미제공 · test는 최종 평가에만 사용',h=1.55)
text(s,5.65,6.55,7.02,.27,'붉은 선: 모델 파라미터로의 역전파 · 공식 fine-tuning 구현 기준',size=11,color='B54040')

s=slide('가중치 fine-tuning의 실제 설정 항목','02 미세조정','',
        'Illustrative configuration only, not a new quality result. Pin checkpoint with extra_classifier_kwargs[model_path]. Use epochs=10, learning_rate=1e-5, context+query=1024, query ratio=.2, one estimator as an initial engineering smoke configuration, not a tuned best setting. Validation supplied separately; no test tuning. Official repository recommends high-memory GPU for its 50k example, so batch/context size must be validated on this 24GB GPU.')
table(s,.65,1.65,[3.9,8.2],['설정','역할'],[
 ['사용 checkpoint','기존 v3 multiclass 가중치를 명시적으로 지정'],['학습률·epoch·조기 종료','갱신 크기와 최적화 비용 제어'],['Context + query 크기','한 학습 과제의 GPU 메모리와 사례 수 제어'],
 ['Query 비율','손실 계산에 사용할 학습 행의 비율'],['학습·검증·추론 estimator 수','품질·메모리·비용의 비교 조건 고정'],['Validation과 test','Checkpoint 선택과 최종 평가의 역할 분리']],row_h=.67,size=18)

s=slide('신규 사례만 사용하는 갱신의 의미','02 미세조정',
        '이번 a·b·c 비교의 중심은 context에 어떤 사례를 유지하는지',
        'User clarified 1000 and 20 were examples, not required sample counts. Ordinary classifier supports a newly represented class through labels in context. A naive FinetunedTabPFNClassifier.fit on a single new class is not equivalent to old-context-plus-new-class training: its train/context label space and validation differ. New-query-only gradients with old context retained need a separate controlled training loop and old-class retention validation; not silently substituted into this experiment.')
card(s,.65,1.65,5.8,3.9,'Context 갱신','신규 라벨 사례를 context에 추가\n기존 사례를 누적하거나 일부 교체\n기존 모델 파라미터의 재학습 없음\n추가·삭제 이후 캐시 재구성 비용 측정',size=21)
card(s,6.85,1.65,5.8,3.9,'신규 사례 중심의 가중치 학습','신규 query에 대한 손실로 θ 갱신\n기존 클래스 context의 보존 필요\n기존·신규 클래스의 별도 유지 평가\n기본 fit과 구분한 추가 학습 설계',size=21)
text(s,.8,5.9,11.7,.5,'1,000개와 20개는 정책을 설명하기 위한 예시',size=20,color=GREY)

s=slide('갱신 비용 · 새 모델을 사용할 수 있는 시점까지','03 연구 방향',
        '가중치 최적화의 유무와 전체 갱신·추론 비용을 분리',
        'Update includes selection/similarity/preprocessing/context cache or model fitting. Initial load distinguished. Serving includes actual input gathering/preprocessing/inference. Hardware and estimator/batch settings fixed within experiment; no speed comparison across incompatible workloads. Use initial + sum updates + serving(Q), not fit time alone. No monetary or energy measurement is claimed.')
table(s,.65,1.7,[2.8,6.1,3.2],['비용','포함 연산','보고 기준'],[
 ['초기 구축','Checkpoint 적재·초기 context 또는 모델 준비','초'],['후속 갱신','삭제·선택·전처리·캐시 재구성 또는 재학습','회당 초 · 누적 초'],['추론','실제 배치의 전처리·모델 계산·후처리','초 · μs/행'],
 ['자원·정보량','GPU·RAM peak, context 수, 공개 라벨 수','GiB · 행']],row_h=.76,size=18)
text(s,.8,5.64,11.8,.65,'총 연산 시간 = 초기 구축 + 후속 갱신 합계 + 평가한 Q개 행의 추론',size=22,bold=True,color=BLUE)

s=slide('평가 시나리오 · 라벨된 공격 클래스의 순차 도입','04 실험 설계',
        '등장 전 공격 라벨 미사용 · 방법별 동일 평가 행 사용',
        'STAGES inherited from exp69; 1014 EXP72 test partition shares full clean test rows. Sequence describes class availability, not raw timestamp replay or network drift. Order is introduction order, not class-support sort, because these are timeline cells. Updated experiment evaluates instant-ready versions and measures update cost separately.')
subtitle(s,'클래스 도입 순서 기준 · 실제 수집 시각의 재생과 구분')
table(s,.65,1.75,[1.65,2.7,2.2,1.65,1.8,2.1],['데이터','초기','도입 1','도입 2','도입 3','도입 4 이후'],[
 ['CIC2018','Benign\nBrute force · DoS','DDoS','Web attacks','Infiltration','Bot'],
 ['ToN','Benign\nScanning · DoS','Injection · DDoS','Password','XSS','Ransomware · Backdoor\nMitM']],row_h=1.12,size=15)
text(s,.85,5.1,11.6,1.1,'라벨 공급 → context 또는 모델 갱신 → 별도 test 평가\n신규 공격의 구분 · 기존 공격의 탐지 · 정상 오탐 · 비용',size=23,color=BLUE)

hist_run=next((ROOT/'tabpfn/results/v7/exp72').glob('*/results.json'))
hist=json.loads(hist_run.read_text())
s=slide('1014 결과 · 누적 10만 사례의 품질과 갱신 비용','05 기존 근거',
        '전체 context 재구성 비용을 줄이기 위한 소규모 context 비교',
        'Exact EXP72 source '+str(hist_run.relative_to(ROOT))+'. Aggregate ideal replay over all original test rows, seed 43. Update is stage 1 onward, excludes initial build. Natural/uniform refer to attack-label training allocation, not test balancing. No new experiment or v3 parameter fine-tuning is represented by this table.')
subtitle(s,'전체 test 순차 재생 · 지연 없음 · seed 43 · 초기 구축 제외 후속 갱신 합계')
rows=[]
for ds,label in [('cic2018','CIC2018'),('toniot','ToN')]:
    for case,clabel in [('natural','자연'),('uniform','균등')]:
        jj=[next(j for j in hist['jobs'] if j['dataset']==ds and j['case']==case and j['method']==m) for m in ['xgb','tabpfn']]
        update=[]
        for m in ['xgb','tabpfn']:
            files=sorted((hist_run.parent/'jobs'/f'{ds}_{case}_{m}').glob('stage_*/fit_metrics.json'))
            update.append(sum(json.loads(f.read_text())['update_seconds'] for f in files if f.parent.name!='stage_0'))
        rows.append([label,clabel,*[f"{j['results']['ideal']['macro_f1']*100:.2f}" for j in jj],*[f'{u:.2f}' for u in update]])
table(s,.65,1.85,[1.75,1.2,2.25,2.25,2.3,2.3],['데이터','배분','XGB F1 (%)','PFN F1 (%)','XGB 갱신 (s)','PFN 갱신 (s)'],rows,row_h=.73,size=17,numeric=(2,3,4,5))
text(s,.8,5.85,11.8,.55,'개선 방향: context 사용량·교체 정책을 제어하며 품질과 비용을 함께 측정',size=21,color=BLUE)

s=slide('추가 비교 · 누적·고정 버퍼·갱신 없음','04 실험 설계','',
        'User examples shown explicitly as examples. Experimental defaults are initial 1000, 100 per new class. Policies all receive same released data, but fixed deliberately does not use later labels; information actually consumed is reported. b protects >=1 row per existing class. Not an incremental KV-cache implementation.')
subtitle(s,'정책 설명용 예시: 기존 1,000개 · 신규 라벨 20개')
table(s,.65,1.8,[1,3.3,2.1,5.6],['정책','Context 구성','예시 행 수','확인할 차이'],[
 ['a','기존 전체 + 신규','1,020','사례 보존과 누적 재준비 비용'],['b','기존 일부 삭제 + 신규','1,000','동일 저장량에서 신규 반영·기존 탐지'],['c','기존 context 고정','1,000','갱신 비용 없는 기준선과 신규 유형 한계']],row_h=.95,size=19,numeric=(2,))
text(s,.85,5.55,11.7,.7,'XGB·다른 비교 모델은 누적 사례로 재학습·재구성\n이번 추가 실행에서 LoCalPFN 제외',size=22,color=BLUE)

s=slide('고정 버퍼의 삭제 기준','04 실험 설계','',
        'FIFO denotes oldest acquisition, not measured semantic obsolescence. Similarity uses old-context StandardScaler and distance to new labelled train cases, never test. All three rules retain at least one example of each old class. Same-stage ties use seeded acquisition order. Nearest-row removal may damage discrimination: hypothesis only.')
steps(s,[('확보 순서 · FIFO','가장 먼저 확보한 사례부터 삭제\n실제 노후화의 대리 기준\n동일 구간 동률 순서 고정'),('신규와 유사','기존 context로 표준화\n신규 사례와 가까운 기존 사례 삭제\n중복 근거 감소 가설'),('무작위','Seed로 재현 가능한 삭제\n단순 정책의 비교 기준\n삭제 후 기존 클래스 유지 확인')],y=1.7,h=3.55)
text(s,.8,5.8,11.8,.55,'공통 조건: 버퍼 크기 고정 · 신규 사례 수 일치 · 기존 클래스 최소 사례 보존',size=20,color=BLUE)

s=slide('EXP76 · 첫 소규모 context 비교 설정','04 실험 설계','',
        'Full protocol docs/research/20261015/exp76_protocol.md. 100 per new class is an investigator setting, not a user-imposed value. 8192 batches differ from 1014; direct timing comparison is not valid. No cross-network transfer claim. New results are exploratory, same previously examined fixed test, single seed.')
table(s,.65,1.65,[3.4,8.7],['항목','첫 실행 설정'],[
 ['초기 / 신규 라벨','초기 1,000개 / 신규 공격 클래스당 100개'],['최종 누적 · CIC2018','1,400개'],['최종 누적 · ToN','1,700개'],['고정 버퍼','1,000개 · FIFO / 유사도 / 무작위'],
 ['비교 모델','TabPFN context 정책 · XGB · BoostPFN · DistPFN'],['평가','기존 test 전체 · 지연 없는 순차 도입 · seed 43'],['측정 조건','RTX 4090 단독 실행 · 배치 8,192개 · 비용 분리']],row_h=.57,size=18)

runfiles=sorted((ROOT/'tabpfn/results/v7/exp76').glob('*/results.json'))
new_path=runfiles[-1] if runfiles else None
new=json.loads(new_path.read_text()) if new_path else {'state':'pending','jobs':[]}
labels={'xgb':'XGBoost 누적','pfn_accumulate':'TabPFN 누적','pfn_fifo':'TabPFN FIFO','pfn_similar':'TabPFN 유사도','pfn_random':'TabPFN 무작위','pfn_fixed':'TabPFN 고정','boostpfn':'BoostPFN 누적','distpfn':'DistPFN 누적'}
for ds,label in [('cic2018','CIC2018'),('toniot','ToN')]:
    jobs={j['method']:j for j in new['jobs'] if j['dataset']==ds}
    s=slide(label+' · 소규모 context의 품질과 비용','06 추가 결과','',
            'Source '+str(new_path.relative_to(ROOT) if new_path else 'pending')+'. All test rows, one seed, instant-ready replay. New label budget and prediction batch differ from 1014. Missing rows mean unfinished, never fabricated. Classwise metrics and support are in the experiment artifacts.')
    subtitle(s,'EXP76 · 전체 test · 초기 구축 제외 후속 갱신 · FPR 분모는 실제 Benign')
    rows=[]
    for method,name in labels.items():
        j=jobs.get(method)
        if not j:rows.append([name,'계산 중','계산 중','계산 중','계산 중'])
        else:rows.append([name,f"{100*j['metrics']['macro_f1']:.2f}",f"{100*j['metrics']['benign_false_positive_rate']:.3f}",f"{j['update_seconds']:.2f}",f"{j['serving_seconds']/j['metrics']['rows']*1e6:.2f}"])
    table(s,.65,1.7,[3.35,2.1,2.1,2.1,2.1],['방법','F1 (%)','FPR (%)','갱신 (s)','추론 (μs/행)'],rows,row_h=.48,size=16,numeric=(1,2,3,4))
    text(s,.8,6.25,11.6,.38,'라벨 예산과 배치 크기가 다른 1014 수치와 직접적인 개선량 비교 제외',size=16,color=GREY)

    if 'pfn_accumulate' in jobs:
        s=slide(label+' · 클래스별 구분 성능과 버퍼 정책','06 추가 결과','',
                'Same full test population as preceding aggregate table. Classes sorted by evaluated support descending. Values are class F1 percentages; support is actual evaluation support, not training budget. All three buffer policies shown, without choosing a test-best policy.')
        subtitle(s,'전체 test의 클래스별 F1 (%) · 평가 표본 수 내림차순 · seed 43')
        methods=['xgb','pfn_accumulate','pfn_fifo','pfn_similar','pfn_random']
        classes=sorted(jobs['pfn_accumulate']['metrics']['classes'],key=lambda c:-c['support'])
        classrows=[]
        for c in classes:
            vals=[]
            for method in methods:
                cc=next((x for x in jobs.get(method,{}).get('metrics',{}).get('classes',[]) if x['name']==c['name']),None)
                vals.append(f"{100*cc['f1']:.2f}" if cc else '계산 중')
            classrows.append([c['name'],f"{c['support']:,}",*vals])
        table(s,.65,1.65,[2.4,1.75,1.55,1.65,1.55,1.55,1.55],['클래스','평가 수','XGB','PFN 누적','FIFO','유사도','무작위'],classrows,row_h=.43,size=14,numeric=(1,2,3,4,5,6))
        text(s,.8,6.5,11.8,.28,'공격별 precision·recall과 혼동행렬은 연결된 실험 결과에 보존',size=14,color=GREY)

s=slide('채택 판단 · 품질을 유지하는 작은 context와 낮은 갱신 비용','07 후속 판단','',
        'Do not claim statistical equivalence without a predeclared margin and repeated data/seeds. Update latency and raw throughput have different meanings. New class F1 and old class errors/FPR must be inspected before choosing a policy; do not select a final model solely from this test ranking. Manuscript changes are proposals only under user steering.')
card(s,.65,1.7,5.8,3.8,'현재 관찰과 대응','소규모 PFN의 갱신 준비 비용 절감\n반복 추론 비용의 별도 경량화 필요\n유사도 삭제의 정상 오탐 증가 확인\n정상·기존 클래스 보존 규칙 보강',size=21)
card(s,6.85,1.7,5.8,3.8,'후속 비교','동일 라벨에서 context 규모만 변경\nEstimator 수·캐시 위치별 품질과 비용\n공격별 precision·recall과 정상 오탐\n조건·seed 확장 이후 최종 방법 판단',size=21)

s=slide('부록 · v3 원문의 전체 구조','부록',note='Original Figure 5 from TabPFN-3 Technical Report, PDF page 6. Cropped only to figure bounds using native PPTX picture crop; no diagram content altered. '+PAPER+'#page=6')
picture=QA/'v3_architecture_page.png'
if picture.exists():
    sh=s.shapes.add_picture(str(picture),Inches(.85),Inches(1.2),width=Inches(4.8),height=Inches(5.5))
    sh.crop_left=.20;sh.crop_right=.19;sh.crop_top=.082;sh.crop_bottom=.285
text(s,6.25,1.65,6.05,4.8,'Feature grouping & cell embedding\n열별 분포 embedding\n행별 feature aggregation\nICL transformer\nMany-class decoder\n\n원문 Figure 5 · PDF 6쪽\n세부 hyperparameter는 Appendix C',size=22,color=NAVY)

s=slide('부록 · v3 fine-tuning 코드 예시','부록',note='Illustrative official API usage, not a reported experiment. Adapt context/query size to GPU and label counts. Verify validation disjointness before use. Model checkpoint is explicitly pinned. This snippet intentionally does not fit only a single newly arriving class. '+FT)
code='from tabpfn.finetuning import FinetunedTabPFNClassifier\n\nft = FinetunedTabPFNClassifier(\n    device="cuda", epochs=10, learning_rate=1e-5,\n    n_finetune_ctx_plus_query_samples=1024,\n    finetune_ctx_query_split_ratio=0.2,\n    n_estimators_finetune=1,\n    n_estimators_validation=1,\n    n_estimators_final_inference=1,\n    eval_metric="log_loss", random_state=43,\n    extra_classifier_kwargs={"model_path": checkpoint_path},\n)\nft.fit(X_train, y_train, X_val=X_val, y_val=y_val)\np = ft.predict_proba(X_test)'
box(s,.65,1.32,12,5.23);text(s,.85,1.44,11.6,4.98,code,size=16,font=MONO)
cite(s,'학습 설정 예시 · 이번 context 정책 실험의 실행 코드와 구분')

prs.save(OUT)
manifest=dict(source_deck=str(BASE.relative_to(ROOT)),source_deck_sha256=hashlib.sha256(BASE.read_bytes()).hexdigest(),
    output=str(OUT.relative_to(ROOT)),slides=records,text_boxes=text_boxes,numeric_cells=numeric_cells,
    historical_source=str(hist_run.relative_to(ROOT)),new_source=str(new_path.relative_to(ROOT)) if new_path else None,
    new_state=new['state'],primary_v3_source=PAPER,finetuning_source=FT,
    manuscript_status='proposal only; manuscript and Overleaf unchanged',
    intentional_overlaps=['card backgrounds under headings/body','table cells','footer backgrounds under footer text'])
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
md=['# 1015 랩미팅 슬라이드 구성','', '1002 원본 보존. 0번 연구 관점 전환에 이어 구조도와 단계별 작동 원리, v3 fine-tuning을 설명하는 한국어 새 날짜본.','']
for r in records:md.extend([f"## {r['page']}. {r['title']}",'',r['notes'],''])
(HERE/'kor/1015_labmeeting_plan.md').write_text('\n'.join(md))
print(OUT,len(prs.slides),'slides')
