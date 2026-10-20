#!/usr/bin/env python3
"""Compact main slides and large-data appendix, with measured provenance."""
from pathlib import Path
import sys,json,hashlib,csv,argparse
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/experiment_paths.py').is_file())
sys.path.insert(0,str(ROOT/'scripts/common'))
from experiment_paths import repo_root
ROOT=repo_root(__file__)
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager,ticker
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

parser=argparse.ArgumentParser();parser.add_argument('--layout-preview',action='store_true');PREVIEW=parser.parse_args().layout_preview
HERE=ROOT/'docs/slides';QA=HERE/'kor/1015_labmeeting_qa';ASSET=QA/('layout_preview_assets' if PREVIEW else 'compact');ASSET.mkdir(exist_ok=True)
BASE=HERE/'kor/1002_labmeeting_ko_draft.pptx';OUT=HERE/'kor/1015_labmeeting_ko.pptx'
if PREVIEW:OUT=QA/'layout_preview.pptx'
SMALL=ROOT/'tabpfn/results/v7/exp77/20261015_230535_initialft10_finalfull_s43_2eb2237d'
FOLLOWUP=ROOT/'tabpfn/results/v7/exp78/20261016_005435_noft_dist_new100_50_20_10_s43_52588487'
LARGE=ROOT/'tabpfn/results/v7/exp79/20261016_022105_saved100k_finalfull_s43_7f90866d'
EXPANDED=ROOT/'tabpfn/results/v7/exp80/20261016_022614_new200_400_noft_s43_c6f6ea24'
MIXED=ROOT/'tabpfn/results/v7/exp81/20261016_025421_mixed_query_final_s43_b6ba1d34'
small=json.loads((SMALL/'results.json').read_text())
assert small['state']=='complete' and json.loads((SMALL/'verification.json').read_text())['passed']
dist=json.loads((FOLLOWUP/'dist_results.json').read_text())
assert dist['state']=='complete' and json.loads((FOLLOWUP/'dist_verification.json').read_text())['passed']
small['jobs'].extend(dist['jobs'])
budget=json.loads((FOLLOWUP/'results.json').read_text())
assert budget['state']=='complete' and json.loads((FOLLOWUP/'verification.json').read_text())['passed']
HISTORICAL=ROOT/'docs/research/20261015/exp73_report_data.json'
historical=json.loads(HISTORICAL.read_text())
HIST_RUN=ROOT/historical['run'];HIST_BASE=ROOT/historical['baseline']
assert json.loads((HIST_RUN/'verification.json').read_text())['state']=='passed'
assert json.loads((HIST_BASE/'verification.json').read_text())['state']=='passed'
if PREVIEW:
    large_jobs=[dict(dataset=k.split('_')[0],case=k.split('_')[1],method=k.split('_')[2],metrics=v['results']['rate_1000']) for k,v in historical['jobs'].items()]
else:
    large=json.loads((LARGE/'results.json').read_text());expanded=json.loads((EXPANDED/'results.json').read_text())
    assert large['state']==expanded['state']=='complete'
    assert json.loads((LARGE/'verification.json').read_text())['passed'] and json.loads((EXPANDED/'verification.json').read_text())['passed']
    large_jobs=large['jobs'];budget['jobs'].extend(expanded['jobs'])
    mixed=json.loads((MIXED/'results.json').read_text())
    assert mixed['state']=='complete' and json.loads((MIXED/'verification.json').read_text())['passed']
    small['jobs']=[j for j in mixed['jobs'] if j['scope']=='small']
    large_jobs=[j for j in mixed['jobs'] if j['scope']=='large']
    for j in budget['jobs']:
        if j['new_per_class']==100 and j['method'] in ['xgb','pfn_pretrained']:
            j['metrics']=next(x for x in small['jobs'] if x['dataset']==j['dataset'] and x['method']==j['method'])['metrics']
NAVY,BLUE,INK,GREY,PALE='00306C','007BC6','21262E','65717D','F0F5FA'
PAPER='https://arxiv.org/pdf/2605.13986';FT='https://github.com/PriorLabs/TabPFN/blob/main/src/tabpfn/finetuning/finetuned_base.py'
records=[];text_boxes=[];chart_data=[];scenario_data=[];historical_counts=[];budget_counts=[];large_scenario_data=[]
prs=Presentation(BASE)
for s in list(prs.slides._sldIdLst):prs.part.drop_rel(s.rId);prs.slides._sldIdLst.remove(s)

font_manager.fontManager.addfont('/home/user/.local/share/fonts/Pretendard/Pretendard-Regular.otf')
plt.rcParams.update({'font.family':'Pretendard','font.size':15,'axes.unicode_minus':False,'svg.fonttype':'none','pdf.fonttype':42})

def save_figure(fig,name):
    for ext in ['png','pdf','svg']:fig.savefig(ASSET/f'{name}.{ext}',dpi=230,facecolor='white')
    plt.close(fig)

def axis_base(ax,labels,title,xmax=100,log=False):
    ax.set_title(title,fontsize=17,loc='left',pad=9,color='#'+NAVY,fontweight='bold')
    ax.set_yticks(np.arange(len(labels)),labels,fontsize=14)
    ax.invert_yaxis();ax.tick_params(axis='both',length=0,pad=5)
    for spine in ax.spines.values():spine.set_visible(False)
    ax.grid(axis='x',color='#E2E8EF',linewidth=.7);ax.set_axisbelow(True)
    if log:
        ax.set_xscale('log');ax.set_xlim(1,12000);ax.set_xticks([1,10,100,1000,10000],['1','10','100','1,000','10,000']);ax.minorticks_off()
    else:
        ax.set_xlim(0,xmax);ax.xaxis.set_major_locator(ticker.MaxNLocator(4))
    ax.tick_params(axis='x',labelsize=11)

def small_plot():
    methods=['xgb','pfn_pretrained']
    labels=['XGB','TabPFN']
    colors=['#80909F','#007BC6']
    fig,axes=plt.subplots(2,2,figsize=(13.1,5.24));fig.subplots_adjust(left=.151,right=.985,top=.915,bottom=.055,wspace=.63,hspace=.44)
    costmax=max(j['update_seconds'] for j in small['jobs'] if j['method'] in methods)*1.25
    for col,(ds,label) in enumerate([('cic2018','CIC2018'),('toniot','ToN')]):
        jobs={j['method']:j for j in small['jobs'] if j['dataset']==ds}
        for row,metric,title,maxval in [(0,'f1',label+' · 최종 Macro-F1 (%)',114),(1,'update',label+' · 후속 갱신 합계 (초)',costmax)]:
            ax=axes[row,col];axis_base(ax,labels,title,maxval)
            ax.tick_params(axis='y',labelsize=12)
            vals=[100*jobs[m]['metrics']['macro_f1'] if metric=='f1' else jobs[m]['update_seconds'] for m in methods]
            ax.barh(np.arange(len(methods)),vals,height=.48,color=colors)
            for i,(m,v) in enumerate(zip(methods,vals)):
                ax.text(v+maxval*.017,i,f'{v:.2f}',va='center',fontsize=12,color='#'+INK)
                chart_data.append(dict(slide=6,dataset=ds,case='final_full_test',method=m,metric=metric,value=v))
            if metric=='f1':ax.set_xticks([0,25,50,75,100])
    save_figure(fig,'small_context_results')

def comparison_plot():
    methods=['xgb','pfn_pretrained','ours','boostpfn','distpfn_pretrained']
    labels=['XGB','TabPFN','제안 모델','BoostPFN','DistPFN']
    colors=['#80909F','#007BC6','#00306C','#7EABB9','#36A2C8']
    fig,axes=plt.subplots(1,2,figsize=(13.1,4.70));fig.subplots_adjust(left=.115,right=.985,top=.88,bottom=.09,wspace=.45)
    for col,(ds,label) in enumerate([('cic2018','CIC2018'),('toniot','ToN')]):
        jobs={j['method']:j for j in small['jobs'] if j['dataset']==ds}
        ax=axes[col];axis_base(ax,labels,label+' · 최종 Macro-F1 (%)',114)
        ax.tick_params(axis='y',labelsize=16)
        vals=[100*jobs[m]['metrics']['macro_f1'] for m in methods]
        ax.barh(np.arange(len(methods)),vals,height=.55,color=colors)
        for y,(method,value) in enumerate(zip(methods,vals)):
            ax.text(value+1.8,y,f'{value:.2f}',va='center',fontsize=15,color='#'+INK)
            chart_data.append(dict(slide=7,dataset=ds,case='final_full_test',method=method,metric='f1',value=value))
        ax.set_xticks([0,25,50,75,100])
    save_figure(fig,'small_sota_results')

def budget_plot():
    fig,axes=plt.subplots(1,2,figsize=(13.1,4.55))
    fig.subplots_adjust(left=.07,right=.98,top=.84,bottom=.22,wspace=.22)
    for ax,(ds,label) in zip(axes,[('cic2018','CIC2018'),('toniot','ToN')]):
        jobs={(j['new_per_class'],j['method']):j for j in budget['jobs'] if j['dataset']==ds}
        levels=[10,20,50,100] if PREVIEW else [10,20,50,100,200,400];xs=np.arange(len(levels))
        for method,label_m,color,offset,va in [('xgb','XGB','#687C8F',-12,'top'),('pfn_pretrained','TabPFN · FT 없음','#007BC6',11,'bottom')]:
            values=[100*jobs[n,method]['metrics']['macro_f1'] for n in levels]
            ax.plot(xs,values,color=color,marker='o',markersize=6,linewidth=2.2,label=label_m)
            for x,n,v in zip(xs,levels,values):
                ax.annotate(f'{v:.2f}',(x,v),xytext=(0,offset),textcoords='offset points',ha='center',va=va,fontsize=11.5,color=color,fontweight='bold')
                chart_data.append(dict(slide=8,dataset=ds,case=f'new_{n}',method=method,metric='f1',value=v))
        totals=[jobs[n,'pfn_pretrained']['label_budget'] for n in levels]
        budget_counts.extend(dict(dataset=ds,new_per_class=n,total=t) for n,t in zip(levels,totals))
        ax.set_title(label,fontsize=19,loc='left',color='#'+NAVY,fontweight='bold',pad=15)
        ax.set_xticks(xs,[f'{n}\n총 {t:,}' for n,t in zip(levels,totals)],fontsize=11.5)
        ax.set_xlabel('신규 공격 클래스당 표본 수 (개)',labelpad=12,fontsize=14)
        ax.set_ylim(50,85);ax.set_yticks([50,60,70,80]);ax.set_xlim(-.42,len(levels)-.58)
        ax.set_ylabel('최종 Macro-F1 (%)' if ds=='cic2018' else '',fontsize=14)
        ax.grid(axis='y',color='#E2E8EF');ax.tick_params(axis='both',length=0,pad=8)
        for spine in ax.spines.values():spine.set_visible(False)
    handles,labels=axes[0].get_legend_handles_labels()
    fig.legend(handles,labels,loc='upper center',bbox_to_anchor=(.55,1.015),ncol=2,frameon=False,fontsize=14)
    save_figure(fig,'label_budget_results')

def historical_plot():
    methods=['xgb','tabpfn','ours','boostpfn','distpfn']
    labels=['XGB','TabPFN','제안 모델','BoostPFN','DistPFN']
    colors=['#80909F','#007BC6','#00306C','#7EABB9','#36A2C8']
    fig,axes=plt.subplots(2,2,figsize=(13.1,5.25))
    fig.subplots_adjust(left=.112,right=.982,top=.93,bottom=.045,wspace=.43,hspace=.50)
    for col,(ds,label) in enumerate([('cic2018','CIC2018'),('toniot','ToN')]):
        for row,(case,case_label) in enumerate([('natural','자연비율'),('uniform','균등 배분')]):
            ax=axes[row,col];axis_base(ax,labels,f'{label} · {case_label}',114)
            ax.title.set_fontsize(16);ax.tick_params(axis='y',labelsize=12)
            jobs={j['method']:j for j in large_jobs if j['dataset']==ds and j['case']==case}
            values=[100*jobs[method]['metrics']['macro_f1'] for method in methods]
            ax.barh(np.arange(len(methods)),values,height=.61,color=colors)
            for y,(method,value) in enumerate(zip(methods,values)):
                ax.text(value+1.7,y,f'{value:.2f}',va='center',fontsize=12,color='#'+INK)
                chart_data.append(dict(slide=11,dataset=ds,case=case,method=method,metric='f1',value=value))
            ax.set_xticks([0,25,50,75,100])
    save_figure(fig,'large_final_results')

def frame(tf,value,size=22,bold=False,color=INK,align=PP_ALIGN.LEFT):
    tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=tf.margin_right=Inches(.03);tf.margin_top=tf.margin_bottom=Inches(.025)
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(3);p.line_spacing=1.08
        r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        ea=OxmlElement('a:ea');ea.set('typeface','Pretendard');r._r.get_or_add_rPr().append(ea)
def text(s,x,y,w,h,value,**kw):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame(sh.text_frame,value,**kw)
    text_boxes.append(dict(page=len(prs.slides),x=x,y=y,w=w,h=h,text=str(value),size=kw.get('size',22)));return sh
def box(s,x,y,w,h,color=PALE):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.adjustments[0]=.03
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
def slide(title='',notes='',title_size=28):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s._element.set('showMasterSp','0');s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
    if title:text(s,.60,.29,12.1,.72,title,size=title_size,bold=True,color=NAVY)
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(7.02),prs.slide_width,Inches(.48));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(NAVY);sh.line.fill.background()
    text(s,12.48,7.09,.4,.25,str(len(prs.slides)),size=10,color='FFFFFF',align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes;records.append(dict(page=len(prs.slides),title=title,notes=notes));return s
def architecture(s):
    sh=s.shapes.add_picture(str(QA/'v3_architecture_page.png'),Inches(.68),Inches(1.08),width=Inches(4.0),height=Inches(5.85))
    sh.crop_left=.20;sh.crop_right=.19;sh.crop_top=.082;sh.crop_bottom=.285
def detail(s,y,heading,body):
    text(s,5.12,y,7.52,.48,heading,size=23,bold=True,color=BLUE)
    text(s,5.14,y+.59,7.48,.85,body,size=21)

small_plot();comparison_plot();budget_plot();historical_plot()

s=slide(notes='No title. Rephrase the user intent into concise presentation language instead of copying their draft. Adaptation to changing networks motivates the research; current class-arrival tests do not establish network-drift robustness.')
text(s,.93,1.75,11.48,1.40,'고정된 환경의 분류 정확도를 넘어\n변화하는 네트워크에서의 적응성 평가',size=31,color=INK)
text(s,.93,4.18,11.48,1.40,'사전학습 TabPFN을 활용한\n소량 라벨·저비용 탐지 모델 갱신',size=31,bold=True,color=BLUE)

s=slide('TabPFN · 표 데이터 예측 규칙을 사전학습한 모델 (v3: arXiv, 2026)',title_size=23,
        notes='Source: '+PAPER+' §2.5, PDF page 10: more than 8 trillion tokens, not datasets. Official TabPFN 8.2.0: '+FT+'. The archived initial-FT ablation in EXP77 uses AdamW, learning rate 1e-5, 10 fixed epochs, context/query cross-entropy; current displayed PFN results use no IDS FT. The v3 decoder weights context one-hot labels, so old-class examples remain necessary in the inference context. Validation is inside the 1000 revealed-label budget; no test tuning.')
for y,heading,body in [
    (1.40,'사전학습','합성 표 데이터의 8조 개 이상 토큰으로 학습\n일부 행의 정답을 맞히며 표 데이터의 예측 규칙 습득'),
    (2.68,'In-context learning','라벨된 사례(context)를 참고해 새 flow(query)의 라벨 예측\n추론 중 모델 가중치는 그대로 유지'),
    (3.96,'Fine-tuning','초기 IDS를 context와 query로 나눠 query의 정답 예측\n예측 오류를 역전파해 모델 가중치를 IDS에 맞게 조정')]:
    text(s,.7,y,2.7,1.0 if heading=='In-context learning' else .5,heading,size=22,bold=True,color=BLUE)
    text(s,3.45,y,9.15,1.0,body,size=21)
box(s,.7,5.45,11.92,1.12)
text(s,.91,5.55,11.5,.36,'갱신: 사전학습 가중치 유지 → 기존 클래스 사례 보존 + 신규 사례 누적',size=20,bold=True,color=NAVY)
text(s,.91,6.06,11.5,.34,'기대: 반복 재학습과 비슷한 탐지 품질을 더 적은 데이터·갱신 비용으로 달성',size=20,color=BLUE)

s=slide('v3 구조',notes='Original Figure 5, '+PAPER+'#page=6. Native PPTX crop excludes the original caption. Labels and source remain in speaker notes. First three stages: feature grouping/cell embedding; target-aware column distribution embedding; row-level feature aggregation. Plain language deliberately omits implementation dimensions.')
architecture(s)
detail(s,1.40,'① 입력 변환 · Cell embedding','패킷 길이·지속시간 등의 값을 숫자 벡터로 변환\n값이 비어 있는 위치도 함께 표시')
detail(s,3.18,'② 특징별 분포 파악 · Feature embedding','여러 flow를 살펴 각 특징의 값 분포를 요약\nContext의 정상·공격 라벨도 함께 반영')
detail(s,4.96,'③ Flow별 정보 통합 · Feature aggregation','한 flow의 여러 특징을 하나의 표현으로 통합\n이 표현을 이용해 flow 사이의 관계 비교')

s=slide('v3 구조',notes='Same Figure 5, '+PAPER+'#page=6; no visible caption by user request. ICL attends to training/context keys; many-class decoder uses attention-weighted one-hot context labels. This is learned contextual relevance, not merely raw-feature nearest-neighbour distance. Context determines available class labels.')
architecture(s)
detail(s,1.74,'④ 사례 참조 · In-context learning','새 flow(query)가 라벨된 기존 사례(context)를 참조\n어떤 사례가 예측에 도움이 되는지 attention으로 계산')
detail(s,4.15,'⑤ 확률 출력 · Many-class decoder','관련성이 큰 사례의 라벨에 더 큰 가중치 부여\n참고한 라벨을 모아 정상·공격별 확률 계산')
text(s,5.14,5.94,7.48,.64,'기존 클래스도 분류하려면\n해당 라벨의 context 유지 필요',size=19,bold=True,color=BLUE)

s=slide('평가시나리오',notes='Displayed XGB and no-FT PFN use the same released IDs as EXP76/77. XGB cumulatively retrains; PFN retains original pretrained weights and cumulative context. Initial and repeated IDS FT ablations remain in the experiment archive and are not shown. Proposed model uses five feature-disjoint roles within the SAME total label budget, original pretrained backbone to keep calibration labels outside weight training. Stage order chronological, within-stage sample counts descending. Main score is the FINAL state on ALL held-out rows, not pooled stage predictions.')
text(s,.66,1.02,12.0,.38,'초기 1,000개 · 신규 공격당 100개 누적 · 단계순, 단계 내 표본 수 내림차순',size=16,color=GREY)
stages={
 'cic2018':[['Benign','Brute force','DoS'],['DDoS'],['Web attacks'],['Infiltration'],['Bot']],
 'toniot':[['Benign','DoS','Scanning'],['DDoS','Injection'],['Password'],['XSS'],['Backdoor','Ransomware'],['MitM']]}
for ds,label,x in [('cic2018','CIC2018',.66),('toniot','ToN',6.90)]:
    text(s,x,1.50,5.75,.4,label,size=22,bold=True,color=BLUE)
    rows=[];total=0;heights=[.43]
    for i,classes in enumerate(stages[ds]):
        counts=[750,125,125] if i==0 else [100]*len(classes);total+=sum(counts)
        rows.append([str(i+1),'\n'.join(classes),'\n'.join(f'{n:,}' for n in counts),f'{total:,}'])
        heights.append(1.08 if i==0 else .74 if len(classes)==2 else .53)
        scenario_data.append(dict(dataset=ds,stage=i+1,classes=classes,new_counts=counts,cumulative=total))
    widths=[.65,2.43,1.22,1.36]
    sh=s.shapes.add_table(len(rows)+1,4,Inches(x),Inches(2.01),Inches(sum(widths)),Inches(sum(heights)));t=sh.table
    for c,w in enumerate(widths):t.columns[c].width=Inches(w)
    for r,values in enumerate([['단계','신규 공개 클래스','클래스별 수','누적 수']]+rows):
        t.rows[r].height=Inches(heights[r])
        for c,v in enumerate(values):
            cell=t.cell(r,c);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if r==0 else PALE if r%2 else 'FFFFFF')
            frame(cell.text_frame,v,size=14 if r==0 else 15.5,bold=r==0,color='FFFFFF' if r==0 else INK,align=PP_ALIGN.RIGHT if c>=2 and r else PP_ALIGN.CENTER if c==0 else PP_ALIGN.LEFT)
text(s,.7,6.68,11.95,.27,'전체 test: CIC2018 3,042,473개 / ToN 2,271,723개 · 최종 모델 · seed 43 혼합 순서',size=14.5,color=GREY)

s=slide('결과',notes='EXP77 final-model full-test results: '+str(SMALL.relative_to(ROOT))+'/results.json. Two displayed methods: XGB cumulative retraining from scratch at every stage and original pretrained TabPFN with cumulative context, no IDS fine-tuning. Initial-FT results are retained in EXP77/EXP78, removed from slides 6 and 7 at user request. Slide 6 alone reports later-stage update cost: GPU-synchronized context selection + model construction + fit summed AFTER the initial stage; initial preparation and all inference excluded. Single seed 43. A separate EXP78 label-budget artifact covers 100/50/20/10 new examples per class, with initial 1000 fixed.')
text(s,.66,1.04,12.0,.36,'최종 상태 전체 test 평가 · XGB: 누적 재학습 · TabPFN: FT 없이 누적 context 갱신',size=16.5,color=GREY)
s.shapes.add_picture(str(ASSET/'small_context_results.png'),Inches(.26),Inches(1.47),width=Inches(12.75),height=Inches(5.10))
text(s,.69,6.63,12.0,.25,'비용: 2단계 이후 갱신 합계 · 초기 준비·추론 제외 · 두 데이터셋 동일 비용 눈금',size=14,color=GREY)

s=slide('제안 모델·SOTA 성능 비교',notes='Same EXP77 released-label budget: CIC2018 1400 / ToN 1700, including all internal validation and calibration. Only final full-test Macro-F1 is displayed; all cost panels and initial-FT baseline removed per user request. Existing K=4 anchor/residual + S1V0 proposed model uses original pretrained weights and global/expert/route/tune/cal roles within that total budget. Both datasets have zero expert calls in the final evaluation, so this is a diagnosed current configuration, not demonstrated ensemble improvement. Candidate expert corrections changed correctness on at most 0/4 calibration rows (CIC2018/ToN), below the existing 30-case requirement. Next diagnostic controls are a single PFN with identical Global context and adjustment of shared anchors/expert-specific allocation; do not retune thresholds on test. BoostPFN is the existing 50-learner port. DistPFN uses original pretrained v3 without IDS FT and current-query-batch posterior adjustment; no query labels. The appendix separately evaluates the saved large-data final model on the same complete test; its training and validation budgets differ from this small-data comparison. Source: '+str((FOLLOWUP/'dist_results.json').relative_to(ROOT)))
text(s,.66,1.02,12.0,.36,'동일 라벨 예산: CIC2018 1,400개 / ToN 1,700개 · 최종 상태 전체 test 평가',size=16,color=GREY)
s.shapes.add_picture(str(ASSET/'small_sota_results.png'),Inches(.22),Inches(1.55),width=Inches(12.8),height=Inches(4.59))
text(s,.69,6.60,12.0,.30,'TabPFN·DistPFN: FT 없음 · test를 8,192개씩 혼합 평가 · 제안 모델 expert 호출 0건',size=14,color=GREY)

s=slide('신규 공격 표본 수에 따른 성능',notes='EXP78 and EXP80. Initial 1000 fixed (Benign750 + two attacks125 each), seed43 once, no IDS FT. EXP80 preserves the original 100 incoming IDs and appends nested 200/400 samples; Web attacks availability 450 limits the shared budget. Final models predict the complete held-out test. EXP81 aligns query composition and audits row-wise prediction reuse. Varying new-class counts also changes class priors. Sources: '+str(FOLLOWUP.relative_to(ROOT))+' and '+str(EXPANDED.relative_to(ROOT)))
text(s,.66,1.04,12.0,.36,'초기 1,000개 고정 · 신규 공격당 10·20·50·100·200·400개 · 최종 상태 전체 test · seed 43' if not PREVIEW else '배치 검토용 · 기존 10·20·50·100개 결과 · 200·400개 추가 실험 대기',size=15.5,color=GREY)
s.shapes.add_picture(str(ASSET/'label_budget_results.png'),Inches(.22),Inches(1.62),width=Inches(12.8),height=Inches(4.45))
box(s,.72,6.18,11.9,.55)
n=100 if PREVIEW else 400
gaps={ds:100*(next(j for j in budget['jobs'] if j['dataset']==ds and j['new_per_class']==n and j['method']=='pfn_pretrained')['metrics']['macro_f1']-next(j for j in budget['jobs'] if j['dataset']==ds and j['new_per_class']==n and j['method']=='xgb')['metrics']['macro_f1']) for ds in ['cic2018','toniot']}
text(s,.90,6.23,11.5,.40,f"신규 공격 {n}개 · XGB 대비 TabPFN 차이: CIC2018 {gaps['cic2018']:+.2f}%p · ToN {gaps['toniot']:+.2f}%p",size=18,bold=True,color=NAVY)

display={'benign':'Benign','brute_force':'Brute force','dos':'DoS','ddos':'DDoS','web_attacks':'Web attacks','infiltration':'Infiltration','bot':'Bot','scanning':'Scanning','injection':'Injection','password':'Password','xss':'XSS','backdoor':'Backdoor','ransomware':'Ransomware','mitm':'MitM'}
for ds,label in [('cic2018','CIC2018'),('toniot','ToN')]:
    s=slide(f'부록 · 학습 10만 개 시나리오 · {label}',notes='Original EXP72/73 stage-wise data. Same attack introduction order as current small-data experiment, but Benign is also added each stage: 15000 CIC2018 /12500 ToN, reaching75000. Small-data Benign750 is initial-only. Natural and uniform each reach25000 attack rows; uniform respects class availability caps. Rows are chronological; classes within each stage sorted by NATURAL incoming count descending. Each numeric cell lists class-specific newly supplied samples in matching order; cumulative totals are separate columns. Separate validation is outside100k; tune40% selects residual temperature, cal60% selects expert gate thresholds. Original LoCalPFN uses validation checkpoint selection; XGB/TabPFN/BoostPFN/DistPFN do not use those labels. Small proposed validation is inside total1400/1700. Sources: '+str(HIST_RUN.relative_to(ROOT))+' and '+str(HIST_BASE.relative_to(ROOT)))
    benign_per_stage=15000 if ds=='cic2018' else 12500
    text(s,.66,1.03,12.0,.36,f'동일 공격 순서 · 소규모: 신규 공격당 100개 · 대규모: Benign 매 단계 {benign_per_stage:,}개 추가',size=15.5,color=GREY)
    text(s,3.45,1.48,4.35,.30,'이번 단계 추가 수 (개)',size=16,bold=True,color=BLUE,align=PP_ALIGN.CENTER)
    text(s,7.80,1.48,4.80,.30,'누적 수 (개)',size=16,bold=True,color=BLUE,align=PP_ALIGN.CENTER)
    previous={case:{} for case in ['natural','uniform']};rows=[];heights=[.40]
    small_stages=[r for r in scenario_data if r['dataset']==ds]
    for i,current in enumerate(small_stages):
        counts={};deltas={}
        for case in previous:
            fit=historical['jobs'][f'{ds}_{case}_tabpfn']['fits'][i]
            counts[case]=dict(zip(fit['class_names'],fit['train_class_counts']))
            deltas[case]={k:v-previous[case].get(k,0) for k,v in counts[case].items() if v-previous[case].get(k,0)>0}
        small_new={name.replace(' ','_').lower():n for name,n in zip(current['classes'],current['new_counts'])}
        names=sorted(set(deltas['natural'])|set(small_new),key=lambda k:(-deltas['natural'].get(k,0),k))
        values={case:[deltas[case].get(k,0) for k in names] for case in deltas};values['small']=[small_new.get(k,0) for k in names]
        totals=dict(small=current['cumulative'],natural=sum(counts['natural'].values()),uniform=sum(counts['uniform'].values()))
        rows.append([str(i+1),'\n'.join(display[k] for k in names)]+['\n'.join(f'{n:,}' for n in values[k]) for k in ['small','natural','uniform']]+[f'{totals[k]:,}' for k in ['small','natural','uniform']])
        heights.append(.72 if len(names)==3 else .49)
        large_scenario_data.append(dict(page=len(prs.slides),dataset=ds,stage=i+1,classes=names,new_counts=values,cumulative=totals))
        previous=counts
    current_counts={name.replace(' ','_').lower():n for r in small_stages for name,n in zip(r['classes'],r['new_counts'])}
    for name in sorted(previous['natural'],key=lambda k:(-previous['natural'][k],k)):
        historical_counts.append(dict(dataset=ds,name=name,natural=previous['natural'][name],uniform=previous['uniform'][name],current=current_counts[name]))
    validations={case:historical['jobs'][f'{ds}_{case}_ours']['fits'][-1]['validation_rows'] for case in previous}
    rows.append(['','별도 검증','','','','0',f"{validations['natural']:,}",f"{validations['uniform']:,}"]);heights.append(.38)
    widths=[.65,2.10,1.45,1.45,1.45,1.60,1.60,1.60]
    sh=s.shapes.add_table(len(rows)+1,8,Inches(.70),Inches(1.86),Inches(sum(widths)),Inches(sum(heights)));table=sh.table
    for c,w in enumerate(widths):table.columns[c].width=Inches(w)
    for r,values in enumerate([['단계','추가 클래스','소규모','10만 자연','10만 균등','소규모','10만 자연','10만 균등']]+rows):
        table.rows[r].height=Inches(heights[r])
        for c,value in enumerate(values):
            cell=table.cell(r,c);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if r==0 else PALE if r%2 else 'FFFFFF')
            cell.margin_top=cell.margin_bottom=Inches(.015);cell.margin_left=cell.margin_right=Inches(.06)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            frame(cell.text_frame,value,size=13 if r==0 else 15,bold=r==0,color='FFFFFF' if r==0 else INK,align=PP_ALIGN.RIGHT if c>=2 and r else PP_ALIGN.CENTER if c==0 else PP_ALIGN.LEFT)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after=Pt(0);paragraph.space_before=Pt(0)
                paragraph.font.size=Pt(13 if r==0 else 15);paragraph.font.name='Pretendard'
                paragraph.line_spacing=Pt(14.5 if r==0 else 16.2)
    text(s,.69,6.40,12.0,.26,'별도 검증: 제안 모델의 잔차 확률 보정·expert 채택 기준 선택 · 소규모는 전체 예산 내 분할',size=14,color=GREY)
    text(s,.69,6.73,12.0,.23,'자연: 학습 공격 비율 · 균등: 공격별 균등, 표본 상한 반영 · 단계순, 단계 내 자연 추가 수 내림차순',size=13.5,color=GREY)

s=slide('부록 · 학습 10만 개의 최종 성능',notes='EXP79/81 inference-only reevaluation: immutable saved EXP72/73 FINAL stage model states, no retraining or threshold retuning. Complete test, mixed query batch8192, identical final-state evaluation convention to main slides. DistPFN adjusts current completed query-batch posterior only, without query labels. No synthetic arrival rate, queue or update delay in these scores. Five methods; LoCalPFN omitted under prior user scope because of slow inference. Separate validation remains available to proposed model, so total label budgets differ from small-data setup beyond training size. BoostPFN v1; other PFN models v3. '+str((MIXED/'results.json').relative_to(ROOT)))
text(s,.66,1.03,12.0,.36,'최종 모델로 전체 test 평가 · Macro-F1 (%) · 학습 10만 개 + 별도 검증' if not PREVIEW else '배치 검토용 · 기존 순차 재생 점수 표시 · 최종 모델 전체 test 재평가 진행 중',size=16,color=GREY)
s.shapes.add_picture(str(ASSET/'large_final_results.png'),Inches(.22),Inches(1.48),width=Inches(12.8),height=Inches(5.13))
text(s,.69,6.67,12.0,.24,'본문과 동일 test·혼합 순서·최종 상태 평가 · 학습·검증 예산 상이 · BoostPFN v1, 나머지 PFN v3',size=13.5,color=GREY)

for page in [5,6,7,8,11]:
    note=' Final evaluation convention: EXP81 seed43 permutation of test row indices independent of labels, query batch8192. PFN probabilities and DistPFN correction recomputed on all mixed batches. Row-wise XGB/proposed/Boost predictions retained only after exact mixed-query audit; any mismatching model is fully reevaluated. Budget10/20/50/200/400 PFN predictions checked with the same audit. Original-order results preserved in EXP77/78/79/80. No model tuning on test. Source: '+str((MIXED/'results.json').relative_to(ROOT))
    prs.slides[page-1].notes_slide.notes_text_frame.text+=note
    records[page-1]['notes']+=note

prs.save(OUT)
with (ASSET/'plotted_values.csv').open('w',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=['slide','dataset','case','method','metric','value']);writer.writeheader();writer.writerows(chart_data)
sources=[SMALL/'results.json',SMALL/'verification.json',FOLLOWUP/'dist_results.json',FOLLOWUP/'dist_verification.json',FOLLOWUP/'results.json',FOLLOWUP/'verification.json',HISTORICAL,HIST_RUN/'verification.json',HIST_BASE/'verification.json']
if not PREVIEW:sources += [LARGE/'results.json',LARGE/'verification.json',EXPANDED/'results.json',EXPANDED/'verification.json',MIXED/'results.json',MIXED/'verification.json',MIXED/'budget_order_audits.json']
manifest=dict(revision='compact_8_expanded_budget_mixed_final_appendix',layout_preview=PREVIEW,slides=records,text_boxes=text_boxes,chart_data=chart_data,scenario_data=scenario_data,historical_counts=historical_counts,budget_counts=budget_counts,large_scenario_data=large_scenario_data,
 source_deck=str(BASE.relative_to(ROOT)),source_deck_sha256=hashlib.sha256(BASE.read_bytes()).hexdigest(),output=str(OUT.relative_to(ROOT)),
 primary_v3_source=PAPER,pretraining_quantity='more than 8 trillion tokens, not dataset count; report section 2.5',
 new_source=str((SMALL/'results.json').relative_to(ROOT)),additional_sources=[str((FOLLOWUP/'dist_results.json').relative_to(ROOT))],budget_source=str((FOLLOWUP/'results.json').relative_to(ROOT)),historical_source=str(HISTORICAL.relative_to(ROOT)),
 expanded_budget_source=str((EXPANDED/'results.json').relative_to(ROOT)),large_final_source=str((LARGE/'results.json').relative_to(ROOT)),
 mixed_final_source=str((MIXED/'results.json').relative_to(ROOT)),query_order='seed43_mixed',query_batch_rows=8192,
 main_slide_count=8,appendix_slides=[9,10,11],historical_evaluation='layout_preview_replay' if PREVIEW else 'final_full_test',
 cost_metric_by_slide={'6':'update_seconds','7':None},cost_excludes=['initial_preparation','all_inference'],cost_axis='linear_shared_between_datasets_on_slide_6',distpfn_initial_ft=False,
 current_results_include_initial_ids_finetuning=False,source_experiment_includes_initial_ids_finetuning=True,evaluation='final_full_test',proposed_design='pretrained weights fixed, cumulative labelled context; initial-FT ablation retained in experiment records',
 original_25_slide_archive='docs/slides/archive/20261015_before_7slide_revision',figure_caption_visible=False,
 source_sha256={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources})
(QA/('layout_preview_manifest.json' if PREVIEW else 'manifest.json')).write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
plan=['# 1015 랩미팅 · 본문 8장 + 부록 3장','','6장은 XGB 누적 재학습과 FT 없는 TabPFN context 갱신의 최종 성능·갱신 시간 비교. 7장은 소규모 제안 모델·SOTA 최종 성능. 8장은 신규 공격당 10·20·50·100·200·400개 결과. 9·10장은 데이터셋별 공격 도입 단계와 클래스별 추가·누적 수 및 별도 검증 용도. 11장은 학습 10만 개의 저장된 최종 모델을 전체 test로 다시 평가한 성능. 날짜형 실험 별칭과 유입률 표기는 발표 화면에서 제거. 초기 FT 및 과거 재생 결과는 원본 기록에 보존.','']
for r in records:plan.extend([f"## {r['page']}. {r['title'] or '(타이틀 없음)'}",'',r['notes'],''])
if not PREVIEW:(HERE/'kor/1015_labmeeting_plan.md').write_text('\n'.join(plan)+'\n')
print(OUT, len(prs.slides),'slides')
