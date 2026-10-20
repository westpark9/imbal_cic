#!/usr/bin/env python3
"""Check compact slides and historical appendix against measured sources."""
from pathlib import Path
import json,re,hashlib,subprocess
import numpy as np
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/experiment_paths.py').is_file())
qa=ROOT/'docs/slides/kor/1015_labmeeting_qa'
m=json.loads((qa/'manifest.json').read_text())
result=json.loads((ROOT/m['new_source']).read_text())
out=ROOT/m['output'];pdf=out.with_suffix('.pdf');prs=Presentation(out)
assert len(prs.slides)==10 and m['evaluation']=='final_full_test'
assert m['main_slide_count']==8 and m['appendix_slides']==[9,10]
budget=json.loads((ROOT/m['budget_source']).read_text())
budget_jobs={(j['dataset'],j['method'],j['new_per_class']):j for j in budget['jobs']}
history=json.loads((ROOT/m['historical_source']).read_text())
assert m['historical_policy']=='rate_1000' and m['historical_evaluation']=='pooled_stage_version_predictions'
assert m['current_results_include_initial_ids_finetuning'] is False
jobs={(j['dataset'],j['method']):j for j in result['jobs']}
for source in m.get('additional_sources',[]):
    extra=json.loads((ROOT/source).read_text())
    jobs.update({(j['dataset'],j['method']):j for j in extra['jobs']})
assert m['cost_metric_by_slide']=={'6':'update_seconds','7':None} and m['distpfn_initial_ft'] is False
assert all(r['metric'] in ('f1','update') for r in m['chart_data'])
assert {r['method'] for r in m['chart_data'] if r['slide']==6}=={'xgb','pfn_pretrained'}
assert all(r['metric']=='f1' for r in m['chart_data'] if r['slide']==7)
assert 'pfn_ft_accumulate' not in {r['method'] for r in m['chart_data']}
assert 'distpfn_ft' not in {r['method'] for r in m['chart_data']}
assert 'distpfn_pretrained' in {r['method'] for r in m['chart_data'] if r['slide']==7}
for ds in ['cic2018','toniot']:
    j=jobs[ds,'distpfn_pretrained'];base=jobs[ds,'pfn_pretrained']
    assert j['initial_finetuning_seconds']==0 and j['update_seconds']==base['update_seconds']
    assert j['restored_base_prediction_mismatches']==0
for source,digest in m['source_sha256'].items():
    assert hashlib.sha256((ROOT/source).read_bytes()).hexdigest()==digest
for row in m['chart_data']:
    if row['slide']==8:
        j=budget_jobs[row['dataset'],row['method'],int(row['case'].removeprefix('new_'))]
        v=100*j['metrics']['macro_f1']
    elif row['slide']==10:
        key=f"{row['dataset']}_{row['case']}_{row['method']}"
        original=json.loads((ROOT/history['run']/'jobs'/key/'results.json').read_text())
        v=100*original['results'][m['historical_policy']]['macro_f1']
    else:
        j=jobs[row['dataset'],row['method']]
        v={'f1':100*j['metrics']['macro_f1'],'update':j['update_seconds']}[row['metric']]
    assert abs(row['value']-v)<1e-10
assert {r['method'] for r in m['chart_data'] if r['slide']==8}=={'xgb','pfn_pretrained'}
for row in m['budget_counts']:
    assert row['total']==budget_jobs[row['dataset'],'pfn_pretrained',row['new_per_class']]['label_budget']
    assert row['total']==1000+row['new_per_class']*(4 if row['dataset']=='cic2018' else 7)
for source,digest in history['source_hashes'].items():
    assert hashlib.sha256((ROOT/source).read_bytes()).hexdigest()==digest
for dataset in ['cic2018','toniot']:
    run=(ROOT/m['new_source']).parent;meta=json.loads((run/'data'/dataset/'manifest.json').read_text())
    source=Path(meta['source'])
    y=np.load(source/'train_y.npy');names=meta['class_names']
    plans=json.loads((run/'data'/dataset/'plans.json').read_text())
    for row in [r for r in m['scenario_data'] if r['dataset']==dataset]:
        plan=plans[row['stage']-1];assert row['cumulative']==len(plan['released'])
        actual={names[c].replace('_',' ').lower():int((y[plan['incoming']]==c).sum()) for c in np.unique(y[plan['incoming']])}
        assert actual=={name.lower():n for name,n in zip(row['classes'],row['new_counts'])}
    rows=[r for r in m['historical_counts'] if r['dataset']==dataset]
    assert [r['natural'] for r in rows]==sorted([r['natural'] for r in rows],reverse=True)
    current={names[c]:int((y[plans[-1]['released']]==c).sum()) for c in np.unique(y[plans[-1]['released']])}
    assert {r['name']:r['current'] for r in rows}==current
    histmeta=json.loads((ROOT/history['baseline']/'data'/dataset/'manifest.json').read_text())
    histdata=Path(histmeta['source']);train_ids=np.load(histdata/'train_ids.npy');train_y=np.load(histdata/'train_y.npy')
    sorter=np.argsort(train_ids)
    for case in ['natural','uniform']:
        job=ROOT/history['baseline']/'jobs'/f'{dataset}_{case}_tabpfn'
        stage=sorted(job.glob('stage_*'))[-1];ids=np.load(stage/'train_ids.npy')
        positions=sorter[np.searchsorted(train_ids[sorter],ids)]
        assert np.array_equal(train_ids[positions],ids)
        actual={name:int((train_y[positions]==c).sum()) for c,name in enumerate(histmeta['class_names'])}
        assert actual=={r['name']:r[case] for r in rows},(dataset,case,actual)
        assert sum(r[case] for r in rows)==100000
        fit=json.loads((stage/'fit_metrics.json').read_text())
        assert dict(zip(fit['class_names'],fit['train_class_counts']))==actual
    # Verify the actual PPTX count table, including the separate validation row.
    tables=[sh.table for sh in prs.slides[8].shapes if sh.has_table]
    table=tables[0 if dataset=='cic2018' else 1]
    for i,row in enumerate(rows,1):
        assert [table.cell(i,c).text for c in range(1,4)]==[f"{row[k]:,}" for k in ['natural','uniform','current']]
    vr=len(rows)+2
    for c,case in enumerate(['natural','uniform'],1):
        stage=sorted((ROOT/history['run']/'jobs'/f'{dataset}_{case}_ours').glob('stage_*'))[-1]
        val=json.loads((stage/'fit_metrics.json').read_text())['validation_rows']
        assert table.cell(vr,c).text==f'{val:,}'
    assert table.cell(vr,3).text=='0'
issues=[];overlaps=[];content_overlaps=[]
for i,slide in enumerate(prs.slides,1):
    texts=[];contents=[]
    for shape in slide.shapes:
        if (shape.has_text_frame and shape.text) or shape.has_table or shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
            contents.append(shape)
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+10000
            assert shape.top+shape.height<=prs.slide_height+10000
        if shape.has_text_frame and shape.text:
            assert not re.search(r'합니다|입니다|습니다|—',shape.text),(i,shape.text)
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+10000
            assert shape.top+shape.height<=prs.slide_height+10000
            texts.append(shape)
        if shape.has_table:
            for row in list(shape.table.rows)[1:]:
                for cell in list(row.cells)[1:]:
                    if re.fullmatch(r'[\d,\s]+',cell.text):
                        assert all(p.alignment==PP_ALIGN.RIGHT for p in cell.text_frame.paragraphs)
    for a,shape in enumerate(texts):
        for other in texts[a+1:]:
            w=min(shape.left+shape.width,other.left+other.width)-max(shape.left,other.left)
            h=min(shape.top+shape.height,other.top+other.height)-max(shape.top,other.top)
            if w>10000 and h>10000:overlaps.append((i,shape.text,other.text))
    for a,shape in enumerate(contents):
        for other in contents[a+1:]:
            w=min(shape.left+shape.width,other.left+other.width)-max(shape.left,other.left)
            h=min(shape.top+shape.height,other.top+other.height)-max(shape.top,other.top)
            if w>10000 and h>10000:content_overlaps.append((i,shape.name,other.name))
assert not overlaps,overlaps
assert not content_overlaps,content_overlaps
assert hashlib.sha256((ROOT/m['source_deck']).read_bytes()).hexdigest()==m['source_deck_sha256']
pages=int(re.search(r'Pages:\s+(\d+)',subprocess.check_output(['pdfinfo',str(pdf)],text=True)).group(1));assert pages==10
text=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
main_text=subprocess.check_output(['pdftotext','-f','1','-l','8','-layout',str(pdf),'-'],text=True)
assert '1014' not in main_text and '미적용 사전 비교' not in text
assert '부록' in text and '1014' in text and '별도 검증' in text
assert '100·50·20·10' in text and '단계별 버전의 예측을 합산' in text
assert '최종' in text and '제안 모델' in text
assert '준비 + 전체 추론' not in text and '비용은 로그눈금' not in text
page7=subprocess.check_output(['pdftotext','-f','7','-l','7','-layout',str(pdf),'-'],text=True)
assert '후속 갱신' not in page7 and '갱신 합계' not in page7
assert not any('초기 FT' in sh.text for slide in list(prs.slides)[5:7] for sh in slide.shapes if sh.has_text_frame)
(qa/'compact/rendered_text.txt').write_text(text)
validation=dict(revision=m['revision'],slides=10,pdf_pages=pages,plot_values_verified=len(m['chart_data']),
    scenario_stages_verified=len(m['scenario_data']),numeric_columns_right_aligned=True,
    initial_ids_finetuning_results=False,evaluation='final_full_test',historical_1014_removed_from_main=True,
    historical_appendix_pages=[9,10],historical_class_counts_verified=len(m['historical_counts']),
    historical_counts_verified_from_original_train_ids=True,historical_evaluation='pooled_stage_version_predictions',
    new_class_budgets=[100,50,20,10],main_budget_slide=8,
    same_label_budget_including_validation=True,unintended_text_overlap_count=len(overlaps),source_deck_unchanged=True,
    unintended_content_overlap_count=len(content_overlaps),
    pptx_sha256=hashlib.sha256(out.read_bytes()).hexdigest(),pdf_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),
    update_cost_on_slide_6_only=True,sota_slide_performance_only=True,initial_ft_removed_from_slides_6_7=True,
    distpfn_without_finetuning=True,visual_review='pending')
for dest in [qa/'validation.json',qa/'compact/validation.json']:dest.write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
