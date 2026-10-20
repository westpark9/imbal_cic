#!/usr/bin/env python3
"""Seven-slide revision with measured-result plots and explicit adaptation scope."""
from pathlib import Path
import sys,json,hashlib,csv
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/experiment_paths.py').is_file())
sys.path.insert(0,str(ROOT/'scripts/common'))
from experiment_paths import repo_root
ROOT=repo_root(__file__)
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager,ticker
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

HERE=ROOT/'docs/slides';QA=HERE/'kor/1015_labmeeting_qa';ASSET=QA/'compact';ASSET.mkdir(exist_ok=True)
BASE=HERE/'kor/1002_labmeeting_ko_draft.pptx';OUT=HERE/'kor/1015_labmeeting_ko.pptx'
SMALL=ROOT/'tabpfn/results/v7/exp77/20261015_230535_initialft10_finalfull_s43_2eb2237d'
FOLLOWUP=ROOT/'tabpfn/results/v7/exp78/20261016_005435_noft_dist_new100_50_20_10_s43_52588487'
small=json.loads((SMALL/'results.json').read_text())
assert small['state']=='complete' and json.loads((SMALL/'verification.json').read_text())['passed']
dist=json.loads((FOLLOWUP/'dist_results.json').read_text())
assert dist['state']=='complete' and json.loads((FOLLOWUP/'dist_verification.json').read_text())['passed']
small['jobs'].extend(dist['jobs'])
NAVY,BLUE,INK,GREY,PALE='00306C','007BC6','21262E','65717D','F0F5FA'
PAPER='https://arxiv.org/pdf/2605.13986';FT='https://github.com/PriorLabs/TabPFN/blob/main/src/tabpfn/finetuning/finetuned_base.py'
records=[];text_boxes=[];chart_data=[];scenario_data=[]
prs=Presentation(BASE)
for s in list(prs.slides._sldIdLst):prs.part.drop_rel(s.rId);prs.slides._sldIdLst.remove(s)

font_manager.fontManager.addfont('/home/user/.local/share/fonts/Pretendard/Pretendard-Regular.otf')
plt.rcParams.update({'font.family':'Pretendard','font.size':15,'axes.unicode_minus':False,'svg.fonttype':'none','pdf.fonttype':42})

def save_figure(fig,name):
    for ext in ['png','pdf','svg']:fig.savefig(ASSET/f'{name}.{ext}',dpi=230,facecolor='white')
    plt.close(fig)

def axis_base(ax,labels,title,xmax=100,log=False):
    ax.set_title(title,fontsize=17,loc='left',pad=9,color='#'+NAVY,fontweight='bold')
    ax.set_yticks(np.arange(len(labels)),labels,fontsize=14)
    ax.invert_yaxis();ax.tick_params(axis='both',length=0,pad=5)
    for spine in ax.spines.values():spine.set_visible(False)
    ax.grid(axis='x',color='#E2E8EF',linewidth=.7);ax.set_axisbelow(True)
    if log:
        ax.set_xscale('log');ax.set_xlim(1,12000);ax.set_xticks([1,10,100,1000,10000],['1','10','100','1,000','10,000']);ax.minorticks_off()
    else:
        ax.set_xlim(0,xmax);ax.xaxis.set_major_locator(ticker.MaxNLocator(4))
    ax.tick_params(axis='x',labelsize=11)

def small_plot():
    methods=['xgb','pfn_pretrained','pfn_ft_accumulate']
    labels=['XGB','PFN · FT 없음','PFN · 초기 FT']
    colors=['#80909F','#A4ADB6','#007BC6']
    fig,axes=plt.subplots(2,2,figsize=(13.1,5.24));fig.subplots_adjust(left=.151,right=.985,top=.915,bottom=.055,wspace=.63,hspace=.44)
    costmax=common_cost_max()
    for col,(ds,label) in enumerate([('cic2018','CIC2018'),('toniot','ToN')]):
        jobs={j['method']:j for j in small['jobs'] if j['dataset']==ds}
        for row,metric,title,maxval in [(0,'f1',label+' · 최종 Macro-F1 (%)',114),(1,'update',label+' · 후속 갱신 합계 (초)',costmax)]:
            ax=axes[row,col];axis_base(ax,labels,title,maxval)
            ax.tick_params(axis='y',labelsize=12)
            vals=[100*jobs[m]['metrics']['macro_f1'] if metric=='f1' else jobs[m]['update_seconds'] for m in methods]
            ax.barh(np.arange(len(methods)),vals,height=.48,color=colors)
            for i,(m,v) in enumerate(zip(methods,vals)):
                ax.text(v+maxval*.017,i,f'{v:.2f}',va='center',fontsize=12,color='#'+INK)
                chart_data.append(dict(slide=6,dataset=ds,case='final_full_test',method=m,metric=metric,value=v))
            if metric=='f1':ax.set_xticks([0,25,50,75,100])
    save_figure(fig,'small_context_results')

def common_cost_max():
    methods=['xgb','pfn_pretrained','pfn_ft_accumulate','ours','boostpfn','distpfn_pretrained']
    return max(j['update_seconds'] for j in small['jobs'] if j['method'] in methods)*1.25

def comparison_plot():
    methods=['xgb','pfn_pretrained','pfn_ft_accumulate','ours','boostpfn','distpfn_pretrained']
    labels=['XGB','PFN · FT 없음','PFN · 초기 FT','제안 모델','BoostPFN','DistPFN · FT 없음']
    colors=['#80909F','#A4ADB6','#007BC6','#00306C','#7EABB9','#36A2C8']
    fig,axes=plt.subplots(2,2,figsize=(13.1,5.17));fig.subplots_adjust(left=.16,right=.985,top=.925,bottom=.065,wspace=.68,hspace=.50)
    for col,(ds,label) in enumerate([('cic2018','CIC2018'),('toniot','ToN')]):
        jobs={j['method']:j for j in small['jobs'] if j['dataset']==ds}
        for row,metric,title in [(0,'f1',label+' · 최종 Macro-F1 (%)'),(1,'update',label+' · 후속 갱신 합계 (초)')]:
            ax=axes[row,col];axis_base(ax,labels,title,114 if metric=='f1' else common_cost_max())
            ax.tick_params(axis='y',labelsize=12)
            vals=[]
            for m in methods:
                job=jobs[m]
                v=100*job['metrics']['macro_f1'] if metric=='f1' else job['update_seconds']
                vals.append(v);chart_data.append(dict(slide=7,dataset=ds,case='final_full_test',method=m,metric=metric,value=v))
            yy=np.arange(len(methods))
            ax.barh(yy,vals,height=.6,color=colors)
            for y,v in zip(yy,vals):ax.text(v+(1.8 if metric=='f1' else common_cost_max()*.017),y,f'{v:,.2f}',va='center',fontsize=12,color='#'+INK)
            if metric=='f1':ax.set_xticks([0,25,50,75,100])
    save_figure(fig,'small_sota_results')

def frame(tf,value,size=22,bold=False,color=INK,align=PP_ALIGN.LEFT):
    tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=tf.margin_right=Inches(.03);tf.margin_top=tf.margin_bottom=Inches(.025)
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(3);p.line_spacing=1.08
        r=p.add_run();r.text=line;r.font.name='Pretendard';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        ea=OxmlElement('a:ea');ea.set('typeface','Pretendard');r._r.get_or_add_rPr().append(ea)
def text(s,x,y,w,h,value,**kw):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame(sh.text_frame,value,**kw)
    text_boxes.append(dict(page=len(prs.slides),x=x,y=y,w=w,h=h,text=str(value),size=kw.get('size',22)));return sh
def box(s,x,y,w,h,color=PALE):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.adjustments[0]=.03
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
def slide(title='',notes='',title_size=28):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s._element.set('showMasterSp','0');s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
    if title:text(s,.60,.29,12.1,.72,title,size=title_size,bold=True,color=NAVY)
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(7.02),prs.slide_width,Inches(.48));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(NAVY);sh.line.fill.background()
    text(s,12.48,7.09,.4,.25,str(len(prs.slides)),size=10,color='FFFFFF',align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes;records.append(dict(page=len(prs.slides),title=title,notes=notes));return s
def architecture(s):
    sh=s.shapes.add_picture(str(QA/'v3_architecture_page.png'),Inches(.68),Inches(1.08),width=Inches(4.0),height=Inches(5.85))
    sh.crop_left=.20;sh.crop_right=.19;sh.crop_top=.082;sh.crop_bottom=.285
def detail(s,y,heading,body):
    text(s,5.12,y,7.52,.48,heading,size=23,bold=True,color=BLUE)
    text(s,5.14,y+.59,7.48,.85,body,size=21)

small_plot();comparison_plot()

s=slide(notes='No title. Rephrase the user intent into concise presentation language instead of copying their draft. Adaptation to changing networks motivates the research; current class-arrival tests do not establish network-drift robustness.')
text(s,.93,1.75,11.48,1.40,'고정된 환경의 분류 정확도를 넘어\n변화하는 네트워크에서의 적응성 평가',size=31,color=INK)
text(s,.93,4.18,11.48,1.40,'사전학습 TabPFN을 활용한\n소량 라벨·저비용 탐지 모델 갱신',size=31,bold=True,color=BLUE)

s=slide('TabPFN · 표 데이터 예측 규칙을 사전학습한 모델 (v3: arXiv, 2026)',title_size=23,
        notes='Source: '+PAPER+' §2.5, PDF page 10: more than 8 trillion tokens, not datasets. Official TabPFN 8.2.0: '+FT+'. EXP77 actually fine-tunes the initial IDS data using AdamW, learning rate 1e-5, 10 fixed epochs, context/query cross-entropy; then freezes weights. The v3 decoder weights context one-hot labels, so old-class examples remain necessary in the inference context. Validation is inside the 1000 revealed-label budget; no test tuning.')
for y,heading,body in [
    (1.40,'사전학습','합성 표 데이터의 8조 개 이상 토큰으로 학습\n일부 행의 정답을 맞히며 표 데이터의 예측 규칙 습득'),
    (2.68,'In-context learning','라벨된 사례(context)를 참고해 새 flow(query)의 라벨 예측\n추론 중 모델 가중치는 그대로 유지'),
    (3.96,'Fine-tuning','초기 IDS를 context와 query로 나눠 query의 정답 예측\n예측 오류를 역전파해 모델 가중치를 IDS에 맞게 조정')]:
    text(s,.7,y,2.7,1.0 if heading=='In-context learning' else .5,heading,size=22,bold=True,color=BLUE)
    text(s,3.45,y,9.15,1.0,body,size=21)
box(s,.7,5.45,11.92,1.12)
text(s,.91,5.55,11.5,.36,'초기 IDS로 1회 fine-tuning → 가중치 고정 → 기존 클래스 유지 + 신규 사례 반영',size=20,bold=True,color=NAVY)
text(s,.91,6.06,11.5,.34,'기대: 반복 재학습과 비슷한 탐지 품질을 더 적은 데이터·갱신 비용으로 달성',size=20,color=BLUE)

s=slide('v3 구조',notes='Original Figure 5, '+PAPER+'#page=6. Native PPTX crop excludes the original caption. Labels and source remain in speaker notes. First three stages: feature grouping/cell embedding; target-aware column distribution embedding; row-level feature aggregation. Plain language deliberately omits implementation dimensions.')
architecture(s)
detail(s,1.40,'① 입력 변환 · Cell embedding','패킷 길이·지속시간 등의 값을 숫자 벡터로 변환\n값이 비어 있는 위치도 함께 표시')
detail(s,3.18,'② 특징별 분포 파악 · Feature embedding','여러 flow를 살펴 각 특징의 값 분포를 요약\nContext의 정상·공격 라벨도 함께 반영')
detail(s,4.96,'③ Flow별 정보 통합 · Feature aggregation','한 flow의 여러 특징을 하나의 표현으로 통합\n이 표현을 이용해 flow 사이의 관계 비교')

s=slide('v3 구조',notes='Same Figure 5, '+PAPER+'#page=6; no visible caption by user request. ICL attends to training/context keys; many-class decoder uses attention-weighted one-hot context labels. This is learned contextual relevance, not merely raw-feature nearest-neighbour distance. Context determines available class labels.')
architecture(s)
detail(s,1.74,'④ 사례 참조 · In-context learning','새 flow(query)가 라벨된 기존 사례(context)를 참조\n어떤 사례가 예측에 도움이 되는지 attention으로 계산')
detail(s,4.15,'⑤ 확률 출력 · Many-class decoder','관련성이 큰 사례의 라벨에 더 큰 가중치 부여\n참고한 라벨을 모아 정상·공격별 확률 계산')
text(s,5.14,5.94,7.48,.64,'기존 클래스도 분류하려면\n해당 라벨의 context 유지 필요',size=19,bold=True,color=BLUE)

s=slide('평가시나리오',notes='EXP77 uses the same released IDs as EXP76. Initial IDS weight fine-tuning: fixed 10 epochs, AdamW, lr 1e-5; initial internal feature-group validation only. Subsequent context policies freeze the initial checkpoint; repeated FT is a separate control. XGB cumulatively retrains. Proposed model uses five feature-disjoint roles within the SAME total label budget, original pretrained backbone to keep calibration labels outside weight training. Stage order chronological, within-stage sample counts descending. Main score is the FINAL state on ALL held-out rows, not pooled stage predictions.')
text(s,.66,1.02,12.0,.38,'초기 FT 비교군: 1단계 10 epoch FT · 이후 가중치 고정 · 단계순, 단계 내 표본 수 내림차순',size=16,color=GREY)
stages={
 'cic2018':[['Benign','Brute force','DoS'],['DDoS'],['Web attacks'],['Infiltration'],['Bot']],
 'toniot':[['Benign','DoS','Scanning'],['DDoS','Injection'],['Password'],['XSS'],['Backdoor','Ransomware'],['MitM']]}
for ds,label,x in [('cic2018','CIC2018',.66),('toniot','ToN',6.90)]:
    text(s,x,1.50,5.75,.4,label,size=22,bold=True,color=BLUE)
    rows=[];total=0;heights=[.43]
    for i,classes in enumerate(stages[ds]):
        counts=[750,125,125] if i==0 else [100]*len(classes);total+=sum(counts)
        rows.append([str(i+1),'\n'.join(classes),'\n'.join(f'{n:,}' for n in counts),f'{total:,}'])
        heights.append(1.08 if i==0 else .74 if len(classes)==2 else .53)
        scenario_data.append(dict(dataset=ds,stage=i+1,classes=classes,new_counts=counts,cumulative=total))
    widths=[.65,2.43,1.22,1.36]
    sh=s.shapes.add_table(len(rows)+1,4,Inches(x),Inches(2.01),Inches(sum(widths)),Inches(sum(heights)));t=sh.table
    for c,w in enumerate(widths):t.columns[c].width=Inches(w)
    for r,values in enumerate([['단계','신규 공개 클래스','클래스별 수','누적 수']]+rows):
        t.rows[r].height=Inches(heights[r])
        for c,v in enumerate(values):
            cell=t.cell(r,c);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if r==0 else PALE if r%2 else 'FFFFFF')
            frame(cell.text_frame,v,size=14 if r==0 else 15.5,bold=r==0,color='FFFFFF' if r==0 else INK,align=PP_ALIGN.RIGHT if c>=2 and r else PP_ALIGN.CENTER if c==0 else PP_ALIGN.LEFT)
text(s,.7,6.68,11.95,.27,'최종 상태로 전체 test 평가 · CIC2018 3,042,473개 / ToN 2,271,723개 · seed 43 단일 실행',size=14.5,color=GREY)

s=slide('결과',notes='EXP77 final-model full-test results: '+str(SMALL.relative_to(ROOT))+'/results.json. Three comparisons: XGB cumulative retraining from scratch at every stage; original pretrained PFN cumulative context; initial-FT PFN cumulative context with frozen weights after stage 1. All plotted metrics independently verified against saved predictions. Slides 6 and 7 use the identical cost definition and linear axis: GPU-synchronized context selection + model construction + fit summed over stages AFTER the initial stage. Initial preparation, initial FT, all inference, serialization, and audit overhead are excluded. Initial FT is disclosed separately. Single seed 43. Other context policies remain in the EXP77 record; classwise and label-budget figures are separate supporting artifacts.')
text(s,.66,1.04,12.0,.36,'최종 상태 전체 test 평가 · XGB: 누적 재학습 · PFN 두 조건: 누적 context 갱신',size=16.5,color=GREY)
s.shapes.add_picture(str(ASSET/'small_context_results.png'),Inches(.26),Inches(1.47),width=Inches(12.75),height=Inches(5.10))
initial_cost=[json.loads((SMALL/'finetuning'/ds/'initial'/'audit.json').read_text())['seconds'] for ds in ['cic2018','toniot']]
text(s,.69,6.63,12.0,.25,f'비용: 2단계 이후 갱신 합계 · 초기 준비·추론 제외 · 초기 FT 별도: CIC2018 {initial_cost[0]:.2f}초 / ToN {initial_cost[1]:.2f}초',size=14,color=GREY)

s=slide('제안 모델·SOTA 비교',notes='Same EXP77 released-label budget: CIC2018 1400 / ToN 1700, including all internal validation and calibration. Existing K=4 anchor/residual + S1V0 proposed model uses the original pretrained backbone; global/expert/route/tune/cal split 40/20/20/8/12% by feature groups. This is a same-total-label-budget method comparison. BoostPFN uses the existing project port, 50 learners, 500 rows/learner. EXP78 DistPFN uses original pretrained v3 with NO FT, cumulative context, and current-query-batch posterior adjustment; no query labels. Restored base predictions exactly match EXP77 no-FT PFN. Context preparation costs are shared with that base, not remeasured or taken from the FT branch. DistPFN correction is inference work and is excluded from the context-update graph. Both slides 6 and 7 show only later-stage cumulative update time with the same linear axis; initial preparation and all inference excluded. Detailed inference measurements remain in research records. LoCalPFN excluded. Source: '+str((FOLLOWUP/'dist_results.json').relative_to(ROOT)))
text(s,.66,1.02,12.0,.36,'동일 라벨 예산: CIC2018 1,400개 / ToN 1,700개 · 비용: 2단계 이후 갱신 합계',size=16,color=GREY)
s.shapes.add_picture(str(ASSET/'small_sota_results.png'),Inches(.22),Inches(1.47),width=Inches(12.8),height=Inches(5.05))
text(s,.69,6.64,12.0,.25,'초기 준비·추론 비용 제외 · DistPFN: FT 없는 PFN에 추론 시 확률 보정 · 제안 모델: 두 데이터 모두 Global 유지',size=13.5,color=GREY)

prs.save(OUT)
with (ASSET/'plotted_values.csv').open('w',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=['slide','dataset','case','method','metric','value']);writer.writeheader();writer.writerows(chart_data)
manifest=dict(revision='compact_7_exp78_noft_dist_update_cost',slides=records,text_boxes=text_boxes,chart_data=chart_data,scenario_data=scenario_data,
 source_deck=str(BASE.relative_to(ROOT)),source_deck_sha256=hashlib.sha256(BASE.read_bytes()).hexdigest(),output=str(OUT.relative_to(ROOT)),
 primary_v3_source=PAPER,pretraining_quantity='more than 8 trillion tokens, not dataset count; report section 2.5',
 new_source=str((SMALL/'results.json').relative_to(ROOT)),additional_sources=[str((FOLLOWUP/'dist_results.json').relative_to(ROOT))],historical_source=None,
 cost_metric_slides_6_7='update_seconds',cost_excludes=['initial_preparation','all_inference'],cost_axis='shared_linear',distpfn_initial_ft=False,
 current_results_include_initial_ids_finetuning=True,evaluation='final_full_test',proposed_design='initial IDS parameter fine-tuning once, then freeze theta and vary inference context',
 original_25_slide_archive='docs/slides/archive/20261015_before_7slide_revision',figure_caption_visible=False,
 source_sha256={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [SMALL/'results.json',SMALL/'verification.json',FOLLOWUP/'dist_results.json',FOLLOWUP/'dist_verification.json']})
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
plan=['# 1015 랩미팅 · 7장 축약본','','초기 IDS 가중치 fine-tuning을 실제 수행한 EXP77의 최종 상태 전체 test 평가. 마지막 장은 같은 소규모 라벨 예산의 제안 모델·SOTA 비교. 1002 원본과 이전 25장·1014 포함 7장 자료는 archive에 보존.','']
for r in records:plan.extend([f"## {r['page']}. {r['title'] or '(타이틀 없음)'}",'',r['notes'],''])
(HERE/'kor/1015_labmeeting_plan.md').write_text('\n'.join(plan)+'\n')
print(OUT, len(prs.slides),'slides')
