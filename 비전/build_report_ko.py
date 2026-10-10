# -*- coding: utf-8 -*-
"""
Homework 1 한국어 Word 보고서 생성기.
homework1_cv.py가 저장한 results_A/B/C.json과 그림(output/)을 읽어
그림 + 정량결과 + 해석(한국어)을 담은 Homework1_Report_KO.docx 를 만든다.

실행:  python build_report_ko.py
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn

from report_updates import Results, finalize_report, BASE
A = Results("A")
B = Results("B")
Cm = Results("C")
os.chdir(BASE)

DOCX = "Homework1_Report_KO.docx"


# ---------- docx 헬퍼 ----------
def set_base_font(doc, font="Malgun Gothic", size=10.5):
    st = doc.styles["Normal"]
    st.font.name = font; st.font.size = Pt(size)
    rpr = st.element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), font); rf.set(qn("w:ascii"), font); rf.set(qn("w:hAnsi"), font)

def h(doc, text, level=1):
    return doc.add_heading(text, level=level)

def para(doc, text, bold=False, italic=False, size=None, align=None):
    p = doc.add_paragraph(); run = p.add_run(text)
    run.bold = bold; run.italic = italic
    if size: run.font.size = Pt(size)
    if align: p.alignment = align
    return p

def bullet(doc, text, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if lead:
        r = p.add_run(lead); r.bold = True
    p.add_run(text); return p

def img(doc, path, width=6.3, caption=None):   # caption 인자는 무시(캡션 제거)
    if path and os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

def quant(doc, lines):                          # '정량 결과' 라벨 없이 수치만
    for ln in lines:
        b = doc.add_paragraph(ln); b.paragraph_format.left_indent = Inches(0.25); b.paragraph_format.space_after = Pt(0)

def interp(doc, text=None):                     # '해석' 섹션 라벨 제거 (no-op)
    return


def build():
    rA = A.run_all();  mA = rA["metrics"]; FA = rA["figures"]; TH = rA.get("a9_thumbs", {})

    doc = Document(); set_base_font(doc)

    # ---------- 표지/개요 ----------
    doc.add_heading("컴퓨터비전 과제 1", level=0)
    para(doc, "점처리 · 공간/주파수 필터링 · Wiener 복원", italic=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    para(doc, '코드는 homework1_cv.py를 참고하며, 실행 결과 그림은 output/ 폴더에 저장된다.')

    h(doc, "Part A. 점처리 및 히스토그램 기반 영상 향상", 1)
    para(doc, "입력: 컬러 RGB 영상(point_processing_input_rgb.png, %dx%d). 점처리는 각 화소를 독립적으로 "
              "변환하고, 히스토그램 기반 기법은 밝기 분포를 재배치해 대비를 개선한다. 모든 향상은 휘도 Y에만 "
              "적용하여 색상(hue/saturation)을 보존한다." % mA["shape"])
    para(doc, "사용 파라미터", bold=True)
    for s in ["RGB→YUV 행렬: 과제 제공(BT.601 계열), 역변환은 그 역행렬",
              "히스토그램: 8-bit, 레벨 L=256",
              "대비 스트레칭: 2nd / 98th 백분위를 r_min, r_max 로 사용",
              "감마: γ = 0.5, 1.0, 2.0 (c=1)",
              "AHE / CLAHE 타일: 8×8, CLAHE clip = 0.01, 0.05 (타일 화소수 대비 비율)",
              "비트평면: 상위 4개(bit 4~7)로 재구성"]:
        bullet(doc, s)

    # A1
    h(doc, "A-1. RGB → YUV 변환", 2)
    para(doc, "주어진 행렬을 직접 적용해 휘도 Y와 색차 U, V를 분리한다. R,G,B∈[0,255]이고 첫 행의 가중치 합이 "
              "1이라 Y∈[0,255](휘도)이며, U·V는 색차로 음수가 될 수 있어 클리핑/8bit 캐스팅 없이 float로 유지한다.")
    img(doc, FA["A1"], caption="그림 A1. 원본 RGB와 Y, U, V 성분")
    quant(doc, ["Y 범위 [%.2f, %.2f], U 범위 [%.2f, %.2f], V 범위 [%.2f, %.2f]"
                % (mA["Y_range"][0], mA["Y_range"][1], mA["U_range"][0], mA["U_range"][1], mA["V_range"][0], mA["V_range"][1]),
                "corr(U, B−Y) = %.3f,  corr(V, R−Y) = %.3f" % (mA["corr_U_BmY"], mA["corr_V_RmY"])])
    interp(doc, "해석 (의미)")
    bullet(doc, "Y 영상엔 얼굴·우주복·헬멧·뒤쪽 성조기 구조가 선명하다. 오렌지색 우주복은 V(빨강 색차)에서 밝게, "
                "U(파랑 색차)에서 어둡게 나타난다(오렌지는 R이 크고 B가 작기 때문).", lead="그림 관찰: ")
    bullet(doc, "사람 눈의 밝기 민감도를 반영한 가중합(G에 0.587로 최대 가중치). 구조·에지가 대부분 담긴다.", lead="Y (휘도): ")
    bullet(doc, "U는 파랑 색차(∝ B−Y), V는 빨강 색차(∝ R−Y). 상관계수가 1.000이라 정확히 비례함을 확인.", lead="U, V (색차): ")

    # A2
    h(doc, "A-2. 히스토그램 계산", 2)
    para(doc, "Y를 8-bit로 양자화(r_k∈{0,…,255})한 뒤 레벨별로 직접 집계한다: h(r_k)=n_k, 정규화 p(r_k)=n_k/MN.")
    img(doc, FA["A2"], caption="그림 A2. 입력 Y, 히스토그램, 정규화 히스토그램")
    quant(doc, ["sum(h) = %.0f ( = M·N ),  sum(p) = %.4f" % (mA["hist_sum"], mA["p_sum"])])
    interp(doc)
    bullet(doc, "코드의 k가 강도 r_k, np.sum(q==k)가 그 강도의 화소 수 n_k 이다.")
    bullet(doc, "정규화 전후 히스토그램 모양이 같은 이유: 모든 막대를 같은 상수 MN으로 나눌 뿐이라 상대 비율은 "
                "그대로이고 세로축 눈금만 바뀐다(확률분포 PDF로 재해석). 분포가 좁게 몰릴수록 대비가 낮다.")

    # A3
    h(doc, "A-3. 히스토그램 평활화 (HE)", 2)
    para(doc, "CDF(r_k)=Σp(r_j)를 변환함수로 사용: s_k=(L−1)·CDF(r_k). 전역적으로 분포를 균일에 가깝게 펴 대비를 "
              "높인다. 평활화된 Y를 원본 U,V와 결합해 RGB로 역변환한다.")
    img(doc, FA["A3a"], caption="그림 A3-1. 원본/평활 Y와 각 히스토그램")
    img(doc, FA["A3b"], caption="그림 A3-2. 원본 vs HE RGB, 변환곡선 s=T(r) (곡선은 참고용, 요구 항목 아님)")
    quant(doc, ["Y 평균 %.1f → %.1f,  표준편차 %.1f → %.1f" % (mA["HE_mean_before"], mA["HE_mean_after"], mA["HE_std_before"], mA["HE_std_after"])])
    interp(doc)
    bullet(doc, "HE 결과 영상은 전반적으로 밝아지고, 어두운 헬멧 그림자와 뒤쪽 성조기·배경 디테일이 더 잘 보인다. "
                "변환곡선이 대각선 위로 올라가(어두운 값을 들어올림) 이 변화를 그대로 반영한다.", lead="그림 관찰: ")
    bullet(doc, "변환함수 기울기 dT/dr=(L−1)·p(r). 히스토그램이 밀집한(p 큰) 구간에서 기울기가 가팔라져 그 구간을 넓게 벌림 → 대비 증가.")
    bullet(doc, "평활 히스토그램이 빗살(comb) 모양이 되는 이유: 밀집 구간은 레벨을 건너뛰어 빈 레벨(gap)이 생기고, "
                "희소 구간은 느린 CDF를 정수로 반올림해 여러 입력이 한 출력으로 합쳐진다(many-to-one).")
    bullet(doc, "평균이 %.1f→%.1f로 올라 전체적으로 밝아진다. 다만 순흑이 많아 검정 배경이 중간 회색으로 들려, "
                "중간톤 세부 대비는 커지되 암부 자체 대비는 줄 수 있다(전역 HE의 한계)." % (mA["HE_mean_before"], mA["HE_mean_after"]))

    # A4
    h(doc, "A-4. 대비 스트레칭 (Contrast Stretching)", 2)
    para(doc, "piecewise-linear 변환으로 [r_min, r_max]를 [0, L−1]로 선형 확장. r_min, r_max는 2/98 백분위(이상치에 강건).")
    img(doc, FA["A4"], caption="그림 A4. 원본/스트레칭 RGB와 전후 히스토그램")
    quant(doc, ["r_min(2%%)=%.1f, r_max(98%%)=%.1f, 이득=(L−1)/(r_max−r_min)=%.3f" % (mA["cs_rmin"], mA["cs_rmax"], mA["cs_gain"])])
    interp(doc)
    bullet(doc, "원본과 스트레칭 결과가 육안으로 거의 구분되지 않고, 전후 히스토그램도 거의 같은 모양이다 — 이미 동적범위가 꽉 차 있기 때문.", lead="그림 관찰: ")
    bullet(doc, "이 영상은 이미 거의 전체 범위를 써서(이득≈%.2f) 전후 차이가 미미하다. 스트레칭은 동적범위가 좁은(뿌연) 영상에서 효과가 크다." % mA["cs_gain"])
    bullet(doc, "스트레칭 vs HE: 스트레칭은 분포 모양을 유지한 채 범위(동적범위)만 선형으로 늘리고, HE는 CDF 기반 "
                "비선형으로 분포 자체를 재배치. 요약하면 '범위를 늘림' vs '분포를 폄'.")

    # A5
    h(doc, "A-5. 감마 보정 (Power-law)", 2)
    para(doc, "s = 255·(r/255)^γ. 일반형 s = c·r^γ 에서 c=1은 추가 스케일 없음. γ=0.5, 1.0, 2.0 비교.")
    img(doc, FA["A5"], caption="그림 A5. 감마별 결과 영상과 변환곡선")
    interp(doc)
    bullet(doc, "γ=0.5에선 헬멧·그림자 등 어두운 부분이 환해져 디테일이 드러나고, γ=2.0에선 전체가 어두워지며 밝았던 "
                "얼굴·배경이 눌린다. 아래 변환곡선의 볼록/오목이 이 밝기 변화를 설명한다.", lead="그림 관찰: ")
    bullet(doc, "밝아짐 + 암부 확장. 기울기 γ·r^(γ−1)가 r→0에서 ∞ → 어두운 값들이 서로 벌어져(대비↑) 그림자 디테일이 "
                "살아난다(압축되는 쪽은 밝은 영역).", lead="γ<1 (0.5): ")
    bullet(doc, "항등 변환(불변).", lead="γ=1: ")
    bullet(doc, "어두워짐. r→0에서 기울기→0이라 암부가 압축(그림자 뭉개짐), 밝은 영역이 확장.", lead="γ>1 (2.0): ")

    # A6
    h(doc, "A-6. 적응형 히스토그램 평활화 (AHE)", 2)
    para(doc, "영상을 8×8 타일로 나눠 각 타일에서 독립적으로 HE를 수행(보간 없음).")
    img(doc, FA["A6"], caption="그림 A6. 원본 Y / 전역 HE / AHE")
    interp(doc)
    bullet(doc, "AHE 결과엔 타일 경계를 따라 생긴 블로킹과, 어두운 배경·매끄러운 우주복 영역의 얼룩덜룩한 잡음이 "
                "뚜렷하게 보인다. 대신 국소 대비(얼굴 음영·질감)는 전역 HE보다 훨씬 살아난다.", lead="그림 관찰: ")
    bullet(doc, "국소 CDF를 쓰므로 전역 HE보다 국소 대비가 크게 향상된다.")
    bullet(doc, "잡음 증폭 이유: 어두운 배경·우주복 매끄러운 부분처럼 거의 균일한 타일은 좁은 밝기 범위를 억지로 "
                "0~255로 펴면서 미세 잡음이 얼룩으로 드러난다. 또 타일 경계 불연속으로 블로킹 아티팩트 발생 → CLAHE로 해결.")

    # A7
    h(doc, "A-7. CLAHE (Contrast Limited AHE)", 2)
    para(doc, "타일 히스토그램을 clip limit로 자르고 초과분을 균등 재분배한 뒤 국소 CDF로 매핑, 타일 중심 간 "
              "쌍선형 보간으로 경계 아티팩트 제거. clip = 0.01, 0.05 비교.")
    img(doc, FA["A7"], caption="그림 A7. 원본/AHE/CLAHE와 히스토그램")
    quant(doc, ["타일 64×64=4096: clip=0.01 → clip_count=40, clip=0.05 → clip_count=204"])
    bullet(doc, "CLAHE(clip=0.01)는 AHE에 있던 블로킹·얼룩 잡음이 사라져 자연스럽고, clip=0.05는 국소 대비가 더 강하다. "
                "히스토그램을 보면 AHE는 뾰족한 스파이크가 많은 반면 CLAHE는 눌려 있다.", lead="그림 관찰: ")
    bullet(doc, "AHE의 타일 내부는 밝기 범위가 좁아 국소 CDF가 가팔라져(대비 이득 폭주) 잡음을 증폭한다. clip이 "
                "히스토그램 봉우리를 잘라 CDF 기울기에 상한을 둬서, 평탄 타일에서 이득이 폭주하지 못하게 한다.", lead="CLAHE가 AHE보다 잡음이 적은 이유: ")
    bullet(doc, "clip limit이 곧 '국소 대비 이득의 상한'(국소 CDF의 최대 기울기)이다. clip이 클수록 CDF가 더 가팔라질 수 "
                "있어 국소 대비가 강해지고(AHE 쪽), 작을수록 CDF가 평평해져 국소 대비가 약해진다(원본 쪽). 즉 clip을 "
                "올리면 그 상한 내에서 국소 대비가 커진다.", lead="clip limit이 국소 대비에 미치는 영향: ")
    bullet(doc, "너무 작으면 ≈ 원본(봉우리가 다 잘려 히스토그램 평평→CDF 직선→항등 매핑, 향상 미미); "
                "너무 크면 ≈ AHE(봉우리 통과→과대비·배경 잡음 증폭). 중간 clip이 균형.", lead="clip이 너무 작/크면: ")

    # A8
    h(doc, "A-8. 비트평면 분해 (Bit-Plane Slicing)", 2)
    para(doc, "Y(x,y)=Σ b_k·2^k 로 8개 비트평면 분해. 상위 4개(bit 4~7)만으로 재구성해 원본과 비교.")
    img(doc, FA["A8a"], caption="그림 A8-1. 8개 비트평면 (좌상단 MSB → 우하단 LSB)")
    img(doc, FA["A8b"], width=5.2, caption="그림 A8-2. 원본 vs 상위 4비트 재구성")
    quant(doc, ["4-MSB 재구성:  MSE=%.2f,  PSNR=%.2f dB" % (mA["bitplane_MSE"], mA["bitplane_PSNR"])])
    interp(doc)
    bullet(doc, "bit 7~5 평면만으로도 얼굴·우주복·헬멧이 또렷이 식별되고, bit 4는 질감을 보강한다. 반면 bit 3부터는 "
                "점점 무작위 점처럼 변하고 bit 0(LSB)은 거의 순수 잡음이다.", lead="그림 관찰: ")
    bullet(doc, "상위 비트(bit 7/6/5)는 큰 밝기 변화 = 시각적 구조/윤곽의 대부분. 하위 비트(bit 1/0)는 가중치가 작아 미세 질감·잡음에 가깝다.")
    bullet(doc, "상위 4비트가 중요한 근거: ① 가중치 합 16+32+64+128=240 으로 최대값 255의 약 94%%, "
                "② 재구성 PSNR≈%.1f dB·육안으로도 거의 동일. 하위 비트 손실로 매끄러운 그라데이션에 약한 윤곽선(false contour) 발생." % mA["bitplane_PSNR"])

    h(doc, 'A-9. 비교 및 논의', 2)
    para(doc, '아래 복잡도는 밝기 레벨마다 전체 화소를 비교하는 현재 구현 기준이다. M×N은 영상 크기, L=256, T=64이다. L과 T를 고정하면 화소 수에 대해 선형이며, 단일 패스 히스토그램 계수는 O(MN+L)로 구현할 수 있다.')
    para(doc, '대비 스트레칭', bold=True)
    para(doc, '전역 처리로 선택한 중앙 밝기 범위를 일정한 이득으로 확장하고 양 끝은 클리핑한다. 이득이 큰 구간에서는 신호와 잡음이 함께 증폭될 수 있다. 현재 히스토그램 구현의 복잡도는 O(LMN+L)이며, 동적범위가 좁은 영상의 전처리에 적합하다.')
    para(doc, '히스토그램 평활화', bold=True)
    para(doc, '전역 CDF에 따라 밝기 구간별 대비를 조절한다. 빈도가 높은 강도 구간의 변동과 잡음도 확대될 수 있으며, 밝기 상승이 전체 표준편차 상승을 보장하지는 않는다. 복잡도는 O(LMN+L)이며, 전역 밝기 분포를 재조정할 때 사용한다.')
    para(doc, 'AHE', bold=True)
    para(doc, '타일별 국소 처리로 불균일한 대비를 개선한다. 평탄한 타일의 미세 변동도 크게 증폭하며, 본 구현은 타일 간 보간이 없어 경계 불연속이 나타난다. 복잡도는 O(LMN+TL)이며, 영역별 대비가 다른 영상에 적합하지만 잡음과 타일 경계를 함께 살펴야 한다.')
    para(doc, 'CLAHE', bold=True)
    para(doc, '타일 히스토그램의 큰 빈도를 제한해 국소 대비 이득을 완화하고, 보간으로 타일 경계 불연속을 줄인다. clip limit가 크면 세부 대비와 잡음 증폭이 함께 커질 수 있다. 복잡도는 O(LMN+TL)이며, 저조도 영상이나 국소 대비를 제어하며 개선할 때 적합하다.')
    para(doc, '감마 보정', bold=True)
    para(doc, '전역 점처리로 밝기에 따라 다른 이득을 적용한다. γ<1은 암부를, γ>1은 밝은 영역을 상대적으로 확장하므로 해당 영역의 잡음도 증폭할 수 있다. 복잡도는 O(MN)이며, 감마 및 톤 조정에 적합하다.')

    rB = B.run_all(); mB = rB["metrics"]; FB = rB["figures"]

    h(doc, "Part B. 공간·주파수 영역 필터링", 1)
    para(doc, "입력: 그레이스케일 f(x,y) (spatial_frequency_filtering_input.png, %dx%d). 각 필터를 공간영역(직접 "
              "합성곱)과 주파수영역(FFT 곱)에서 모두 수행해 비교한다. 푸리에 스펙트럼은 중심화 로그 크기 "
              "log(1 + |fftshift(F)|)로 표시한다." % (mB["shape"][0], mB["shape"][1]))
    para(doc, "사용 파라미터", bold=True)
    for s in ["블러 커널 h_b = (1/9)·ones(3,3)",
              "샤프닝 커널 h_s = [[0,-1,0],[-1,5,-1],[0,-1,0]]",
              "주파수 필터링: (M+m-1, N+n-1)로 zero-pad → FFT가 '선형'(순환 아님) 합성곱을 수행; 공간과 동일한 zero 경계",
              "Unsharp masking: σ = 1.0, 3.0 ; k = 1.0, 2.0 ; 가우시안 반경 = ⌈3σ⌉"]:
        bullet(doc, s)

    # B1
    h(doc, "B-1. 블러 필터링", 2)
    para(doc, "윗줄은 요구된 4개 표시: (1) 입력 f, (2) log|F|, (3) 임펄스응답 h_b, (4) log|H_b|. 아랫줄은 공간결과 "
              "g_s=h_b*f, 주파수결과 g_f=F⁻¹{H_b F}, 출력 스펙트럼 log(1+|G_f|), 그리고 차이 |g_s − g_f|.")
    img(doc, FB["B1"], caption="그림 B1. 블러: 입력/응답(윗줄), 결과+차이(아랫줄)")
    quant(doc, ["공간 vs 주파수:  MSE = %.2e,  PSNR = %.1f dB,  SSIM = %.6f" % (mB["blur_mse"], mB["blur_psnr"], mB["blur_ssim"])])
    interp(doc)
    bullet(doc, "입력(동전 영상)의 에지·질감이 g_s·g_f에서 똑같이 부드럽게 뭉개지고, 오른쪽 아래 차이맵 |g_s−g_f|는 "
                "완전히 검정(최댓값 ~1e-13)이라 두 결과가 사실상 같음을 한눈에 보여준다.", lead="그림 관찰: ")
    bullet(doc, "입력 스펙트럼 |F|는 중앙(DC·저주파)이 가장 밝고 가장자리(고주파)로 갈수록 어둡다 — 영상에 매끄러운 "
                "넓은 영역이 많아 에너지가 저주파에 몰리기 때문. log|H_b|는 중앙이 밝고 바깥으로 사라지는 저역통과 응답이고, "
                "log(1+|G_f|)=|H_b F|는 |F|에서 고주파가 더 깎여 중앙에 더 집중된다(블러 영상은 고주파 에너지가 적다).", lead="스펙트럼 읽기: ")
    bullet(doc, "합성곱 정리로 두 방법이 같은 선형 합성곱을 계산하므로 결과가 부동소수점 반올림만 빼면 동일"
                "(MSE ~1e-26, PSNR >300 dB, SSIM=1). 차이맵 |g_s−g_f| 최댓값 ~1e-13이 확인.", lead="왜 공간=주파수: ")

    # B2
    h(doc, "B-2. 샤프닝 필터링", 2)
    para(doc, "h_s = 원본 + 라플라시안 기반 에지강조(계수합 1이라 평균밝기 보존). 레이아웃은 B-1과 동일.")
    img(doc, FB["B2"], caption="그림 B2. 샤프닝: 입력/응답(윗줄), 결과+차이(아랫줄)")
    quant(doc, ["공간 vs 주파수:  MSE = %.2e,  PSNR = %.1f dB,  SSIM = %.6f" % (mB["sharp_mse"], mB["sharp_psnr"], mB["sharp_ssim"])])
    interp(doc)
    bullet(doc, "결과에서 동전 경계·각인 질감이 또렷해지고 에지 주변이 밝게 강조된다(오버슈트). 여기서도 차이맵은 ≈0으로 공간·주파수 결과가 일치한다.", lead="그림 관찰: ")
    bullet(doc, "샤프닝도 동일(반올림 제외). 출력은 에지에서 오버/언더슈트로 [0,255]를 벗어나 표시할 때만 클리핑(지표는 원본 float).")

    # B3
    h(doc, "B-3. 주파수 영역 분석", 2)
    para(doc, "각 필터에 대해 |F|, |H|, |G|=|HF|를 비교한다. 중심화 스펙트럼에서 중앙=저주파, 가장자리=고주파.")
    img(doc, FB["B3_blur"], width=6.6)
    img(doc, FB["B3_sharpen"], width=6.6)
    bullet(doc, "블러 |H_b|는 중앙이 밝고 가장자리로 어둡다. 샤프닝 |H_s|는 반대로 중앙이 어둡고 모서리로 밝다.", lead="그림 관찰: ")
    bullet(doc, "중앙=저주파, 가장자리=고주파. |H_b|가 중앙(저주파)에서 밝고(≈1) 가장자리(고주파)에서 어두우므로, "
                "|G|=|H_b||F|에서 저주파는 통과하고 고주파는 감쇠된다 → 저역통과이고, 그래서 영상이 흐려진다. 추가 그래프 "
                "없이 이 2D 스펙트럼만으로 바로 읽을 수 있다.", lead="블러가 저역통과인 이유: ")
    bullet(doc, "|H_s|는 중앙이 어둡고 가장자리가 밝으므로 |G|=|H_s||F|에서 고주파가 증폭된다 → 고역통과이고, 그래서 "
                "에지·디테일이 또렷해진다.", lead="샤프닝이 고주파를 강조하는 이유: ")
    bullet(doc, "에지는 급격한 밝기 변화 = 강한 고주파. 고주파를 키우면 에지가 또렷해지고, 깎으면 뭉개진다.", lead="에지와 고주파: ")

    # B4
    h(doc, "B-4. 임펄스 응답 검증", 2)
    para(doc, "임펄스 δ에 필터를 적용하면 임펄스응답이 복원된다: h*δ = h.")
    img(doc, FB["B4_blur"], width=6.8, caption="그림 B4-1. 블러: δ, δ의 평탄한 FFT, 출력(중앙 9×9), h_b, |H_b|")
    img(doc, FB["B4_sharpen"], width=6.8, caption="그림 B4-2. 샤프닝: δ, δ의 평탄한 FFT, 출력(중앙 9×9), h_s, |H_s|")
    quant(doc, ["max |출력중앙 − h|:  블러 = %.1e,  샤프닝 = %.1e" % (mB["impulse_err_blur"], mB["impulse_err_sharpen"])])
    bullet(doc, "출력 h*δ는 (중앙 9×9로 잘라 표시 — h가 3×3뿐이라 31×31 출력의 나머지는 모두 0이다) 정확히 h다: "
                "블러는 균일한 밝은 블록, 샤프닝은 중앙이 밝고 상하좌우가 어두운 십자. δ의 FFT 패널이 한 가지 색으로 "
                "균일한 것은 |FFT(δ)|가 모든 주파수에서 1이기 때문이다(그 한 색은 상수값에 대한 컬러맵 색일 뿐).", lead="그림 관찰: ")
    bullet(doc, "단위 임펄스 δ를 넣으면 출력이 h*δ = h가 되므로, h는 말 그대로 '시스템의 임펄스에 대한 응답'이다 "
                "(그래서 impulse response). δ가 모든 주파수를 똑같이 포함하므로 이 한 번의 입력이 LTI 시스템 전체"
                "(공간응답 h, 주파수응답 H)를 드러낸다. 측정 오차는 ~0.", lead="왜 h를 임펄스 응답이라 부르나: ")

    # B5
    h(doc, "B-5. 가우시안 기반 선명화 (Unsharp Masking)", 2)
    para(doc, "가우시안 저역통과로 f_L을 얻고 고주파 f_H = f − f_L을 더해 선명화: g = (1+k)f − k·f_L.")
    para(doc, "4×5 그림 두 장: 각 변형이 '이미지 행 + |F| 스펙트럼 행'을 차지하고, 열은 원본 f, 가우시안 커널, 블러 f_L, "
              "고주파 f_H, 선명화 g. 첫 장은 σ를 다르게(k 고정), 둘째 장은 k를 다르게(σ 고정).")
    img(doc, FB["B5_sigma"], width=6.9)
    img(doc, FB["B5_k"], width=6.9)
    img(doc, FB["B5_vs_fixed"], width=6.2)
    bullet(doc, "σ가 클수록 커널이 넓어지고 그 스펙트럼은 더 좁은 중앙 블롭이 된다(더 많은 대역을 '저주파'로 제거) → "
                "f_H가 더 굵은 윤곽을 담고, σ=3의 f_H 스펙트럼이 σ=1보다 넓은 대역을 덮는다. k를 키우면 선명화 영상의 "
                "에지 테두리(halo)가 밝아진다.", lead="그림 관찰: ")
    bullet(doc, "f_L은 저주파(저역통과)라 f_H = f − f_L은 저주파가 상쇄되고 고주파(에지·디테일)만 남는다(가우시안 고역통과, f_H 스펙트럼 열에서 확인).", lead="왜 블러를 빼면 고주파: ")
    bullet(doc, "σ가 클수록 더 넓은 대역을 '저주파'로 제거 → f_H가 더 넓은 대역(굵은 에지 포함). 작을수록 미세 디테일만.", lead="σ의 영향: ")
    bullet(doc, "k는 고주파를 더하는 강도. 클수록 선명하나 너무 크면 에지 오버슈트(halo/링잉)·잡음 증폭·포화.", lead="k의 영향: ")
    bullet(doc, "고정 3×3 커널 h_s는 아주 작은 블러·고정 강도의 unsharp 특수 경우라 가장 미세한 디테일만 선명화할 수 있다. "
                "Unsharp masking은 σ(대역)·k(강도)를 독립 조절해, σ를 크게(=3) 하면 고정 h_s가 못 하는 '굵은 구조 선명화'도 "
                "가능하다(비교 그림 참조).", lead="고정 커널 h_s와 비교: ")

    # B6
    h(doc, "B-6. 논의 — 합성곱 정리", 2)
    para(doc, "g(x,y) = h(x,y) * f(x,y)   ⟺   G(u,v) = H(u,v) F(u,v).")
    bullet(doc, "영상을 2차원 사인파들의 합(푸리에 성분)으로 본다. LTI 시스템은 각 사인파에 따로 작용하는데, 입력이 "
                "단일 주파수 e^{j2π(ux+vy)}이면 출력은 '같은 주파수'에 복소수 H(u,v)만 곱해진 것이다(사인파가 시스템의 "
                "고유함수). 따라서 필터링은 각 푸리에 성분 F(u,v)에 H(u,v)를 곱하는 것(G=HF)과 같다. 즉 공간의 합성곱이 "
                "주파수영역에서는 '성분별 곱'에 대응한다 — 큰 커널에선 FFT(O(N²logN))가 직접합성곱(O(N²k²))보다 유리.", lead="왜 합성곱이 곱에 대응하나: ")
    para(doc, "실무에서 작은 차이가 생기는 이유", bold=True)
    bullet(doc, "FFT 곱은 본질적으로 순환(circular) 합성곱 — 영상이 가장자리에서 반대편으로 감겨(wrap-around) 오른쪽 "
                "끝이 왼쪽에 샌다. (M+m-1, N+n-1)로 zero-padding하면 그 감김이 패딩 영역에만 생겨 선형이 되고 'same'으로 자른다.", lead="순환 합성곱 / zero-padding: ")
    bullet(doc, "공간 합성곱은 영상 밖 화소를 뭘로 볼지(경계조건) 정해야 한다(여기선 0). 두 방법이 다른 경계(zero/replicate/wrap)를 "
                "쓰면 테두리 화소만 달라진다. 본 과제는 양쪽 0이라 테두리까지 일치.", lead="경계 처리: ")
    bullet(doc, "FFT/IFFT는 유한정밀 부동소수점이라 수학적으로 같아도 반올림이 ~1e-12~1e-13 누적. 그래서 MSE가 정확히 0이 아니라 ~1e-26.", lead="수치 정밀도: ")

    # =================================================================
    # PART C
    # =================================================================
    rC = Cm.run_all(); inputs = rC["inputs"]; primary = inputs[0]

    def k_table(rows, bestP, bestS):
        tbl = doc.add_table(rows=1, cols=4); tbl.style = "Table Grid"
        for j, ht in enumerate(["K", "MSE", "PSNR (dB)", "SSIM"]):
            cpar = tbl.rows[0].cells[j]; cpar.text = ""
            rr = cpar.paragraphs[0].add_run(ht); rr.bold = True; rr.font.size = Pt(9)
        for K, m_, ps, ss in rows:
            cells = tbl.add_row().cells; note = ""
            if K == bestP: note += "  (PSNR 최고)"
            if K == bestS: note += "  (SSIM 최고)"
            for j, v in enumerate(["%.0e%s" % (K, note), "%.5f" % m_, "%.2f" % ps, "%.4f" % ss]):
                cells[j].text = ""; rr = cells[j].paragraphs[0].add_run(v); rr.font.size = Pt(9)
        doc.add_paragraph()

    h(doc, "Part C. Wiener 필터 (영상 복원)", 1)
    para(doc, "열화 모델: g = h*f + n (h는 점확산함수 PSF, n은 가산 잡음). PSF는 15×15 가우시안(σ=2.5)이며 "
              "sum(h)=1로 정규화하고, 입력 f는 [0,1]로 정규화한다. Wiener 필터를 직접 설계하고 정규화 상수 K를 스윕하며, "
              "입력 영상 %d장에 대해 실험한다." % len(inputs))
    para(doc, "사용 파라미터", bold=True)
    for s in ["PSF: 15×15 가우시안, σ=2.5, sum=1 정규화",
              "잡음: 평균 0 가우시안, RMS=0.03 (실제 RMS를 정확히 맞춤)",
              "Wiener: F̂ = conj(H)/(|H|² + K)·G ; K ∈ {1e-6, 1e-4, 1e-3, 1e-2, 1e-1}",
              "지표(원본 대비, 범위 [0,1]): MSE, PSNR, SSIM"]:
        bullet(doc, s)

    h(doc, "C-1. 블러 영상 생성", 2)
    para(doc, "g_b = h*f (주파수영역 곱으로 구현, H는 zero-phase PSF의 FFT).")
    img(doc, primary["figs"]["blur"])
    quant(doc, ["PSF 합 = %.6f (정규화)" % rC["psf_sum"]])
    bullet(doc, "PSF는 가운데가 밝은 작은 가우시안 점으로 보이고, 블러 영상(카메라맨)은 전체 윤곽이 뿌옇게 번져 "
                "삼각대·배경 건물 경계가 흐려진다.", lead="그림 관찰: ")

    h(doc, "C-2. 평균 0 가우시안 잡음 추가", 2)
    para(doc, "g = g_b + n. 생성한 잡음을 재스케일해 실제 RMS를 목표값에 정확히 맞춘다.")
    img(doc, primary["figs"]["noise"])
    quant(doc, ["실제 잡음 RMS = %.5f (목표 0.03000)" % primary["noise_rms"]])
    bullet(doc, "오른쪽(블러+잡음)은 매끄러운 하늘·잔디 영역에 오돌토돌한 그레인이 뚜렷이 얹혀 있어, 왼쪽 블러 영상과 구별된다.", lead="그림 관찰: ")

    h(doc, "C-3. Wiener 필터로 복원", 2)
    para(doc, "F̂(u,v) = [ H*(u,v) / (|H(u,v)|² + K) ]·G(u,v), 그리고 복원 = IFFT(F̂). K는 잡음/신호 전력비를 "
              "근사하며, K=0이면 역필터(inverse filter).")
    img(doc, primary["figs"]["restore"])
    bullet(doc, "복원(K=1e-2) 영상은 가운데 열화 영상의 흐림이 걷혀 카메라맨·삼각대 윤곽이 또렷해지고 잡음도 억제되어, 원본에 상당히 가까워진다.", lead="그림 관찰: ")

    h(doc, "C-4. K의 효과", 2)
    para(doc, "기본 입력(%s)에 대한 K 스윕:" % primary["label"])
    img(doc, primary["figs"]["sweep"], width=6.8)
    k_table(primary["table"], primary["best_psnr_K"], primary["best_ssim_K"])
    bullet(doc, "K=1e-6은 화면이 잡음으로 완전히 덮여 피사체가 안 보이고, K=1e-4는 겨우 윤곽만, K=1e-3부터 선명해지며 "
                "잔여 잡음, K=1e-2가 가장 깨끗하고 선명, K=1e-1은 더 매끈하지만 약간 흐릿(배경 건물이 뭉개짐)하다.", lead="그림 관찰: ")
    bullet(doc, "|H|²이 0에 가까운 고주파에서 1/|H|²이 폭발해 역필터에 가까워져 잡음을 극단적으로 증폭한다"
                "(K=1e-6에서 PSNR ~5 dB — 선명해 보여도 잡음에 파묻힘).", lead="K가 너무 작으면: ")
    bullet(doc, "분모가 K에 지배되어 필터가 H*/K에 가까워지고, 역필터링이 약해져 흐릿한(과평활) 복원이 된다(잡음은 적지만 해상도 손실).", lead="K가 너무 크면: ")
    bullet(doc, "K는 잡음 억제 ↔ 디블러 선명도의 트레이드오프를 조절한다. K ≈ S_n/S_f(잡음대신호 전력비)일 때 최적에 가깝다.", lead="K의 역할: ")

    h(doc, "C-5. 다른 입력들에 대한 결과", 2)
    para(doc, "같은 열화 + Wiener 복원을 다른 입력 영상들에도 반복해, 최적 K가 영상 내용에 따라 어떻게 달라지는지 확인한다.")
    for r in inputs[1:]:
        h(doc, "입력: %s" % r["label"], 3)
        img(doc, r["figs"]["deg"], width=5.8)
        img(doc, r["figs"]["sweep"], width=6.4)
        k_table(r["table"], r["best_psnr_K"], r["best_ssim_K"])
    best_list = ", ".join("%s: PSNR@%.0e·SSIM@%.0e" % (r["label"], r["best_psnr_K"], r["best_ssim_K"]) for r in inputs)
    bullet(doc, "최적 K는 영상마다 다르다(%s). 더 매끄러운 영상일수록 큰 K를 더 잘 견딘다 — 이는 K ≈ S_n/S_f 에 부합하며, "
                "최적 정규화는 잡음뿐 아니라 신호 자체의 스펙트럼에 달려 있다." % best_list, lead="입력별 비교: ")

    doc.add_paragraph()
    finalize_report(doc, "ko")
    doc.save(DOCX); print("saved", DOCX)


if __name__ == "__main__":
    build()
