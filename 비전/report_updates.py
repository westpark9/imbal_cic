"""Shared scientific corrections and measured tables for both report languages."""
from pathlib import Path
import json
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent

def load_result(part):
    def relocate(v):
        if isinstance(v,dict): return {k:relocate(x) for k,x in v.items()}
        if isinstance(v,list): return tuple(relocate(x) for x in v)
        if isinstance(v,str) and v.endswith('.png'):
            return str(BASE/'output'/v.replace('\\','/').split('/')[-1])
        return v
    return relocate(json.loads((BASE/f'results_{part}.json').read_text(encoding='utf-8')))

class Results:
    def __init__(self, part): self.part=part
    def run_all(self): return load_result(self.part)

def insert_para(doc, anchor, text, bold=False):
    p=doc.add_paragraph()
    p.add_run(text).bold=bold
    anchor.addnext(p._p)
    return p._p

def insert_table(doc, anchor, headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    anchor.addnext(t._tbl)
    return t._tbl

def insert_picture(doc, anchor, path, width=6.4):
    p=doc.add_paragraph(); p.add_run().add_picture(str(path),width=Inches(width))
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    anchor.addnext(p._p)
    return p._p

def find(doc, prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))

KO_EDITS = [
('입력: 컬러 RGB 영상', '입력은 512×512 RGB 영상(point_processing_input_rgb.png)이다. 휘도 Y에 향상 기법을 적용하고 색차 U, V를 유지한다. 이는 색차 성분을 유지하는 처리이며, 밝기 변경과 최종 RGB 클리핑 때문에 색상과 채도가 완전히 보존된다고 보장할 수는 없다.'),
('U, V (색차):', 'U, V (색차): U는 B−Y, V는 R−Y에 거의 비례한다. 주어진 변환계수는 반올림된 값이므로 상관계수 1.000은 표시 정밀도에서의 결과이다.'),
('평균이 ', 'HE 후 평균은 115.4에서 129.5로 증가하지만 전체 표준편차는 75.1에서 71.4로 감소한다. 이처럼 평균 밝기 상승과 전체 대비 상승은 같은 뜻이 아니다. 검정 입력은 약 28의 어두운 회색으로 올라가며, 빈도가 높은 밝기 구간과 낮은 구간의 대비 변화가 서로 다르다.'),
('s = 255·(r/255)^γ.', '정규화한 r_n=r/255에 s_n=c·r_n^γ, c=1을 적용하고 s=255·s_n으로 되돌린다. γ=0.5, 1.0, 2.0을 비교한다. 기울기는 ds/dr=γ(r/255)^(γ−1)이며, 정규화 좌표에서는 ds_n/dr_n=γ·r_n^(γ−1)이다.'),
('그림 관찰: CLAHE(clip=0.01)', '그림 관찰: clip=0.01은 독립 타일 AHE보다 경계 불연속과 과도한 암부 변화를 줄이고, clip=0.05는 더 강한 국소 대비와 배경 변동을 보인다. 아래에는 원본·AHE·두 CLAHE 결과의 전역 히스토그램을 같은 세로축 범위로 표시했다. 전역 출력 히스토그램의 봉우리는 타일별 clip limit의 제한 대상이 아니다.'),
('CLAHE가 AHE보다 잡음이 적은 이유:', 'CLAHE가 잡음 증폭을 줄이는 이유: 좁은 밝기 구간에 몰린 타일 히스토그램의 빈도를 제한하여 국소 CDF가 지나치게 가팔라지는 것을 완화한다. 보간은 타일 경계 불연속을 줄인다. 보간과 clipping의 역할은 서로 다르며, 결과 영상만으로 잡음량을 직접 측정했다고 해석하지 않는다.'),
('clip limit이 국소 대비에 미치는 영향:', 'clip limit의 의미: clip_count=floor(clip×타일 화소수)이며 여기서는 40 또는 204이다. 초과량 E를 모든 256개 bin에 한 번 균등 재분배하므로 최종 bin 값은 min(h,clip_count)+E/256이 된다. 따라서 설정 count는 재분배 후의 엄격한 상한이 아니다. 반올림 전 CDF 매핑의 레벨당 증가량은 255·h_final/4096이다.'),
('clip이 너무 작/크면:', 'clip이 작으면 타일 히스토그램이 더 균일해져 매핑이 항등변환에 가까워지고 향상이 약해진다. 매우 크면 clipping의 영향이 사라져 보간된 AHE에 접근한다. 본 실험의 AHE는 보간하지 않으므로 두 영상이 완전히 같아지는 것은 아니다.'),
('그림 관찰: 입력(동전 영상)', '그림 관찰: 동전의 에지와 질감이 두 결과에서 동일하게 부드러워진다. 차이맵은 오차를 10⁻¹² 단위로 확대하여 표시하므로 무늬가 보일 수 있다. 최대 오차 약 3.13×10⁻¹³은 영상 강도 범위 0~255에 비해 매우 작다.'),
('스펙트럼 읽기:', '스펙트럼 읽기: F, H, G는 실제 선형 합성곱에 사용한 동일한 514×514 패딩 격자에서 얻었으며 G=HF이다. F와 G는 같은 로그 표시 범위를 쓴다. 3×3 평균 필터는 DC 이득이 1이고 고주파를 전반적으로 감쇠하지만, Gaussian과 달리 영점과 측엽이 있어 응답이 단조롭게 감소하지는 않는다.'),
('h_s = 원본 +', 'h_s는 원본에 라플라시안 기반 성분을 더하는 고역강조 커널이다. 계수합이 1이므로 DC 이득은 1이고 영상 내부의 상수 영역을 보존한다. 그러나 zero-padding과 same 크롭을 적용한 전체 영상의 평균은 경계 때문에 변할 수 있다.'),
('블러가 저역통과인 이유:', '블러가 저역통과인 이유: 중심 정렬 좌표에서 H_b(ωx,ωy)=(1+2cosωx)(1+2cosωy)/9이다. DC에서 1이며 영점과 측엽을 포함하면서 고주파를 감쇠한다. 이로 인해 빠른 밝기 변화인 에지와 세부 질감이 약해진다.'),
('샤프닝이 고주파를 강조하는 이유:', '샤프닝이 고주파를 강조하는 이유: 중심 정렬 응답 H_s(ωx,ωy)=5−2cosωx−2cosωy는 DC에서 1, (π,π)에서 9이다. DC를 제거하지 않으므로 순수 고역통과가 아니라 고역강조 필터이다. 스펙트럼 중앙의 어두운 색은 상대적으로 낮은 이득을 뜻하며 이득 0을 뜻하지 않는다.'),
('그림 관찰: σ가 클수록', '그림 관찰: σ가 클수록 공간 커널이 넓어지고 저역통과 응답의 통과 대역은 좁아진다. f_L에는 더 낮은 주파수가 주로 남고, f_H=f−f_L에는 상대적으로 낮은 주파수까지 포함한 넓은 대역이 남아 굵은 윤곽이 나타난다. k 증가 시 에지 halo와 범위 밖 출력이 늘어난다.'),
('σ의 영향:', 'σ의 영향: σ 증가 → Gaussian 통과 대역 축소 → 잔차 응답 1−H_G의 강조 범위가 더 낮은 주파수까지 확장된다. 따라서 더 큰 공간 규모의 윤곽이 강조된다. 유한 영상의 zero 경계에서는 잔차에 경계 효과도 포함된다.'),
('고정 커널 h_s와 비교:', '고정 커널 h_s와 비교: 고정 3×3 커널은 주파수별 이득이 고정되어 있다. Gaussian unsharp masking은 σ로 공간 규모를, k로 강도를 조절한다. 고정 커널도 다양한 윤곽에 반응하므로 굵은 구조를 전혀 강조할 수 없다고 단정하지 않는다.'),
('순환 합성곱 / zero-padding:', '순환 합성곱과 패딩: FFT 곱은 순환 합성곱을 계산한다. 두 배열을 최소 (M+m−1, N+n−1)로 패딩하면 선형 합성곱의 전체 지지영역을 담을 수 있어 순환 겹침이 발생하지 않는다. 역변환한 전체 결과에서 커널 반경부터 M×N 크기로 잘라 공간영역 same 결과와 정렬한다.'),
('F̂(u,v) =', '복원식은 F̂(u,v)=[H*(u,v)/(|H(u,v)|²+K)]·G(u,v)이고, 역 FFT의 실수부를 복원 영상으로 사용한다. H*는 복소켤레이다. K는 주파수별 잡음/신호 전력비를 상수로 근사한 정규화 항이다. K=0이면 H≠0인 곳에서 1/H가 되지만 H=0에서는 역복원이 정의되지 않는다.'),
('K가 너무 작으면:', 'K가 너무 작으면: |H|²에 비해 K가 충분히 작은 주파수에서 필터는 1/H에 접근한다. 흐림으로 약해진 성분을 되돌리는 동시에 잡음을 증폭하며, 작은 |H|에서는 그 영향이 크다. K=10⁻⁶의 원시 복원 PSNR은 −16.37 dB이고 [0,1] 클리핑 후에는 5.04 dB이다.'),
('K의 역할:', 'K의 역할: 작은 K는 역복원을 강하게 하고 큰 K는 잡음 증폭을 더 억제한다. 지나치게 큰 K는 디블러 효과와 세부 구조를 약화시키며 DC 이득도 1/(1+K)가 되어 밝기가 감소할 수 있다. 최적값은 평가 지표와 영상, 잡음 실현에 따라 달라진다.'),
('입력별 비교:', '입력별 비교: 이번 5개 후보와 고정된 잡음 시드에서 cameraman과 input 2의 clipping 후 PSNR 최고는 K=10⁻², SSIM 최고는 K=10⁻¹이다. rocket은 두 지표 모두 K=10⁻¹에서 최고이다. 이는 영상 내용과 평가 지표에 따라 선호하는 복원 강도가 다름을 보여주며, 모든 매끄러운 영상에 같은 K가 최적이라고 일반화할 수는 없다.'),
]

EN_EDITS = [
('Input: a color RGB image', 'The input is a 512x512 RGB image (point_processing_input_rgb.png). Enhancement is applied to Y while retaining U and V. This retains chrominance coordinates; it does not guarantee unchanged hue and saturation after changing luminance and clipping the reconstructed RGB values.'),
('U is proportional to', 'U and V are approximately proportional to B-Y and R-Y, respectively. The supplied matrix coefficients are rounded, so correlations displayed as 1.000 do not prove exact proportionality.'),
('the mean rises', 'The mean increases from 115.4 to 129.5, whereas the global standard deviation decreases from 75.1 to 71.4. Increased brightness therefore does not imply increased global contrast. Black input maps to approximately 28, a dark gray, and contrast changes differently across populated and sparse intensity intervals.'),
('s = 255*(r/255)^gamma.', 'Normalize r_n=r/255, apply s_n=c*r_n^gamma with c=1, and convert back using s=255*s_n. The tested gamma values are 0.5, 1.0 and 2.0. In original intensity units, ds/dr=gamma*(r/255)^(gamma-1).'),
('Observed in the figure: CLAHE', 'Observed in the figure: clip=0.01 reduces boundary discontinuities and excessive dark-region changes relative to independent-tile AHE; clip=0.05 produces stronger local contrast and background variation. Histograms of the original, AHE and both CLAHE outputs use the same count scale. A cap on local input histograms does not cap peaks in the global output histogram.'),
('Why CLAHE amplifies less noise', 'Why CLAHE limits noise amplification: clipping reduces concentration in each local histogram and moderates steep local CDF mappings. Interpolation reduces discontinuities across tile boundaries. These are different mechanisms; output appearance alone is not a direct measurement of noise.'),
('How the clip limit affects', 'Meaning of the clip limit: clip_count=floor(clip*tile_pixel_count), giving 40 or 204. Excess E is redistributed once across all 256 bins, so h_final=min(h,clip_count)+E/256. The configured count is therefore not a strict upper bound after redistribution. Before LUT rounding, the mapping increment is 255*h_final/4096.'),
('Clip limit too small', 'A small limit makes the local histogram more uniform and the mapping closer to identity, reducing enhancement. A sufficiently large limit removes clipping and approaches interpolated AHE. The AHE baseline here has no interpolation, so it need not become exactly the same image.'),
('Observed in the figure: the input (coins)', 'Observed in the figure: coin edges and texture are softened equally in both outputs. The difference map displays error in units of 1e-12, so round-off patterns remain visible. Its maximum, approximately 3.13e-13, is negligible relative to the intensity range 0 to 255.'),
('Reading the spectra:', 'Reading the spectra: F, H and G come from the same 514x514 padded grid used in filtering, with G=H F. F and G share a log display scale. The box filter has unit DC gain and generally attenuates high frequencies, but unlike a Gaussian response it has zeros and sidelobes rather than a monotonic radial decay.'),
('Same layout for h_s', 'The sharpening kernel adds a Laplacian-based component to the original. Its coefficients sum to one, giving unit DC gain and preserving constant interior regions. With zero boundaries and same cropping, the mean of the complete finite image can nevertheless change.'),
('Why blur is low-pass:', 'Why blur is low-pass: in centered kernel coordinates, H_b(wx,wy)=(1+2*cos(wx))*(1+2*cos(wy))/9. Its DC gain is one, and high frequencies are attenuated, with zeros and sidelobes. This weakens abrupt intensity changes and fine texture.'),
('Why sharpen emphasizes high freq:', 'Why sharpening emphasizes high frequencies: H_s(wx,wy)=5-2*cos(wx)-2*cos(wy) has gain one at DC and nine at (pi,pi). It is a high-frequency emphasis filter, not a pure high-pass filter. The dark center of its autoscaled spectrum represents lower relative gain, not zero gain.'),
('Observed in the figure: larger sigma', 'Observed in the figure: increasing sigma broadens the spatial kernel and narrows its low-pass passband. The residual f_H=f-f_L consequently includes a wider frequency range extending toward lower frequencies, producing broader outlines. Increasing k strengthens halos and increases out-of-range output values.'),
('Effect of sigma:', 'Effect of sigma: a larger sigma narrows the Gaussian passband. The residual response 1-H_G then emphasizes a range extending to lower frequencies, corresponding to larger spatial structures. Zero boundaries also introduce finite-image boundary effects into the residual.'),
('Comparison with the fixed kernel h_s:', 'Comparison with the fixed kernel: the 3x3 kernel has fixed frequency-dependent gain. Gaussian unsharp masking allows the spatial scale sigma and strength k to be adjusted separately. The fixed kernel can still respond to broad structures; claiming that it cannot sharpen them at all would be too strong.'),
('Zero-padding / circular convolution:', 'Padding and circular convolution: the FFT product computes circular convolution. Padding both arrays to at least (M+m-1,N+n-1) accommodates the full support of linear convolution, preventing circular overlap. Cropping the inverse transform from the kernel-radius offset aligns the same-size spatial result.'),
('F_hat(u,v) =', 'The restoration is F_hat=[conj(H)/(|H|^2+K)]*G, followed by the real part of its inverse FFT. K approximates the frequency-dependent noise-to-signal power ratio by a constant. At K=0 the response is 1/H where H is nonzero; inversion is undefined at zeros of H.'),
('K too small:', 'K too small: at frequencies where K is negligible relative to |H|^2, the response approaches 1/H. Undoing attenuation also amplifies noise, especially at small |H|. At K=1e-6 the raw restoration PSNR is -16.37 dB, whereas clipping to [0,1] gives 5.04 dB.'),
('Role of K:', 'Role of K: smaller K allows stronger inversion, whereas larger K suppresses noise amplification. Excessively large K weakens deblurring and fine detail; DC gain is also 1/(1+K), so brightness can decrease. The preferred value depends on the metric, image and noise realization.'),
('Across inputs:', 'Across these five candidates and the fixed noise realization, cameraman and input 2 have their highest clipped-output PSNR at K=1e-2 and SSIM at K=1e-1. Rocket has both maxima at K=1e-1. These observations demonstrate image- and metric-dependent preferences, not a universal rule for all smooth images.'),
]

def finalize_report(doc, lang):
    ko=lang=='ko'; a=load_result('A'); b=load_result('B'); c=load_result('C')
    edits=KO_EDITS if ko else EN_EDITS
    for prefix,replacement in edits:
        matches=[p for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(matches)!=1: raise ValueError((lang,prefix,len(matches)))
        matches[0].text=replacement
    for p in doc.paragraphs:
        if not p.text: continue
        t=p.text
        t=t.replace('상관계수가 1.000이라 정확히 비례함을 확인.','반올림된 상관계수는 1.000이다.')
        t=t.replace('중간 회색','어두운 회색').replace('mid-gray','dark gray')
        t=t.replace('아래에는 원본·AHE·두 CLAHE','그림에는 원본·AHE·두 CLAHE')
        t=t.replace('bit 0(LSB)은 거의 순수 잡음이다.','bit 0(LSB)은 미세한 변동을 보인다. 이 그림만으로 신호와 잡음을 구분할 수는 없다.')
        t=t.replace('bit 0 (LSB) is almost pure noise.','bit 0 (LSB) shows fine variation, which cannot be classified as pure noise from this display alone.')
        t=t.replace('기울기 γ·r^(γ−1)','정규화 좌표의 기울기 γ·r_n^(γ−1)').replace('The slope gamma*r^(gamma-1)','In normalized coordinates, the slope gamma*r_n^(gamma-1)')
        t=t.replace('샤프닝=고역통과','샤프닝=고역강조').replace('sharpening as high-pass','sharpening as high-frequency emphasis')
        t=t.replace('정확한 PDF','확률질량함수').replace('확률분포 PDF로 재해석','이산 확률질량함수로 재해석')
        t=t.replace('모든 중간 계산은 부동소수점(float64)으로 수행한다.','필요한 중간 계산은 float64로 수행하며, 히스토그램 색인과 비트평면 연산에는 정수를 사용한다.')
        t=t.replace('All intermediate computations use floating-point (float64).','Intermediate arithmetic uses float64 where needed; histogram indexing and bit-plane operations use integers.')
        if t!=p.text: p.text=t

    if ko:
        find(doc,'스트레칭 vs HE:').text='스트레칭 vs HE: 선택한 중앙 범위에서는 일정한 이득으로 선형 확장하지만 양 끝에서는 클리핑이 발생한다. HE는 누적분포에 따라 강도 구간별 이득이 달라지는 비선형 변환이다.'
        find(doc,'그림 관찰: K=1e-6').text='그림 관찰: K=10⁻⁶은 강한 잡음으로 구조 식별이 어렵고, 10⁻⁴에서는 피사체 윤곽과 큰 잡음이 함께 보인다. 10⁻²는 선명도와 잡음의 균형에서 PSNR이 가장 높다. 10⁻¹은 더 매끈하며 일부 세부 구조가 약해지고 SSIM이 가장 높다.'
    else:
        find(doc,'stretching vs. HE:').text='Stretching applies a constant linear gain inside the chosen central range but clips the tails. Histogram equalization is a nonlinear CDF mapping with different gains across intensity intervals.'
        find(doc,'Observed in the figure: K=1e-6').text='Observed in the figure: K=1e-6 severely obscures structure with noise; at 1e-4 the subject is visible amid substantial noise. K=1e-2 gives the highest PSNR among the tested values. K=1e-1 is smoother, loses some fine detail, and has the highest SSIM.'

    # Histogram percentile method is explicitly linked to the implemented CDF.
    p=find(doc,'A-4.')
    insert_para(doc,p._p, '백분위수는 양자화된 Y의 히스토그램 CDF가 0.02와 0.98에 처음 도달하는 레벨로 선택한다. 이번 입력에서는 r_min=0, r_max=235이다.' if ko else 'The percentile levels are the first levels at which the histogram CDF of quantized Y reaches 0.02 and 0.98. They are r_min=0 and r_max=235 for this input.')

    # Concise measurements accompany each experiment; no extra tables.
    stats={row[0]:row[1:] for row in a['metrics']['enhancement_stats']}
    texts={
        'A-5.': ('휘도 평균: 원본 %.3f, γ=0.5 %.3f, γ=1 %.3f, γ=2 %.3f.' if ko else 'Mean luminance: original %.3f; gamma=0.5 %.3f; gamma=1 %.3f; gamma=2 %.3f.') % tuple(stats[n][0] for n in ['Original','Gamma 0.5','Gamma 1.0','Gamma 2.0']),
        'A-6.': ('64×64 타일 표준편차의 평균: 원본 %.3f, HE %.3f, AHE %.3f. 이 값에는 구조·대비·잡음이 함께 포함된다.' if ko else 'Mean standard deviation across 64x64 tiles: original %.3f; HE %.3f; AHE %.3f. This includes structure, contrast and noise.') % tuple(stats[n][2] for n in ['Original','HE','AHE']),
        'A-7.': ('타일 표준편차의 평균: AHE %.3f, CLAHE clip=0.01 %.3f, clip=0.05 %.3f. 이는 잡음만의 측정값이 아니다.' if ko else 'Mean tile standard deviation: AHE %.3f; CLAHE clip=0.01 %.3f; clip=0.05 %.3f. This is not a noise-only measurement.') % tuple(stats[n][2] for n in ['AHE','CLAHE 0.01','CLAHE 0.05']),
    }
    for prefix,text in texts.items():
        insert_para(doc,find(doc,prefix)._p,text)

    p=find(doc,'Part B.')
    anchor=insert_para(doc,p._p,'구현과 지표 조건' if ko else 'Implementation and metric settings',True)
    anchor=insert_para(doc,anchor,'m×n 커널을 패딩 배열의 좌상단에 두고 (M+m−1,N+n−1) FFT를 계산한다. 역 FFT 결과의 (m//2,n//2)부터 원본 크기로 잘라 커널 중심을 정렬한다. B-1~B-3의 F,H,G는 이 패딩 격자의 실제 배열이며, G는 크롭 전 스펙트럼이다.' if ko else 'Place the mxn kernel in the top-left of the padded array and compute FFTs of size (M+m-1,N+n-1). Crop the inverse FFT from (m//2,n//2) to the input size to align the kernel center. The F,H,G panels in B-1 through B-3 use these actual arrays; G is the spectrum before cropping.')
    insert_para(doc,anchor,'SSIM은 11×11 Gaussian 창(σ=1.5), C1=(0.01L)², C2=(0.03L)², zero-padding, 경계 포함 전체 평균을 사용한다. PSNR의 peak 및 SSIM의 L은 B에서 255, C에서 1이다. B 지표는 클리핑 전 float 결과끼리 비교한다.' if ko else 'SSIM uses an 11x11 Gaussian window (sigma=1.5), C1=(0.01L)^2, C2=(0.03L)^2, zero-padding, and the mean over all pixels including borders. The PSNR peak and SSIM L are 255 in B and 1 in C. Part B compares raw floating-point outputs before display clipping.')
    p=find(doc,'B-2.')
    insert_para(doc,p._p,('전체 평균: 입력 %.3f, zero 경계 샤프닝 %.3f.' if ko else 'Whole-image mean: input %.3f; sharpening with zero boundaries %.3f.') % (b['metrics']['mean_input'],b['metrics']['mean_sharpen']))
    p=find(doc,'B-4.')
    insert_para(doc,p._p,'31×31 배열 중앙을 표시 좌표의 원점으로 두어 단위 임펄스를 배치한다. 배열 인덱스상으로는 이동된 임펄스이므로 출력도 같은 위치로 이동한 h이다. 이동은 Fourier 위상에 영향을 주지만 크기는 모든 주파수에서 1이다. 공간·FFT 임펄스 출력도 함께 수치 검증한다.' if ko else 'The center of the 31x31 array is the displayed coordinate origin of the impulse. In array-index coordinates this is a shifted impulse, producing the correspondingly shifted h. The shift changes Fourier phase but not its constant unit magnitude. Spatial and FFT impulse outputs are also checked numerically.')
    p=find(doc,'σ의 영향:' if ko else 'Effect of sigma:')
    prev=p._p.getprevious()
    prev=insert_para(doc,prev,'공간·주파수 영역 검증' if ko else 'Spatial and frequency verification',True)
    rows=b['metrics']['unsharp_comparison']
    prev=insert_para(doc,prev,'σ=1·3, k=1·2의 네 조합에서 공간영역과 주파수영역 결과를 비교했다. Gaussian 커널 크기는 2⌈3σ⌉+1로 7×7 또는 19×19이다. 네 조합의 MSE 범위는 %.2e~%.2e, PSNR은 %.2f~%.2f dB이며, SSIM은 표시 정밀도에서 모두 1이다.' % (min(r[3] for r in rows),max(r[3] for r in rows),min(r[4] for r in rows),max(r[4] for r in rows)) if ko else 'The four combinations of sigma=1 or 3 and k=1 or 2 were evaluated in both domains. Gaussian kernel size is 2*ceil(3*sigma)+1, giving 7x7 or 19x19. MSE ranges from %.2e to %.2e, PSNR from %.2f to %.2f dB, and SSIM is one at displayed precision for every pair.' % (min(r[3] for r in rows),max(r[3] for r in rows),min(r[4] for r in rows),max(r[4] for r in rows)))
    prev=insert_picture(doc,prev,b['figures']['B5_domains'],5.8)
    insert_para(doc,prev,'모든 조합의 SSIM은 표시 정밀도에서 1이며, 차이는 부동소수점 반올림 수준이다. σ=3에서 k를 1에서 2로 늘리면 범위 밖 비율이 2.673%에서 7.665%로 증가한다. 이는 강한 선명화가 포화와 halo를 증가시킨다는 해석을 뒷받침한다.' if ko else 'SSIM is one at the displayed precision for every pair, and differences are at round-off scale. At sigma=3, increasing k from 1 to 2 raises the out-of-range fraction from 2.673% to 7.665%, supporting the interpretation of increased saturation and halos with stronger sharpening.')

    p=find(doc,'Part C.')
    anchor=insert_para(doc,p._p,'비너 필터의 목적' if ko else 'Purpose of the Wiener filter',True)
    anchor=insert_para(doc,anchor,'비너 필터(Wiener filter)는 흐려지고 잡음이 더해진 영상에서 원본을 추정하는 복원 방법이다. Gaussian blur는 에지·미세 질감을 약하게 만든다. 이를 단순히 역으로 증폭하면 잡음도 함께 커진다. 비너 필터는 흐림의 역복원과 잡음 증폭 억제를 함께 고려하며, 여기서는 K로 그 균형을 조절한다. 원본은 성능 평가에만 사용하고 복원 계산에는 열화 영상 g, 알려진 PSF의 H, 선택한 K만 사용한다.' if ko else 'The Wiener filter estimates an original image from an observation that has been blurred and corrupted by noise. Gaussian blur attenuates edges and fine texture; directly inverting this attenuation also amplifies noise. The filter balances deblurring and suppression of noise amplification through K. The original is used for evaluation only; restoration uses g, the known PSF response H, and the selected K.')
    anchor=insert_para(doc,anchor,'열화·복원은 동일한 주기적 경계(순환 합성곱) 모델을 사용한다. PSF를 영상 크기로 패딩하고 중심을 배열 원점으로 이동해 H를 만든다. B의 zero 경계 선형 합성곱과는 경계조건이 다르다. 잡음 시드는 0이고 세 영상에 같은 난수 생성기를 순차 사용하며, 가산 후 g는 클리핑하지 않는다.' if ko else 'Degradation and restoration use the same periodic boundary model (circular convolution). Pad the PSF to the image size and shift its center to the array origin before forming H. This differs from the zero-boundary linear convolution in B. Noise uses seed 0 with one generator sequentially across the three images. The noisy observation g is not clipped.')
    p=find(doc,'C-4.')
    insert_para(doc,p._p,'평가 기준: 복원값을 [0,1]로 클리핑한 뒤 원본과 비교한다. 아래 K 비교표와 C-5의 표는 모두 이 기준을 사용한다.' if ko else 'Evaluation: clip the restored values to [0,1] before comparison with the original. The K table below and the tables in C-5 all use this convention.')

    format_document(doc,ko)

def format_document(doc, ko):
    font='Malgun Gothic' if ko else 'Times New Roman'
    for sec in doc.sections:
        sec.top_margin=sec.bottom_margin=Inches(.65)
        sec.left_margin=sec.right_margin=Inches(.65)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
        st=doc.styles[name]; st.font.name=font; st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font)
        if name=='Normal': st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(6)
        if name.startswith('Heading'): st.paragraph_format.keep_with_next=True
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for p in doc.paragraphs:
        for border in p._p.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        if p.style.name=='Title':
            p.paragraph_format.space_after=Pt(10)
        p.paragraph_format.widow_control=True
        if p._p.xpath('.//w:drawing'): p.paragraph_format.keep_together=True
        p.paragraph_format.keep_together=True
        if p.text in ['대비 스트레칭','히스토그램 평활화','AHE','CLAHE','감마 보정','Contrast stretching','Histogram equalization','Gamma correction','공간·주파수 영역 검증','Spatial and frequency verification']:
            p.paragraph_format.keep_with_next=True
        if not ko:
            for r in p.runs:
                r.font.name=font
                rf=r._r.get_or_add_rPr().get_or_add_rFonts()
                for key in ['ascii','hAnsi','eastAsia','cs']:
                    rf.set(qn('w:'+key),font)
                for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                    rf.attrib.pop(qn('w:'+key),None)
    for t in doc.tables:
        t.autofit=False
        widths=[7.2/len(t.columns)]*len(t.columns)
        for j,w in enumerate(widths): t.columns[j].width=Inches(w)
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr()
            ns=OxmlElement('w:cantSplit');pr.append(ns)
            if i==0:
                rep=OxmlElement('w:tblHeader');pr.append(rep)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j])
                cp=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2' if i==0 else ('F7F8FA' if i%2==0 else 'FFFFFF'));cp.append(shade)
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                cp.append(borders)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    if len(t.columns)!=8:
                        p.paragraph_format.keep_with_next=i<len(t.rows)-1
                    for r in p.runs:
                        r.font.size=Pt(8.5 if len(t.columns)==8 else 9)
                        r.font.name=font
                        if i==0:r.bold=True
