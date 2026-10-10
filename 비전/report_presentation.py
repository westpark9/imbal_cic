"""Question-only bullets and the A-9 comparison table for both reports."""
from copy import deepcopy
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = Path(__file__).resolve().parent


def find(doc, prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f'Expected one paragraph for {prefix!r}, got {len(matches)}')
    return matches[0]


def after(p, text=''):
    el = OxmlElement('w:p')
    p._p.addnext(el)
    new = Paragraph(el, p._parent)
    new.add_run(text)
    return new


def remove(p):
    p._p.getparent().remove(p._p)


def set_font(run, font, size=None):
    run.font.name = font
    if size is not None:
        run.font.size = Pt(size)
    rf = run._r.get_or_add_rPr().get_or_add_rFonts()
    for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + k), font)
    for k in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + k), None)


def preserve_english_edits(doc):
    """Keep the user's abbreviated opening sections when rebuilding from source."""
    for start, end in [('Part A.', 'A-1.'), ('Part B.', 'B-1.'), ('Part C.', 'Parameters')]:
        first = find(doc, start)
        node = first._p.getnext()
        while node is not None:
            p = Paragraph(node, doc._body)
            if p.text.startswith(end):
                break
            nxt = node.getnext()
            node.getparent().remove(node)
            node = nxt
    for p in list(doc.paragraphs):
        if p.text == 'Point Processing, Spatial & Frequency-Domain Filtering, Wiener Restoration':
            nxt = p._p.getnext()
            remove(p)
            if nxt is not None and not Paragraph(nxt, doc._body).text and not nxt.xpath('.//w:drawing'):
                nxt.getparent().remove(nxt)
        elif p.text.startswith('See homework1_cv.py'):
            p.text = p.text.removeprefix('See ')
    chroma = [p for p in doc.paragraphs if p.text.startswith('U and V are approximately')]
    if chroma:
        luma = find(doc, 'Y (luminance):')
        luma.add_run(' ' + chroma[0].text)
        remove(chroma[0])


def question_style(doc, font):
    """Use an actual Word bullet whose glyph is an ASCII hyphen."""
    if 'Homework Answer' in doc.styles:
        # Word expects abstract numbering definitions before concrete instances.
        root = doc.part.numbering_part.element
        first_num = root.find(qn('w:num'))
        for abstract in list(root.findall(qn('w:abstractNum'))):
            first_num.addprevious(abstract)
        return
    style = doc.styles.add_style('Homework Answer', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles['Normal']
    style.font.name = font
    style.font.size = Pt(10.5)
    style.paragraph_format.left_indent = Inches(.17)
    style.paragraph_format.first_line_indent = Inches(-.17)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_together = True
    style.paragraph_format.widow_control = True
    numbering = doc.part.numbering_part.element
    aid = max([int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))] + [-1]) + 1
    nid = max([int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))] + [0]) + 1
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(aid))
    lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), '0')
    for tag, val in [('start', '1'), ('numFmt', 'bullet'), ('lvlText', '-'), ('suff', 'space'), ('lvlJc', 'left')]:
        el = OxmlElement('w:' + tag); el.set(qn('w:val'), val); lvl.append(el)
    rpr = OxmlElement('w:rPr'); rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), font); rf.set(qn('w:hAnsi'), font); rpr.append(rf); lvl.append(rpr)
    abstract.append(lvl)
    first_num = numbering.find(qn('w:num'))
    if first_num is None: numbering.append(abstract)
    else: first_num.addprevious(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(nid))
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(aid)); num.append(ref); numbering.append(num)
    npr = style.element.get_or_add_pPr().get_or_add_numPr()
    npr.get_or_add_ilvl().val = 0
    npr.get_or_add_numId().val = nid


def format_answers(doc, ko):
    font = 'Malgun Gothic' if ko else 'Times New Roman'
    question_style(doc, font)
    if any(p.style.name == 'Homework Answer' for p in doc.paragraphs):
        for p in doc.paragraphs:
            if p.style.name == 'Homework Answer':
                for num in p._p.xpath('./w:pPr/w:numPr'): num.getparent().remove(num)
                p._p.get_or_add_pPr().append(deepcopy(doc.styles['Homework Answer'].element.pPr.numPr))
        return
    # Explanations, observations and parameter lists remain ordinary prose.
    for p in doc.paragraphs:
        if p.style.name.startswith('List Bullet'):
            p.style = doc.styles['Normal']
            p.paragraph_format.left_indent = None
            p.paragraph_format.first_line_indent = None
        for num in p._p.xpath('./w:pPr/w:numPr'):
            num.getparent().remove(num)

    def answer(prefix, number, question, text=None):
        p = find(doc, prefix) if isinstance(prefix, str) else prefix
        body = p.text if text is None else text
        if text is None and isinstance(prefix, str) and ': ' in body and (prefix.endswith(':') or prefix.startswith(('Why ', '왜 ', 'CLAHE가 잡음', '샤프닝이 고주파', '블러가 저역통과'))):
            body = body.split(': ', 1)[1]
        p.clear()
        p.style = doc.styles['Homework Answer']
        p._p.get_or_add_pPr().append(deepcopy(doc.styles['Homework Answer'].element.pPr.numPr))
        p.paragraph_format.left_indent = Inches(.17)
        p.paragraph_format.first_line_indent = Inches(-.17)
        p.paragraph_format.keep_with_next = False
        lead = ('질문' if ko else 'Question ') + str(number) + ': ' + question + ' '
        r = p.add_run(lead); r.bold = True; set_font(r, font)
        r = p.add_run(body); r.bold = False; set_font(r, font)
        return p

    if ko:
        luma, chroma = find(doc, 'Y (휘도):'), find(doc, 'U, V (색차):')
        merged = luma.text + ' ' + chroma.text
        remove(chroma)
        answer(luma, 1, 'Y, U, V는 각각 무엇을 의미하는가?', merged)
    else:
        answer('Y (luminance):', 1, 'What do Y, U and V represent?')
    answer('스트레칭 vs HE:' if ko else 'Stretching applies', 1,
           '대비 스트레칭은 HE와 어떻게 다른가?' if ko else 'How does contrast stretching differ from HE?')
    for i, (kp, ep, kq, eq) in enumerate([
        ('γ<1', 'gamma < 1', 'γ<1이면 왜 밝아지는가?', 'Why does gamma < 1 brighten the image?'),
        ('γ=1:', 'gamma = 1:', 'γ=1이면 왜 변하지 않는가?', 'Why is the image unchanged at gamma = 1?'),
        ('γ>1', 'gamma > 1', 'γ>1이면 왜 어두워지는가?', 'Why does gamma > 1 darken the image?'),
    ], 1):
        answer(kp if ko else ep, i, kq if ko else eq)
    answer('국소 CDF를' if ko else 'using a local CDF', 1,
           'AHE가 전역 HE보다 국소 대비를 더 잘 개선하는 이유는?' if ko else 'Why does AHE improve local contrast better than global HE?')
    answer('잡음 증폭 이유:' if ko else 'noise amplification:', 2,
           'AHE가 잡음을 증폭하는 이유는?' if ko else 'Why can AHE amplify noise?')
    for i, (kp, ep, kq, eq) in enumerate([
        ('CLAHE가 잡음', 'Why CLAHE limits', 'CLAHE는 어떻게 잡음 증폭을 줄이는가?', 'How does CLAHE reduce noise amplification?'),
        ('clip limit의 의미:', 'Meaning of the clip limit:', 'clip limit은 국소 대비에 어떤 영향을 주는가?', 'How does the clip limit affect local contrast?'),
        ('clip이 작으면', 'A small limit makes', 'clip limit이 너무 작거나 크면 어떻게 되는가?', 'What happens if the clip limit is too small or too large?'),
    ], 1):
        answer(kp if ko else ep, i, kq if ko else eq)
    p = find(doc, '상위 비트(bit' if ko else 'high-order planes')
    old = p.text
    split = ' 하위 비트' if ko else ' Low-order planes'
    left, right = old.split(split, 1)
    second = after(p, split.strip() + right)
    answer(p, 1, '시각적 구조를 가장 많이 담는 비트평면은?' if ko else 'Which bit planes contain most of the visual structure?', left)
    answer(second, 2, '미세 디테일이나 잡음이 주로 나타나는 비트평면은?' if ko else 'Which bit planes mainly contain fine details or noise?')
    answer('왜 공간=주파수:' if ko else 'Why spatial and frequency match:', 1,
           '공간영역과 주파수영역 결과는 동일한가?' if ko else 'Are the spatial- and frequency-domain results identical?')
    p = find(doc, '블러가 저역통과인 이유:' if ko else 'Why blur is low-pass:')
    second = after(p, '빠른 밝기 변화에 해당하는 고주파 성분, 즉 에지와 미세 질감이 전반적으로 감쇠한다. 상자 필터의 응답에는 영점과 측엽이 있으므로 감쇠는 주파수에 따라 단조롭지 않다.' if ko else 'High-frequency components associated with abrupt intensity changes, edges and fine texture are generally attenuated. The box-filter response has zeros and sidelobes, so attenuation is not monotonic with frequency.')
    answer(p, 1, '블러 필터가 저역통과로 동작하는 이유는?' if ko else 'Why does the blur filter behave as a low-pass filter?')
    answer(second, 2, '블러 필터가 감쇠시키는 주파수 성분은?' if ko else 'Which frequency components does the blur filter attenuate?')
    answer('샤프닝이 고주파' if ko else 'Why sharpening emphasizes', 3,
           '샤프닝 필터가 고주파를 강조하는 이유는?' if ko else 'Why does the sharpening filter emphasize high frequencies?')
    answer('에지와 고주파:' if ko else 'Edges and high freq:', 4,
           '에지는 고주파 성분과 어떤 관계인가?' if ko else 'How are edges related to high-frequency components?')
    answer('왜 h를 임펄스' if ko else 'Why h is called', 1,
           'h를 임펄스 응답이라고 부르는 이유는?' if ko else 'Why is h called the impulse response?')
    p = answer('왜 블러를 빼면' if ko else 'Why subtracting the blur', 1,
           'Gaussian 블러를 빼면 왜 고주파 정보가 추출되는가?' if ko else 'Why does subtracting the Gaussian blur extract high-frequency information?')
    sigma = find(doc, 'σ의 영향:' if ko else 'Effect of sigma:')
    # Keep the five discussion answers together after the verification figure.
    sigma._p.addprevious(p._p)
    answer(sigma, 2, 'σ는 추출되는 고주파 성분에 어떤 영향을 주는가?' if ko else 'How does sigma affect the extracted high-frequency components?')
    p = find(doc, 'k의 영향:' if ko else 'Effect of k:')
    second = after(p, 'k가 너무 크면 에지 오버슈트(halo/링잉), 잡음 증폭, 표시 시 클리핑과 포화가 증가한다.' if ko else 'Excessively large k causes edge overshoot (halos/ringing), noise amplification and more clipping or saturation when displayed.')
    answer(p, 3, 'k는 선명화 강도에 어떤 영향을 주는가?' if ko else 'How does k affect sharpening strength?',
           'k는 고주파 잔차를 원본에 더하는 강도이다. k가 클수록 잔차가 강하게 더해져 에지와 디테일이 더 강조된다.' if ko else 'k scales the high-frequency residual added to the original. Increasing k strengthens edge and detail emphasis.')
    answer(second, 4, 'k가 너무 크면 어떻게 되는가?' if ko else 'What happens when k is too large?')
    answer('고정 커널 h_s와 비교:' if ko else 'Comparison with the fixed kernel:', 5,
           '앞의 고정 샤프닝 커널과 어떻게 다른가?' if ko else 'How does this compare with the fixed sharpening kernel?')
    answer('왜 합성곱이 곱에' if ko else 'Why convolution corresponds', 1,
           '공간 합성곱이 주파수영역의 곱과 대응하는 이유는?' if ko else 'Why does spatial convolution correspond to frequency-domain multiplication?')
    starts = ['순환 합성곱과 패딩:', '경계 처리:', '수치 정밀도:'] if ko else ['Padding and circular convolution:', 'Boundary handling:', 'Numerical precision:']
    parts = [find(doc, s) for s in starts]
    combined = ' '.join(p.text for p in parts)
    label = find(doc, '실제로 작은 차이가' if ko else 'Why the two can differ slightly') if not ko else next(p for p in doc.paragraphs if p._p is parts[0]._p.getprevious())
    answer(label, 2, '실제 구현에서 작은 차이가 발생하는 이유는?' if ko else 'Why can small differences occur in practice?', combined)
    for p in parts:
        remove(p)
    for i, (kp, ep, kq, eq) in enumerate([
        ('K가 너무 작으면:', 'K too small:', 'K가 너무 작으면 어떻게 되는가?', 'What happens when K is too small?'),
        ('K가 너무 크면:', 'K too large:', 'K가 너무 크면 어떻게 되는가?', 'What happens when K is too large?'),
        ('K의 역할:', 'Role of K:', 'K는 잡음 억제와 선명도의 균형을 어떻게 조절하는가?', 'How does K balance noise suppression and sharpness?'),
    ], 1):
        last = answer(kp if ko else ep, i, kq if ko else eq)
    observed = find(doc, '그림 관찰: K=10⁻⁶' if ko else 'Observed in the figure: K=1e-6')
    observed.text = ('그림 관찰: K=10⁻⁶은 강한 잡음으로 구조 식별이 어렵고, 10⁻⁴에서는 피사체 윤곽과 큰 잡음이 함께 보인다. 10⁻¹은 더 매끈하지만 일부 세부 구조가 약해진다.' if ko else 'Observed in the figure: K=1e-6 severely obscures structure with noise; at 1e-4 the subject is visible amid substantial noise. K=1e-1 is smoother but loses some fine detail.')
    best = after(last, '이번 다섯 후보 중 cameraman의 PSNR은 K=10⁻²에서 가장 높다. 잡음 증폭과 과도한 평활화 사이에서 화소 오차가 가장 작은 선택이다. SSIM 기준으로는 K=10⁻¹이 가장 높으므로 최적값은 평가 지표에 따라 달라진다.' if ko else 'Among the five tested values for cameraman, K=1e-2 gives the highest PSNR, balancing noise amplification and excessive smoothing to minimize pixel error. K=1e-1 gives the highest SSIM, so the preferred value depends on the evaluation metric.')
    answer(best, 4, '가장 좋은 복원을 주는 K는 무엇이며 그 이유는?' if ko else 'Which K gives the best restoration, and why?')
    for r in observed.runs:
        set_font(r, font)


TABLE_EN = [
    ('Contrast stretching', 'stretch', 'Global', 'Constant gain within the chosen range; tails clipped.', 'Signal and noise both increase where gain is large.', 'O(LMN+L)', 'Narrow dynamic range; preprocessing'),
    ('Histogram equalization', 'he', 'Global', 'CDF-dependent gain across intensity intervals.', 'Can amplify noise in populated intervals; higher mean does not imply higher global contrast.', 'O(LMN+L)', 'Redistribute a global intensity distribution'),
    ('AHE', 'ahe', 'Local', 'Strong local enhancement using tile CDFs.', 'Amplifies small variations in flat tiles; no interpolation here, so seams can appear.', 'O(LMN+TL)', 'Spatially varying contrast; consider noise and seams'),
    ('CLAHE', 'clahe', 'Local', 'Clipping moderates local gain; interpolation reduces seams.', 'Larger clip limits can increase detail and noise. Redistribution means the count cap is not strict.', 'O(LMN+TL)', 'Controlled local enhancement; low-light images'),
    ('Gamma correction', 'gamma', 'Global', 'Gamma < 1 expands dark tones; gamma > 1 expands bright tones.', 'Existing noise can increase in the expanded intensity range.', 'O(MN)', 'Gamma and tone adjustment'),
]
TABLE_KO = [
    ('대비 스트레칭', 'stretch', '전역', '선택 범위 안에서 일정한 이득; 양 끝은 클리핑.', '이득이 큰 구간에서는 신호와 잡음이 함께 증가.', 'O(LMN+L)', '좁은 동적범위 영상의 전처리'),
    ('히스토그램 평활화', 'he', '전역', 'CDF에 따라 밝기 구간별 이득이 달라짐.', '밀집 구간의 잡음을 증폭할 수 있음. 평균 증가가 전체 대비 증가를 뜻하지는 않음.', 'O(LMN+L)', '전역 밝기 분포 재분배'),
    ('AHE', 'ahe', '국소', '타일별 CDF로 국소 대비를 강하게 향상.', '평탄한 타일의 미세 변동도 증폭. 보간이 없어 타일 경계가 나타남.', 'O(LMN+TL)', '공간적으로 대비가 다른 영상; 잡음·경계 고려'),
    ('CLAHE', 'clahe', '국소', 'clipping으로 이득 완화, 보간으로 경계 완화.', 'clip이 크면 디테일과 잡음이 모두 증가 가능. 재분배 후 count는 엄격한 상한이 아님.', 'O(LMN+TL)', '제어된 국소 향상, 저조도 영상'),
    ('감마 보정', 'gamma', '전역', 'γ<1은 암부, γ>1은 밝은 구간을 확장.', '확장되는 밝기 구간의 기존 잡음도 증폭 가능.', 'O(MN)', '감마 및 톤 조정'),
]


def restore_a9_table(doc, ko):
    heading, following = find(doc, 'A-9.'), find(doc, 'Part B.')
    # A second call must not add another table or section boundary.
    existing = heading._p.getnext()
    if existing is not None and existing.tag == qn('w:tbl'):
        return
    node = heading._p.getnext()
    while node is not following._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt
    # A wide comparison table gets one landscape section, then return to portrait.
    portrait = deepcopy(doc.sections[0]._sectPr)
    for el in portrait.findall(qn('w:type')):
        portrait.remove(el)
    typ = OxmlElement('w:type'); typ.set(qn('w:val'), 'nextPage'); portrait.insert(0, typ)
    landscape = deepcopy(portrait)
    size = landscape.find(qn('w:pgSz'))
    width, height = size.get(qn('w:w')), size.get(qn('w:h'))
    size.set(qn('w:w'), height); size.set(qn('w:h'), width); size.set(qn('w:orient'), 'landscape')
    boundary = OxmlElement('w:p'); heading._p.addprevious(boundary)
    Paragraph(boundary, doc._body)._p.get_or_add_pPr().append(portrait)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.page_break_before = False
    table = doc.add_table(rows=1, cols=8)
    table.style = 'Table Grid'
    table.autofit = False
    heading._p.addnext(table._tbl)
    headers = (['방법', '결과 영상', '범위', '대비 향상', '잡음 민감도', '히스토그램', '계산 복잡도', '적용 분야'] if ko else
               ['Method', 'Result image', 'Scope', 'Contrast improvement', 'Noise sensitivity', 'Histogram', 'Complexity', 'Suitable applications'])
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    thumbs = json.loads((BASE / 'results_A.json').read_text(encoding='utf-8'))['a9_thumbs']
    for name, key, scope, contrast, noise, cost, app in (TABLE_KO if ko else TABLE_EN):
        cells = table.add_row().cells
        for index, text in [(0, name), (2, scope), (3, contrast), (4, noise), (6, cost), (7, app)]:
            cells[index].text = text
        ip, hp = thumbs[key]
        cells[1].paragraphs[0].add_run().add_picture(str(BASE / ip), width=Inches(1.02))
        cells[5].paragraphs[0].add_run().add_picture(str(BASE / hp), width=Inches(1.16))
    widths = [.85, 1.15, .55, 1.6, 1.8, 1.3, 1., 1.45]
    font = 'Malgun Gothic' if ko else 'Times New Roman'
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EDF2' if i == 0 else 'FFFFFF'); pr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: e.set(qn('w:' + k), v)
                borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', 80), ('bottom', 80), ('left', 60), ('right', 60)]:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), str(value)); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or j in [1, 2, 5, 6] else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_font(r, font, 8.5)
                    r.bold = i == 0 or j == 0
                    r.font.color.rgb = RGBColor(0, 0, 0)
    note = OxmlElement('w:p'); table._tbl.addnext(note)
    p = Paragraph(note, doc._body)
    p.add_run('복잡도는 레벨별로 화소를 탐색하는 현재 구현 기준이다. 영상 크기는 M×N, L=256, T=64이며 L과 T를 고정하면 화소 수에 선형이다. 단일 패스 히스토그램은 O(MN+L)로 구현할 수 있다. 표시 예: CLAHE clip=0.01, γ=0.5.' if ko else 'Complexity describes this implementation, which scans pixels for each intensity level: image MxN, L=256, T=64. With L and T fixed, cost is linear in pixel count. Single-pass histogram counting can use O(MN+L). Displayed examples: CLAHE clip=0.01 and gamma=0.5.')
    p.paragraph_format.space_before = Pt(6)
    for r in p.runs: set_font(r, font, 9)
    p._p.get_or_add_pPr().append(landscape)
    following.paragraph_format.page_break_before = False


def apply_presentation(doc, lang, rebuilding=False):
    ko = lang == 'ko'
    if rebuilding and not ko:
        preserve_english_edits(doc)
    format_answers(doc, ko)
    restore_a9_table(doc, ko)
    for p in doc.paragraphs:
        if p.text in ['Spatial and frequency verification', '공간·주파수 영역 검증', '비너 필터의 목적']:
            p.paragraph_format.keep_with_next = True


if __name__ == '__main__':
    for lang, name in [('en', 'Homework1_Report.docx'), ('ko', 'Homework1_Report_KO.docx')]:
        path = BASE / name
        doc = Document(path)
        apply_presentation(doc, lang)
        doc.save(path)
        print(name, 'question answers:', sum(p.style.name == 'Homework Answer' for p in doc.paragraphs), 'tables:', len(doc.tables))
