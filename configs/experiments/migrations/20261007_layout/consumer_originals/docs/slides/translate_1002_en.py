#!/usr/bin/env python3
"""kor/1002_labmeeting_ko_draft.pptx → eng/1002_labmeeting_en_draft.pptx.

The Korean slide order, charts and numbers are retained. Text and speaker notes use
the checked English translation file, with explicit font and spacing adjustments.
The four data-review matrices are redrawn from the same CSVs with English annotations.
Untranslated Korean text or changed numeric cells fail validation.
"""
import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
IN = HERE / 'kor/1002_labmeeting_ko_draft.pptx'
OUT = HERE / 'eng/1002_labmeeting_en_draft.pptx'
KO_ASSETS = HERE / 'kor/1002_intro_assets'
EN_ASSETS = HERE / 'eng/1002_intro_assets'
QA = HERE / 'eng/1002_operational_qa'
DATASETS = [('cse_cic_ids2018', 'CIC2018'), ('ton_iot', 'ToN-IoT'),
            ('bot_iot', 'BoT-IoT'), ('unsw_nb15', 'UNSW-NB15')]
LABELS = dict(benign='Benign', bot='Bot', brute_force='Brute force', ddos='DDoS', dos='DoS',
              infiltration='Infiltration', web_attacks='Web attack', backdoor='Backdoor',
              injection='Injection', mitm='MitM', password='Password', ransomware='Ransomware',
              scanning='Scanning', xss='XSS', reconnaissance='Recon', theft='Theft',
              analysis='Analysis', exploits='Exploits', fuzzers='Fuzzers', generic='Generic',
              shellcode='Shellcode', worms='Worms')
SHORT = dict(benign='Ben', bot='Bot', brute_force='Brute', ddos='DDoS', dos='DoS',
             infiltration='Inf', web_attacks='Web', backdoor='Back', injection='Inj',
             mitm='MitM', password='Pass', ransomware='Ran', scanning='Scan', xss='XSS',
             reconnaissance='Recon', theft='Theft', analysis='Anal', exploits='Expl',
             fuzzers='Fuzz', generic='Gen', shellcode='Shell', worms='Worm')


def redraw_matrices():
    """Same figure as intro_1002.prepare_data, drawn from the saved Korean assets with an English annotation."""
    EN_ASSETS.mkdir(parents=True, exist_ok=True)
    for path in (Path.home() / '.local/share/fonts/Pretendard').glob('*.otf'):
        font_manager.fontManager.addfont(path)
    plt.rcParams.update({'font.family': 'Pretendard', 'svg.fonttype': 'none', 'axes.unicode_minus': False})
    prov = json.loads((KO_ASSETS / 'provenance.json').read_text())
    cmap = LinearSegmentedColormap.from_list('KENTECH', ['#F2F7FB', '#80B6DA', '#1875B4', '#00306C'])
    out = {}
    for ds, title in DATASETS:
        m = pd.read_csv(KO_ASSETS / f'{ds}_pair_percent.csv', index_col=0)
        classes = list(m.index)
        assert classes == prov[ds]['class_order'] and list(m.columns) == classes
        matrix = m.values
        fig = plt.figure(figsize=(5.90, 2.25), facecolor='white')
        ax = fig.add_axes([0.215, 0.035, 0.765, 0.71])
        ax.imshow(matrix, cmap=cmap, vmin=0, vmax=100, aspect='auto', interpolation='nearest')
        ax.set_xticks(range(len(classes)), [SHORT[c] for c in classes], fontsize=9.3)
        ax.set_yticks(range(len(classes)), [LABELS[c] for c in classes], fontsize=9.5)
        ax.tick_params(axis='both', length=0, pad=4)
        ax.xaxis.tick_top()
        for spine in ax.spines.values():
            spine.set_visible(False)
        for i in range(len(classes)):
            for j in range(len(classes)):
                if i == j:
                    ax.add_patch(Rectangle((j - .5, i - .5), 1, 1, facecolor='#E9EDF1', edgecolor='white', linewidth=.6))
                    continue
                v = matrix[i, j]
                if v:
                    label = f'{v:.1f}' if v >= 1 else '·'
                    ax.text(j, i, label, ha='center', va='center', fontsize=8.5 if len(classes) >= 10 else 9.3,
                            color='white' if v >= 55 else '#163F60', weight='bold' if v >= 50 else 'normal')
        ax.set_xticks(np.arange(-.5, len(classes), 1), minor=True)
        ax.set_yticks(np.arange(-.5, len(classes), 1), minor=True)
        ax.grid(which='minor', color='white', linewidth=.6)
        ax.tick_params(which='minor', bottom=False, left=False, top=False)
        fig.text(.015, .945, title, color='#00306C', fontsize=14.5, weight='bold', va='center')
        fig.text(.985, .945, f"Conflicting rows {prov[ds]['conflict_fraction'] * 100:.2f}%",
                 color='#5D6773', fontsize=10.3, va='center', ha='right')
        fig.savefig(EN_ASSETS / f'{ds}_matrix.png', dpi=260)
        fig.savefig(EN_ASSETS / f'{ds}_matrix.svg')
        plt.close(fig)
        out[ds] = EN_ASSETS / f'{ds}_matrix.png'
    (EN_ASSETS / 'README.md').write_text(
        'English redraw of kor/1002_intro_assets/*_matrix.png from the saved *_pair_percent.csv and provenance.json; '
        'numbers unchanged, only the annotation text differs.\n')
    return out


def replace_picture(slide, pic, path):
    sp = pic._element
    parent = sp.getparent()
    idx = parent.index(sp)
    rid = sp.blip_rId
    new = slide.shapes.add_picture(str(path), pic.left, pic.top, pic.width, pic.height)
    parent.remove(new._element)
    parent.insert(idx, new._element)
    parent.remove(sp)
    slide.part.drop_rel(rid)


def main():
    import hashlib
    from slide_style import set_text, Inches, PP_ALIGN

    def sha(path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    config = json.loads((HERE / '1002_english_translation.json').read_text())
    translations, notes = config['translations'], config['notes']
    korean_hashes = {p.name: sha(p) for p in (HERE / 'kor').glob('1002_labmeeting_ko_draft.*')}
    prs = Presentation(IN)
    assert len(prs.slides) == 10
    assert prs.slides[3].shapes[0].text == 'ToN · 신규 공격 반영 성능과 비용'
    hangul = re.compile('[가-힣]')
    numeric = re.compile(r'^[+−-]?[\d,]+(?:\.\d+)?%?$')
    untranslated = []
    replaced = 0
    table_numbers = []
    bodies = {
        '공통 context로\n기본 분류기 구성',
        '입력 표현·예측 확률\n정답과의 차이·오류 크기',
        '실패 특성에 따라\nK개 군집으로 분할',
        '공통 anchor\n+ 군집별 특화 block',
        'Global·expert의\n교정·훼손 사례 학습',
        '기본 클래스 예측', 'Expert 선택\n호출 여부 판단',
        '선택한 context로 분류', 'Expert 예측의\n채택 여부 판단',
        '채택 시 expert 예측',
    }

    def translate_frame(frame, slide_number, shape=None):
        nonlocal replaced
        old = frame.text
        if old not in translations:
            if hangul.search(old):
                untranslated.append((slide_number, old))
            return
        runs = [r for p in frame.paragraphs for r in p.runs if r.text]
        assert runs, (slide_number, old)
        source = runs[0]
        size = source.font.size.pt if source.font.size else 13
        if slide_number == 7 and old in bodies:
            size = 10.5
        if slide_number == 5 and old.startswith('행: 원래 라벨'):
            size = 11.3
        if slide_number == 5 and old.startswith('빈칸:'):
            size = 11
            assert shape is not None
            shape.top, shape.height = Inches(6.23), Inches(.52)
        if slide_number == 8 and old.startswith('정적 평가에서의'):
            size = 16
        if slide_number == 10 and old.startswith('연구의 판단 기준:'):
            size = 20
            assert shape is not None
            shape.height = Inches(.75)
        ink = str(source.font.color.rgb) if source.font.color.type is not None else '21262E'
        align = frame.paragraphs[0].alignment or PP_ALIGN.LEFT
        margins = (frame.margin_left, frame.margin_right, frame.margin_top, frame.margin_bottom)
        anchor = frame.vertical_anchor
        set_text(frame, translations[old], size=size, bold=bool(source.font.bold), ink=ink,
                 align=align, font=source.font.name or 'Pretendard', inset=0)
        frame.margin_left, frame.margin_right, frame.margin_top, frame.margin_bottom = margins
        frame.vertical_anchor = anchor
        replaced += 1

    for i, slide in enumerate(prs.slides, 1):
        table_index = 0
        for sh in slide.shapes:
            if sh.has_text_frame:
                translate_frame(sh.text_frame, i, sh)
            if sh.has_table:
                for ri, row in enumerate(sh.table.rows):
                    for ci, cell in enumerate(row.cells):
                        if numeric.fullmatch(cell.text):
                            table_numbers.append((i, table_index, ri, ci, cell.text))
                        translate_frame(cell.text_frame, i)
                table_index += 1
        assert str(i) in notes
        slide.notes_slide.notes_text_frame.text = notes[str(i)]
    assert not untranslated, untranslated

    # Translate the four embedded scientific figures from the exact source CSVs.
    pngs = redraw_matrices()
    slots = {(0.62, 1.47): 'cse_cic_ids2018', (6.81, 1.47): 'ton_iot',
             (0.62, 3.89): 'bot_iot', (6.81, 3.89): 'unsw_nb15'}
    audit = prs.slides[4]
    pictures = [sh for sh in audit.shapes if sh.shape_type == 13]
    assert len(pictures) == 4
    used = []
    for pic in pictures:
        slot = min(slots, key=lambda k: abs(k[0] - pic.left / 914400) + abs(k[1] - pic.top / 914400))
        used.append(slots[slot])
        replace_picture(audit, pic, pngs[slots[slot]])
    assert len(set(used)) == 4
    for slide in prs.slides:
        frames = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                frames.append(sh.text_frame)
            if sh.has_table:
                frames.extend(c.text_frame for row in sh.table.rows for c in row.cells)
        frames.append(slide.notes_slide.notes_text_frame)
        for frame in frames:
            for p in frame.paragraphs:
                for r in p.runs:
                    r._r.get_or_add_rPr().set('lang', 'en-US')
    prs.core_properties.title = 'TabPFN for operational intrusion detection'
    prs.core_properties.subject = 'New attack adaptation, data review, static performance and cost'
    prs.core_properties.comments = 'English translation of the latest 10-slide Korean deck, including completed CIC2018 and ToN EXP70 results'
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)

    check = Presentation(OUT)
    actual_numbers, dump, issues, titles = [], [], [], []
    for i, slide in enumerate(check.slides, 1):
        entries, text_shapes = [], []
        table_index = 0
        for sh in slide.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > check.slide_width + 20 or sh.top + sh.height > check.slide_height + 20:
                issues.append((i, 'outside slide', sh.name))
            if sh.top < Inches(6.9) and sh.top + sh.height > Inches(6.9):
                issues.append((i, 'footer overlap', sh.name))
            if sh.has_text_frame and sh.text.strip():
                entries.append(sh.text)
                text_shapes.append(sh)
                if sh.top < Inches(.96):
                    titles.append(sh.text)
            if sh.has_table:
                for ri, row in enumerate(sh.table.rows):
                    entries.append(' | '.join(c.text for c in row.cells))
                    for ci, cell in enumerate(row.cells):
                        if numeric.fullmatch(cell.text):
                            actual_numbers.append((i, table_index, ri, ci, cell.text))
                            if ri and ci:
                                assert all(p.alignment == PP_ALIGN.RIGHT for p in cell.text_frame.paragraphs)
                table_index += 1
        for ai, a in enumerate(text_shapes):
            for b in text_shapes[ai + 1:]:
                if (min(a.left + a.width, b.left + b.width) - max(a.left, b.left) > Inches(.005)
                    and min(a.top + a.height, b.top + b.height) - max(a.top, b.top) > Inches(.005)):
                    issues.append((i, 'text boxes overlap', a.text[:25], b.text[:25]))
        joined = '\n'.join(entries)
        assert not hangul.search(joined), (i, 'untranslated slide text')
        assert not hangul.search(slide.notes_slide.notes_text_frame.text), (i, 'untranslated notes')
        assert 'TBD' not in joined
        dump.append(f'## {i}\n{joined}\n')
    assert actual_numbers == table_numbers, 'Numeric table cells changed during translation'
    assert not issues, issues
    assert korean_hashes == {p.name: sha(p) for p in (HERE / 'kor').glob('1002_labmeeting_ko_draft.*')}
    QA.mkdir(exist_ok=True)
    (QA / 'slide_text.md').write_text('\n'.join(dump))
    (QA / 'validation.json').write_text(json.dumps(dict(
        slides=len(check.slides), source=str(IN), source_sha256=sha(IN),
        sha256=sha(OUT), titles=titles, text_replacements=replaced,
        numeric_cells_unchanged=len(actual_numbers), numeric_cells_right_aligned=True,
        korean_unchanged=True, untranslated_text=[], geometry_issues=issues,
        translated_matrix_sources={ds: sha(KO_ASSETS / f'{ds}_pair_percent.csv') for ds, _ in DATASETS},
        translator='docs/slides/translate_1002_en.py', translation_file='docs/slides/1002_english_translation.json'),
        ensure_ascii=False, indent=2) + '\n')
    print('Saved', OUT, 'slides=', len(check.slides), 'translated frames=', replaced,
          'unchanged numeric cells=', len(actual_numbers))


if __name__ == '__main__':
    main()
