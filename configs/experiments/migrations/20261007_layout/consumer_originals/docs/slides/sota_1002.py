"""Measured SOTA and cost slides, backed by the same data as the HTML report."""
import json
from pathlib import Path
from pptx.enum.text import PP_ALIGN
from intro_1002 import footer

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/'docs/research/20260930/exp61_report_data.json'
METHODS=['global_raw','xgb','boostpfn','localpfn','distpfn','xgb_full']
NAMES=dict(global_raw='Global TabPFN v3',xgb='XGBoost',boostpfn='BoostPFN',localpfn='LoCalPFN · FT',distpfn='DistPFN · v3',xgb_full='XGBoost · 전체 train')


def highlight(h,shape,row,cols,fill=None):
    for col in cols:
        cell=shape.table.cell(row,col)
        if fill:cell.fill.fore_color.rgb=h.color(fill)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:r.font.bold=True;r.font.color.rgb=h.color(h.BLUE)


def build(h):
    data=json.loads(DATA.read_text())
    ds=data['datasets'];a=ds['cic2018']['results'];b=ds['toniot']['results']
    s=h.slide('정제 데이터의 SOTA 분류 성능','SOTA 비교')
    h.text(s,.62,1.04,12.09,.34,'Seed 43 · 동일 C0 100,000행 · CIC2018 test 304만 행 / ToN test 227만 행',size=14,ink=h.GREY,inset=0)
    rows=[]
    for m in METHODS:
        rows.append([NAMES[m],f'{a[m]["macro_f1"]:.4f}',f'{b[m]["macro_f1"]:.4f}', '100,000' if m!='xgb_full' else '866만 / 467만'])
    t=h.table(s,.62,1.60,[3.90,2.30,2.30,3.59],['방법','CIC2018 Macro-F1','ToN Macro-F1','학습 풀 행 수 · CIC / ToN'],rows,row_h=.49,font=15)
    highlight(h,t,5,[1],h.HEAD);highlight(h,t,2,[2],h.HEAD)
    for cell in t.table.rows[6].cells:cell.fill.fore_color.rgb=h.color('E8EDF2')
    h.text(s,.62,5.18,12.09,.35,'전체 train XGB: 추가 학습량을 허용한 참고 기준 · Global: expert·S/V 적용 전 기본 분류기',size=12,ink=h.GREY,inset=0)
    for x,title,value in [(.62,'CIC2018 · 클래스별 차이','Infiltration 최고 F1 0.3055 · DistPFN\nWeb attack 최고 F1 0.3321 · XGB'),(6.81,'ToN · 클래스별 차이','MitM 최고 F1 0.2363 · XGB\nScanning 최고 F1 0.3281 · BoostPFN')]:
        h.box(s,x,5.79,5.90,.98,fill=h.CARD)
        h.text(s,x+.14,5.84,5.62,.24,title+' · 공통 100k',size=12,bold=True,ink=h.NAVY,inset=0)
        h.text(s,x+.14,6.15,5.62,.49,value,size=12,ink=h.INK,inset=0)
    footer(h,s,'동일 10만 행 기준 최고 Macro-F1 · CIC2018 DistPFN 0.7839 · ToN XGB 0.6983')
    s.notes_slide.notes_text_frame.text=(
        'EXP61 완료 결과, 2026-09-30 갱신\n'
        '모순 제거 fixed split 및 Global C0의 정확히 같은 100,000개 학습 ID 사용\n'
        'Test: CIC2018 3,042,473행, ToN 2,271,723행, 전체 test 평가\n'
        'XGB 전체 train: CIC2018 8,666,430행, ToN 4,669,119행, 별도 학습량 기준\n'
        'Global/DistPFN backbone v3, BoostPFN/LoCalPFN backbone v1\n'
        'LoCalPFN은 validation AUC로 미세조정 checkpoint 선택, validation CIC 30,194 / ToN 41,744행\n'
        'BoostPFN T=50, context 500, LoCalPFN context 1000, FT 21 epochs × 30 steps\n'
        '동일 train ID 비교이며 backbone과 validation 사용까지 동일하다는 의미는 아님\n'
        '클래스별 최고값은 공통 100k 방법 내 비교, 전체 train XGB 제외\n'
        '현재 제안 expert/S/V 전체 시스템의 성능은 이 표에 포함하지 않음\n'
        '단일 seed, 현재 test는 이미 관찰한 development holdout\n'+str(DATA.relative_to(ROOT)))

    s=h.slide('SOTA 학습·추론 비용','SOTA 비교')
    h.text(s,.62,1.04,12.09,.35,'RTX 4090 24GB · 최대 2개 작업 병렬 · 자원 경합을 포함한 경과 시간',size=14,ink=h.GREY,inset=0)
    rows=[]
    for m in METHODS:
        rows.append([NAMES[m],f'{a[m]["fit_seconds"]:,.1f}',f'{a[m]["predict_seconds"]:,.1f}',f'{b[m]["fit_seconds"]:,.1f}',f'{b[m]["predict_seconds"]:,.1f}',f'{a[m]["peak_gpu_sampled_gib"]:.2f} / {b[m]["peak_gpu_sampled_gib"]:.2f}'])
    h.table(s,.62,1.64,[3.05,1.65,1.80,1.65,1.80,2.14],['방법','CIC 학습 (초)','CIC 추론 (초)','ToN 학습 (초)','ToN 추론 (초)','GPU GiB\nCIC / ToN'],rows,row_h=.50,font=13.3)
    h.text(s,.62,5.39,12.09,.40,'LoCalPFN 학습에 validation 포함 · 전체 test 추론에 kNN 검색 포함',size=13,ink=h.GREY,inset=0)
    h.text(s,.62,5.91,12.09,.40,'DistPFN은 Global fit·추론 공유 · prior 보정 추가 시간 약 0.19초',size=13,ink=h.GREY,inset=0)
    h.text(s,.62,6.41,12.09,.30,'GPU: 1초 간격 프로세스별 최대값 · 단독 실행 속도 비교와 구분',size=11.8,ink=h.GREY,inset=0)
    footer(h,s,'LoCalPFN 전체 test 추론 · CIC2018 4.81시간 · ToN 3.10시간')
    s.notes_slide.notes_text_frame.text=(
        '같은 EXP61 실행에서 측정한 fit, full-test prediction, 프로세스별 GPU 사용량\n'
        '2개 worker가 GPU를 공유하므로 모델 간 단독 실행 속도 배수로 해석하지 않음\n'
        'LoCalPFN validation은 fit에 포함: CIC 3,129.97초 / ToN 5,694.41초\n'
        'DistPFN fit/predict는 Global 계산 공유, 합산하여 두 번 계산하지 않음\n'
        'Global은 별도 저장된 backbone 예측 기준, DistPFN의 prior 후처리만 약 0.19초 추가\n'
        '추론 throughput과 단계별 RAM, 전체 input-loading peak는 HTML 상세표 및 cost.json 제공\n'
        '모델 RAM 범위는 전처리, 모델 로딩, fit, validation, predict이며 원본 PKL 로딩 제외\n'
        '단일 GPU 실제 전체 wall time은 약 6.65시간, 병렬 job 시간 합과 구분\n'+str(DATA.relative_to(ROOT)))

    s=h.slide('제안 모델의 최종 성능과 비용','최종 모델 비교')
    h.text(s,.62,1.04,12.09,.35,'검증된 expert·S/V 구성의 최종 비교',size=14,ink=h.GREY,inset=0)
    h.table(s,.62,1.65,[3.49,2.20,2.20,2.10,2.10],['구성','CIC2018','ToN','학습 정보','호출 비용'],
        [[m,'TBD','TBD','TBD','TBD'] for m in ['Expert · S/V 없음','Expert · S만','Expert · V만','Expert · S+V']],row_h=.64,font=14)
    h.box(s,.62,5.27,12.09,1.08,fill=h.CARD)
    h.text(s,.80,5.39,11.73,.33,'비교 기준: 5장의 SOTA 성능 · 6장의 측정 비용',size=16,bold=True,ink=h.NAVY,inset=0)
    h.text(s,.80,5.89,11.73,.27,'Expert·route의 추가 학습 행 수와 조건부 추론 비용을 포함한 비교',size=13,ink=h.GREY,inset=0)
    s.notes_slide.notes_text_frame.text='현재 seed 43, K=4 bank의 신규 S/V ablation 및 최종 제안 모델 결과 TBD\nEXP61 기준선 성능은 5장, 비용은 6장에 작성 완료\nExpert와 route의 추가 학습 라벨은 Global C0 100k 예산과 별도로 보고'
