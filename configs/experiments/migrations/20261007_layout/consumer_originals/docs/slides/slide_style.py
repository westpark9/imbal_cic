#!/usr/bin/env python3
"""Reusable KENTECH drawing helpers. Importing does not modify any deck."""
import glob
import hashlib
import json
import numpy as np
import re
import shutil
import time
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
HERE = ROOT / 'docs/slides'
OUTPUT = HERE / 'kor/1002_labmeeting_ko_draft.pptx'
QA = HERE / 'kor/1002_draft_qa'
ARCHIVE = HERE / 'archive/20261002_before_claim_first'
NAVY, BLUE, CYAN = '00306C', '007BC6', '00B8EE'
CARD, HEAD, LINE = 'F3F6FA', 'E7F0F8', 'D9E2EC'
INK, GREY, MUTED = '21262E', '5D6773', '8393A5'
GOOD, BAD = '1B7F5A', 'B0413E'
FONT, NUMBER_FONT = 'Pretendard', 'DejaVu Sans Mono'
NUM = re.compile(r'^[+−-]?[\d,]+(?:\.\d+)?%?$')


def color(value):
    return RGBColor.from_string(value)


def set_text(frame, value, size=16, bold=False, ink=INK, align=PP_ALIGN.LEFT, middle=True, inset=.10, font=FONT):
    frame.clear(); frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Inches(inset); frame.margin_top = frame.margin_bottom = Inches(.035)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if middle else MSO_ANCHOR.TOP
    for i, value_line in enumerate(str(value).split('\n')):
        p = frame.paragraphs[0] if not i else frame.add_paragraph()
        p.alignment = align; p.space_before = p.space_after = Pt(0); p.line_spacing = 1.10
        etree.SubElement(p._p.get_or_add_pPr(), qn('a:buNone'))
        r = p.add_run(); r.text = value_line; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color(ink)
        rp = r._r.get_or_add_rPr(); rp.set('lang', 'ko-KR')
        for name in ['a:ea', 'a:cs']:
            etree.SubElement(rp, qn(name)).set('typeface', font)


def text(s, x, y, w, h, value, **kwargs):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); set_text(sh.text_frame, value, **kwargs); return sh


def box(s, x, y, w, h, fill=CARD, border=None, rounded=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        sh.adjustments[0] = .035
    sh.fill.solid(); sh.fill.fore_color.rgb = color(fill); sh.shadow.inherit = False
    if border:
        sh.line.color.rgb = color(border); sh.line.width = Pt(.8)
    else:
        sh.line.fill.background()
    return sh


def table(s, x, y, widths, headers, rows, row_h=.48, font=13, first_col_bold=True):
    sh = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * (len(rows) + 1)))
    t = sh.table; t.first_row = t.horz_banding = False
    for c, width in enumerate(widths):
        t.columns[c].width = Inches(width)
    for r, values in enumerate([headers] + rows):
        t.rows[r].height = Inches(row_h)
        for c, value in enumerate(values):
            value = str(value); cell = t.cell(r, c); cell.fill.solid()
            cell.fill.fore_color.rgb = color(NAVY if r == 0 else ('FFFFFF' if r % 2 else CARD))
            numeric = bool(NUM.fullmatch(value)) and c > 0 and r > 0
            align = PP_ALIGN.RIGHT if numeric else (PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
            set_text(cell.text_frame, value, size=font - .35 if numeric else font, bold=(r == 0 or (c == 0 and first_col_bold)),
                     ink='FFFFFF' if r == 0 else INK, align=align, font=NUMBER_FONT if numeric else FONT, inset=.07)
            cell.margin_left = cell.margin_right = Inches(.07); cell.margin_top = cell.margin_bottom = Inches(.025)
    return sh


def chrome(s, title, section):
    text(s, .62, .30, 12.09, .65, title, size=28, bold=True, ink=NAVY, inset=0)
    box(s, 0, 6.95, 13.333333, .55, fill=NAVY, rounded=False)
    text(s, 9.05, 7.03, 3.00, .30, section, size=10, ink='D6E4F3', align=PP_ALIGN.RIGHT)
    text(s, 12.23, 7.01, .48, .34, '0', size=11, bold=True, ink='FFFFFF', align=PP_ALIGN.RIGHT)


def new_slide(prs, title, section):
    layout = next(l for l in prs.slide_layouts if l.name == 'TITLE_ONLY')
    s = prs.slides.add_slide(layout)
    for sh in list(s.shapes):
        s.shapes._spTree.remove(sh._element)
    chrome(s, title, section); return s


def rebuild(s, title, section):
    for sh in list(s.shapes):
        s.shapes._spTree.remove(sh._element)
    chrome(s, title, section); return s


def subtitle(s, value, size=13):
    return text(s, .62, 1.02, 12.09, .35, value, size=size, ink=GREY, inset=0)


def footer(s, value):
    text(s, .62, 7.035, 8.20, .30, value, size=11, bold=True, ink='FFFFFF', inset=0)


def card(s, x, y, w, h, title):
    box(s, x, y, w, h); box(s, x, y, w, .50, fill=HEAD, rounded=False)
    text(s, x + .17, y + .065, w - .34, .36, title, size=17, bold=True, ink=NAVY, inset=0)


def notes(s, value):
    s.notes_slide.notes_text_frame.text = value


def highlight(sh, row, cols, fill=HEAD, ink=BLUE):
    for c in cols:
        cell = sh.table.cell(row, c); cell.fill.fore_color.rgb = color(fill)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.color.rgb = color(ink)


def latest(pattern):
    paths = [p for p in sorted(glob.glob(str(ROOT / pattern))) if Path(p + '/results.json').exists()]
    return json.loads(Path(paths[-1] + '/results.json').read_text()) if paths else None


