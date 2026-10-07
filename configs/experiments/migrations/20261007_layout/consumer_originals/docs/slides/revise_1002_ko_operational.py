#!/usr/bin/env python3
"""Simplify the Korean introduction and put completed arrival evidence next to it.

Edit the immediately preceding user deck. Results are frozen in the QA folder;
--refresh-results imports only datasets completed by both compared methods.
"""
import argparse
import csv
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from slide_style import *

parser = argparse.ArgumentParser()
parser.add_argument('--refresh-results', action='store_true')
args = parser.parse_args()
ARCHIVE = HERE / 'archive/20261002_before_simplified_intro'
BASE = ARCHIVE / 'docs/slides/kor/1002_labmeeting_ko_draft.pptx'
QA = HERE / 'kor/1002_operational_qa'
QA.mkdir(exist_ok=True)
RUN = ROOT / 'tabpfn/results/20261002_exp70_class_arrival_100k_s43'
SOURCE = QA / 'result_snapshot.json'
LAST = {'cic2018': 4, 'toniot': 5}
LABELS = dict(ddos='DDoS', web_attacks='Web attack', infiltration='Infiltration',
              bot='Bot', injection='Injection', password='Password', xss='XSS',
              ransomware='Ransomware', backdoor='Backdoor', mitm='MitM')
INTRO = {'cic2018': {'ddos': 1, 'web_attacks': 2, 'infiltration': 3, 'bot': 4},
         'toniot': {'injection': 1, 'ddos': 1, 'password': 2, 'xss': 3,
                    'ransomware': 4, 'backdoor': 4, 'mitm': 5}}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


# Preserve whichever English revision is current when this Korean-only editor runs.
english_before = {str(p.relative_to(ROOT)): sha(p)
                  for p in (HERE / 'eng').glob('1002_labmeeting_en_draft.*')}


if args.refresh_results or not SOURCE.exists():
    raw = {n: (RUN / n).read_bytes() for n in ['summary.csv', 'class_metrics.csv']}
    summary = list(csv.DictReader(io.StringIO(raw['summary.csv'].decode())))
    classes = list(csv.DictReader(io.StringIO(raw['class_metrics.csv'].decode())))
    complete = []
    for ds, last in LAST.items():
        valid = True
        for method in ['xgb', 'tabpfn']:
            for stage in range(last + 1):
                job = RUN / 'jobs' / (f'{ds}_{method}_initial' if stage == 0 else
                                      f'{ds}_{method}_n100000/stage_{stage}')
                marker = job / 'COMPLETE.json'
                if not marker.exists() or json.loads(marker.read_text())['state'] != 'complete':
                    valid = False
                if len([r for r in summary if r['dataset'] == ds and r['method'] == method
                        and int(r['stage']) == stage]) != 1:
                    valid = False
        if valid:
            complete.append(ds)
    assert 'cic2018' in complete, 'CIC2018 completion must be verified before insertion'
    diagnostics = {}
    metric_hashes = {}
    if 'toniot' in complete:
        for method in ['xgb', 'tabpfn']:
            path = RUN / f'jobs/toniot_{method}_n100000/stage_5/metrics.json'
            d = json.loads(path.read_text())
            metric_hashes[str(path.relative_to(RUN))] = sha(path)
            names = [c['name'] for c in d['classes']]
            m = next(c for c in d['classes'] if c['name'] == 'mitm')
            diagnostics[method] = dict(
                mitm_precision=m['precision'], mitm_recall=m['recall'], mitm_fp=m['fp'],
                benign_to_mitm=d['confusion_matrix'][names.index('benign')][names.index('mitm')])
    snapshot = dict(captured_utc=datetime.now(timezone.utc).isoformat(),
                    source_run=str(RUN.relative_to(ROOT)), complete_datasets=complete,
                    source_sha256={k: hashlib.sha256(v).hexdigest() for k, v in raw.items()},
                    diagnostics=diagnostics, diagnostic_source_sha256=metric_hashes,
                    summary=[r for r in summary if r['dataset'] in complete],
                    class_metrics=[r for r in classes if r['dataset'] in complete])
    SOURCE.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + '\n')
snapshot = json.loads(SOURCE.read_text())
prs = Presentation(BASE)
old = list(prs.slides)
assert len(old) == 13


def retitle(slide, title, section):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.top < Inches(.96) and sh.text.strip():
            set_text(sh.text_frame, title, size=28, bold=True, ink=NAVY, inset=0)
        elif sh.left > Inches(9) and sh.top > Inches(7) and sh.width > Inches(1):
            set_text(sh.text_frame, section, size=10, ink='D6E4F3', align=PP_ALIGN.RIGHT)


def replace_exact(slide, old_text, value, **style):
    matches = [sh for sh in slide.shapes if sh.has_text_frame and sh.text == old_text]
    assert len(matches) == 1, (old_text, len(matches))
    set_text(matches[0].text_frame, value, **style)


# 1. Three plain paragraphs, following the user's requested rationale.
s = rebuild(old[0], 'IDS 운영을 위한 TabPFN 활용', '01 연구 배경')
for y, heading, body in [
    (1.53, '기존 평가', '날짜별 공격 시나리오로 수집된 IDS 데이터\n'
                       '전체 공격 유형을 포함해 학습하고, baseline 대비 F1 향상에 집중'),
    (3.04, '실제 운영', '새 공격 라벨의 추가·환경 변화에 따라 모델 갱신·재학습 필요\n'
                       '고정 test의 F1만으로는 실제 운영 적합성을 판단하기 어려움'),
    (4.55, 'TabPFN 활용', '사전학습된 tabular foundation model을 추론에 바로 활용\n'
                          '소규모 데이터의 예측 능력과 가중치 재학습 없는 context 갱신 활용'),
]:
    text(s, .62, y, 2.05, .44, heading, size=21, bold=True, ink=BLUE, inset=0)
    text(s, 2.92, y, 9.79, .95, body, size=20, ink=INK, inset=0, middle=False)
footer(s, '새 공격 사례를 context에 반영하는 탐지기 갱신 방식 검증')
notes(s, '사용자 요청: 복잡한 연구 질문 도식 대신 기존 F1 중심 평가, 실제 운영의 갱신 요구, '
      'TabPFN 활용 이유를 간단히 연결. 전체 데이터라는 말은 test까지 학습했다는 뜻이 아니라 '
      '전체 공격 유형을 포함한 기존 train을 뜻한다. 단순 분류 자체가 권장되지 않는다고 단정하지 않고 '
      '고정 test F1만으로 실사용 적합성을 판단하기 어렵다고 표현. 새로운 라벨·환경 변화는 갱신 필요성을 '
      '설명하는 운영 배경이며 EXP70이 실제 환경 변화를 재현한다는 뜻이 아니다. '
      'TabPFN backbone을 재학습하지 않아도 전처리·cache·추론 계산 비용은 발생. '
      '소규모 데이터 능력의 문헌 근거: Hollmann et al., Nature 2025, '
      'https://www.nature.com/articles/s41586-024-08328-6 . 이 일반적 근거를 현재 IDS 성능·비용 우위로 '
      '단정하지 않는다. 날짜별 시나리오 근거: https://www.unb.ca/cic/datasets/ids-2018.html '
      '및 EXP69의 NF-v3 timestamp 검증 기록.')

# 2. Move the actual evaluation scenario directly after the motivation.
scenario = old[6]
retitle(scenario, '새 공격 사례 추가에 따른 탐지기 갱신 실험', '02 신규 공격 반영')
replace_exact(scenario,
              'EXP70 실행 중 · XGBoost 재학습과 Global TabPFN context 확장 · 동일한 누적 사례 제공',
              'XGBoost 누적 재학습과 Global TabPFN context 확장 · 동일 사례 · Seed 43',
              size=13, ink=GREY, inset=0)
replace_exact(scenario,
              '최종 각 100,000행 · 등장 전 클래스 학습 제외 · 기존 test 유지 · Seed 43',
              '최종 각 100,000행 · 미래 클래스 학습 제외 · 정제 test 중 도입한 클래스 전체 평가',
              size=15, bold=True, ink=NAVY, inset=0)
replace_exact(scenario,
              '원본의 클래스 최초 출현 순서 보존 · 확인 대상: 새 공격 분류 / 기존 탐지 유지 / 갱신·추론 비용',
              '클래스 최초 출현 순서 보존 · 새 공격 성능·기존 탐지 유지·갱신 및 추론 비용 비교',
              size=13, ink=GREY, inset=0)
notes(scenario, 'EXP70: 전체 raw 시간 순서 재현이 아니라 원본의 클래스 최초 출현 순서를 보존한 '
      '통제된 신규 공격 도입. 이미 도입한 클래스의 전체 기존 clean test로 평가하며 future 클래스는 제외. '
      '단계 간 전체 Macro-F1은 클래스 수 변화의 영향을 받으므로 그대로 향상으로 해석하지 않는다. '
      'Global 갱신 비교이며 expert/S/V 완전판 결과와 구별. '
      '프로토콜: docs/research/20261002/exp70_protocol.md . '
      '완료 결과는 이 장 뒤에 데이터셋별로 삽입. 각 데이터셋의 두 방법 전체 단계 완료 확인 후 채택.')


def summary_row(ds, method, stage):
    rows = [r for r in snapshot['summary'] if r['dataset'] == ds and r['method'] == method
            and int(r['stage']) == stage]
    assert len(rows) == 1
    return rows[0]


def class_row(ds, method, stage, name):
    rows = [r for r in snapshot['class_metrics'] if r['dataset'] == ds and r['method'] == method
            and int(r['stage']) == stage and r['name'] == name]
    assert len(rows) == 1
    return rows[0]


result_slides = []
result_checks = {}
for ds in snapshot['complete_datasets']:
    label = 'CIC2018' if ds == 'cic2018' else 'ToN'
    s = new_slide(prs, f'{label} · 신규 공격 반영 성능과 비용', '02 신규 공격 반영')
    subtitle(s, '동일 누적 사례 · Global TabPFN v3 · Seed 43 · RTX 4090 · GPU 작업 순차 실행')
    text(s, .62, 1.55, 6.34, .39, '각 신규 공격의 도입 직후 F1 (%)', size=18, bold=True, ink=NAVY, inset=0)
    text(s, .62, 1.99, 6.34, .31, '정제 test 표본 수 내림차순 · 괄호는 도입 단계', size=11.5, ink=GREY, inset=0)
    ordered = sorted(INTRO[ds], key=lambda c: -int(class_row(ds, 'xgb', INTRO[ds][c], c)['support']))
    rows = []
    for c in ordered:
        stage = INTRO[ds][c]
        pair = [class_row(ds, method, stage, c) for method in ['xgb', 'tabpfn']]
        assert pair[0]['support'] == pair[1]['support']
        rows.append([f'{LABELS[c]} ({stage})', f"{int(pair[0]['support']):,}",
                     *[f"{100*float(r['f1']):.2f}" for r in pair]])
    table(s, .62, 2.42, [2.34, 1.22, 1.39, 1.39],
          ['신규 공격', 'Test 수', 'XGBoost', 'TabPFN'], rows, row_h=.44, font=12.0)
    text(s, 7.31, 1.55, 5.40, .39, '최종 100,000행에서의 성능·비용', size=18, bold=True, ink=NAVY, inset=0)
    final = [summary_row(ds, method, LAST[ds]) for method in ['xgb', 'tabpfn']]
    assert final[0]['eval_rows'] == final[1]['eval_rows']
    text(s, 7.31, 1.99, 5.40, .31, f"전체 test {int(final[0]['eval_rows']):,}행 · 시간 단위: 초",
         size=11.5, ink=GREY, inset=0)
    metrics = [('Macro-F1 (%)', 'macro_f1', 100), ('정상 오탐률 (%)', 'benign_false_alarm_rate', 100),
               ('학습·구성 시간', 'fit_seconds', 1), ('추론 시간', 'predict_seconds', 1)]
    cost_rows = [[title, *[f"{factor*float(r[key]):,.2f}" for r in final]] for title, key, factor in metrics]
    table(s, 7.31, 2.42, [2.40, 1.50, 1.50], ['지표', 'XGBoost', 'TabPFN'], cost_rows, row_h=.49, font=12.4)
    if ds == 'cic2018':
        detail = 'TabPFN: 가중치 고정, 전처리·cache 구성\nXGBoost: 누적 데이터로 매 단계 재학습'
        conclusion = 'Web attack·Infiltration의 도입 직후 F1 개선 · 큰 context의 추론 비용 절감 필요'
    else:
        d = snapshot['diagnostics']['tabpfn']
        detail = (f"TabPFN MitM · recall {100*d['mitm_recall']:.2f}% · precision {100*d['mitm_precision']:.2f}%\n"
                  f"오탐 {d['mitm_fp']:,}건 중 Benign {d['benign_to_mitm']:,}건")
        conclusion = 'MitM 탐지율은 높으나 정상 오탐 다수 · context 구성과 추론 비용 개선 필요'
    text(s, 7.31, 5.11, 5.40, .67, detail, size=12.5, ink=GREY, inset=0)
    text(s, .62, 6.35, 12.09, .39, conclusion, size=16, bold=True, ink=NAVY, inset=0)
    footer(s, '신규 공격 성능과 계산 비용을 함께 평가')
    notes(s, f'EXP70 완료 결과. {snapshot["source_run"]}/summary.csv 및 class_metrics.csv. '
          '왼쪽은 신규 공격별 최초 도입 단계에서의 클래스 F1로 모든 신규 공격을 포함하며 '
          'test support 내림차순. 오른쪽은 최종 단계의 전체 test와 비용. 단계·평가 집합이 다른 '
          '수치를 직접 비교하지 않는다. 기존 공격의 보존은 별도 저장된 previous-class 및 transition '
          '지표로 확인해야 하며 최종 정상 오탐률만으로 모두 유지됐다고 해석하지 않는다. '
          '학습·구성 시간은 fit_context_or_retrain 경과 시간이며 모델 생성·적재 단계와 구별. '
          'TabPFN cache CPU 보관, ensemble 4, subsampling 없음. 비용은 현재 배치 조건의 실측이며 '
          '실시간 서비스 처리량의 직접 검증은 아님. 단일 seed와 개발 holdout 결과. '
          'ToN MitM 오탐은 최종 단계 confusion matrix에서 확인. 미탐보다 precision이 낮은 현상이 '
          '직접 관찰되지만 그 원인이 context 비율·다양성·feature 표현 중 무엇인지는 별도 비교 필요. '
          '신규 공격 라벨 확보 후의 분류로, 라벨 없는 OOD 탐지와 구별. '
          f'결과 snapshot: docs/slides/kor/1002_operational_qa/result_snapshot.json, {snapshot["captured_utc"]}.')
    result_checks[ds] = dict(class_order=ordered, introduction_rows=rows,
                             final_rows=cost_rows, eval_rows=int(final[0]['eval_rows']))
    result_slides.append(s)

# Direct action title for the existing data audit.
retitle(old[2], '데이터 검토: 동일 feature의 상충 라벨 삭제', '03 데이터 검토')
retitle(old[3], '정제 이후 클래스 분포와 고정 test', '03 데이터 검토')
# Remove exactly the lower-left tail explanation card (shape + textbox).
remove = [sh for sh in old[3].shapes if sh.left < Inches(1) and
          Inches(5.65) <= sh.top <= Inches(5.9)]
assert len(remove) == 2
for sh in remove:
    old[3].shapes._spTree.remove(sh._element)
replace_exact(old[3], '평가에서의 클래스 희소성과 context의 학습 사례 수 구분',
              '상충 라벨 제거 전후의 클래스 분포와 고정 평가 표본',
              size=11, bold=True, ink='FFFFFF', inset=0)

retitle(old[4], 'Context 기반 Global·expert 모델 구조', '04 현재 모델')
retitle(old[5], '정제 데이터의 현재 성능 · SOTA 비교', '04 현재 모델')
replace_exact(old[5],
              '현재 성능 수준의 확인 · 신규 공격 반영 효과와 갱신·추론 비용은 순차 도입으로 검증',
              '정적 평가에서의 현재 성능 · 순차 도입의 Global 비교와 별도로 해석',
              size=17, bold=True, ink=NAVY, inset=0)
old[5].notes_slide.notes_text_frame.text = old[5].notes_slide.notes_text_frame.text.replace(
    '운영 비용은 부록의 측정 조건을 함께 설명한다.', '운영 비용은 다음 장의 측정 조건과 함께 설명한다.')

# Former cost appendix now follows static performance in the main body.
retitle(old[10], '정제 데이터의 SOTA 비용', '04 현재 모델')
cost_table = next(sh.table for sh in old[10].shapes if sh.has_table)
for col, title in [(1, 'CIC 학습·구성\n초'), (3, 'ToN 학습·구성\n초')]:
    set_text(cost_table.cell(0, col).text_frame, title, size=11, bold=True,
             ink='FFFFFF', align=PP_ALIGN.CENTER, inset=.04)
old[10].notes_slide.notes_text_frame.text += ('\n현재 정적 평가의 비용을 본문에 유지. '
    'TabPFN 학습·구성은 backbone 재학습이 아닌 전처리·cache 계산을 포함. '
    'EXP70 순차 GPU 비용과 기존 EXP61 병렬 작업 비용을 같은 실행 조건으로 비교하지 않는다.')
retitle(old[7], '후속 검증 · 새로운 공격을 반영하는 시스템', '05 후속 검증')
if 'toniot' in snapshot['complete_datasets']:
    replace_exact(old[7],
                  '새 공격 사례를 반영하는 방식의 효용 확인 · 성능·기존 탐지 유지·운영 비용의 연결',
                  'CIC2018 일부 신규 공격에 이득 · ToN 정상 오탐과 두 데이터셋의 계산 비용 개선 필요',
                  size=13, ink=GREY, inset=0)
    replace_exact(old[7],
                  'Context 구성 개선 → Expert·S/V의 역할과 갱신 검증 → 운영 조건별 효과 확인',
                  '정상 사례 다양성·context 규모 조절 → 오탐·계산 비용 비교 → Expert·S/V 갱신 검증',
                  size=16, bold=True, ink=BLUE, inset=0)
    old[7].notes_slide.notes_text_frame.text += (
        '\nEXP70 완료 후: CIC2018 최종 Macro-F1은 TabPFN이 높고 ToN은 XGB가 높다. '
        '두 데이터셋 모두 현재 100k 배치 조건의 fit 및 추론 시간은 XGB가 짧다. '
        '따라서 가중치 재학습 불필요를 실제 비용 우위로 대체하지 않는다. '
        '후속 context 구성에서는 정상 대표성을 유지하며 오탐과 처리 비용을 함께 줄이는지 검증한다. '
        '이는 다음 검증 가설이며 이번 결과가 개선 방법의 효능까지 입증한 것은 아니다.')

# Drop old slides 2, 9, 10, 12, 13; keep the system-level closing after SOTA cost.
order = [old[0], scenario, *result_slides, old[2], old[3], old[4], old[5], old[10], old[7]]
slide_ids = {id(prs.slides[i]): el for i, el in enumerate(prs.slides._sldIdLst)}
chosen = [slide_ids[id(s)] for s in order]
for el in list(prs.slides._sldIdLst):
    prs.slides._sldIdLst.remove(el)
for el in chosen:
    prs.slides._sldIdLst.append(el)
live_rel_ids = {el.rId for el in chosen}
for rel in list(prs.part.rels.values()):
    if rel.reltype.endswith('/slide') and rel.rId not in live_rel_ids:
        prs.part.drop_rel(rel.rId)
for i, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if sh.has_text_frame and abs(sh.left - Inches(12.23)) < Inches(.05) and abs(sh.top - Inches(7.01)) < Inches(.05):
            set_text(sh.text_frame, str(i), size=11, bold=True, ink='FFFFFF', align=PP_ALIGN.RIGHT)
prs.core_properties.title = 'IDS 운영을 위한 TabPFN 활용'
prs.core_properties.comments = ('Korean only. Simplified motivation, arrival evidence near the opening, '
                                'data cleaning, static model performance/cost, system-level follow-up.')
prs.save(OUTPUT)

issues, dump, titles = [], [], []
numeric = 0
for i, slide in enumerate(prs.slides, 1):
    chunks, text_shapes = [], []
    for sh in slide.shapes:
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > prs.slide_width + 20 or sh.top + sh.height > prs.slide_height + 20:
            issues.append((i, 'outside', sh.name))
        if sh.top < Inches(6.9) and sh.top + sh.height > Inches(6.9):
            issues.append((i, 'footer overlap', sh.name))
        if sh.has_text_frame and sh.text.strip():
            chunks.append(sh.text); text_shapes.append(sh)
            if sh.top < Inches(.96):
                titles.append(sh.text)
        if sh.has_table:
            for ri, row in enumerate(sh.table.rows):
                chunks.append(' | '.join(c.text for c in row.cells))
                for ci, cell in enumerate(row.cells):
                    if ri and ci and NUM.fullmatch(cell.text):
                        numeric += 1
                        assert all(p.alignment == PP_ALIGN.RIGHT for p in cell.text_frame.paragraphs)
    for ai, a in enumerate(text_shapes):
        for b in text_shapes[ai + 1:]:
            if (min(a.left + a.width, b.left + b.width) - max(a.left, b.left) > Inches(.005)
                    and min(a.top + a.height, b.top + b.height) - max(a.top, b.top) > Inches(.005)):
                issues.append((i, 'text overlap', a.text[:20], b.text[:20]))
    joined = '\n'.join(chunks)
    assert not re.search(r'합니다|입니다|습니다|[—–]|상대사례|부록|Tail: 고정 test', joined), (i, 'stale wording')
    dump.append(f'## {i}\n{joined}\n')
assert not issues, issues
assert len(prs.slides) == 8 + len(result_slides)
assert all(sha(ROOT / p) == value for p, value in english_before.items())
(QA / 'slide_text.md').write_text('\n'.join(dump))
(QA / 'validation.json').write_text(json.dumps(dict(
    slides=len(prs.slides), appendix=0, titles=titles, geometry_issues=issues,
    numeric_cells_right_aligned=numeric, english_unchanged=True,
    removed_previous_pages=[2, 9, 10, 12, 13], result_checks=result_checks,
    result_snapshot=str(SOURCE.relative_to(ROOT)), generator=str(Path(__file__).relative_to(ROOT)),
    sha256=sha(OUTPUT)), ensure_ascii=False, indent=2) + '\n')
print(OUTPUT, 'slides=', len(prs.slides), 'completed results=', snapshot['complete_datasets'])
