"""Validate saved predictions and write the Korean EXP86 report automatically."""
from pathlib import Path
import sys
_root=next(p for p in Path(__file__).resolve().parents if (p/'scripts/common/experiment_paths.py').is_file())
sys.path.insert(0,str(_root/'scripts/common'))
from experiment_paths import bootstrap,repo_root
bootstrap(_root)
ROOT=repo_root(__file__)
import argparse,json
from datetime import datetime
from zoneinfo import ZoneInfo
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from exp61_utils import write,sha
from lablog import write_record,build_report


def main():
    p=argparse.ArgumentParser();p.add_argument('--run',required=True,type=Path);a=p.parse_args()
    out=a.run.resolve();cfg=json.loads((out/'config.json').read_text())
    saved=json.loads((out/'results.json').read_text()) if (out/'results.json').exists() else {'state':'preparing','jobs':[]}
    jobs=saved['jobs'];date=cfg['report_date'];now=datetime.now(ZoneInfo('Asia/Seoul')).isoformat()
    rows=[];perclass=[];checks=[]
    for j in jobs:
        names=j['class_names'];benign=names.index('benign');att=[c for c in range(len(names)) if c!=benign]
        path=Path(j['predictions_path']);pred=np.load(path,mmap_mode='r')
        meta=json.loads((out/'data'/j['dataset']/'manifest.json').read_text())
        y=np.load(Path(meta['test_base'])/'test_y.npy',mmap_mode='r')
        C=len(names);cm=np.bincount(np.asarray(y,dtype='int64')*C+pred,minlength=C*C).reshape(C,C)
        assert cm.tolist()==j['phase_b']['confusion'] and cm[:,benign].sum()==0
        cm[benign,:]=0;assert cm.tolist()==j['phase_a']['confusion']
        assert len(j['train_source_ids'])==j['n']*len(att) and len(set(j['train_source_ids']))==len(j['train_source_ids'])
        rows.append(dict(dataset=j['dataset'],seed=j['seed'],n=j['n'],method=j['method'],
            attack_f1=j['phase_a']['macro_f1'],attack_recall=np.mean([j['phase_a']['classes'][c]['recall'] for c in att]),
            normal_included_attack_f1=np.mean([j['phase_b']['classes'][c]['f1'] for c in att]),
            all_macro_f1=j['phase_b']['macro_f1'],feature_disjoint_attack_f1=j['feature_disjoint_attack']['macro_f1'],
            preparation_seconds=j['preparation_seconds'],inference_seconds=j['inference_seconds'],
            benign_fpr=j['phase_b']['benign_false_positive_rate'],overlap_attack_rows=j['exact_feature_overlap_attack_rows'],
            train_unique_features=j['training_unique_feature_rows'],xgb_tree_splits=j['xgb_tree_splits']))
        for c in att:
            perclass.append(dict(dataset=j['dataset'],seed=j['seed'],n=j['n'],method=j['method'],class_name=names[c],
                support=j['phase_a']['classes'][c]['support'],f1=j['phase_a']['classes'][c]['f1'],
                recall=j['phase_a']['classes'][c]['recall'],precision=j['phase_a']['classes'][c]['precision'],
                normal_included_f1=j['phase_b']['classes'][c]['f1'],normal_included_precision=j['phase_b']['classes'][c]['precision']))
        checks.append(dict(job=path.parent.name,confusion_recomputed=True,benign_predictions=0))
    frame=pd.DataFrame(rows);pc=pd.DataFrame(perclass)
    text=["# EXP86 · 극소수 공격 사례의 TabPFN·XGB 분류 비교", "## 목적과 실행 상태",
          f"연구 버전 v7의 EXP86은 갱신·버퍼 삭제 이전의 소수 공격 분류 능력을 확인한다. 실행 시작일은 {date}, 현재 기록 시각은 {now}다. 상태는 {saved['state']}, 완료 조건은 {len(jobs)}/100이다. 실행 경로는 `{out.relative_to(ROOT)}`이다.",
          "## 공통 설정과 평가 의미",
          "CIC2018과 ToN-IoT를 독립적으로 평가하며 공격당 1·2·5·10·20행, seed 43–47을 사용한다. 모든 공격 클래스에서 같은 수를 뽑고 정상 학습 사례는 0개다. 전체 정제 train에서 클래스당 20개를 추출한 뒤 앞 n개를 사용하므로 표본 집합이 중첩된다. 초기 1000행이나 기존 IDS fine-tuning 결과는 사용하지 않는다.",
          "같은 seed·n의 XGB와 PFN은 동일한 행 ID를 사용한다. PFN은 기존 v3 checkpoint, 4 estimators, fit_with_cache, CPU cache이며 가중치를 학습하지 않는다. XGB는 300 trees, depth 8, learning rate 0.05, subsample/colsample 0.8, min_child_weight 1, CUDA hist의 기존 설정을 고정한다. 극소 표본에서 이 설정으로 split이 불가능한지 실제 split 수를 기록한다. 이 문제가 나타나면 해당 n의 결과를 최적화된 XGB 대비 우위로 주장하지 않는다.",
          "Phase A는 고정 test 중 공격만의 Macro-F1·클래스별 F1/recall이다. Phase B는 같은 모델과 예측을 정상 포함 모집단에서 집계하며 정상으로부터 발생한 공격 클래스 false positive가 포함된다. 정상 출력 클래스가 없어 정상 recall 0 및 정상 FPR 100%는 구조적으로 결정된다. 따라서 B는 실제 정상 탐지 능력의 증거가 아니다.",
          "시험 순서는 seed 43의 라벨과 무관한 혼합 순서, batch 8192로 고정한다. 두 phase는 같은 query 배치에서 얻은 예측을 사용하므로 B의 영향은 평가 모집단 변경이며 query 구성 변경의 인과 효과가 아니다. 기존 특징은 행별 zero fill과 고정 clipping을 사용하고 모델의 데이터 의존 전처리는 제공된 소수 행에서만 적합한다. 기존 정제 과정과 이미 검토한 고정 test의 영향을 포함하므로 탐색적 결과로 해석한다.",
          "준비시간은 매 조건의 새 모델 생성과 XGB fit 또는 PFN cache 준비를 포함한다. 순차 후속 갱신 비용은 측정하지 않았으며 추론과 파일 저장·검증은 준비시간에서 제외한다. 추론시간은 정상 포함 전체 test를 처리한 시간이다. 같은 장비라도 첫 조건의 라이브러리 초기화 등 cold-start 영향은 별도로 해석한다."]
    if rows:
        text.extend(['## 실제 제공·평가 표본 수',
            '| 데이터 | 공격 클래스 | 학습/context 행 수 (n=1/2/5/10/20) | 공격 test | 정상 test | 전체 test |',
            '|---|---:|---|---:|---:|---:|'])
        for ds in ['cic2018','toniot']:
            sample=next((j for j in jobs if j['dataset']==ds),None)
            if sample is None:continue
            b=sample['class_names'].index('benign');cm=np.asarray(sample['phase_b']['confusion']);total=int(cm.sum());normal=int(cm[b].sum());c=len(cm)-1
            sizes='/'.join(str(c*n) for n in [1,2,5,10,20])
            text.append(f'| {ds} | {c} | {sizes} | {total-normal:,} | {normal:,} | {total:,} |')
        frame.to_csv(out/'exp86_values.csv',index=False);pc.to_csv(out/'exp86_per_class.csv',index=False)
        keys=['dataset','n','method'];summary=frame.groupby(keys).agg(['mean','std','count']).reset_index()
        summary.to_csv(out/'exp86_summary.csv',index=False)
        piv=frame.pivot(index=['dataset','seed','n'],columns='method',values='attack_f1').dropna()
        if set(['xgb','pfn'])<=set(piv):
            piv['pfn_minus_xgb']=piv.pfn-piv.xgb;piv.to_csv(out/'exp86_paired.csv')
        # Check matched rows for every completed pair.
        index={(j['dataset'],j['seed'],j['n'],j['method']):j for j in jobs}
        for (ds,seed,n,method),j in index.items():
            other=index.get((ds,seed,n,'xgb' if method=='pfn' else 'pfn'))
            if other:assert other['train_source_ids']==j['train_source_ids']
        figures=ROOT/'docs/lablog/figures';figures.mkdir(exist_ok=True)
        colors={'xgb':'#222222','pfn':'#0072B2'}
        for filename,metric,label in [('quality','attack_f1','Attack-only Macro-F1 (%)'),
                ('phase_b','normal_included_attack_f1','Attack-class F1, benign included (%)'),
                ('feature_disjoint','feature_disjoint_attack_f1','Feature-disjoint attack Macro-F1 (%)')]:
            fig,axes=plt.subplots(1,2,figsize=(10,3.5))
            for ax,ds in zip(axes,['cic2018','toniot']):
                for method in ['xgb','pfn']:
                    data=frame[(frame.dataset==ds)&(frame.method==method)]
                    if data.empty:continue
                    z=data.groupby('n')[metric].agg(['mean','std'])
                    ax.errorbar(z.index,z['mean']*100,yerr=z['std'].fillna(0)*100,color=colors[method],marker='o',capsize=3,label=method.upper())
                ax.set(title=ds,xlabel='Labeled rows per attack class',ylabel=label,xticks=[1,2,5,10,20],ylim=(0,100));ax.grid(alpha=.2);ax.legend()
            fig.tight_layout();fig.savefig(figures/f'exp86_{filename}.pdf');plt.close(fig)
            text.extend([f"## {label}",f"![{label}; mean ± sample SD, incomplete cells use available seeds](../../figures/exp86_{filename}.pdf)"])
        fig,axes=plt.subplots(2,2,figsize=(10,6))
        for col,ds in enumerate(['cic2018','toniot']):
            for row,(metric,label) in enumerate([('preparation_seconds','Preparation (seconds)'),('inference_seconds','Full-test inference (seconds)')]):
                ax=axes[row,col]
                for method in ['xgb','pfn']:
                    z=frame[(frame.dataset==ds)&(frame.method==method)].groupby('n')[metric].agg(['mean','std'])
                    if len(z):ax.errorbar(z.index,z['mean'],yerr=z['std'].fillna(0),color=colors[method],marker='o',capsize=3,label=method.upper())
                ax.set(title=ds,xlabel='Rows per attack class',ylabel=label,xticks=[1,2,5,10,20]);ax.grid(alpha=.2);ax.legend()
        fig.tight_layout();fig.savefig(figures/'exp86_cost.pdf');plt.close(fig)
        text.extend(['## 비용','![초기 준비와 전체 추론을 분리한 비용](../../figures/exp86_cost.pdf)',
            '## 수치 요약','평균과 표본 표준편차(ddof=1)는 고정 분할에서 선택 표본·모델 seed를 바꾼 반복이다. 신뢰구간이나 통계적 동등성을 뜻하지 않는다. 부분 완료에서는 완료 seed 수를 함께 표시하며 서로 다른 반복 수의 곡선을 확정 비교하지 않는다.',
            '| 데이터 | n | 모델 | 반복 수 | 공격 F1 (%) | 준비 (초) | 추론 (초) |','|---|---:|---|---:|---:|---:|---:|'])
        for (ds,n,method),g in frame.groupby(keys):
            fmt=lambda col,mul=1:f'{g[col].mean()*mul:.2f} ± {g[col].std()*mul:.2f}' if len(g)>1 else f'{g[col].mean()*mul:.2f} (1회)'
            text.append(f'| {ds} | {n} | {method} | {len(g)} | {fmt("attack_f1",100)} | {fmt("preparation_seconds")} | {fmt("inference_seconds")} |')
        fig,axes=plt.subplots(2,2,figsize=(12,8))
        for row,ds in enumerate(['cic2018','toniot']):
            order=pc[pc.dataset==ds].groupby('class_name').support.max().sort_values(ascending=False).index
            for col,method in enumerate(['xgb','pfn']):
                ax=axes[row,col];data=pc[(pc.dataset==ds)&(pc.method==method)]
                z=data.pivot_table(index='class_name',columns='n',values='f1',aggfunc='mean').reindex(index=order,columns=[1,2,5,10,20])*100
                ax.imshow(z.to_numpy(),vmin=0,vmax=100,cmap='Blues',aspect='auto')
                ax.set(xticks=range(5),xticklabels=z.columns,yticks=range(len(z)),yticklabels=z.index,title=f'{ds}: {method.upper()} F1 (%)',xlabel='Rows per class')
                for i in range(len(z)):
                    for k in range(5):
                        v=z.iloc[i,k]
                        if np.isfinite(v):ax.text(k,i,f'{v:.1f}',ha='center',va='center',color='white' if v>60 else 'black',fontsize=8)
        fig.tight_layout();fig.savefig(figures/'exp86_per_class.pdf');plt.close(fig)
        text.extend(['## 클래스별 결과','![공격 test 표본 수 내림차순; 클래스별 F1 평균](../../figures/exp86_per_class.pdf)',
            '## 근거에 따른 해석과 후속 실험'])
        for ds,g in frame.groupby('dataset'):
            available=sorted(set(g.n))
            text.append(f'{ds}에서 현재 집계한 n은 {available}이다. 동일한 seed·n의 PFN−XGB 차이는 exp86_paired.csv에, 클래스별 수치는 exp86_per_class.csv에 보존한다.')
            for n in [1,20]:
                pair=piv.loc[(ds,slice(None),n),:] if len(piv) and ds in piv.index.get_level_values(0) and n in piv.loc[ds].index.get_level_values('n') else pd.DataFrame()
                if len(pair):text.append(f'n={n}의 대응 PFN−XGB 차이는 {pair.pfn_minus_xgb.mean()*100:.2f} pp, 완료된 대응 반복은 {len(pair)}회다.')
            for method in ['xgb','pfn']:
                for n in [5,20]:
                    point=g[(g.method==method)&(g.n==n)]
                    if point.empty:continue
                    weakest=pc[(pc.dataset==ds)&(pc.method==method)&(pc.n==n)].groupby('class_name').f1.mean().sort_values().head(3)
                    detail=', '.join(f'{name} {score*100:.2f}%' for name,score in weakest.items())
                    text.append(f'{ds} {method.upper()} n={n}: 공격 Macro-F1 평균 {point.attack_f1.mean()*100:.2f}% (완료 {len(point)} seeds). 클래스별 F1이 낮은 3종은 {detail}다. 평균 상승과 약한 클래스의 개선 여부를 분리해 본다.')
        text.extend(['## 정상 포함 시 오분류 유입 (n=20)',
            '이 표는 정상 query가 어느 공격으로 분류되는지를 나타낸다. 모든 모델이 정상 출력 라벨을 가지지 않으므로 비율 합은 100%이며, 정상 탐지 능력의 순위로 해석하지 않는다. 정상에서 유입된 false positive 때문에 Phase B의 공격 precision/F1은 Phase A와 달라진다.',
            '| 데이터 | 모델 | 공격 클래스 | 정상 중 해당 공격으로 분류 (%) | 완료 seeds |',
            '|---|---|---|---:|---:|'])
        for ds in ['cic2018','toniot']:
            for method in ['xgb','pfn']:
                selected=[j for j in jobs if j['dataset']==ds and j['method']==method and j['n']==20]
                if not selected:continue
                names=selected[0]['class_names'];b=names.index('benign')
                rates=np.array([np.asarray(j['phase_b']['confusion'])[b]/sum(j['phase_b']['confusion'][b]) for j in selected]).mean(axis=0)
                supports=np.array(selected[0]['phase_b']['confusion']).sum(axis=1)
                for c in np.argsort(-supports):
                    if c!=b:text.append(f'| {ds} | {method} | {names[c]} | {rates[c]*100:.2f} | {len(selected)} |')
        text.append('행 순서는 각 데이터의 고정 공격 test 표본 수 내림차순이다. 위 비율의 분모는 해당 데이터의 정상 test 수다.')
        stumps=frame[(frame.method=='xgb')&(frame.xgb_tree_splits==0)]
        if len(stumps):
            text.append(f'XGB의 실제 split 수가 0인 조건이 {len(stumps)}개 발견되었다. 극소 표본에서 min_child_weight와 subsampling의 제약이 영향을 줄 수 있으므로, 해당 점수는 고정된 기존 XGB 설정의 결과다. 다음 비교에서는 작은 표본에 적합한 min_child_weight/subsample 설정을 사전에 정한 대조군으로 추가해 모델 차이와 설정 제약을 분리해야 한다. 이 대조군은 이번 승인된 100조건에 자동 추가하지 않는다.')
        text.append('PFN의 공격 분류가 유지되는 조건은 정상·기존 공격을 함께 둔 순차 context 갱신으로 연결한다. 처음부터 약한 클래스는 무작위 대 대표성·다양성 선택 및 특징 구분 가능성을 먼저 검토한다. 두 모델 모두 잘하는 경우에는 동일 품질에서의 비용과 보존 사례 수를 확인한다. 이는 후속 제안이며 새 실험 실행이나 모델 수정은 아직 수행하지 않았다.')
        text.append('특징이 정확히 같은 학습 사례가 있는 test 행을 제외한 진단은 n·seed에 따라 평가 모집단이 달라진다. 따라서 전체 고정 test 결과를 대체하거나 그 곡선의 변화만으로 인과관계를 주장하지 않는다. feature hash는 기존 정제 파이프라인의 값이며 근접 중복·동일 세션 누수까지 제거했다는 의미는 아니다.')
    else:
        text.extend(['## 대기 또는 실패 상태','아직 완료된 모델 조건이 없다. 실행 경로의 progress.json, stdout.log 및 존재할 경우 FAILURE.json으로 현재 상태와 원인을 확인한다.'])
    text.extend(['## 재현 근거',f'설정·환경·데이터 ID·query 순서·예측·혼동행렬·비용·소스 사본은 `{out.relative_to(ROOT)}`에 있다. 원고와 슬라이드에는 결과를 자동 채택하지 않는다.'])
    target=write_record(ROOT/'docs/lablog/evidence'/date/'exp86_record', '\n\n'.join(text).replace('|\n\n|', '|\n|'),
        summary=f'극소수 공격 1/2/5/10/20행, 5 seeds, PFN·XGB; {saved["state"]} {len(jobs)}/100조건; 공격/정상 포함 집계 분리',compile_pdf=False)
    validation=build_report(target)
    write(out/'report_validation.json',dict(**validation,checked_at=now,validated_jobs=len(jobs),checks=checks))
    if saved['state']=='complete' and len(jobs)==100:
        # Close the approved plan's status after the numerical report is verified.
        import openpyxl
        workbook=ROOT/'docs/lablog/evidence/references/tabpfn_papers_survey.xlsx'
        wb=openpyxl.load_workbook(workbook);sheet=wb['20261021_선행실험안']
        sheet.cell(2,3,f'EXP86 완료: 100/100 조건, 완료 확인 {now}; 상세 결과는 20261021_exp86.tex')
        sheet.cell(18,3,f'100조건 완료 및 최종 보고서 임시 컴파일 검증 완료 ({now})')
        tmp=workbook.with_name(workbook.stem+'.exp86.tmp.xlsx');wb.save(tmp);tmp.replace(workbook)
        write(out/'workbook_completion.json',dict(updated_at=now,workbook=str(workbook.relative_to(ROOT)),cells=['20261021_선행실험안!C2','20261021_선행실험안!C18'],sha256=sha(workbook)))
    print(json.dumps(dict(report=str(target),jobs=len(jobs),state=saved['state']),ensure_ascii=False),flush=True)


if __name__=='__main__':main()
